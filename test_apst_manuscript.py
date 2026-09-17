import os
import docx

def test_apst_manuscript():
    docx_path = '12_manuscript/APST_Long_Term_Rainfall_Occurrence_Regimes_Northeastern_Thailand.docx'
    assert os.path.exists(docx_path), f"File {docx_path} does not exist"
    size = os.path.getsize(docx_path)
    assert size > 0, "DOCX file is empty"

    doc = docx.Document(docx_path)
    assert len(doc.paragraphs) > 0, "DOCX paragraph count is 0"

    # Title check
    title = doc.paragraphs[0].text.strip()
    expected_title = "Spatial Heterogeneity and Temporal Stability of Daily Rainfall Occurrence Regimes in Northeastern Thailand"
    assert title == expected_title, f"Title mismatch. Got: {title}"

    # Tables check
    assert len(doc.tables) == 4, f"Expected 4 tables, got {len(doc.tables)}"

    # Images check
    img_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.target_ref)
    assert img_count == 5, f"Expected 5 embedded figures, got {img_count}"

    # Text validation: check sections & no TODOs
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Abstract" in full_text
    assert "1. Introduction" in full_text
    assert "2. Materials and Methods" in full_text
    assert "3. Results and Discussion" in full_text
    assert "4. Conclusions" in full_text
    assert "References" in full_text

    for word in ["TODO", "DRAFT", "PLACEHOLDER"]:
        assert word not in full_text, f"Found draft marker: {word}"

    print("ALL MANUSCRIPT TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    test_apst_manuscript()
