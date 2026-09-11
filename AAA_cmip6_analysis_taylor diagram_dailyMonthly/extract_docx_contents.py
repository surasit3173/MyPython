from pathlib import Path
from zipfile import ZipFile
import json

from docx import Document


FILES = [
    Path(r"D:\วารสารบูรพา\ID1558_Manuscript_BSJ (11).docx"),
    Path(r"D:\วารสารบูรพา\ID1558_correction_plan (1).docx"),
    Path(r"D:\วารสารบูรพา\Evaluation of Quantile Delta Mapping for Bias Correction of CMIP6 Daily Rainfall Using Independent Temporal Cross-Validation in Prachuap Khiri Khan Province.docx"),
]


def iter_block_items(doc):
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "paragraph", child
        elif tag == "tbl":
            yield "table", child


def para_text(p_elm):
    return "".join(node.text or "" for node in p_elm.iter() if node.tag.rsplit("}", 1)[-1] == "t").strip()


def table_from_elm(doc, tbl_elm):
    for i, table in enumerate(doc.tables, start=1):
        if table._element is tbl_elm:
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip().replace("\n", " | ") for cell in row.cells])
            return i, rows
    return None, []


def image_count(path):
    with ZipFile(path) as zf:
        return len([n for n in zf.namelist() if n.startswith("word/media/")])


def extract(path):
    doc = Document(path)
    blocks = []
    para_i = 0
    for kind, elm in iter_block_items(doc):
        if kind == "paragraph":
            text = para_text(elm)
            if not text:
                continue
            para_i += 1
            style = ""
            try:
                for p in doc.paragraphs:
                    if p._element is elm:
                        style = p.style.name if p.style else ""
                        break
            except Exception:
                style = ""
            blocks.append({"type": "p", "n": para_i, "style": style, "text": text})
        else:
            table_i, rows = table_from_elm(doc, elm)
            blocks.append({"type": "table", "n": table_i, "rows": rows})

    return {
        "file": str(path),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": image_count(path),
        "blocks": blocks,
    }


def main():
    out_dir = Path("docx_extracts")
    out_dir.mkdir(exist_ok=True)
    index = []
    for path in FILES:
        data = extract(path)
        safe = path.stem[:80].replace(" ", "_")
        out = out_dir / f"{safe}.json"
        out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        index.append({"file": str(path), "out": str(out), "paragraphs": data["paragraphs"], "tables": data["tables"], "images": data["images"]})
    (out_dir / "index.json").write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(index, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
