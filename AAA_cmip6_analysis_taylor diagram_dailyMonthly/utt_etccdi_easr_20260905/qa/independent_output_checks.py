"""Read-only checks of a completed Uttaradit analysis output directory.

This intentionally does not import the production module.  It verifies table
contracts and independently recomputes two raw daily-data quantities used in
the reported change table.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESULTS = (
    Path(sys.argv[1]).resolve()
    if len(sys.argv) == 2
    else ROOT / "outputs" / "full_run" / "results"
)
EXPECTED_STATIONS = {
    "351001", "351002", "351003", "351004", "351005", "351006", "351007",
    "351008", "351009", "351010", "351011", "351012", "351201",
}
EXPECTED_INDICES = {
    "PRCPTOT", "SDII", "Rx1day", "Rx5day", "CDD", "CWD",
    "R10mm", "R20mm", "R50mm", "R95p", "R99p",
}


def read_csv(name: str) -> pd.DataFrame:
    path = RESULTS / name
    assert path.is_file() and path.stat().st_size > 0, f"missing/empty: {name}"
    return pd.read_csv(path)


def daily_series(path: str, station: str) -> pd.Series:
    raw = pd.read_csv(path)
    dates = pd.to_datetime(dict(year=raw.YEAR, month=raw.MONTH, day=raw.DAY), errors="raise")
    values = pd.to_numeric(raw[station], errors="raise").to_numpy(dtype=float)
    return pd.Series(values, index=dates).sort_index()


def main() -> None:
    meta = json.loads((RESULTS / "analysis_metadata.json").read_text(encoding="utf-8"))
    assert meta["run_mode"] == "full_analysis"
    assert meta["observed_trend_rows"] == 13 * 11
    assert meta["future_trend_rows"] == 7 * 2 * 13 * 11
    assert meta["future_change_rows"] == 7 * 2 * 13 * 11

    manifest = read_csv("cmip6_file_manifest.csv")
    assert len(manifest) == 21
    assert not manifest.duplicated(["Model", "Dataset_period"]).any()
    assert set(manifest.Status) == {"OK"}

    coordinates = read_csv("station_coordinates_used.csv")
    assert set(coordinates.station.astype(str)) == EXPECTED_STATIONS
    assert len(coordinates) == 13
    assert coordinates[["latitude", "longitude"]].apply(np.isfinite).all().all()

    thresholds = read_csv("station_thresholds.csv")
    assert set(thresholds.Station.astype(str)) == EXPECTED_STATIONS
    assert (thresholds[["P95_threshold_mm", "P99_threshold_mm"]] > 0).all().all()
    assert (thresholds.P99_threshold_mm >= thresholds.P95_threshold_mm).all()

    observed_indices = read_csv("observed_indices_1981_2014.csv")
    assert len(observed_indices) == 13 * 34 * 11
    assert set(observed_indices.Station.astype(str)) == EXPECTED_STATIONS
    assert set(observed_indices.Index) == EXPECTED_INDICES
    assert observed_indices.Year.between(1981, 2014).all()
    assert set(observed_indices.Annual_status) == {"OK"}
    assert (observed_indices.Value >= 0).all()

    future_indices = read_csv("future_indices_2021_2050.csv")
    assert len(future_indices) == 7 * 2 * 13 * 30 * 11
    assert future_indices.Year.between(2021, 2050).all()
    assert set(future_indices.Annual_status) == {"OK"}
    assert (future_indices.Value >= 0).all()

    for name, expected_n in [("observed_trend_1981_2014.csv", 34), ("future_trend_2021_2050.csv", 30)]:
        trends = read_csv(name)
        assert set(trends.Index) == EXPECTED_INDICES
        assert (trends.N == expected_n).all(), f"unexpected N in {name}"
        for col in ("p_MK", "p_MMK2004"):
            values = trends[col].to_numpy(dtype=float)
            assert np.isfinite(values).all() and ((0 <= values) & (values <= 1)).all()
        assert (trends.Sen_CI95_low <= trends.Sen_slope).all()
        assert (trends.Sen_slope <= trends.Sen_CI95_high).all()
        assert (trends.n_effective > 0).all()

    changes = read_csv("future_change_2021_2050.csv")
    assert len(changes) == 7 * 2 * 13 * 11
    assert set(changes.Baseline_Period) == {"1995-2014"}
    assert set(changes.Future_Period) == {"2021-2050"}
    assert np.isfinite(changes[["Baseline_mean", "Future_mean", "Absolute_change"]]).all().all()
    assert np.allclose(changes.Absolute_change, changes.Future_mean - changes.Baseline_mean)

    agreement = read_csv("future_model_agreement.csv")
    assert len(agreement) == 2 * 13 * 11
    assert (agreement.n_models_available == 7).all()
    assert agreement[["sign_agreement_pct", "sig_agreement_pct"]].apply(
        lambda col: col.between(0, 100).all()
    ).all()

    # Independent raw-data check: model-consistent PRCPTOT annual means.
    model, scenario, station = "ACCESS-ESM1-5", "SSP2-4.5", "351001"
    historical_path = manifest.loc[
        (manifest.Model == model) & (manifest.Dataset_period == "historical"), "File"
    ].item()
    future_path = manifest.loc[
        (manifest.Model == model) & (manifest.Dataset_period == "ssp245"), "File"
    ].item()
    hist = daily_series(historical_path, station)
    future = daily_series(future_path, station)
    # ETCCDI PRCPTOT sums wet days (RR >= 1 mm), not trace precipitation.
    hist_wet = hist.loc["1995":"2014"].where(hist.loc["1995":"2014"] >= 1.0, 0.0)
    future_wet = future.loc["2021":"2050"].where(future.loc["2021":"2050"] >= 1.0, 0.0)
    expected_baseline = hist_wet.groupby(hist_wet.index.year).sum().mean()
    expected_future = future_wet.groupby(future_wet.index.year).sum().mean()
    row = changes.loc[
        (changes["Model"] == model)
        & (changes["Scenario"] == scenario)
        & (changes["Station"].astype(str) == station)
        & (changes["Index"] == "PRCPTOT")
    ].iloc[0]
    assert np.isclose(row.Baseline_mean, expected_baseline)
    assert np.isclose(row.Future_mean, expected_future)

    workbook = load_workbook(RESULTS / "Uttaradit_Trend_Analysis.xlsx", read_only=True, data_only=True)
    required_sheets = {
        "Metadata", "Station Coordinates", "Station Thresholds", "Observed Indices",
        "Observed Trend", "Future Indices", "Future Trend", "Future Change",
        "Model Summary", "QC Observed", "QC Models", "File Manifest",
    }
    assert required_sheets.issubset(set(workbook.sheetnames))
    workbook.close()
    print("INDEPENDENT_OUTPUT_CHECKS_OK")


if __name__ == "__main__":
    main()
