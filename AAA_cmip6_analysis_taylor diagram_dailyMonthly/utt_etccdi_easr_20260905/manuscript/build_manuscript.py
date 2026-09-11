"""Build a transparent research draft from the supplied EASR template."""
from pathlib import Path
from copy import deepcopy
import shutil
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r'C:\MyPython\CMIP6Uttaradit\paper3\EASR+Template+2026.docx')
OUT = ROOT / 'manuscript' / 'EASR_Uttaradit_research_draft.docx'
R = ROOT / 'outputs' / 'full_with_figures' / 'results'
shutil.copy2(TEMPLATE, OUT)
doc = Document(OUT)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
for name in ['Normal', 'Els-body-text', 'Els-1storder-head', 'Els-2ndorder-head', 'Els-Abstract-text', 'Els-Abstract-head']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
if 'Title' not in doc.styles:
    doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
doc.styles['Title'].font.name = 'Times New Roman'
doc.styles['Title'].font.size = Pt(11)
doc.styles['Title'].font.bold = True
doc.styles['Title'].font.color.rgb = RGBColor(0,0,0)
doc.styles['Title'].paragraph_format.line_spacing = 2

def p(text, style='Els-body-text'):
    obj = doc.add_paragraph(text, style)
    num = OxmlElement('w:numPr')
    nid = OxmlElement('w:numId'); nid.set(qn('w:val'),'0'); num.append(nid)
    obj._p.get_or_add_pPr().append(num)
    obj.paragraph_format.first_line_indent = Inches(.2) if style == 'Els-body-text' else Inches(0)
    return obj

def h(text):
    obj = p(text, 'Els-1storder-head')
    obj.paragraph_format.keep_with_next = True
    obj.runs[0].bold = True

def sub(text):
    obj = p(text, 'Els-2ndorder-head')
    obj.paragraph_format.keep_with_next = True
    obj.runs[0].italic = True

def fig(name, caption, width=5.8):
    obj = doc.add_paragraph()
    obj.paragraph_format.keep_with_next = True
    obj.paragraph_format.line_spacing = 1
    inline_shape = obj.add_run().add_picture(str(ROOT/'manuscript'/'figures'/name), width=Inches(width))
    inline_shape._inline.docPr.set('title', caption.split(' ', 2)[0] + ' ' + caption.split(' ', 2)[1])
    inline_shape._inline.docPr.set('descr', caption)
    p(caption, 'Normal')

def table(headers, rows):
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = True
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = str(text)
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing = 1
                para.paragraph_format.keep_with_next = i < len(t.rows)-1
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs:
                    run.font.name='Times New Roman'; run.font.size=Pt(10); run.bold=(i==0)
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom']:
        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'6'); borders.append(edge)
    t._tbl.tblPr.append(borders)
    return t

p('Research Article', 'Normal').runs[0].bold=True
p('Observed and Near-Term CMIP6 Changes in Precipitation Extremes over Uttaradit, Thailand', 'Title')
p('Surasit Punyawansiri1,*', 'Normal')
p('1Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok 10300, Thailand', 'Normal')
p('*Corresponding author: Surasit.pu@ku.th', 'Normal')
p('Revised research draft requiring confirmation of data provenance before submission', 'Normal')
h('Abstract')
p('This study examines observed precipitation extremes at 13 stations in Uttaradit, Thailand, and near-term changes in seven available CMIP6 model series under SSP2-4.5 and SSP5-8.5. Eleven annual indices were calculated from observed 1981-2014 records and model historical (1995-2014) and future (2021-2050) series. Standard Mann-Kendall (MK) tests and Sen slopes provide the primary station-level trend description. The full-lag Yue-Wang modification was retained only as a sensitivity analysis after a Monte Carlo audit showed inflated null-rejection rates for short series. Projected changes were summarized from seven model-specific station medians, with their median, IQR and positive/negative support reported separately. Of 143 standard-MK tests, 30 were significant before and 14 after Benjamini-Hochberg adjustment; no province-wide field inference is made. Median PRCPTOT changes were 2.9% and 0.9%, and median Rx1day changes were 6.1% and 12.3%, under SSP2-4.5 and SSP5-8.5, respectively. R99p retains negative model support under SSP5-8.5 despite a positive median. A power-2 IDW map is used only as a descriptive spatial visualization, with leave-one-out error reported separately. Conclusions remain conditional on supplied data whose observation provenance, bias-correction procedure, units and native calendar require documentary verification before submission.', 'Els-Abstract-text')
p('Keywords: Precipitation extremes, Uttaradit, CMIP6, Mann-Kendall, Sen slope, Climate projections', 'Normal')
h('1. Introduction')
p('Annual rainfall totals alone cannot describe the intensity, timing and persistence of precipitation extremes. Standardized indices based on daily observations provide complementary information on heavy rainfall and dry or wet spells [1,2]. For engineering applications, distinguishing these measures helps frame questions about drainage, storage and drought exposure. An index change does not itself quantify infrastructure failure or flood discharge, which require additional hydrological and design information.')
p('CMIP6 provides coordinated climate simulations for examining historical conditions and future scenarios [3,4]. Local interpretation nevertheless depends on model choice, spatial sampling and preprocessing of supplied series. Trend inference can also be sensitive to serial correlation treatment, especially for short annual records. The full-lag Yue-Wang effective-sample-size formulation [5] is retained as a sensitivity analysis, while conventional MK results and Sen slopes provide the primary station-level description.')
p('The primary question is how observed precipitation extremes vary across Uttaradit and how the available CMIP6 models describe near-term changes under SSP2-4.5 and SSP5-8.5. Observed trends provide context, future mean change is the primary projection result, and within-window future trend is a secondary diagnostic. The analysis distinguishes change in the 2021-2050 mean relative to 1995-2014 from monotonic trend within 2021-2050; these quantities need not have the same sign.')
h('2. Materials and methods')
sub('2.1 Study network and supplied data')
p('The network comprises stations 351001-351012 and 351201. Coordinates were read from the supplied station workbook. Observed records cover 1981-2014. Twenty-one supplied precipitation CSV files represent historical, SSP2-4.5 and SSP5-8.5 data from ACCESS-ESM1-5, CanESM5, CESM2, EC-Earth3, FGOALS-g3, MIROC6 and MRI-ESM2-0. These seven models were included because they were the supplied series satisfying the target-station, scenario and temporal-coverage requirements; they are not treated as statistically independent. The filenames designate the model series as bias-corrected; the original correction algorithm, training period, validation and observational agency were not established from the files inspected. These details must be supplied and checked before submission. Rainfall values are treated as daily millimetres, an input assumption requiring confirmation from the data provider.')
fig('Figure1_station_network.png', 'Figure 1 Locations of the 13 supplied stations and the Uttaradit polygon selected by province code 53 from the supplied GIS file.', 4.3)
sub('2.2 Quality control and annual indices')
p('Dates, station identifiers and coordinates were validated before calculation. Duplicate dates, negative rainfall and non-numeric rainfall are rejected. Each annual series is reindexed to a Gregorian calendar; a supplied record without February 29 therefore has that date missing, rather than having a synthetic value inserted. The 90% completeness denominator is 365 or 366 Gregorian days, as applicable. No missing day is imputed. An unrecorded day breaks a run and prevents a five-day accumulation spanning that gap. This is a study-specific annual completeness policy rather than a claim of full ETCCDI software equivalence. The native model calendar, rainfall variable, units and any upstream conversion must be confirmed from provider documentation before submission.')
p('The core ETCCDI-type indices are PRCPTOT, SDII, Rx1day, Rx5day, CDD, CWD, R10mm, R20mm, R95p and R99p. Wet days satisfy RR >= 1 mm. PRCPTOT sums wet-day precipitation and SDII divides that sum by the number of wet days. Rx1day is the annual daily maximum; Rx5day is the maximum complete five-day running sum within the calendar year. CDD and CWD are the longest within-year runs below and at or above 1 mm, respectively. Spells and Rx5day windows are bounded by calendar years in this implementation. R50mm counts days at or above 50 mm and is a study-specific additional threshold index. R95p and R99p sum rainfall strictly exceeding observed station-specific wet-day percentiles [1,2].')
p('Percentiles use the full observed 1981-2014 wet-day reference sample and the NumPy linear quantile convention; they are then frozen for model historical and future series. No in-base-period bootstrap is applied; the resulting percentile indices are therefore described as fixed-reference ETCCDI-based indices. Threshold uncertainty is not propagated. Annual completeness >= 90% does not guarantee that a within-year missing day has no effect on Rx5day, CDD or CWD. The threshold CSV preserves the reference values and period.')
sub('2.3 Trend estimation')
p('The Mann-Kendall statistic S sums the signs of all later-minus-earlier annual differences. Its variance includes the standard correction for tied values. A continuity-corrected normal statistic yields a two-sided p-value. Constant series return S = 0, Z = 0 and p = 1. At least ten valid annual values are required. Standard MK classifications and Sen slopes, the median of pairwise changes divided by actual differences in years [6], are the primary station-level results. A conventional 95% order-statistic slope interval is calculated with SciPy Theil-Sen routines; it is not adjusted for serial dependence.')
p('For the observed analysis, the 143 station-index standard-MK p-values were treated as one multiplicity family and adjusted using the Benjamini-Hochberg procedure at q = 0.05. BH-FDR controls the expected proportion of false discoveries within this family; it is not a spatial field-significance test. Future within-window classifications are secondary diagnostics and are reported without an additional FDR claim.')
p('As a prespecified sensitivity analysis, the retained Yue-Wang effective-sample-size procedure [5] removes the Sen trend before computing sample autocorrelations with a common mean and full-series sum-of-squares denominator. The variance multiplier is f = 1 + 2 sum[(1 - k/n) rho(k)], for lags k = 1,...,n - 1; modified variance is f Var(S), and effective sample size is n/f. All lags are included in the retained implementation. A Monte Carlo audit generated Gaussian AR(1) annual series with n = 30 and 34, rho = 0.0, 0.1, 0.2, 0.3 and 0.5, null, positive and negative linear trends, 1,000 replicates per condition and nominal alpha = 0.05. Standard and full-lag Yue-Wang rejection rates were compared; materially inflated null rejection rendered the full-lag procedure unsuitable for primary inference (Table S9). Standard-MK results remain descriptive station-level evidence and do not assert serial independence or regional field significance.')
sub('2.4 Future changes and spatial presentation')
p('For each model, station and index, absolute change is the 2021-2050 annual mean minus the 1995-2014 historical annual mean. Relative change divides this difference by the historical mean and multiplies by 100; zero baselines do not produce finite percentages. Observed trends use 1981-2014, whereas projected changes use each model\'s own 1995-2014 historical baseline; this study does not perform a historical model-performance evaluation. For a single aggregation order, each model is first summarized by its median across the 13 stations; the central projection is the median of these seven model-specific station medians. Model spread is their IQR, and directional support is the number of positive and negative models. These summaries are descriptive; no ensemble significance test is claimed.')
p('Figure 2 presents the analytical framework. Future mean change compares the 2021-2050 and 1995-2014 means, whereas within-window trend describes monotonic change inside 2021-2050. Figure 5 uses power-2 IDW strictly as a descriptive spatial depiction, with the grid extended across the full province polygon and overlaid with all station values. The northern tip above the northernmost station is extrapolated beyond the station convex hull and must not be read as station-supported detail. Leave-one-out performance was evaluated as MAE of 14.14 mm (SSP2-4.5) and 18.05 mm (SSP5-8.5); the IDW surface is not interpreted as independently validated gridded information or spatial-significance evidence.')
fig('Figure5_mean_change_vs_trend.png', 'Figure 2 Analytical framework for observed trends, near-term projected change and within-window temporal diagnostics. The three layers distinguish observed variability, projected future-period mean change and the secondary temporal diagnostic. Future mean change and within-window trend answer different questions.', 5.5)
h('3. Results')
sub('3.1 Data coverage and observed trends')
p('All 442 observed station-years and 5,460 future station-model-scenario-years met the implemented annual completeness criterion. This does not imply every day was present. The output contains 4,862 observed annual index values and 60,060 future values. Observed trend series have 34 annual values and future trend series have 30. The separate output checker confirmed expected dimensions, bounded p-values, ordered slope intervals and model-consistent PRCPTOT means recomputed from raw daily input.')
obs=pd.read_csv(R/'observed_trend_1981_2014.csv')
fdr=pd.read_csv(ROOT/'manuscript'/'supplementary_data'/'Table_S5_observed_MK_FDR.csv')

def observed_direction_counts(index):
    values=obs.loc[obs.Index==index,'Trend_MK'].value_counts()
    return {
        'increasing': int(values.get('Increasing',0)),
        'decreasing': int(values.get('Decreasing',0)),
        'nonsignificant': int(values.get('No significant trend',0)),
    }

prcptot_counts=observed_direction_counts('PRCPTOT')
r50mm_counts=observed_direction_counts('R50mm')
sdii_counts=observed_direction_counts('SDII')
n_observed_tests=len(obs)
n_raw_significant=int(obs['Significant_MK'].astype(str).str.strip().str.lower().eq('true').sum())
n_fdr_significant=int(fdr['MK_significant_FDR_BH'].astype(str).str.strip().str.lower().eq('true').sum())
if prcptot_counts['increasing']==0 and prcptot_counts['decreasing']==0:
    prcptot_summary=(
        f'For PRCPTOT, no station showed a significant increasing or decreasing standard-MK trend; '
        f'all {prcptot_counts["nonsignificant"]} stations were nonsignificant.'
    )
else:
    prcptot_summary=(
        f'For PRCPTOT, {prcptot_counts["increasing"]} stations had a significant increasing trend, '
        f'{prcptot_counts["decreasing"]} had a significant decreasing trend and '
        f'{prcptot_counts["nonsignificant"]} were nonsignificant.'
    )
p(
    f'Observed standard-MK classifications were spatially mixed (Figure 3). Of {n_observed_tests} '
    f'station-index tests, {n_raw_significant} were significant at unadjusted alpha = 0.05, whereas '
    f'{n_fdr_significant} remained significant after Benjamini-Hochberg false-discovery-rate adjustment. '
    f'{prcptot_summary} R50mm had '
    f'{r50mm_counts["increasing"]} increasing, {r50mm_counts["decreasing"]} decreasing and '
    f'{r50mm_counts["nonsignificant"]} nonsignificant classifications. SDII had '
    f'{sdii_counts["increasing"]} increasing, {sdii_counts["decreasing"]} decreasing and '
    f'{sdii_counts["nonsignificant"]} nonsignificant classifications. These station-level results do not demonstrate '
    'a province-wide trend. The full-lag Yue-Wang sensitivity analysis classified 66 tests as '
    'significant; Table S9 reports its null-rejection rates for every simulated n and AR(1) condition, '
    'and it is not used for primary inference.'
)
p('Table 1 Observed standard-MK direction counts based on unadjusted alpha = 0.05 across the 13 stations, 1981-2014.', 'Normal')
order=['PRCPTOT','SDII','Rx1day','Rx5day','CDD','CWD','R10mm','R20mm','R50mm','R95p','R99p']
rows=[]
for idx in order:
    v=obs.loc[obs.Index==idx,'Trend_MK'].value_counts()
    rows.append([idx,v.get('Increasing',0),v.get('Decreasing',0),v.get('No significant trend',0)])
table(['Index','Increasing','Decreasing','Not significant'],rows)
fig('Figure2_observed_mmk_direction.png','Figure 3 Observed precipitation-index trends across the 13 stations, 1981-2014. Panel (a) shows standardized Sen slopes for within-index pattern comparison; raw slopes and units are provided in the Supplementary Material. Panel (b) shows the direction of trends that remain significant after Benjamini-Hochberg FDR adjustment; nonsignificant station-index combinations are displayed separately and do not encode magnitude. S001-S012 and S201 denote the last three digits of station IDs.')
sub('3.2 Projected mean changes and within-window trends')
p('The model-first descriptive summaries show small PRCPTOT medians with inter-model distributions spanning both positive and negative changes (Table 2; Figure 4). Rx1day medians are 6.1% and 12.3%, with positive changes in seven and five of seven models under SSP2-4.5 and SSP5-8.5, respectively. R50mm is a study-specific additional threshold index rather than an ETCCDI core index. R99p has positive medians but broad spread, including negative model changes under SSP5-8.5. Table 2 reports IQR and positive/negative model support so central values are not interpreted as model consensus. Large percentage changes in infrequent exceedances must be read with their absolute baseline frequencies.')
chg=pd.read_csv(R/'future_change_2021_2050.csv')
model_station=chg.groupby(['Scenario','Model','Station','Index'],as_index=False).Relative_change_pct.median()
model_summary=model_station.groupby(['Scenario','Model','Index'],as_index=False).Relative_change_pct.median()
summ=model_summary.groupby(['Scenario','Index']).Relative_change_pct.median().unstack('Scenario').reindex(order)
def uncertainty_cell(scenario, idx):
    vals=model_summary.loc[(model_summary.Scenario==scenario)&(model_summary.Index==idx),'Relative_change_pct'].dropna()
    q1,q3=vals.quantile([.25,.75])
    pos=int((vals>0).sum()); neg=int((vals<0).sum())
    return f'{summ.loc[idx,scenario]:.2f} [{q1:.2f}, {q3:.2f}]; +{pos}/7, -{neg}/7'
p('Table 2 Multi-model summary of relative change (%) for 2021-2050 against each model\'s 1995-2014 baseline. Entries are median [IQR] (%) of seven model-specific station medians, followed by positive and negative model support.', 'Normal')
table(['Index','SSP2-4.5 median [IQR] (%); + / - models','SSP5-8.5 median [IQR] (%); + / - models'],[[idx,uncertainty_cell('SSP2-4.5',idx),uncertainty_cell('SSP5-8.5',idx)] for idx in order])
fig('Figure3_future_median_relative_change.png','Figure 4 Seven model-specific station-median relative changes, 2021-2050 relative to 1995-2014. Open points identify models, the black point is the median and the horizontal line is the IQR. R99p is shown in a separate lower strip to avoid compressing the remaining indices. This figure displays structural model spread; it is not an ensemble significance test.',5.5)
p('Figure 5 provides a descriptive power-2 IDW depiction of the station-median PRCPTOT change, with every station labelled and overlaid. The common colour scale permits scenario comparison, but the surface is a deterministic interpolation subject to the leave-one-out errors reported in Table S10; it is not independently validated gridded information and carries no spatial-significance claim. Within the 2021-2050 window, standard MK identified 79 of 1,001 tests under SSP2-4.5 and 71 of 1,001 under SSP5-8.5 as significant. Under SSP2-4.5, decreasing classifications outnumbered increasing classifications (74 versus 5); under SSP5-8.5, increasing classifications outnumbered decreasing classifications (57 versus 14). These secondary diagnostics can coexist with a future mean above the earlier baseline because they measure different quantities (Figure 2).')
fig('Figure4_prcptot_idw_change.png','Figure 5 Descriptive power-2 IDW depiction of seven-model-median absolute PRCPTOT change, 2021-2050 relative to 1995-2014. Station labels and symbols identify the values used to interpolate; both panels share a colour scale and the grid covers the full province polygon. The northern tip above the northernmost station is extrapolated beyond the station convex hull and is not station-supported detail. The surface is descriptive only and is not interpreted as independently validated gridded information; LOOCV MAE is 14.14 and 18.05 mm (Table S10).')
h('4. Discussion')
p('Observed extremes are heterogeneous across the station network. The standard-MK results do not support a single province-wide direction for PRCPTOT or most other indices, and the unadjusted classifications must not be interpreted as field significance. The descriptive projection pattern separates rainfall accumulation from heavy-event metrics: small median PRCPTOT changes coexist with larger central changes for Rx1day and selected threshold indices. This can motivate separate assessment of water supply and intense rainfall exposure, but it cannot determine drainage capacity, return-period rainfall or flood damage from these indices alone. Likewise, CDD is a precipitation-spell measure and does not incorporate evapotranspiration or soil water storage.')
p('This study contributes by separating three quantities that can be conflated in local precipitation-projection studies: observed station-level trend, projected future-period mean change and monotonic trend within the future period. The multi-model analysis further separates a central projected change from inter-model spread and directional support. Model choice is a material source of uncertainty. The multi-model median is a descriptive central value, not evidence that the seven models are independent or agree in sign. The IQR and positive/negative model support in Table 2 should accompany any interpretation of the median. Large relative changes in R50mm, R95p and R99p require particular care because percentage changes can be amplified when the historical frequency or baseline amount is low.')
p('The discrepancy between standard and modified test classifications is a substantive inference sensitivity. The full-lag multiplier falls below one for 128 of 143 observed series and 1,992 of 2,002 future series; consequently this sensitivity analysis often becomes more permissive. In the present simulation audit at sample sizes of 30 and 34 years, it had inflated null rejection rates and is therefore not used for inferential conclusions. Any alternative autocorrelation-adjusted procedure would require separate validation before it could replace standard MK as the primary test. The reported conventional Sen intervals retain their independence assumptions.')
p('Additional limits arise upstream. Before submission, data-provider documentation must establish the observational agency and quality control, rainfall variable and units, native model calendars and any conversion, spatial sampling, bias-correction method, reference observations, calibration period and validation evidence. A filename alone does not verify preprocessing performance. Fixed observed percentiles aid comparability to a common historical threshold but omit threshold-estimation uncertainty and differ from bootstrap treatments used in some ETCCDI implementations. The analysis is therefore a reproducible, model-conditioned assessment of the supplied series, not a validated engineering design standard or deterministic province-wide forecast.')
h('5. Conclusions')
p('Observed precipitation extremes are heterogeneous across the 13 stations; standard-MK station tests do not establish a province-wide trend. Descriptive near-term projections show larger median relative changes for selected heavy-rainfall measures, including Rx1day and the study-specific R50mm index, although model spread is substantial for several tail-related indices. These changes should be interpreted jointly with model spread, positive/negative support and the absolute baselines of infrequent indices. Future mean change and within-window trend answer different questions and may have different signs. Finally, statistical classifications are sensitive to the selected trend-test treatment, and all projection inferences remain model-conditioned until data provenance, bias correction, units and calendar treatment are documented and verified.')
h('6. Acknowledgements and declarations')
p('[To be completed by the authors: data-provider acknowledgement, funding, conflicts of interest, author contributions and any required AI-use disclosure. No declarations have been inferred.]', 'Normal')
p('Data and code availability: The revised code, regression tests, input-file manifest, QC tables and result workbook accompany this draft locally. [Authors must supply data-use permissions and a public repository or justified access statement.]', 'Normal')
h('7. References')
refs=[
'Zhang X, Alexander L, Hegerl GC, Jones P, Klein Tank A, Peterson TC, et al. Indices for monitoring changes in extremes based on daily temperature and precipitation data. WIREs Clim Change. 2011;2:851-870. https://doi.org/10.1002/wcc.147',
'ETCCDI. Definitions of the 27 core indices [Internet]. [cited 2026 Sep 5]. Available from: https://etccdi.pacificclimate.org/list_27_indices.shtml',
'Eyring V, Bony S, Meehl GA, Senior CA, Stevens B, Stouffer RJ, et al. Overview of the Coupled Model Intercomparison Project Phase 6 experimental design and organization. Geosci Model Dev. 2016;9:1937-1958. https://doi.org/10.5194/gmd-9-1937-2016',
'O\'Neill BC, Tebaldi C, van Vuuren DP, Eyring V, Friedlingstein P, Hurtt G, et al. The Scenario Model Intercomparison Project for CMIP6. Geosci Model Dev. 2016;9:3461-3482. https://doi.org/10.5194/gmd-9-3461-2016',
'Yue S, Wang CY. The Mann-Kendall test modified by effective sample size to detect trend in serially correlated hydrological series. Water Resour Manage. 2004;18:201-218. https://doi.org/10.1023/B:WARM.0000043140.61082.60',
'Sen PK. Estimates of the regression coefficient based on Kendall\'s tau. J Am Stat Assoc. 1968;63:1379-1389. https://doi.org/10.1080/01621459.1968.10480934',
'Shepard D. A two-dimensional interpolation function for irregularly-spaced data. In: Proceedings of the 23rd ACM National Conference. New York: ACM; 1968. p. 517-524. https://doi.org/10.1145/800186.810616']
for i,text in enumerate(refs,1): p(f'[{i}] {text}', 'Normal')
doc.core_properties.author='Surasit Punyawansiri'
doc.core_properties.last_modified_by=''
doc.core_properties.title='Observed and Near-Term CMIP6 Changes in Precipitation Extremes over Uttaradit, Thailand'
doc.save(OUT)
print(OUT)
