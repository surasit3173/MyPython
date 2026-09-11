from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out_path = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_response_to_reviewers_revised_TH.docx")

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

styles = doc.styles
styles["Normal"].font.name = "TH Sarabun New"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
styles["Normal"].font.size = Pt(14)
for name in ["Heading 1", "Heading 2"]:
    styles[name].font.name = "TH Sarabun New"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    styles[name].font.bold = True

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_width(cell, width_twips):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.first_child_found_in("w:tcW")
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width_twips))
    tcW.set(qn("w:type"), "dxa")

def set_cell_no_wrap(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    if tcPr.first_child_found_in("w:noWrap") is None:
        tcPr.append(OxmlElement("w:noWrap"))

def add_para(text="", style=None, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.font.name = "TH Sarabun New"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    r.font.size = Pt(14)
    r.bold = bold
    return p

title = add_para("คำชี้แจงการแก้ไขต้นฉบับตามข้อเสนอแนะของผู้ทรงคุณวุฒิ", bold=True)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.size = Pt(18)
subtitle = add_para("ต้นฉบับ ID1558: Evaluation of Quantile Delta Mapping for Bias Correction of CMIP6 Daily Rainfall Using Independent Temporal Holdout in Prachuap Khiri Khan Province")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_para("")
add_para("เรียน บรรณาธิการและผู้ทรงคุณวุฒิ", bold=True)
add_para("คณะผู้วิจัยขอขอบพระคุณบรรณาธิการและผู้ทรงคุณวุฒิสำหรับข้อเสนอแนะที่เป็นประโยชน์ต่อการปรับปรุงต้นฉบับอย่างยิ่ง ข้อเสนอแนะดังกล่าวช่วยให้คณะผู้วิจัยทบทวนกรอบการประเมิน วิธีวิเคราะห์ การตีความผล และความถูกต้องของเอกสารอ้างอิงอย่างรอบคอบมากขึ้น ต้นฉบับฉบับแก้ไขได้ปรับจากการประเมินในช่วงข้อมูลเดียวกันไปสู่การตรวจสอบด้วยช่วงเวลาอิสระ โดยประมาณฟังก์ชัน QDM จากช่วง ค.ศ. 1981-2000 และประเมินผลกับช่วง ค.ศ. 2001-2014 ที่กันไว้โดยไม่ใช้ข้อมูลตรวจวัดในช่วงดังกล่าวเพื่อปรับเทียบใหม่")
add_para("การตอบรายข้อต่อไปนี้เรียบเรียงโดยยึดข้อเท็จจริงจากต้นฉบับฉบับแก้ไขล่าสุด และระบุหัวข้อหรือส่วนของบทความที่เกี่ยวข้องเพื่อให้ตรวจสอบได้โดยตรง")

add_para("สรุปการแก้ไขหลัก", style="Heading 1")
summary_items = [
    "ปรับชื่อเรื่องและกรอบการศึกษาให้สะท้อน independent temporal holdout แทนการประเมินจากข้อมูลเต็มช่วง",
    "แก้ปัญหา data leakage โดยแยกช่วงประมาณฟังก์ชัน QDM ออกจากช่วงประเมินผลอย่างชัดเจน",
    "เปลี่ยนการวิเคราะห์หลักเป็นระดับคู่สถานี-แบบจำลอง 60 คู่ และใช้ MME รายเดือนเป็นการวิเคราะห์รอง",
    "เพิ่มตัวชี้วัดด้าน distributional fidelity, heavy rainfall, time-paired error และ temporal correspondence",
    "เพิ่ม year-block bootstrap 2,000 ซ้ำและควบคุม false discovery rate ด้วย Benjamini-Yekutieli เพื่อรองรับการพึ่งพาเชิงพื้นที่ของคู่ประเมิน",
    "เพิ่มการแยกองค์ประกอบ MSE เพื่ออธิบายกลไกของการเปลี่ยนแปลง RMSE โดยไม่ตีความว่า QDM สร้างความสอดคล้องเชิงเวลาที่ไม่มีในแบบจำลองเดิม",
    "ตรวจและแก้ไขรายการอ้างอิงที่ไม่สอดคล้องกับ DOI เดิม และตัดรายการที่ตรวจสอบไม่ได้ออก",
]
for item in summary_items:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

rows = [
    ("1", "ชื่อเรื่องควรสะท้อนว่าฉบับเดิมเน้น unweighted multi-model ensemble ในช่วง historical และควรระบุ validation design ที่ใช้จริง",
     "ปรับชื่อเรื่องเป็นการประเมิน QDM ด้วย independent temporal holdout และปรับโครงสร้างการวิเคราะห์ให้ผลหลักอยู่ที่คู่สถานี-แบบจำลอง 60 คู่ ส่วน MME เป็นการวิเคราะห์รองที่คำนวณระดับรายเดือนเท่านั้น",
     "ชื่อเรื่อง; 2.3; 2.6; Results; 3.6",
     "คณะผู้วิจัยเห็นด้วยและได้ปรับชื่อเรื่องรวมทั้งกรอบการนำเสนอใหม่ เพื่อให้สะท้อนวัตถุประสงค์และการออกแบบการตรวจสอบจริงของงานวิจัย ฉบับแก้ไขไม่ได้ใช้ MME เป็นหน่วยหลักของข้อสรุปอีกต่อไป แต่รายงานผลรายแบบจำลองและรายคู่สถานี-แบบจำลองก่อน แล้วจึงนำเสนอ MME เป็นผลรองเพื่อประกอบการตีความ"),
    ("2", "ตัวเลขในบทคัดย่อเดิมไม่สอดคล้องกับตารางและบทคัดย่อภาษาอังกฤษ",
     "เขียนบทคัดย่อไทยและอังกฤษใหม่ โดยรายงานค่าจากช่วงตรวจสอบอิสระและใช้ตัวเลขชุดเดียวกับตารางผล เช่น KS, PSS, RMSE, NSE และจำนวนคู่ที่ดีขึ้นหรือแย่ลง",
     "บทคัดย่อ; Abstract; 3.2; 3.3; ตารางที่ 3-4",
     "คณะผู้วิจัยขอขอบพระคุณที่ชี้ข้อคลาดเคลื่อนดังกล่าว และได้ตรวจสอบตัวเลขทั้งหมดใหม่ให้ตรงกับผลวิเคราะห์ฉบับแก้ไข โดยตัดค่าที่ไม่สอดคล้องจากฉบับเดิมออกและแทนด้วยผลจากช่วง ค.ศ. 2001-2014 ซึ่งเป็นช่วงตรวจสอบอิสระ"),
    ("3", "ข้อสรุปเดิมเกินหลักฐาน เพราะใช้ข้อมูลช่วงเดียวกันในการปรับและประเมิน และไม่มี out-of-sample validation หรือการประเมิน extreme/hydrological performance",
     "ออกแบบการวิเคราะห์ใหม่โดยแยก calibration period กับ holdout period และเพิ่มตัวชี้วัด q95, q99, Rx1day สำหรับฝนหนัก พร้อมหลีกเลี่ยงการอ้าง hydrological performance เพราะไม่ได้ขับแบบจำลองอุทกวิทยา",
     "2.3; 2.5; 3.2; 4.4; 4.5; Conclusions",
     "คณะผู้วิจัยเห็นด้วย จึงปรับกรอบการศึกษาเป็น independent temporal holdout และลดระดับข้อสรุปให้จำกัดอยู่กับคุณสมบัติของข้อมูลฝนที่ประเมินจริง ไม่อ้างว่าผลหลังปรับแก้ยืนยันสมรรถนะทางอุทกวิทยาหรือความคงที่ภายใต้ภูมิอากาศอนาคต"),
    ("4", "ภาษาอังกฤษและหน่วยในบทคัดย่อเดิมไม่ถูกต้อง โดย RMSE ควรใช้หน่วย mm month-1 และควรระวังคำว่า decreased",
     "ปรับ Abstract ใหม่ให้รายงานทิศทางผลตามข้อมูลจริง แยก annual-pooled fitting กับ calendar-month fitting และใช้หน่วย mm month-1 สำหรับ RMSE",
     "Abstract; 3.3; ตารางที่ 3-4",
     "คณะผู้วิจัยได้แก้ภาษาอังกฤษและหน่วยของตัวชี้วัดแล้ว โดยรายงานว่า calendar-month fitting ลด RMSE รายเดือน ขณะที่ annual-pooled fitting เพิ่ม RMSE รายเดือน เพื่อหลีกเลี่ยงการใช้ถ้อยคำที่สื่อว่าทุกกรณีดีขึ้น"),
    ("5", "ข้อสรุปภาษาอังกฤษเดิมเป็นประโยคไม่สมบูรณ์และสับสนระหว่าง RMSE กับ NSE",
     "เขียน Conclusions ภาษาอังกฤษใหม่ แยก RMSE และ NSE ออกจากกัน และตีความการเปลี่ยนแปลงของ RMSE/NSE ร่วมกับองค์ประกอบของ MSE",
     "Abstract: Conclusions; Conclusions; 3.3; 3.5",
     "คณะผู้วิจัยได้ตัดข้อความที่สับสนในฉบับเดิมออกทั้งหมด และเขียนข้อสรุปใหม่ให้ชัดว่า RMSE เป็นตัวชี้วัดขนาดความคลาดเคลื่อน ส่วน NSE เป็นดัชนีประสิทธิภาพเชิงอนุกรมที่ควรตีความควบคู่กับองค์ประกอบของความคลาดเคลื่อน"),
    ("6", "ข้ออ้างเรื่องช่องว่างการวิจัยควรอ้างอิงงานล่าสุดในไทยและเอเชียตะวันออกเฉียงใต้ และไม่ควรสรุปจากงานเพียงบางรายการ",
     "ปรับบทนำให้ใช้ถ้อยคำระมัดระวังขึ้น และเพิ่มงานที่เกี่ยวข้องกับ CMIP6/QDM/ฝนในภูมิภาค เช่น Supharatid et al., Kuinkel et al., Humphries et al. และ Arif Rafhida et al.",
     "Introduction; References",
     "คณะผู้วิจัยเห็นด้วยและได้ปรับการนำเสนอช่องว่างงานวิจัยให้เป็นการสังเคราะห์เชิงประเด็น แทนการกล่าวอ้างแบบกว้างเกินหลักฐาน พร้อมเพิ่มวรรณกรรมร่วมสมัยในประเทศไทยและภูมิภาคเอเชียตะวันออกเฉียงใต้เท่าที่เกี่ยวข้องโดยตรง"),
    ("7", "ต้องรายงานแหล่งข้อมูล สัดส่วน missing รายสถานี เกณฑ์ outlier วิธีจัดการ missing homogeneity test และข้อมูล relocation/instrument changes",
     "เพิ่มรายละเอียดแหล่งข้อมูล การจัดการ NoData/ค่าติดลบ เกณฑ์ความครบถ้วนอย่างน้อย 95% การทดสอบ Pettitt ที่ระดับ 0.05 และผลว่าข้อมูลครบถ้วน 100% และไม่พบจุดเปลี่ยนอย่างมีนัยสำคัญในสถานีทั้ง 12 แห่ง พร้อมรายงานลักษณะรายสถานีในตารางที่ 1",
     "2.2; ตารางที่ 1; 4.5; Acknowledgments",
     "คณะผู้วิจัยได้เพิ่มรายละเอียดการควบคุมคุณภาพข้อมูลและระบุข้อจำกัดอย่างตรงไปตรงมาว่าการตรวจสอบนี้ไม่ใช้แทนเมตาดาทาเกี่ยวกับการย้ายสถานีหรือการเปลี่ยนเครื่องมือของหน่วยงานเจ้าของข้อมูล เพื่อไม่ให้ตีความเกินจากข้อมูลที่มี"),
    ("8", "Nearest-neighbor เป็นการเลือก/รีกริดจุดกริด ไม่ใช่ downscaling และควรรายงานระยะสถานี-กริดกับจำนวนสถานีที่ใช้กริดเดียวกัน",
     "แก้คำอธิบายเป็น nearest-grid-cell extraction และระบุชัดว่าไม่ใช่การลดมาตราส่วนเชิงพื้นที่ พร้อมรายงานจำนวนอนุกรมกริดเซลล์ที่แตกต่างจริงในตารางที่ 2 และอธิบายว่าบางสถานีใช้กริดเดียวกัน",
     "2.2; ตารางที่ 2; 4.5",
     "คณะผู้วิจัยเห็นด้วยและได้ยกเลิกถ้อยคำที่สื่อว่า nearest-neighbor สร้างข้อมูลระดับ sub-grid เพิ่มเติม ฉบับแก้ไขใช้ถ้อยคำว่าเป็นการดึงค่าจากกริดเซลล์ที่ใกล้ที่สุดและตีความคู่ประเมิน 60 คู่เป็นชุดประเมิน ไม่ใช่หน่วยเชิงพื้นที่อิสระ 60 หน่วย"),
    ("9", "สมมติฐานของ ensemble ไม่มีหลักฐานรองรับ ควรรายงานผลราย GCM genealogy และ uncertainty ก่อนสรุป ensemble",
     "ปรับให้ผลรายแบบจำลองเป็นหลัก และรายงานความแตกต่างระหว่างแบบจำลอง/สถานีใน Results ส่วน MME เป็นการวิเคราะห์รอง ไม่ใช้แทนความไม่แน่นอนเชิงโครงสร้าง และไม่อ้างว่าลด structural uncertainty",
     "2.6; 3.6; ตารางที่ 5; รูปที่ 6-7; 4.4",
     "คณะผู้วิจัยเห็นด้วยในหลักการ จึงปรับการนำเสนอให้ไม่ซ่อนความแตกต่างของแบบจำลองไว้ในค่าเฉลี่ยเดียว ส่วนประเด็น genealogy ไม่ได้วิเคราะห์เชิงปริมาณในงานนี้ จึงระบุขอบเขตการตีความของ MME อย่างระมัดระวังแทนการอ้างความเป็นอิสระของแบบจำลอง"),
    ("10", "เกิด data leakage หากใช้ช่วงเดียวกัน train QDM และประเมินผล ควรใช้ calibration period และ independent validation period หรือ blocked cross-validation",
     "วิเคราะห์ใหม่โดยใช้ ค.ศ. 1981-2000 สำหรับประมาณฟังก์ชัน QDM และ ค.ศ. 2001-2014 สำหรับประเมินผล โดยไม่ใช้ observations ของช่วงตรวจสอบในการ recalibration",
     "2.3; รูปที่ 2; 2.4; 2.5; Results",
     "คณะผู้วิจัยเห็นด้วยและได้แก้กรอบการวิเคราะห์หลักใหม่ทั้งหมด ฟังก์ชัน QDM ถูกประมาณจากช่วงสอบเทียบเพียงช่วงเดียว จากนั้นตรึงไว้และนำไปใช้กับช่วงตรวจสอบอิสระ ทำให้ผลที่รายงานสะท้อนการประเมินนอกช่วงที่ใช้สร้างความสัมพันธ์"),
    ("11", "สูตร QDM กล่าวถึง projection distribution และ relative delta แต่ต้นฉบับวิเคราะห์ historical 1981-2014 จึงต้องระบุว่าช่วง projection คือช่วงใด",
     "นิยามช่วง ค.ศ. 2001-2014 เป็น target/holdout period ในการประยุกต์ QDM และระบุว่าการแจกแจงของแบบจำลองในช่วงเป้าหมายใช้เพื่อกำหนดตำแหน่งควอนไทล์ตามหลัก QDM โดยไม่ใช้ข้อมูลสังเกตในช่วงตรวจสอบเพื่อประมาณฟังก์ชันใหม่",
     "2.3; 2.4",
     "คณะผู้วิจัยได้ปรับคำอธิบายสมการ QDM ให้ตรงกับการวิเคราะห์ historical holdout โดยหลีกเลี่ยงถ้อยคำที่ทำให้เข้าใจว่าเป็น future projection และระบุบทบาทของช่วงเป้าหมายอย่างชัดเจน"),
    ("12", "เหตุผลเรื่อง full-period evaluation ไม่ยอมรับเป็นการประเมินประสิทธิภาพ เพราะทำให้ optimistic และไม่วัด generalization",
     "ตัดกรอบ full-period evaluation เดิมออก และประเมินผลเฉพาะช่วง holdout 2001-2014",
     "2.3; 2.5; Results",
     "คณะผู้วิจัยเห็นด้วยและได้ยกเลิกเหตุผลดังกล่าวในฉบับแก้ไข การประเมินประสิทธิภาพทั้งหมดทำในช่วงตรวจสอบอิสระเท่านั้น เพื่อให้การตีความสอดคล้องกับความสามารถในการถ่ายทอดข้ามช่วงเวลา"),
    ("13", "การประมาณ quantile ด้วยข้อมูลเต็มช่วงอาจใช้เพื่อ fitting แต่ไม่ใช่เหตุผลให้ละเว้น validation ควรแยก fitting ออกจาก evaluation",
     "แยกขั้นตอน fitting และ evaluation อย่างชัดเจน พร้อมตรวจสอบว่า observations ของช่วงตรวจสอบไม่เข้าสู่ขั้นตอนสร้างฟังก์ชัน QDM",
     "2.3; 2.4; รูปที่ 2",
     "คณะผู้วิจัยได้ปรับกระบวนการให้แยกการประมาณฟังก์ชัน การนำฟังก์ชันไปใช้ และการประเมินผลออกจากกัน โดยเพิ่มคำอธิบายใน Methodology และรูปผังการไหลของข้อมูลเพื่อให้ตรวจสอบได้"),
    ("14", "การเฉลี่ยค่าฝนรายวันจาก free-running GCMs ที่ phase ไม่ตรงกันทำลาย wet-day/extreme structure ควรประเมินรายโมเดลแล้วสรุป ensemble ภายหลัง",
     "ยกเลิกการสร้าง MME รายวัน และสร้าง MME หลังสะสมเป็นรายเดือนเท่านั้น พร้อมระบุว่า MME ไม่ใช่อนุกรม weather realization",
     "2.6; 3.6; ตารางที่ 5",
     "คณะผู้วิจัยเห็นด้วยและได้แก้ไขวิธีวิเคราะห์ตามข้อเสนอแนะ โดยประเมินแบบจำลองรายตัวก่อน จากนั้นจึงสรุป MME รายเดือนเป็นผลรองเพื่อหลีกเลี่ยงการเฉลี่ยเหตุการณ์ฝนรายวันที่ไม่ตรงเฟสกัน"),
    ("15", "ensemble spread ที่แคบลงไม่ใช่หลักฐานว่า structural uncertainty ลดลง และอาจประเมิน uncertainty ต่ำเกินจริง",
     "ตัดการตีความ ensemble spread contraction เป็นการลด uncertainty และรายงาน MME เฉพาะในฐานะค่าเฉลี่ยรายเดือนที่ช่วยลด RMSE บางกรณี แต่ไม่ได้เพิ่ม temporal correspondence อย่างจำเป็น",
     "2.6; 3.6; 4.4",
     "คณะผู้วิจัยได้ปรับการตีความตามข้อเสนอแนะ โดยไม่ใช้การหดตัวของการกระจายหลังการปรับแก้เป็นหลักฐานของการลดความไม่แน่นอนเชิงโครงสร้าง และระบุข้อจำกัดของ MME อย่างชัดเจน"),
    ("16", "การอธิบายผลของ S9 ด้วย orographic rainfall เป็นข้อสันนิษฐาน เพราะ S9 เป็น Inland Plain และสูงเพียง 14.45 m",
     "ตัดข้อสรุปเชิงสาเหตุเกี่ยวกับ orographic rainfall และใช้การอธิบายจากลักษณะระเบียนที่สังเกตได้แทน เช่น SDII และ lag-1 autocorrelation ของสถานี",
     "3.1; 3.6; 4.5",
     "คณะผู้วิจัยเห็นด้วยและได้ลดระดับการตีความเชิงสาเหตุ โดยหลีกเลี่ยงการอ้างอิทธิพลภูเขาโดยไม่มีหลักฐานจาก exposure หรือ grid mapping ที่เพียงพอ"),
    ("17", "ไม่มี statistical test, p-value หรือ confidence interval จึงไม่ควรใช้คำว่าอย่างมีนัยสำคัญ หรือควรเพิ่ม dependence-aware block bootstrap/permutation test",
     "เพิ่ม year-block bootstrap 2,000 ซ้ำสำหรับผลต่างแบบจับคู่และควบคุม FDR ด้วย Benjamini-Yekutieli ภายใต้ครอบครัว 120 สมมติฐาน พร้อมรายงาน confidence limits และจำนวนคู่ที่ดีขึ้น/แย่ลง/ไม่ต่างในตารางที่ 4",
     "2.5; ตารางที่ 4; 3.3",
     "คณะผู้วิจัยเห็นด้วยและได้เพิ่มการอนุมานทางสถิติที่คำนึงถึงการพึ่งพาเชิงเวลาและการทดสอบหลายครั้ง โดยใช้ปีปฏิทินเป็นบล็อกในการ bootstrap และใช้ Benjamini-Yekutieli เป็นเกณฑ์หลักเนื่องจากคู่ประเมินไม่เป็นอิสระต่อกันเชิงพื้นที่"),
    ("18", "Taylor diagram และกราฟอนุกรมเวลาเป็นหลักฐานเชิงพรรณนา ไม่ได้ยืนยันความน่าเชื่อถือ โดยเฉพาะเมื่อเป็น in-sample และ daily/model phases ไม่ตรงกัน",
     "ตัดการพึ่งพา Taylor diagram เป็นหลักฐานยืนยัน และแทนด้วยการรายงานตัวชี้วัดเชิงปริมาณ ตารางผล bootstrap และการแยก MSE/temporal correlation",
     "3.3-3.5; ตารางที่ 3-4; รูปที่ 4-5",
     "คณะผู้วิจัยเห็นด้วยและได้ปรับการนำเสนอให้กราฟทำหน้าที่สนับสนุนการอธิบายเชิงพรรณนาเท่านั้น ส่วนข้อสรุปหลักอาศัยผลตัวชี้วัด การอนุมานแบบ bootstrap และการแยกองค์ประกอบของความคลาดเคลื่อน"),
    ("19", "ข้อสรุปต้องลดระดับ เพราะไม่มีการทดสอบนัยสำคัญและไม่มี independent validation",
     "เขียน Conclusions ใหม่ภายใต้กรอบ independent temporal holdout และใช้ผลจาก bootstrap/BY ประกอบเฉพาะตัวชี้วัดที่ทดสอบ",
     "Conclusions; 4.1-4.5",
     "คณะผู้วิจัยได้ลดระดับข้อสรุปให้สอดคล้องกับหลักฐาน โดยสรุปว่า QDM มีความสามารถในการถ่ายทอดบางคุณสมบัติข้ามช่วงเวลาอย่างมีเงื่อนไข แต่ไม่สามารถกู้คืนความสอดคล้องเชิงเวลาที่ไม่มีในแบบจำลองเดิม"),
    ("20", "รายการอ้างอิงหนึ่งตรวจไม่พบตาม DOI ที่อ้าง",
     "ตัดรายการที่ตรวจสอบไม่ได้ออก และคงเฉพาะรายการที่ DOI และข้อมูลบรรณานุกรมตรวจสอบได้",
     "References",
     "คณะผู้วิจัยได้ตรวจสอบรายการอ้างอิงใหม่และนำรายการที่ไม่สามารถยืนยันความตรงกันระหว่างชื่อบทความกับ DOI ออกจากต้นฉบับ เพื่อให้เอกสารอ้างอิงมีความน่าเชื่อถือ"),
    ("21", "DOI ของ Vrac (2018) ไม่สอดคล้องกับชื่อบทความ ต้องตรวจและแก้ทั้งชื่อเรื่อง วารสาร หน้า และ DOI",
     "ตัดรายการ Vrac ที่มีความไม่สอดคล้องออกจาก reference list และไม่ใช้อ้างอิงในเนื้อหาฉบับแก้ไข",
     "References",
     "คณะผู้วิจัยเห็นด้วยและเลือกตัดรายการดังกล่าวออกแทนการคงรายการที่ยังไม่แน่นอน โดยคำอธิบายวิธี QDM อ้างอิง Cannon et al. (2015) เป็นแหล่งวิธีวิทยาหลักร่วมกับวรรณกรรมที่ตรวจสอบได้"),
    ("22", "DOI ของ Wang et al. ไม่ตรงกับบทความ precipitation bias correction ที่อ้าง",
     "ตัดรายการ Wang et al. ที่ไม่ตรงกับ DOI ออก และใช้แหล่งอ้างอิงด้านวิธีวิทยา/ภูมิภาคที่ตรวจสอบได้แทน",
     "Introduction; References",
     "คณะผู้วิจัยได้แก้ไขตามข้อเสนอแนะ โดยตัดรายการที่ DOI ไม่ตรงกับบทความออกจากรายการอ้างอิงและไม่ใช้เป็นหลักฐานในเนื้อหาฉบับแก้ไข"),
]

add_para("คำชี้แจงรายข้อ", style="Heading 1")
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
table.autofit = False
hdr = table.rows[0].cells
headers = ["ข้อ", "ประเด็นผู้ทรงคุณวุฒิ", "การแก้ไขในบทความฉบับใหม่", "ตำแหน่งในบทความ", "คำชี้แจงที่เสนอ"]
widths = [900, 2200, 2750, 1450, 3150]
for cell, h, w in zip(hdr, headers, widths):
    set_cell_shading(cell, "D9EAF7")
    set_cell_width(cell, w)
    if h == "ข้อ":
        set_cell_no_wrap(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(h)
    run.bold = True
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(11)

for row in rows:
    cells = table.add_row().cells
    for idx, (cell, text, w) in enumerate(zip(cells, row, widths)):
        set_cell_width(cell, w)
        if idx == 0:
            set_cell_no_wrap(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        if idx == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
        run.font.size = Pt(10.5)

add_para("")
add_para("ประเด็นที่ควรตรวจขั้นสุดท้ายก่อนส่ง", style="Heading 1")
final_checks = [
    "ใช้คำเรียกกรอบการตรวจสอบให้สม่ำเสมอทั้งบทความ เช่น ใช้ “การตรวจสอบด้วยช่วงเวลาอิสระ” สำหรับแนวคิด และ “ช่วงตรวจสอบอิสระ” สำหรับช่วงข้อมูล",
    "ตรวจข้อกำหนดของวารสารว่าต้องส่งต้นฉบับพร้อม line numbers หรือฉบับสะอาด หากไม่กำหนด ควรถอดเลขบรรทัดก่อนส่ง",
    "ตรวจการวางรูปและคำบรรยายภาพ โดยเฉพาะรูปที่ 7 ซึ่งในฉบับเรนเดอร์แยกภาพกับ caption คนละหน้า",
    "ตรวจความสอดคล้องของคำว่า calibration, holdout, validation และ fitting configuration ในภาษาไทยและอังกฤษ",
    "ตรวจว่า references ทุกฉบับที่อยู่ในรายการอ้างอิงยังมีการอ้างในเนื้อหา ตาราง หรือหมายเหตุตารางตามความจำเป็น",
]
for item in final_checks:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

add_para("")
add_para("โดยสรุป คำชี้แจงฉบับนี้ยืนยันว่าต้นฉบับได้ปรับแก้ประเด็นหลักของผู้ทรงคุณวุฒิแล้ว โดยเฉพาะการแก้ data leakage การแยกการประเมินรายแบบจำลอง การลดระดับข้อสรุป การเพิ่มการอนุมานทางสถิติ และการแก้รายการอ้างอิงที่มีปัญหา ทั้งนี้ยังควรตรวจรูปแบบเอกสารขั้นสุดท้ายก่อนส่งเข้าสู่ระบบวารสาร")

doc.save(out_path)
print(out_path)
