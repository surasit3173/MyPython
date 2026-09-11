"""Create the EnNRJ-ready Uttaradit precipitation-extremes manuscript.

The document is authored from the validated Uttaradit analysis outputs and
the written requirements in the 2025 Environment and Natural Resources
Journal (EnNRJ) author/style guides.  The supplied EnNRJ DOCX is a guidelines
export rather than a manuscript template, so page geometry and styles are
constructed explicitly here.
"""
from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


HERE = Path(__file__).resolve().parent
RESULTS = (
    HERE.parent
    / "utt_etccdi_easr_20260905"
    / "outputs"
    / "full_with_figures"
    / "results"
)
FIGURES = HERE / "figures_submission"
ELEVATION_FILE = Path(
    r"C:\MyPython\CMIP6Uttaradit\Data_Uttaradit\station_coordinates_Uttaradit.csv"
)
OUT = HERE / "EnNRJ_Uttaradit_Precipitation_Extremes_Manuscript.docx"

TITLE = "Projected Precipitation Extremes over Uttaradit Province under CMIP6 SSP Scenarios"
INDICES = [
    "PRCPTOT", "SDII", "Rx1day", "Rx5day", "CDD", "CWD",
    "R10mm", "R20mm", "R50mm", "R95p", "R99p",
]
MODELS = [
    "ACCESS-ESM1-5", "CanESM5", "CESM2", "EC-Earth3",
    "FGOALS-g3", "MIROC6", "MRI-ESM2-0",
]
SCENARIOS = ["SSP2-4.5", "SSP5-8.5"]


# ---------------------------------------------------------------------------
# Word-format helpers
# ---------------------------------------------------------------------------

def set_run_font(run, size: float = 12, *, bold: bool = False,
                 italic: bool = False, font: str = "Times New Roman") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def format_paragraph(paragraph, *, align=None, before: float = 0,
                     after: float = 0, line_spacing: float = 1.5,
                     keep_with_next: bool = False,
                     first_line_cm: float | None = None) -> None:
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    pf.keep_with_next = keep_with_next
    pf.widow_control = True
    if first_line_cm is not None:
        pf.first_line_indent = Cm(first_line_cm)


def add_text(doc: Document, text: str, *, size: float = 12,
             bold: bool = False, italic: bool = False, align=None,
             before: float = 0, after: float = 0,
             line_spacing: float = 1.5, keep_with_next: bool = False,
             first_line_cm: float | None = 0.75):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size, bold=bold, italic=italic)
    format_paragraph(
        p, align=align, before=before, after=after,
        line_spacing=line_spacing, keep_with_next=keep_with_next,
        first_line_cm=first_line_cm,
    )
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, 12, bold=(level == 1), italic=(level == 2))
    format_paragraph(
        p, before=6 if level == 1 else 3, after=0,
        line_spacing=1.5, keep_with_next=True, first_line_cm=0,
    )
    return p


def add_caption(doc: Document, label: str, text: str, *, above: bool):
    p = doc.add_paragraph()
    r1 = p.add_run(label + " ")
    set_run_font(r1, 10, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, 10)
    format_paragraph(
        p, before=3 if above else 0, after=0 if above else 6,
        line_spacing=1.0, keep_with_next=above, first_line_cm=0,
    )
    return p


def add_equation(doc: Document, expression: str, number: int) -> None:
    """Insert editable OMML math with centered expression and right number."""
    p = doc.add_paragraph()
    format_paragraph(
        p, align=WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2,
        line_spacing=1.0, first_line_cm=0,
    )
    ppr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    centre = OxmlElement("w:tab")
    centre.set(qn("w:val"), "center")
    centre.set(qn("w:pos"), "4520")
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:pos"), "9040")
    tabs.extend([centre, right])
    ppr.append(tabs)
    tab1 = p.add_run("\t")
    set_run_font(tab1, 11)

    omath = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mrpr = OxmlElement("m:rPr")
    msty = OxmlElement("m:sty")
    msty.set(qn("m:val"), "p")
    mrpr.append(msty)
    mt = OxmlElement("m:t")
    mt.set(qn("xml:space"), "preserve")
    mt.text = expression
    mr.extend([mrpr, mt])
    omath.append(mr)
    p._p.append(omath)

    tab2 = p.add_run(f"\t({number})")
    set_run_font(tab2, 11)


def set_cell_text(cell, text: object, *, bold: bool = False,
                  size: float = 9, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    set_run_font(r, size, bold=bold)
    format_paragraph(p, align=align, line_spacing=1.0, first_line_cm=0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.margin_top = Cm(0.04)
    cell.margin_bottom = Cm(0.04)


def _set_border(element, edge: str, value: str, size: str = "6") -> None:
    props = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element
    borders = props.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    old = borders.find(qn(f"w:{edge}"))
    if old is not None:
        borders.remove(old)
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), value)
    if value != "nil":
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
    borders.append(node)


def apply_minimal_table_rules(table) -> None:
    tblpr = table._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    for edge in ("left", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tblpr.append(borders)
    for cell in table.rows[0].cells:
        _set_border(cell._tc, "bottom", "single", "6")


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    trpr.append(node)


def set_table_widths(table, widths_inches: Sequence[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths_inches):
            cell.width = Inches(width)
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(int(width * 1440)))
            tcw.set(qn("w:type"), "dxa")


def add_table_note(doc: Document, text: str) -> None:
    p = add_text(
        doc, "Note: " + text, size=9, line_spacing=1.0,
        before=0, after=6, first_line_cm=0,
    )
    p.paragraph_format.keep_with_next = False


def add_figure(doc: Document, image_path: Path, label: str,
               caption: str, width_inches: float = 6.05) -> None:
    p = doc.add_paragraph()
    format_paragraph(
        p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=0,
        line_spacing=1.0, keep_with_next=True, first_line_cm=0,
    )
    p.add_run().add_picture(str(image_path), width=Inches(width_inches))
    add_caption(doc, label, caption, above=False)


def add_page_number(section) -> None:
    footer = section.footer
    footer.distance = Cm(0.9)
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.extend([begin, instr, separate])
    set_run_font(run, 10)
    run2 = p.add_run("1")
    set_run_font(run2, 10)
    run2._r.append(end)


def add_line_numbering(section) -> None:
    sectpr = section._sectPr
    old = sectpr.find(qn("w:lnNumType"))
    if old is not None:
        sectpr.remove(old)
    node = OxmlElement("w:lnNumType")
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:distance"), "360")
    node.set(qn("w:restart"), "continuous")
    sectpr.append(node)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    add_line_numbering(section)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------------------
# Validated analytical summaries
# ---------------------------------------------------------------------------

def bh_fdr_mask(pvalues: pd.Series, q: float = 0.05) -> pd.Series:
    p = pd.to_numeric(pvalues, errors="coerce")
    mask = pd.Series(False, index=p.index)
    valid = p.dropna().sort_values()
    if valid.empty:
        return mask
    ranks = np.arange(1, len(valid) + 1)
    passed = valid.to_numpy(float) <= q * ranks / len(valid)
    if passed.any():
        cutoff = float(valid.iloc[np.flatnonzero(passed)[-1]])
        mask.loc[p.index[p <= cutoff]] = True
    return mask


def observed_summary() -> pd.DataFrame:
    raw = pd.read_csv(RESULTS / "observed_trend_1981_2014.csv", dtype={"Station": str})
    raw["BH"] = bh_fdr_mask(raw["p_MK"], 0.05)
    units = {
        "PRCPTOT": "mm", "SDII": "mm/day", "Rx1day": "mm",
        "Rx5day": "mm", "CDD": "days", "CWD": "days",
        "R10mm": "days", "R20mm": "days", "R50mm": "days",
        "R95p": "mm", "R99p": "mm",
    }
    rows = []
    for index in INDICES:
        sub = raw.loc[raw["Index"] == index].copy()
        sig = pd.to_numeric(sub["p_MK"], errors="coerce") < 0.05
        z = pd.to_numeric(sub["Z_MK"], errors="coerce")
        bh = sub["BH"]
        rows.append({
            "Index": index,
            "Unit": units[index],
            "MK_inc": int((sig & (z > 0)).sum()),
            "MK_dec": int((sig & (z < 0)).sum()),
            "MK_NS": int((~sig).sum()),
            "BH_inc": int((bh & (z > 0)).sum()),
            "BH_dec": int((bh & (z < 0)).sum()),
            "Median_slope_decade": float(pd.to_numeric(sub["Sen_slope"], errors="coerce").median() * 10.0),
        })
    return pd.DataFrame(rows)


def future_summary() -> pd.DataFrame:
    raw = pd.read_csv(RESULTS / "future_change_2021_2050.csv", dtype={"Station": str})
    model_first = (
        raw.groupby(["Scenario", "Model", "Index"], as_index=False, observed=True)
        ["Relative_change_pct"].median()
    )
    rows = []
    for scenario in SCENARIOS:
        for index in INDICES:
            vals = (
                model_first.loc[
                    (model_first["Scenario"] == scenario) &
                    (model_first["Index"] == index)
                ]
                .set_index("Model")
                .reindex(MODELS)["Relative_change_pct"]
                .to_numpy(float)
            )
            q1, med, q3 = np.nanquantile(vals, [0.25, 0.50, 0.75])
            finite = vals[np.isfinite(vals)]
            rows.append({
                "Scenario": scenario,
                "Index": index,
                "Median": med,
                "Q1": q1,
                "Q3": q3,
                "Min": np.min(finite),
                "Max": np.max(finite),
                "Positive": int(np.sum(finite > 0)),
                "Negative": int(np.sum(finite < 0)),
                "Zero": int(np.sum(np.isclose(finite, 0.0, atol=1e-12))),
            })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------

def add_station_table(doc: Document) -> None:
    coords = pd.read_csv(RESULTS / "station_coordinates_used.csv", dtype={"station": str})
    elevations = pd.read_csv(ELEVATION_FILE, dtype={"Station_ID": str})
    elevations = elevations.rename(columns={"Station_ID": "station", "Elevation": "elevation"})
    data = coords.merge(elevations[["station", "elevation"]], on="station", how="left").sort_values("station")
    if len(data) != 13 or data["elevation"].isna().any():
        raise ValueError("Station table does not contain the validated 13 gauges with elevation metadata.")
    add_caption(doc, "Table 1.", "Rain-gauge stations used in the Uttaradit analysis.", above=True)
    headers = ["Station ID", "Latitude (°N)", "Longitude (°E)", "Elevation (m MSL)"]
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    repeat_header(table.rows[0])
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, size=9)
    for row in data.itertuples(index=False):
        cells = table.add_row().cells
        values = [row.station, f"{row.latitude:.4f}", f"{row.longitude:.4f}", f"{row.elevation:.2f}"]
        for j, value in enumerate(values):
            set_cell_text(cells[j], value, size=9)
        prevent_row_split(table.rows[-1])
    set_table_widths(table, [1.22, 1.55, 1.60, 1.70])
    apply_minimal_table_rules(table)
    add_table_note(
        doc,
        "Station IDs and coordinates are the inputs used in the analysis. Elevation is supplied station metadata and was not used to calculate the indices, trends, changes, or IDW weights. Station names were not created because the analytical input identifies gauges by ID only.",
    )


def add_model_table(doc: Document) -> None:
    rows = [
        ("ACCESS-ESM1-5", "CSIRO and ARCCSS", "Australia", "r1i1p1f1 / gn"),
        ("CanESM5", "CCCma", "Canada", "r1i1p1f1 / gn"),
        ("CESM2", "NCAR", "United States", "r11i1p1f1 / gn"),
        ("EC-Earth3", "EC-Earth Consortium", "Europe", "r1i1p1f1 / gr"),
        ("FGOALS-g3", "IAP, Chinese Academy of Sciences", "China", "r1i1p1f1 / gn"),
        ("MIROC6", "MIROC Consortium", "Japan", "r1i1p1f1 / gn"),
        ("MRI-ESM2-0", "Meteorological Research Institute", "Japan", "r1i1p1f1 / gn"),
    ]
    add_caption(doc, "Table 2.", "CMIP6 global climate models used in the study.", above=True)
    headers = ["Model", "Institution/consortium", "Country/region", "Member/grid", "Data used"]
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    repeat_header(table.rows[0])
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, size=8.4)
    for model, institution, country, member in rows:
        cells = table.add_row().cells
        values = [model, institution, country, member, "Baseline; SSP2-4.5; SSP5-8.5"]
        for j, value in enumerate(values):
            set_cell_text(
                cells[j], value, size=8.2,
                align=WD_ALIGN_PARAGRAPH.LEFT if j in (0, 1, 4) else WD_ALIGN_PARAGRAPH.CENTER,
            )
        prevent_row_split(table.rows[-1])
    set_table_widths(table, [1.06, 1.58, 0.92, 0.92, 1.64])
    apply_minimal_table_rules(table)
    add_table_note(
        doc,
        "Historical data for 1995–2014 provided the model-consistent baseline, and SSP2-4.5 and SSP5-8.5 data for 2021–2050 provided the future period. Daily files were supplied as bias-corrected series; the correction procedure was not re-estimated or independently validated in this study.",
    )


def add_index_table(doc: Document) -> None:
    rows = [
        ("PRCPTOT", "Annual total wet-day precipitation", "Sum of RR on days with RR≥1 mm", "mm"),
        ("SDII", "Simple daily intensity index", "PRCPTOT divided by the number of wet days", "mm/day"),
        ("Rx1day", "Maximum 1-day precipitation", "Annual maximum 1-day precipitation", "mm"),
        ("Rx5day", "Maximum 5-day precipitation", "Annual maximum consecutive 5-day precipitation total", "mm"),
        ("CDD", "Consecutive dry days", "Maximum consecutive days with RR<1 mm", "days"),
        ("CWD", "Consecutive wet days", "Maximum consecutive days with RR≥1 mm", "days"),
        ("R10mm", "Heavy precipitation days", "Number of days with RR≥10 mm", "days"),
        ("R20mm", "Very heavy precipitation days", "Number of days with RR≥20 mm", "days"),
        ("R50mm", "Extremely heavy precipitation days", "Number of days with RR≥50 mm", "days"),
        ("R95p", "Very wet-day precipitation", "Annual precipitation total from wet days exceeding the fixed station P95 threshold", "mm"),
        ("R99p", "Extremely wet-day precipitation", "Annual precipitation total from wet days exceeding the fixed station P99 threshold", "mm"),
    ]
    add_caption(doc, "Table 3.", "Definitions of precipitation-extreme indices.", above=True)
    headers = ["Index", "Name", "Operational definition", "Unit"]
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    repeat_header(table.rows[0])
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, size=8.5)
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(
                cells[j], value, size=8.1,
                align=WD_ALIGN_PARAGRAPH.LEFT if j in (1, 2) else WD_ALIGN_PARAGRAPH.CENTER,
            )
        prevent_row_split(table.rows[-1])
    set_table_widths(table, [0.72, 1.68, 3.03, 0.72])
    apply_minimal_table_rules(table)
    add_table_note(
        doc,
        "RR denotes daily precipitation. A wet day was RR≥1 mm. Station-specific P95 and P99 thresholds were estimated from observed wet days during 1981–2014 and then held fixed. R50mm is a study-specific fixed-threshold index.",
    )


def add_observed_table(doc: Document, summary: pd.DataFrame) -> None:
    add_caption(
        doc, "Table 4.",
        "Observed precipitation-index trends at 13 stations during 1981–2014.",
        above=True,
    )
    headers = ["Index", "Unit", "MK inc.", "MK dec.", "MK NS", "BH inc.", "BH dec.", "Median Sen slope/decade"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    repeat_header(table.rows[0])
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, size=8.0)
    for row in summary.itertuples(index=False):
        cells = table.add_row().cells
        values = [
            row.Index, row.Unit, row.MK_inc, row.MK_dec, row.MK_NS,
            row.BH_inc, row.BH_dec, f"{row.Median_slope_decade:.2f}",
        ]
        for j, value in enumerate(values):
            set_cell_text(cells[j], value, size=8.0)
        prevent_row_split(table.rows[-1])
    set_table_widths(table, [0.95, 0.72, 0.60, 0.60, 0.60, 0.60, 0.60, 1.45])
    apply_minimal_table_rules(table)
    add_table_note(
        doc,
        "MK inc., MK dec., and MK NS are counts from the standard two-sided Mann–Kendall test at α=0.05. BH inc. and BH dec. are counts retained after Benjamini–Hochberg false-discovery-rate control across all 143 station–index tests. Sen slopes retain the index unit per decade.",
    )


def add_future_table(doc: Document, summary: pd.DataFrame) -> None:
    add_caption(
        doc, "Table 5.",
        "Model-first relative changes in annual precipitation indices for 2021–2050 relative to 1995–2014.",
        above=True,
    )
    headers = ["Scenario", "Index", "Median (%)", "IQR (%)", "Range (%)", "Models +/−/0"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    repeat_header(table.rows[0])
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, size=8.4)
    for row in summary.itertuples(index=False):
        cells = table.add_row().cells
        values = [
            row.Scenario, row.Index, f"{row.Median:.1f}",
            f"{row.Q1:.1f} to {row.Q3:.1f}", f"{row.Min:.1f} to {row.Max:.1f}",
            f"{row.Positive}/{row.Negative}/{row.Zero}",
        ]
        for j, value in enumerate(values):
            set_cell_text(
                cells[j], value, size=8.2,
                align=WD_ALIGN_PARAGRAPH.LEFT if j in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER,
            )
        prevent_row_split(table.rows[-1])
    set_table_widths(table, [0.93, 0.69, 0.82, 1.20, 1.20, 0.92])
    apply_minimal_table_rules(table)
    add_table_note(
        doc,
        "For each model, station-level relative changes were summarized by the station median; the displayed median, IQR, and range were then calculated across the seven model-level values. +/−/0 denotes the number of model-level station medians above/below/equal to zero. Undefined R99p relative changes at station 351011 (zero baseline) were excluded, so the within-model R99p median used 12 stations. The ensemble statistics are descriptive and do not assume model independence.",
    )


# ---------------------------------------------------------------------------
# Manuscript text
# ---------------------------------------------------------------------------

ABSTRACT = (
    "Projected changes in precipitation extremes were assessed for Uttaradit Province, northern Thailand, using daily series at 13 rain gauges from seven bias-corrected Coupled Model Intercomparison Project Phase 6 (CMIP6) models under SSP2-4.5 and SSP5-8.5. Eleven annual indices were calculated for 2021–2050 and compared with each model’s 1995–2014 baseline. A model-first scheme summarized station medians within each model before estimating the seven-model median, interquartile range, range, and sign agreement. Station-referenced inverse-distance weighting described spatial contrasts in annual total wet-day precipitation (PRCPTOT). Under SSP2-4.5, median changes were +2.9% for PRCPTOT, +6.1% for maximum 1-day precipitation (Rx1day), +4.2% for maximum 5-day precipitation (Rx5day), and +26.7% for extremely wet-day precipitation (R99p). All seven models indicated increases in Rx1day and R99p, and six indicated higher Rx5day. Under SSP5-8.5, corresponding medians were +0.9%, +12.3%, +2.9%, and +14.4%, respectively, although model ranges crossed zero. PRCPTOT changes were weakly positive at most stations, with localized negative areas under both scenarios. The results suggest possible intensification of short-duration and upper-tail rainfall without a comparably robust increase in annual total. Wide inter-model spread, model non-independence, and the descriptive nature of the spatial interpolation limit deterministic interpretation."
)

KEYWORDS = [
    "precipitation extremes", "climate projections", "CMIP6",
    "SSP scenarios", "Uttaradit Province", "model uncertainty",
]

HIGHLIGHTS = [
    "Seven CMIP6 models projected near-term extremes at 13 rain gauges.",
    "Rx1day increased in all models under SSP2-4.5.",
    "R99p median change reached 26.7%, but model spread was wide.",
    "PRCPTOT change was weakly positive with local negative pockets.",
    "IDW maps depict station contrast, not gridded climate forecasts.",
]


REFERENCES = [
    "Alexander LV, Zhang X, Peterson TC, Caesar J, Gleason B, Klein Tank AMG, et al. Global observed changes in daily climate extremes of temperature and precipitation. J Geophys Res Atmos. 2006;111:D05109. doi:10.1029/2005JD006290.",
    "Benjamini Y, Hochberg Y. Controlling the false discovery rate: a practical and powerful approach to multiple testing. J R Stat Soc Series B Stat Methodol. 1995;57(1):289–300. doi:10.1111/j.2517-6161.1995.tb02031.x.",
    "Cannon AJ, Sobie SR, Murdock TQ. Bias correction of GCM precipitation by quantile mapping: how well do methods preserve changes in quantiles and extremes? J Clim. 2015;28(17):6938–6959. doi:10.1175/JCLI-D-14-00754.1.",
    "Chaiwino W, Chaisee K, Oonariya C, Wongsaijai B. Analysis of long-term rainfall trend and extreme in upper northern Thailand. Sci Rep. 2025;15:33380. doi:10.1038/s41598-025-18217-1.",
    "de Oliveira-Júnior JF, Mendes D, Porto HD, Cardoso KRA, Ferreira Neto JA, da Silva EBC, et al. Analysis of drought and extreme precipitation events in Thailand: trends, climate modeling, and implications for climate change adaptation. Sci Rep. 2025;15:4501. doi:10.1038/s41598-025-86826-x.",
    "Eyring V, Bony S, Meehl GA, Senior CA, Stevens B, Stouffer RJ, et al. Overview of the Coupled Model Intercomparison Project Phase 6 (CMIP6) experimental design and organization. Geosci Model Dev. 2016;9:1937–1958. doi:10.5194/gmd-9-1937-2016.",
    "Hamed KH, Rao AR. A modified Mann–Kendall trend test for autocorrelated data. J Hydrol. 1998;204(1–4):182–196. doi:10.1016/S0022-1694(97)00125-X.",
    "Humphries UW, Waqas M, Hlaing PT, Dechpichai P, Wangwongchai A. Assessment of CMIP6 GCMs for selecting a suitable climate model for precipitation projections in Southern Thailand. Results Eng. 2024;23:102417. doi:10.1016/j.rineng.2024.102417.",
    "Intergovernmental Panel on Climate Change. Climate change 2021: the physical science basis. Contribution of Working Group I to the Sixth Assessment Report of the Intergovernmental Panel on Climate Change. Cambridge: Cambridge University Press; 2021. doi:10.1017/9781009157896.",
    "Knutti R, Sedláček J, Sanderson BM, Lorenz R, Fischer EM, Eyring V. A climate model projection weighting scheme accounting for performance and interdependence. Geophys Res Lett. 2017;44(4):1909–1918. doi:10.1002/2016GL072012.",
    "Kuinkel D, Promchote P, Upreti KR, Wang S-Y. Projected changes in precipitation extremes in Southern Thailand using CMIP6 models. Theor Appl Climatol. 2024;155:8703–8716. doi:10.1007/s00704-024-05150-y.",
    "Madolli MJ, Gade SA, Gupta V, Chakraborty A, et al. A systematic review on rainfall patterns of Thailand: insights into variability and its relationship with ENSO and IOD. Earth-Sci Rev. 2025;264:105102. doi:10.1016/j.earscirev.2025.105102.",
    "Mann HB. Nonparametric tests against trend. Econometrica. 1945;13(3):245–259. doi:10.2307/1907187.",
    "Maraun D. Bias correcting climate change simulations—a critical review. Curr Clim Change Rep. 2016;2:211–220. doi:10.1007/s40641-016-0050-x.",
    "Maraun D, Shepherd TG, Widmann M, Zappa G, Walton D, Gutiérrez JM, et al. Towards process-informed bias correction of climate change simulations. Nat Clim Chang. 2017;7:764–773. doi:10.1038/nclimate3418.",
    "Nontikansak P, Shrestha S, Shanmugam MS, Loc HH, Virdis SGP. Rainfall extremes under climate change in the Pasak River Basin, Thailand. J Water Clim Change. 2022;13(10):3729–3746. doi:10.2166/wcc.2022.232.",
    "O’Neill BC, Tebaldi C, van Vuuren DP, Eyring V, Friedlingstein P, Hurtt G, et al. The Scenario Model Intercomparison Project (ScenarioMIP) for CMIP6. Geosci Model Dev. 2016;9:3461–3482. doi:10.5194/gmd-9-3461-2016.",
    "Riahi K, van Vuuren DP, Kriegler E, Edmonds J, O’Neill BC, Fujimori S, et al. The Shared Socioeconomic Pathways and their energy, land use, and greenhouse gas emissions implications: an overview. Glob Environ Change. 2017;42:153–168. doi:10.1016/j.gloenvcha.2016.05.009.",
    "Sen PK. Estimates of the regression coefficient based on Kendall’s tau. J Am Stat Assoc. 1968;63(324):1379–1389. doi:10.1080/01621459.1968.10480934.",
    "Shepard D. A two-dimensional interpolation function for irregularly-spaced data. In: Proceedings of the 23rd ACM National Conference. New York: ACM; 1968. p. 517–524. doi:10.1145/800186.810616.",
    "Sillmann J, Kharin VV, Zhang X, Zwiers FW, Bronaugh D. Climate extremes indices in the CMIP5 multimodel ensemble: part 1. Model evaluation in the present climate. J Geophys Res Atmos. 2013;118(4):1716–1733. doi:10.1002/jgrd.50203.",
    "Sillmann J, Kharin VV, Zwiers FW, Zhang X, Bronaugh D. Climate extremes indices in the CMIP5 multimodel ensemble: part 2. Future climate projections. J Geophys Res Atmos. 2013;118(6):2473–2493. doi:10.1002/jgrd.50188.",
    "Tebaldi C, Debeire K, Eyring V, Fischer E, Fyfe J, Friedlingstein P, et al. Climate model projections from the Scenario Model Intercomparison Project (ScenarioMIP) of CMIP6. Earth Syst Dyn. 2021;12:253–293. doi:10.5194/esd-12-253-2021.",
    "Try S, Qin X. Evaluation of future changes in climate extremes over Southeast Asia using downscaled CMIP6 GCM projections. Water. 2024;16:2207. doi:10.3390/w16152207.",
    "Waqas M, Humphries UW. Artificial intelligence-driven precipitation downscaling and projections over Thailand using CMIP6 climate models. Big Earth Data. 2025;9(3):583–614. doi:10.1080/20964471.2025.2547500.",
    "Yue S, Wang CY. The Mann–Kendall test modified by effective sample size to detect trend in serially correlated hydrological series. Water Resour Manag. 2004;18:201–218. doi:10.1023/B:WARM.0000043140.61082.60.",
    "Zhang B, Song S, Wang H, Guo T, Ding Y. Evaluation of the performance of CMIP6 models in simulating extreme precipitation and its projected changes in global climate regions. Nat Hazards. 2024;121(2):1737–1763. doi:10.1007/s11069-024-06850-4.",
    "Zhang X, Alexander L, Hegerl GC, Jones P, Tank AMGK, Peterson TC, et al. Indices for monitoring changes in extremes based on daily temperature and precipitation data. Wiley Interdiscip Rev Clim Change. 2011;2(6):851–870. doi:10.1002/wcc.147.",
]


def add_front_matter(doc: Document) -> None:
    p = add_text(
        doc, TITLE, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
        before=18, after=12, first_line_cm=0,
    )
    p.paragraph_format.keep_with_next = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Surasit Punyawansiri")
    set_run_font(r, 12, bold=True)
    r = p.add_run("1,*")
    set_run_font(r, 12, bold=True)
    r.font.superscript = True
    format_paragraph(p, after=6, first_line_cm=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("1")
    set_run_font(r, 11, italic=True)
    r.font.superscript = True
    r = p.add_run("Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok, Thailand")
    set_run_font(r, 11, italic=True)
    format_paragraph(p, after=6, first_line_cm=0)
    add_text(
        doc, "*Corresponding author: Surasit.pu@ku.th", size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0,
    )

    doc.add_page_break()
    add_heading(doc, "Abstract", 1)
    add_text(doc, ABSTRACT, first_line_cm=0)
    p = doc.add_paragraph()
    r = p.add_run("Keywords: ")
    set_run_font(r, 12, bold=True)
    r = p.add_run("; ".join(KEYWORDS))
    set_run_font(r, 12)
    format_paragraph(p, first_line_cm=0)

    add_heading(doc, "Highlights", 1)
    for item in HIGHLIGHTS:
        p = doc.add_paragraph()
        r = p.add_run("• " + item)
        set_run_font(r, 12)
        format_paragraph(p, first_line_cm=0)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
    doc.add_page_break()


def add_introduction(doc: Document) -> None:
    add_heading(doc, "1. Introduction", 1)
    add_text(
        doc,
        "Precipitation extremes influence flood generation, soil erosion, reservoir operation, agricultural water demand, and drought exposure. A warmer atmosphere can hold more moisture, but regional circulation and monsoon variability determine whether this thermodynamic potential appears as higher event intensity, altered wet-day frequency, or longer dry and wet spells. Consequently, changes in annual precipitation totals alone do not describe the full hydroclimatic response (Alexander et al., 2006; Intergovernmental Panel on Climate Change, 2021; Zhang et al., 2011). A coordinated set of intensity, duration, frequency, and percentile-based indices is therefore required to identify potentially important changes in both rainfall accumulation and distribution.",
    )
    add_text(
        doc,
        "The Coupled Model Intercomparison Project Phase 6 (CMIP6) and Scenario Model Intercomparison Project provide a common experimental framework for assessing future climate under alternative forcing pathways (Eyring et al., 2016; O’Neill et al., 2016; Riahi et al., 2017; Tebaldi et al., 2021). Multi-model ensembles allow structural differences among models to be displayed, but their members are neither independent nor equally credible at every location. Ensemble summaries should therefore be interpreted as descriptive evidence conditioned on the selected models, scenarios, and processing chain, rather than as calibrated probabilities (Knutti et al., 2017; Sillmann et al., 2013a; Sillmann et al., 2013b). This issue is particularly important for precipitation tails, where model spread and small baseline values can produce large relative changes.",
    )
    add_text(
        doc,
        "Local applications also depend on downscaling and bias correction. These procedures can improve the representation of observed distributions, but they may alter temporal dependence, wet-day occurrence, and event sequencing if not evaluated for the target statistic (Cannon et al., 2015; Maraun, 2016; Maraun et al., 2017). Regional studies over Southeast Asia and Thailand indicate spatially heterogeneous future rainfall extremes, substantial inter-model spread, and sensitivity to model or method selection (Humphries et al., 2024; Kuinkel et al., 2024; Nontikansak et al., 2022; Try and Qin, 2024; Waqas and Humphries, 2025; Zhang et al., 2024). Recent observational syntheses likewise show that Thai rainfall variability cannot be represented by one national trend (Chaiwino et al., 2025; de Oliveira-Júnior et al., 2025; Madolli et al., 2025). Province-scale, station-referenced evidence remains useful for planning when its uncertainty and spatial support are stated explicitly.",
    )
    add_text(
        doc,
        "Uttaradit Province lies in upper northern Thailand and contains a rain-gauge network spanning low-elevation and upland settings. The available projections support a near-term comparison for 2021–2050, a period directly relevant to infrastructure and water-management decisions. However, a defensible provincial assessment must preserve station-level information, apply a consistent historical baseline within each model, and distinguish spatial visualization from independently validated gridded projection.",
    )
    add_text(
        doc,
        "This study evaluated projected changes in 11 annual precipitation-extreme indices at 13 stations in Uttaradit Province under SSP2-4.5 and SSP5-8.5. The objectives were to (1) quantify model-first changes for 2021–2050 relative to 1995–2014, including interquartile range, full model range, and sign agreement; (2) describe the spatial pattern of station-level PRCPTOT changes using inverse-distance weighting (IDW); and (3) provide concise observational and future-window trend diagnostics without conflating them with baseline-to-future climatological change. The primary contribution is a transparent, projection-focused synthesis in which ensemble spread and limits of spatial inference are treated as results rather than ancillary qualifications.",
    )


def add_methodology(doc: Document) -> None:
    add_heading(doc, "2. Methodology", 1)
    add_heading(doc, "2.1 Study area and data", 2)
    add_text(
        doc,
        "The analysis used 13 rain gauges distributed across Uttaradit Province (Figure 1 and Table 1). Station coordinates ranged from 17.2300°N to 18.0200°N and from 100.0500°E to 101.0700°E; supplied station elevations ranged from 54.57 to 427.85 m MSL. Coordinates used for analysis were read directly from the validated station input. Elevation was retained as descriptive metadata and did not enter the statistical or IDW calculations.",
    )
    add_station_table(doc)
    add_figure(
        doc, FIGURES / "Figure1_Uttaradit_study_area.png", "Figure 1.",
        "Location of Uttaradit Province and the 13 rain gauges used in the analysis. The inset identifies the province within Thailand. Station labels are gauge IDs; no station names were inferred.",
        5.75,
    )
    add_text(
        doc,
        "The observed input comprised daily precipitation from 1 January 1981 to 31 December 2014. Seven supplied bias-corrected CMIP6 daily precipitation series were analyzed (Table 2). For each model, 1995–2014 was used as the model-consistent historical baseline and 2021–2050 as the future period. The two scenarios were SSP2-4.5 and SSP5-8.5. The file inventory included one ensemble member and native or regridded grid label per model. Because the supplied daily files did not contain complete metadata for the preceding bias-correction implementation, that procedure was neither reconstructed nor independently validated; all conclusions are conditional on the supplied corrected series.",
    )
    add_model_table(doc)

    add_heading(doc, "2.2 Precipitation indices and quality control", 2)
    add_text(
        doc,
        "Eleven annual precipitation indices represented wet-day amount and intensity, annual maxima, spell duration, fixed-threshold frequency, and percentile exceedance totals (Table 3). These follow the Expert Team on Climate Change Detection and Indices framework except R50mm, which is a study-specific fixed-threshold count (Alexander et al., 2006; Zhang et al., 2011). A wet day was defined as daily precipitation RR≥1 mm. The station-specific 95th- and 99th-percentile thresholds were estimated from all observed wet days during 1981–2014 and held fixed for observed and model calculations. This fixed-reference construction prevents the percentile threshold itself from changing between periods.",
    )
    add_index_table(doc)
    add_text(
        doc,
        "Daily records were converted to numeric values after explicit removal of predefined missing-value flags. Invalid dates, duplicate dates, negative precipitation, duplicate normalized station IDs, and nonnumeric values were rejected rather than silently repaired. Each series was reindexed to the complete Gregorian calendar. No missing precipitation was imputed; a missing date interrupted a consecutive spell and invalidated any five-day window containing it. An annual index was retained only when at least 90% of calendar days were valid. All 442 observed station-years and all 5,460 future station-model-scenario-years passed this criterion. The minimum completeness was 100% for observations and 99.73% for model series. The resulting data contained 4,862 observed and 60,060 future annual-index values. Observed station thresholds ranged from 12.700 to 43.370 mm for P95 and from 19.723 to 89.553 mm for P99.",
    )

    add_heading(doc, "2.3 Trend inference and multiplicity", 2)
    add_text(
        doc,
        "The standard two-sided Mann–Kendall (MK) test was the primary trend inference for the 1981–2014 observed annual series (Mann, 1945). For n observations, the score S was calculated by Equation 1, with ties included in the variance in Equation 2. The continuity-corrected standardized statistic in Equation 3 was evaluated against the standard normal distribution at α=0.05.",
    )
    add_equation(doc, "S = Σₖ₌₁ⁿ⁻¹ Σₗ₌ₖ₊₁ⁿ sgn(xₗ − xₖ)", 1)
    add_equation(doc, "Var(S) = [n(n−1)(2n+5) − Σₜ t(t−1)(2t+5)] / 18", 2)
    add_equation(doc, "Z = (S−1)/√Var(S), S>0;  Z = 0, S=0;  Z = (S+1)/√Var(S), S<0", 3)
    add_text(
        doc,
        "Trend magnitude was estimated with Sen’s median pairwise slope (Equation 4), where t denotes year (Sen, 1968). Because 13 stations and 11 indices produced 143 simultaneous observed tests, the Benjamini–Hochberg false-discovery-rate (BH-FDR) procedure was applied as one family at q=0.05. Ordered p-values were compared with the criterion in Equation 5, where m=143 and k is the largest passing rank (Benjamini and Hochberg, 1995). Both unadjusted MK counts and BH-FDR-retained counts were reported so that multiplicity control is transparent.",
    )
    add_equation(doc, "β̂ = medianₖ<ₗ [(xₗ − xₖ)/(tₗ − tₖ)]", 4)
    add_equation(doc, "p₍ₖ₎ ≤ (k/m)q", 5)
    add_text(
        doc,
        "Serial dependence was examined using the Yue–Wang detrended effective-sample-size formulation as a sensitivity diagnostic, not as the primary decision rule (Yue and Wang, 2004). After removal of the Sen trend, the full-lag factor and adjusted variance were calculated by Equation 6. This implementation differs from the rank-autocorrelation modification of Hamed and Rao (1998). Factors below one occurred frequently in the short annual series and can reduce the variance; therefore, sensitivity-test significance was not interpreted as confirmatory evidence.",
    )
    add_equation(doc, "n/n* = 1 + 2Σₕ₌₁ⁿ⁻¹(1−h/n)ρₕ;   Var*(S) = Var(S)(n/n*)", 6)

    add_heading(doc, "2.4 Projection-change aggregation", 2)
    add_text(
        doc,
        "For model m, scenario s, index i, and station j, the relative change compared the future-period mean with the same model’s historical mean (Equation 7). Using a model-consistent baseline avoids combining observed and modeled climatologies in the denominator. If the baseline mean was zero, relative change was undefined and recorded as missing; no arbitrary denominator or offset was introduced. This occurred for R99p at station 351011 in seven models under both scenarios (14 station-level rows). Those values were excluded from the relevant station-median calculation, leaving 12 stations for each model’s R99p summary.",
    )
    add_equation(doc, "Δₘ,ₛ,ᵢ,ⱼ = 100( X̄F,ₘ,ₛ,ᵢ,ⱼ − X̄B,ₘ,ᵢ,ⱼ ) / X̄B,ₘ,ᵢ,ⱼ", 7)
    add_text(
        doc,
        "Aggregation was model-first. Equation 8 took the median across stations within each model. Equation 9 then took the median across the seven model-level values; the interquartile range (IQR), full range, and counts of positive, negative, and zero model medians were calculated from the same seven values. This order gives every model one contribution to the ensemble summary, regardless of the magnitude of its within-province spatial variation. The seven-model median is a descriptive measure of central tendency and not an estimate of event probability or formal model independence.",
    )
    add_equation(doc, "Dₘ,ₛ,ᵢ = medianⱼ₌₁,…,J (Δₘ,ₛ,ᵢ,ⱼ)", 8)
    add_equation(doc, "D̃ₛ,ᵢ = medianₘ₌₁,…,M (Dₘ,ₛ,ᵢ),   M=7", 9)
    add_text(
        doc,
        "A standard MK test was also applied separately to each 2021–2050 model–station–scenario–index annual series. These 2,002 tests were retained only as unadjusted within-window diagnostics. They answer whether an index changed monotonically during 2021–2050, whereas Equations 7–9 answer whether its future climatological level differed from the 1995–2014 baseline. No field-significance claim was made from the diagnostic counts.",
    )

    add_heading(doc, "2.5 Spatial depiction", 2)
    add_text(
        doc,
        "For PRCPTOT, the relative change was first summarized as the median across seven models at each station and scenario. IDW with power p=2 was then applied using great-circle distances in kilometers (Equation 10), where zⱼ is the station value and dⱼ is its distance to an interpolation point (Shepard, 1968). The surface was evaluated on a regular grid and clipped to the Uttaradit provincial polygon. The same symmetric percentage scale was used for both scenarios. Values outside the station convex hull, including the far northern part of the polygon, are spatial extrapolations. The maps are descriptive station-referenced depictions, not dynamically downscaled or independently validated continuous-resolution climate projections.",
    )
    add_equation(doc, "ẑ(x) = [Σⱼ₌₁ᴶ zⱼ dⱼ(x)⁻ᵖ] / [Σⱼ₌₁ᴶ dⱼ(x)⁻ᵖ],   p=2", 10)


def add_results_discussion(doc: Document, obs: pd.DataFrame, fut: pd.DataFrame) -> None:
    add_heading(doc, "3. Results and Discussion", 1)
    add_heading(doc, "3.1 Data readiness and historical context", 2)
    add_text(
        doc,
        "Quality control retained the complete intended station network and all required model–scenario files. The standard MK test identified significant trends in 30 of 143 observed station–index series; BH-FDR retained 14. No station showed a significant PRCPTOT trend by either decision rule, and the station-median Sen slope was −6.75 mm/decade (Table 4). The BH-FDR results were spatially heterogeneous: retained changes included three decreases in CDD, three R50mm changes in opposing directions, and mixed R95p and R99p directions. The observed period therefore provides local context but does not establish a uniform historical trajectory from which future changes can be extrapolated.",
    )
    add_observed_table(doc, obs)

    add_heading(doc, "3.2 Projected changes across precipitation indices", 2)
    add_text(
        doc,
        "The model-first ensemble showed that projected changes in event maxima and upper-tail totals were generally larger than the change in annual wet-day precipitation (Table 5 and Figure 2). Under SSP2-4.5, the median PRCPTOT change was +2.9% (IQR −4.0% to +4.9%; range −5.4% to +18.1%), with four model medians positive and three negative. By contrast, Rx1day increased by a median of 6.1% (IQR +0.5% to +15.4%) and R99p by 26.7% (IQR +13.7% to +33.5%); all seven model-level medians were positive for both indices. Rx5day increased by 4.2%, with six of seven models positive. These sign patterns provide stronger ensemble evidence for intensified short-duration and upper-tail precipitation than for a uniform increase in annual total.",
    )
    add_text(
        doc,
        "Under SSP5-8.5, median PRCPTOT change remained weakly positive at +0.9% (IQR −5.0% to +5.5%; range −9.9% to +15.3%), again with a four-to-three split in sign. Rx1day had the largest scenario amplification among the headline indices, with a median increase of 12.3% compared with 6.1% under SSP2-4.5. Nevertheless, its model range extended from −10.6% to +23.4%, and only five models were positive. Median changes in Rx5day and R99p were +2.9% and +14.4%, respectively, but their IQRs crossed zero and sign agreement weakened to five of seven and four of seven models. The higher forcing scenario therefore did not yield a uniformly larger median response across all indices.",
    )
    add_text(
        doc,
        "The remaining indices reinforced the mixed character of the projections. Under SSP2-4.5, CDD had a median change of −2.1% with six models negative, while R10mm and R20mm had small negative medians of −0.9% and −2.6%. R50mm increased by a median of 24.4%, and R95p by 4.5%, but their ranges crossed zero. Under SSP5-8.5, median changes were −2.4% for SDII, −3.0% for R10mm, −3.1% for R20mm, +17.5% for R50mm, and +4.3% for R95p. R50mm and R99p displayed particularly broad ranges because rare-event totals or counts can have small baseline means. Their percentage changes should therefore be read together with ensemble spread rather than treated as typical deterministic outcomes.",
    )
    add_future_table(doc, fut)
    add_figure(
        doc, FIGURES / "Figure2_future_model_changes.png", "Figure 2.",
        "Model-first relative changes in annual precipitation indices for 2021–2050 relative to 1995–2014 under (a) SSP2-4.5 and (b) SSP5-8.5. Each open circle represents one model after its station median was calculated; the black point is the seven-model median and the thick horizontal segment is the interquartile range. A common x-axis permits direct scenario comparison.",
    )

    add_heading(doc, "3.3 Spatial pattern of projected PRCPTOT change", 2)
    add_text(
        doc,
        "Station-level ensemble medians showed broadly weak positive PRCPTOT changes under both scenarios (Figure 3). Under SSP2-4.5, 12 of 13 stations were positive and one was negative; station medians ranged from −1.0% to +3.4%. Under SSP5-8.5, 10 stations were positive and three were negative, with a range from −1.7% to +4.8%. The IDW field under SSP2-4.5 was consequently positive across most of the province, with a localized negative pocket near station 351011. Under SSP5-8.5, weak positive values persisted across much of the network, while negative pockets occurred toward the southwestern and eastern station locations and near-neutral values occurred around station 351201.",
    )
    add_text(
        doc,
        "These fields communicate broad spatial contrast rather than resolved climate dynamics. Station density is lower toward the provincial margins, and the entire northern tip lies outside the gauge convex hull. The continuous colors in those areas are therefore generated by the IDW distance rule, not by independent observations or model-grid validation. Topography may contribute to real rainfall gradients, but the present analysis did not test an elevation–change relationship. The maps should be used to identify where station-referenced ensemble medians differ, not to assign precise changes at ungauged locations.",
    )
    add_figure(
        doc, FIGURES / "Figure3_prcptot_idw_fields.png", "Figure 3.",
        "Spatial distribution of projected changes in annual total wet-day precipitation (PRCPTOT) over Uttaradit Province for 2021–2050 relative to 1995–2014 under (a) SSP2-4.5 and (b) SSP5-8.5. IDW with power 2 was applied to station-level seven-model medians, and the same percentage scale was used in both panels. White areas outside the provincial boundary are masked. Areas within the polygon but outside the station convex hull are extrapolations rather than independently validated gridded projections.",
    )

    add_heading(doc, "3.4 Within-window trend diagnostics", 2)
    add_text(
        doc,
        "The unadjusted standard MK diagnostic identified significant trends in 79 of 1,001 SSP2-4.5 annual series: five were increasing and 74 decreasing. For SSP5-8.5, 71 of 1,001 series were significant, comprising 57 increases and 14 decreases. These counts concern the slope within 2021–2050 and do not contradict the baseline-to-future changes in Table 5. For example, an index may have a future-period mean above its 1995–2014 baseline while declining within the future window. Because no multiplicity correction was applied to these 2,002 diagnostic tests and each series spans only 30 years, their significance counts were not used to rank scenarios or claim ensemble robustness.",
    )

    add_heading(doc, "3.5 Implications and limitations", 2)
    add_text(
        doc,
        "The clearest projected signal is a possible redistribution toward heavier individual events without a comparably robust increase in annual wet-day precipitation. Under SSP2-4.5, unanimous positive model signs for Rx1day and R99p and six-of-seven agreement for Rx5day contrast with the split PRCPTOT response. Under SSP5-8.5, the larger Rx1day median suggests stronger short-duration intensity, but agreement is weaker for most other indices. Similar regional assessments have found that precipitation-total and extreme-event signals can differ and that their magnitudes depend on model and processing choices (Kuinkel et al., 2024; Nontikansak et al., 2022; Try and Qin, 2024; Zhang et al., 2024). For water management, this pattern means that flood and erosion hazard may intensify even where annual totals change little; realized impacts will also depend on event timing, antecedent moisture, catchment response, drainage capacity, and exposure.",
    )
    add_text(
        doc,
        "Four limitations constrain interpretation. First, the seven models are a selected, non-independent ensemble, so the median and sign counts do not define probabilities. Second, the supplied bias-corrected series could not be independently audited against complete correction metadata; extremes may be sensitive to choices in wet-day adjustment and temporal sequencing (Cannon et al., 2015; Maraun, 2016; Maraun et al., 2017). Third, relative changes in rare-event indices are unstable when baseline values are small or zero, as shown by the undefined R99p changes at one station and the broad R50mm/R99p ranges. Fourth, IDW represents geographic proximity only and cannot resolve an orographic gradient, convective organization, or model physics absent from the station values. These limitations favor adaptation strategies that are robust to a range of plausible changes rather than optimized to one ensemble median.",
    )


def add_back_matter(doc: Document) -> None:
    add_heading(doc, "4. Conclusions", 1)
    add_text(
        doc,
        "For 2021–2050 relative to 1995–2014, the seven-model ensemble projected modest and uncertain changes in PRCPTOT but larger median increases in several precipitation extremes. Under SSP2-4.5, median changes were +2.9% for PRCPTOT, +6.1% for Rx1day, +4.2% for Rx5day, and +26.7% for R99p; model sign agreement was strongest for Rx1day and R99p. Under SSP5-8.5, the corresponding medians were +0.9%, +12.3%, +2.9%, and +14.4%, with wider sign disagreement for most indices. The higher forcing pathway therefore intensified the ensemble-median Rx1day response but did not uniformly amplify all indices. PRCPTOT changes were weakly positive at most stations, with localized negative pockets; IDW mapped these station-referenced contrasts but did not provide a validated gridded forecast. Planning should consider the possibility of stronger short-duration and upper-tail rainfall together with substantial inter-model uncertainty and limited information at ungauged locations.",
    )

    add_heading(doc, "5. Acknowledgments", 1)
    add_text(
        doc,
        "The author acknowledges the Office of Water Management and Hydrology, Royal Irrigation Department, for institutional support.",
    )

    add_heading(doc, "6. Author Contributions", 1)
    add_text(
        doc,
        "S.P.: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Visualization, Writing—original draft, and Writing—review and editing.",
    )

    add_heading(doc, "7. Declaration of Competing Interest", 1)
    add_text(doc, "The authors declare no conflict of interest.")

    add_heading(doc, "8. References", 1)
    for reference in REFERENCES:
        p = add_text(
            doc, reference, size=12, line_spacing=1.5,
            first_line_cm=0, after=0,
        )
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)


def validate_inputs() -> None:
    required = [
        RESULTS / "station_coordinates_used.csv",
        RESULTS / "observed_trend_1981_2014.csv",
        RESULTS / "future_change_2021_2050.csv",
        ELEVATION_FILE,
        FIGURES / "Figure1_Uttaradit_study_area.png",
        FIGURES / "Figure2_future_model_changes.png",
        FIGURES / "Figure3_prcptot_idw_fields.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing manuscript input(s):\n" + "\n".join(missing))
    if len(ABSTRACT.split()) > 250:
        raise ValueError(f"Abstract exceeds 250 words: {len(ABSTRACT.split())}")
    if len(KEYWORDS) > 6:
        raise ValueError("More than six keywords.")
    too_long = [(item, len(item)) for item in HIGHLIGHTS if len(item) > 85]
    if too_long:
        raise ValueError(f"Highlights exceed 85 characters: {too_long}")


def main() -> None:
    validate_inputs()
    obs = observed_summary()
    fut = future_summary()

    # Numerical locks from the audited source output.
    locks = {
        ("SSP2-4.5", "PRCPTOT"): 2.91,
        ("SSP2-4.5", "Rx1day"): 6.07,
        ("SSP2-4.5", "R99p"): 26.75,
        ("SSP5-8.5", "PRCPTOT"): 0.88,
        ("SSP5-8.5", "Rx1day"): 12.32,
        ("SSP5-8.5", "R99p"): 14.39,
    }
    for key, expected in locks.items():
        actual = float(fut.loc[(fut["Scenario"] == key[0]) & (fut["Index"] == key[1]), "Median"].iloc[0])
        if not np.isclose(actual, expected, atol=0.01):
            raise AssertionError(f"Numerical lock failed for {key}: {actual} != {expected}")
    if int(obs[["MK_inc", "MK_dec"]].to_numpy().sum()) != 30:
        raise AssertionError("Observed standard-MK count is not 30.")
    if int(obs[["BH_inc", "BH_dec"]].to_numpy().sum()) != 14:
        raise AssertionError("Observed BH-FDR count is not 14.")

    doc = Document()
    configure_document(doc)
    add_front_matter(doc)
    add_introduction(doc)
    add_methodology(doc)
    add_results_discussion(doc, obs, fut)
    add_back_matter(doc)

    doc.core_properties.title = TITLE
    doc.core_properties.author = "Surasit Punyawansiri"
    doc.core_properties.subject = "CMIP6 precipitation-extreme projections for Uttaradit Province"
    doc.core_properties.keywords = "; ".join(KEYWORDS)
    doc.core_properties.comments = "Prepared to the written 2025 EnNRJ manuscript requirements."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    print(f"Abstract words: {len(ABSTRACT.split())}")
    print("Highlight characters: " + ", ".join(str(len(item)) for item in HIGHLIGHTS))
    print(f"References: {len(REFERENCES)}")


if __name__ == "__main__":
    main()
