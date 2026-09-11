"""Programmatic release checks for the authored EnNRJ manuscript."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
DOCX = HERE / "EnNRJ_Uttaradit_Precipitation_Extremes_Manuscript.docx"
QA_JSON = HERE / "qa" / "manuscript_quality_report.json"
FIGURES = HERE / "figures_submission"
EXPECTED_TITLE = "Projected Precipitation Extremes over Uttaradit Province under CMIP6 SSP Scenarios"


def words(text: str) -> int:
    return len(re.findall(r"\b[\w%+−–-]+\b", text, flags=re.UNICODE))


def extract_document_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def main() -> None:
    doc = Document(DOCX)
    text = extract_document_text(doc)
    paragraphs = [p.text.strip() for p in doc.paragraphs]

    # Geometry and running elements.
    section = doc.sections[0]
    cm = 360000
    geometry = {
        "page_width_cm": round(section.page_width / cm, 2),
        "page_height_cm": round(section.page_height / cm, 2),
        "left_margin_cm": round(section.left_margin / cm, 2),
        "right_margin_cm": round(section.right_margin / cm, 2),
        "top_margin_cm": round(section.top_margin / cm, 2),
        "bottom_margin_cm": round(section.bottom_margin / cm, 2),
    }
    xml = DOCX.read_bytes()
    with zipfile.ZipFile(DOCX) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )

    # Front matter limits.
    abstract_i = paragraphs.index("Abstract")
    keyword_i = next(i for i, value in enumerate(paragraphs) if value.startswith("Keywords:"))
    abstract_text = " ".join(paragraphs[abstract_i + 1:keyword_i])
    highlight_i = paragraphs.index("Highlights")
    intro_i = paragraphs.index("1. Introduction")
    highlights = [p.removeprefix("• ") for p in paragraphs[highlight_i + 1:intro_i] if p]

    refs_i = paragraphs.index("8. References")
    main_start = intro_i
    main_text = " ".join(paragraphs[main_start:refs_i])
    references = [p for p in paragraphs[refs_i + 1:] if p]

    table_shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
    expected_shapes = [(14, 4), (8, 5), (12, 4), (12, 8), (23, 6)]

    # Artwork quality.
    figure_dpi = {}
    for path in sorted(FIGURES.glob("Figure*.tif")):
        with Image.open(path) as image:
            dpi = image.info.get("dpi", (0, 0))
            figure_dpi[path.name] = {
                "pixels": list(image.size),
                "dpi": [round(float(dpi[0]), 1), round(float(dpi[1]), 1)],
            }

    forbidden = [
        "Phetchaburi", "primary-coordinate", "coordinate sensitivity",
        "Figure 4", "Figure 5", "maximum 3 consecutive",
    ]
    forbidden_hits = [term for term in forbidden if term.lower() in text.lower()]

    checks = {
        "file_exists": DOCX.exists(),
        "exact_title": paragraphs[0] == EXPECTED_TITLE,
        "a4_portrait": abs(geometry["page_width_cm"] - 21.0) <= 0.02 and abs(geometry["page_height_cm"] - 29.7) <= 0.02,
        "margins": geometry["left_margin_cm"] == 2.54 and geometry["right_margin_cm"] == 2.54 and geometry["top_margin_cm"] == 1.9 and geometry["bottom_margin_cm"] == 1.9,
        "line_numbering": "w:lnNumType" in document_xml,
        "page_number_field": "PAGE" in footer_xml,
        "abstract_within_250": words(abstract_text) <= 250,
        "keyword_count_six": paragraphs[keyword_i].count(";") + 1 == 6,
        "highlight_count_3_to_5": 3 <= len(highlights) <= 5,
        "highlights_within_85_characters": all(len(item) <= 85 for item in highlights),
        "main_text_within_4000_words": words(main_text) <= 4000,
        "five_tables_with_expected_shapes": table_shapes == expected_shapes,
        "three_embedded_figures": len(doc.inline_shapes) == 3,
        "ten_native_math_objects": document_xml.count("<m:oMath>") == 10,
        "references_at_least_15": len(references) >= 15,
        "forbidden_terms_absent": not forbidden_hits,
        "zero_baseline_disclosed": "zero baseline" in text.lower() and "station 351011" in text,
        "idw_limit_disclosed": "not dynamically downscaled or independently validated" in text,
        "all_separate_tiffs_at_least_600_dpi": len(figure_dpi) == 3 and all(v["dpi"][0] >= 599 and v["dpi"][1] >= 599 for v in figure_dpi.values()),
    }

    # Caption and cross-reference continuity.
    for number in range(1, 6):
        checks[f"table_{number}_caption_and_citation"] = text.count(f"Table {number}") >= 2
    for number in range(1, 4):
        checks[f"figure_{number}_caption_and_citation"] = text.count(f"Figure {number}") >= 2

    # Verify the author-date reference surnames have at least one textual citation.
    surnames = [
        "Alexander", "Benjamini", "Cannon", "Chaiwino", "de Oliveira-Júnior",
        "Eyring", "Hamed", "Humphries", "Intergovernmental Panel on Climate Change",
        "Knutti", "Kuinkel", "Madolli", "Mann", "Maraun", "Nontikansak",
        "O’Neill", "Riahi", "Sen", "Shepard", "Sillmann", "Tebaldi", "Try",
        "Waqas", "Yue", "Zhang",
    ]
    missing_citations = [name for name in surnames if name not in main_text]
    checks["all_reference_families_cited"] = not missing_citations

    report = {
        "document": str(DOCX),
        "geometry": geometry,
        "abstract_words": words(abstract_text),
        "keywords": paragraphs[keyword_i],
        "highlight_characters": [len(item) for item in highlights],
        "main_text_words_excluding_references": words(main_text),
        "reference_count": len(references),
        "table_shapes": table_shapes,
        "embedded_figure_count": len(doc.inline_shapes),
        "native_math_count": document_xml.count("<m:oMath>"),
        "figure_files": figure_dpi,
        "forbidden_hits": forbidden_hits,
        "missing_citation_families": missing_citations,
        "checks": checks,
        "all_checks_passed": all(checks.values()),
    }
    QA_JSON.parent.mkdir(parents=True, exist_ok=True)
    QA_JSON.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if not report["all_checks_passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
