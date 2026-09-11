from __future__ import annotations

import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "manuscript_source.docx"
WORKFLOW = BASE / "figure2_workflow_revised.png"
OUTPUT = BASE / "Manuscript_MBCn_Engineering_Access_Ready.docx"

TITLE = (
    "Multivariate bias correction improves inter-station dependence in multi-site daily "
    "precipitation without altering single-site performance: an evaluation of MBCn against "
    "quantile delta mapping over a coastal gauge network in Thailand"
)

ABSTRACT = (
    "Bias correction of climate-model precipitation is usually evaluated station by station, "
    "although distributed hydrological applications also require realistic dependence among "
    "gauges. We compared N-dimensional multivariate bias correction (MBCn), quantile delta "
    "mapping (QDM), and uncorrected output for 12 daily rain gauges on the western Gulf of "
    "Thailand coast and seven CMIP6 models. Methods were calibrated on 1981-2000 and evaluated "
    "independently on 2001-2014. Because the final MBCn step reinstates QDM values at each gauge, "
    "both methods produced identical results for six marginal diagnostics across all 84 "
    "gauge-model combinations. MBCn improved all six dependence diagnostics. Relative to QDM, "
    "it reduced energy distance from 0.852 to 0.571 and Spearman correlation-matrix RMSE from "
    "0.286 to 0.187; all seven models improved on every diagnostic (one-sided exact Wilcoxon "
    "p = 0.0078; rank-biserial effect size = 1.00). Most dependence gains occurred within 10-15 "
    "iterations, results were insensitive to random seed, and projected wet-day quantile changes "
    "remained comparable to QDM. QDM is therefore sufficient when only station-level distributions "
    "are required, whereas MBCn is better supported when spatial dependence matters."
)

REPLACEMENTS = {
    "Article type: Research article": "Article type: Research article",
    "Statistical bias correction is applied to climate-model rainfall at almost every impact study": ABSTRACT,
    "Global climate models remain too coarse to be used directly at the scale of a catchment": (
        "Global climate models remain too coarse for direct use at catchment and station scales, "
        "and their daily precipitation contains systematic errors in both occurrence and intensity "
        "[1]. Distribution-based bias correction is therefore standard in hydrological climate-impact "
        "studies [2], [3]. Quantile delta mapping (QDM) is widely used because it adjusts the full "
        "distribution at a site while preserving the model-projected relative change at each quantile "
        "[4]. This is usually sufficient when the analysis treats each gauge independently."
    ),
    "What QDM does not do, in common with almost every method in routine use": (
        "QDM does not directly control relationships among gauges because each site is corrected "
        "independently. This limitation matters when corrected fields drive distributed hydrological "
        "models, joint flood analyses, or assessments of spatially extensive droughts and storms. "
        "Accurate marginal distributions can coexist with an unrealistic joint structure, and such "
        "errors propagate into analyses of compound and spatially concurrent extremes [5]."
    ),
    "A separate line of work targets the joint distribution explicitly": (
        "Multivariate methods address the joint distribution explicitly. Examples include rank-based "
        "resampling [6], optimal transport [7], and N-dimensional multivariate bias correction (MBCn) "
        "[8]. MBCn repeatedly rotates the multistation field, quantile-maps each rotated component to "
        "the corresponding observations, and reverses the rotation; a final step reinstates the QDM "
        "marginal values at each gauge. Comparative studies show that multivariate adjustment can "
        "improve dependence while introducing method-specific trade-offs [9], and regional applications "
        "confirm the relevance of evaluating spatial as well as marginal behaviour [10]."
    ),
    "That closing step determines how any comparison between the two methods must be interpreted": (
        "The final rank-substitution step determines how QDM and MBCn should be compared. Because MBCn "
        "restores the QDM values at every gauge, the two methods must be identical for distribution-based "
        "single-site diagnostics. Their added value can therefore be judged only with diagnostics that "
        "measure dependence among gauges."
    ),
    "Rainfall extremes and their spatial organisation are among the more consequential": (
        "The spatial organisation of rainfall extremes is both consequential and uncertain under climate "
        "change [11]. CMIP6 studies have evaluated precipitation change across mainland Southeast Asia "
        "[12]-[14] and Thailand [15]-[18], while recent work has compared correction methods in multistation "
        "settings [19]. However, few studies combine an independent validation period with several complementary "
        "dependence diagnostics in one gauge-network experiment. We therefore ask: once single-site performance "
        "is fixed by construction, how much does MBCn improve the joint behaviour of a coastal gauge network, "
        "and does it compromise numerical stability or the projected quantile-change signal?"
    ),
    "Twelve daily rain gauges in Phetchaburi and Prachuap Khiri Khan provinces": (
        "Twelve daily rain gauges in Phetchaburi and Prachuap Khiri Khan provinces, on the western shore "
        "of the Gulf of Thailand, provide the observations (Table 1 and Figure 1). Each record is complete "
        "from 1 January 1981 to 31 December 2014. Mean annual precipitation ranges from approximately 944 "
        "to 1,376 mm, and wet-day frequency varies markedly across the network. This spatial contrast provides "
        "a demanding test of whether correction preserves local distributions while improving network dependence."
    ),
    "Coordinates and elevations come from the project metadata archive": (
        "Coordinates and elevations were obtained from the project metadata archive. Figure 1 locates the "
        "gauges within the provincial boundary. The background shading is an inverse-distance interpolation "
        "of mean annual precipitation for orientation only; it is not used in any analysis."
    ),
    "Seven CMIP6 models supply the simulated rainfall": (
        "Seven CMIP6 models supplied daily precipitation for the gauge locations (Table 2). Ensemble-member "
        "and grid labels were verified against the archive filenames. CESM2 used r11i1p1f1; the other six "
        "models used r1i1p1f1."
    ),
    "The record was split once and that split governed every method": (
        "A single split was used for every method: 1981-2000 for calibration and 2001-2014 for independent "
        "validation. Validation observations were not used during fitting. All comparisons were restricted "
        "to the same 12 gauges."
    ),
    "QDM operates on one gauge and one calendar month at a time": (
        "QDM was applied separately to each gauge and calendar month. Observed and simulated quantile functions "
        "were estimated in the calibration period, and relative model-projected changes were preserved at each "
        "quantile [4]. MBCn began from the QDM-corrected field and applied iterative multivariate adjustment [8]. "
        "Each of 25 passes used a seeded random orthogonal rotation across the 12 station dimensions, mapped each "
        "rotated component to the corresponding rotated observations, and reversed the rotation. Final rank "
        "substitution restored the QDM values at every station, retaining QDM marginal distributions while "
        "adjusting inter-station ranks (Figure 2). Related multistation comparisons also highlight the need to "
        "evaluate correction methods beyond marginal performance [19]."
    ),
    "Every estimated quantity, including the QDM quantile maps": (
        "All QDM mappings, rotations, and rotated-space mappings were fitted using 1981-2000 data only. The "
        "2001-2014 observations were reserved for evaluation. A wet day was defined as precipitation of at least "
        "1 mm. Thresholds for extreme-event diagnostics were estimated separately from each series' wet days so "
        "that these diagnostics primarily reflected dependence rather than residual intensity differences."
    ),
    "Free-running climate models are not synchronised with observed weather": (
        "Because free-running climate simulations are not synchronized with observed daily weather, single-site "
        "performance was assessed distributionally. Six diagnostics were computed for each of 84 gauge-model "
        "combinations: percentage bias in total precipitation, wet-day-frequency error, the ratio of simulated to "
        "observed wet-day standard deviation, Perkins skill score [20], Kolmogorov-Smirnov distance, and RMSE "
        "between matched wet-day quantiles. Day-matched criteria were not used because they would penalize timing "
        "that free-running simulations are not designed to reproduce."
    ),
    "Six diagnostics address the joint behaviour of the network": (
        "Six diagnostics evaluated joint behaviour. Energy distance compared the 12-dimensional daily precipitation "
        "vectors with observations [21] using a fixed-seed subsample of at most 1,500 validation days. Three metrics "
        "measured RMSE across the 66 off-diagonal station pairs for Spearman correlation, Pearson correlation, and "
        "binary wet-day correlation. Two metrics targeted concurrent extremes: RMSE of pairwise joint exceedance "
        "above each series' 95th wet-day percentile and the Wasserstein-1 distance between distributions of daily "
        "network wet extent. Each diagnostic was computed separately for each climate model."
    ),
    "The central comparison rests on the seven paired model-level differences": (
        "Seven paired model-level differences between MBCn and QDM were summarized with a one-sided exact Wilcoxon "
        "signed-rank test, the matched-pairs rank-biserial effect size, and the number of models favouring MBCn. "
        "The test describes consistency within this ensemble rather than inference to a population of climate models. "
        "Holm-Bonferroni adjustment was applied to the prespecified family of Spearman, Pearson, and wet-day-correlation "
        "endpoints; the other diagnostics were treated as exploratory."
    ),
    "Two properties were verified before any conclusion was drawn": (
        "Robustness was evaluated by varying the number of MBCn passes from 0 (QDM) to 100 and repeating the default "
        "25-pass correction with 20 random seeds. Fidelity to the model-projected signal was evaluated by comparing "
        "raw and corrected percentage changes in five wet-day quantiles for every gauge, model, pathway, and future "
        "period. This check addresses quantile changes only and does not validate the projections themselves."
    ),
    "Both corrections repair the familiar deficiencies of the raw output": (
        "The raw output produced rain too frequently and underestimated wet-day variability (Table 3 and Figure 3). "
        "After correction, Kolmogorov-Smirnov distance decreased from 0.221 to 0.101 and Perkins skill score increased "
        "from 0.82 to 0.91. Both corrections retained an approximately 17% dry bias in total precipitation and a "
        "wet-day-frequency bias of about -6 percentage points; this residual bias should be considered in water-balance applications."
    ),
    "QDM and MBCn returned the same value to three decimal places": (
        "QDM and MBCn were identical to three decimal places for every marginal score across all 84 gauge-model "
        "combinations. This confirms the expected algorithmic behaviour: final rank substitution returns the same "
        "set of daily QDM values at each gauge. The multivariate stage can therefore affect only how values are "
        "arranged jointly across the network, not their station-level distributions."
    ),
    "Once attention moves to the joint structure the two methods separate": (
        "MBCn outperformed QDM for every dependence diagnostic (Table 4). Energy distance decreased from 0.852 for "
        "QDM to 0.571 for MBCn. Spearman correlation-matrix RMSE decreased from 0.484 in raw output to 0.286 with "
        "QDM and to 0.187 with MBCn. Pearson correlation, wet-day correlation, joint 95th-percentile exceedance, and "
        "wet-extent distribution showed the same ordering."
    ),
    "The weight of this result comes from its consistency": (
        "Every model improved under MBCn for all six diagnostics (Table 5). With seven concordant non-zero paired "
        "differences, the one-sided exact Wilcoxon p-value was 0.0078, the smallest attainable value at n = 7, and "
        "the matched-pairs rank-biserial effect size was 1.00. Holm-Bonferroni-adjusted p-values for the three "
        "prespecified correlation endpoints were 0.0234. These values summarize a unanimous within-ensemble pattern; "
        "the models are not a random sample, and the six diagnostics are related views of the same dependence improvement."
    ),
    "Nothing in the algorithm guarantees this outcome": (
        "The consistency across models is empirical rather than guaranteed. Figure 4 shows the Spearman matrices for "
        "MIROC6, selected because its MBCn error was closest to the ensemble median. Raw output imposed uniformly high "
        "correlation; QDM reduced it only partly; and MBCn recovered more of the observed spatial contrast. Figure 5 "
        "shows all seven paired model results, each of which favours MBCn."
    ),
    "Figure 6 avoids the shelter of an average": (
        "Figure 6 relates pairwise Spearman correlation to gauge separation. Observed correlation declines modestly "
        "with distance, whereas raw output remains near 0.96 across most pairs. QDM reduces this over-coherence, and "
        "MBCn most closely follows the observed distance relationship. Coarse-grid extraction contributes to the raw "
        "over-coherence when gauges share or occupy adjacent grid cells. The binned relationship and bootstrap bands "
        "are descriptive because the 66 gauge pairs are not independent."
    ),
    "Correlation matrices summarise dependence but reveal little": (
        "Two event-oriented diagnostics complement the correlation analysis. Figure 7 compares observed and simulated "
        "probabilities that both gauges in a pair exceed their own 95th wet-day percentiles on the same date. Raw output "
        "is widely dispersed, QDM retains a positive bias, and MBCn lies closest to the 1:1 line without a systematic tilt."
    ),
    "Figure 8 restates the question in a form a hydrologist can act on": (
        "Figure 8 shows the pooled distribution of the number of gauges wet on each day. MBCn is closest to the observed "
        "distribution, consistent with the per-model wet-extent scores in Tables 4 and 5. The inset values in Figure 8 "
        "are calculated from the pooled distributions, whereas Table 4 reports the mean of seven separately calculated "
        "model-level Wasserstein-1 distances; the two summaries therefore need not be numerically identical."
    ),
    "Two checks establish that the multivariate stage is well behaved": (
        "Figure 9 shows the convergence analysis for three representative climate models. Most improvement occurred "
        "within 10-15 passes; diagnostics then oscillated within a narrow range, placing the 25-pass default on the "
        "plateau. Across 20 random seeds, correlation-based scores changed little (coefficient of variation approximately "
        "4%). The larger spread in energy distance reflects its fixed-size random subsampling and does not appear in "
        "the full-sample correlation diagnostics."
    ),
    "Figure 10 addresses the objection most likely to be raised": (
        "Projected wet-day quantile changes were similarly preserved by both methods (Figure 10). Across all gauges, "
        "models, pathways, and future periods, mean absolute deviation from the raw change signal was 9.4 percentage "
        "points for QDM and 9.2 for MBCn; the corresponding R-squared values were 0.485 and 0.465. Thus, MBCn achieved "
        "its dependence gains without a material additional loss of the QDM change-preservation property [4]. This "
        "check does not cover sequence-dependent statistics, such as wet- and dry-spell lengths, which MBCn may alter."
    ),
    "Taken together the results support a decision rule rather than a ranking": (
        "The results support a context-specific choice rather than a universal ranking. Recent comparisons likewise "
        "show that no bias-correction method dominates every temporal and dependence property [22]; method selection "
        "should follow the structure required by the downstream application. QDM is sufficient and less computationally "
        "demanding when only single-site distributions are required, including many intensity-duration-frequency "
        "applications in Thailand [23]. MBCn is better supported when gauges are used jointly in distributed hydrological "
        "modelling, spatial hazard mapping, or multivariate risk analysis."
    ),
    "A broader methodological point follows": (
        "More broadly, station-wise evaluation cannot distinguish methods that share marginal distributions but differ "
        "in joint behaviour. At least one dependence diagnostic should accompany conventional marginal scores whenever "
        "corrected variables or locations will be analysed together."
    ),
    "Four limitations bound these conclusions": (
        "Four limitations bound the conclusions. One realization per model does not represent internal variability. "
        "The single coastal network may not represent inland or more strongly monsoonal settings. The stationary "
        "calibration assumes that historical dependence remains an appropriate target under future climate. Finally, "
        "the residual dry bias in both corrections limits water-balance applications. These limitations constrain "
        "generalization but do not affect the paired comparison, for which both methods used identical inputs and periods."
    ),
    "Compared side by side on twelve complete coastal gauge records": (
        "Across 12 coastal gauges and seven CMIP6 models, QDM and MBCn were identical for all six station-level "
        "diagnostics because MBCn restored the QDM marginal values. In contrast, MBCn improved all six dependence "
        "diagnostics for every model, reducing energy distance and correlation-structure errors while converging rapidly, "
        "remaining insensitive to random seed, and preserving projected wet-day quantile changes comparably to QDM. "
        "QDM therefore remains appropriate when only marginal fidelity is needed; MBCn is the better-supported choice "
        "when inter-station dependence is part of the scientific or hydrological question. Dependence diagnostics should "
        "become routine whenever corrected locations are used jointly. Evaluation across larger, topographically diverse "
        "networks and multi-realization ensembles is the next step for testing the generality of these findings."
    ),
    "The authors thank the Royal Irrigation Department": (
        "The author thanks the Royal Irrigation Department for access to the daily rain-gauge records and the World "
        "Climate Research Programme and participating modelling groups for producing and distributing CMIP6 output."
    ),
    "Conceptualization: X.X.": (
        "Surasit Punyawansiri: conceptualization, methodology, software, validation, formal analysis, investigation, "
        "resources, data curation, writing - original draft, writing - review and editing, visualization, supervision, "
        "and project administration. The author has read and approved the submitted manuscript."
    ),
    "The authors declare no conflict of interest": "The author declares no conflict of interest.",
    "The bias-correction, evaluation, and figure code is deterministic": (
        "The daily rain-gauge records are available from the Royal Irrigation Department subject to its data-sharing "
        "conditions. CMIP6 output is publicly available through the Earth System Grid Federation. Analysis code and "
        "underlying numerical data for the figures are available from the corresponding author on reasonable request."
    ),
}

CAPTION_REPLACEMENTS = {
    "Figure 2.": (
        "Figure 2. MBCn workflow. QDM [4] first establishes each gauge's marginal distribution; the iterative "
        "rotation, mapping, and inverse-rotation stage [8] adjusts inter-station dependence; final rank substitution "
        "restores the QDM values. All transformations are fitted using calibration data only."
    ),
    "Figure 6.": (
        "Figure 6. Inter-station Spearman correlation versus gauge separation. MBCn follows the observed distance "
        "relationship more closely than QDM, whereas raw output remains uniformly high. Bootstrap bands are exploratory "
        "because gauge pairs are not independent."
    ),
    "Figure 8.": (
        "Figure 8. Pooled distribution of the daily number of simultaneously wet gauges for observations and each method. "
        "Inset Wasserstein-1 values are computed from the pooled distributions; Tables 4 and 5 use separately computed "
        "per-model scores."
    ),
    "Figure 9.": (
        "Figure 9. Dependence diagnostics versus MBCn iteration count for three representative climate models; iteration "
        "0 is QDM. Most improvement occurs within approximately 10-15 passes, placing the default of 25 on the plateau."
    ),
    "Figure 10.": (
        "Figure 10. Raw versus corrected projected change in wet-day quantiles for every gauge, model, pathway, and future "
        "period. QDM and MBCn preserve the model-projected quantile-change signal to a similar degree."
    ),
}

REFERENCES = [
    '[1] D. Maraun, "Bias correcting climate change simulations - a critical review," Curr. Clim. Change Rep., vol. 2, no. 4, pp. 211-220, 2016, doi: 10.1007/s40641-016-0050-x.',
    '[2] C. Teutschbein and J. Seibert, "Bias correction of regional climate model simulations for hydrological climate-change impact studies: Review and evaluation of different methods," J. Hydrol., vols. 456-457, pp. 12-29, 2012, doi: 10.1016/j.jhydrol.2012.05.052.',
    '[3] L. Gudmundsson, J. B. Bremnes, J. E. Haugen, and T. Engen-Skaugen, "Technical Note: Downscaling RCM precipitation to the station scale using statistical transformations - a comparison of methods," Hydrol. Earth Syst. Sci., vol. 16, no. 9, pp. 3383-3390, 2012, doi: 10.5194/hess-16-3383-2012.',
    '[4] A. J. Cannon, S. R. Sobie, and T. Q. Murdock, "Bias correction of GCM precipitation by quantile mapping: How well do methods preserve changes in quantiles and extremes?" J. Clim., vol. 28, no. 17, pp. 6938-6959, 2015, doi: 10.1175/JCLI-D-14-00754.1.',
    '[5] E. Bevacqua et al., "Advancing research on compound weather and climate events via large ensemble model simulations," Nat. Commun., vol. 14, Art. no. 2145, 2023, doi: 10.1038/s41467-023-37847-5.',
    '[6] M. Vrac, "Multivariate bias adjustment of high-dimensional climate simulations: The Rank Resampling for Distributions and Dependences (R2D2) bias correction," Hydrol. Earth Syst. Sci., vol. 22, no. 6, pp. 3175-3196, 2018, doi: 10.5194/hess-22-3175-2018.',
    '[7] Y. Robin, M. Vrac, P. Naveau, and P. Yiou, "Multivariate stochastic bias corrections with optimal transport," Hydrol. Earth Syst. Sci., vol. 23, no. 2, pp. 773-786, 2019, doi: 10.5194/hess-23-773-2019.',
    '[8] A. J. Cannon, "Multivariate quantile mapping bias correction: An N-dimensional probability density function transform for climate model simulations of multiple variables," Clim. Dyn., vol. 50, nos. 1-2, pp. 31-49, 2018, doi: 10.1007/s00382-017-3580-6.',
    '[9] B. François, M. Vrac, A. J. Cannon, Y. Robin, and D. Allard, "Multivariate bias corrections of climate simulations: Which benefits for which losses?" Earth Syst. Dyn., vol. 11, no. 2, pp. 537-562, 2020, doi: 10.5194/esd-11-537-2020.',
    '[10] D. Dieng et al., "Multivariate bias-correction of high-resolution regional climate change simulations for West Africa: Performance and climate change implications," J. Geophys. Res. Atmos., vol. 127, no. 8, Art. no. e2021JD034836, 2022, doi: 10.1029/2021JD034836.',
    '[11] S. I. Seneviratne et al., "Weather and climate extreme events in a changing climate," in Climate Change 2021: The Physical Science Basis, Cambridge, U.K.: Cambridge Univ. Press, 2021, pp. 1513-1766, doi: 10.1017/9781009157896.013.',
    '[12] S. Supharatid, J. Nafung, and T. Aribarg, "Projected changes in temperature and precipitation over mainland Southeast Asia by CMIP6 models," J. Water Clim. Change, vol. 13, no. 1, pp. 337-356, 2022, doi: 10.2166/wcc.2021.015.',
    '[13] X. Qin and C. Dai, "Comparison of different quantile delta mapping schemes in frequency analysis of precipitation extremes over mainland Southeast Asia under climate change," J. Hydrol., vol. 606, Art. no. 127421, 2022, doi: 10.1016/j.jhydrol.2021.127421.',
    '[14] S. Try and X. Qin, "Evaluation of future changes in climate extremes over Southeast Asia using downscaled CMIP6 GCM projections," Water, vol. 16, no. 15, Art. no. 2207, 2024, doi: 10.3390/w16152207.',
    '[15] U. W. Humphries, M. Waqas, P. T. Hlaing, P. Dechpichai, and A. Wangwongchai, "Assessment of CMIP6 GCMs for selecting a suitable climate model for precipitation projections in Southern Thailand," Results Eng., vol. 23, Art. no. 102417, 2024, doi: 10.1016/j.rineng.2024.102417.',
    '[16] S. Rojpratak and S. Supharatid, "Regional extreme precipitation index: Evaluations and projections from the multi-model ensemble CMIP5 over Thailand," Weather Clim. Extremes, vol. 37, Art. no. 100475, 2022, doi: 10.1016/j.wace.2022.100475.',
    '[17] D. Khadka, M. S. Babel, M. Collins, S. Shrestha, S. G. P. Virdis, and A. S. Chen, "Projected changes in the near-future mean climate and extreme climate events in Northeast Thailand," Int. J. Climatol., vol. 42, no. 4, pp. 2470-2492, 2022, doi: 10.1002/joc.7377.',
    '[18] D. Kuinkel, P. Promchote, K. R. Upreti, S.-Y. Wang, N. Dahal, and B. Pokharel, "Projected changes in precipitation extremes in Southern Thailand using CMIP6 models," Theor. Appl. Climatol., vol. 155, no. 9, pp. 8703-8716, 2024, doi: 10.1007/s00704-024-05150-y.',
    '[19] Y. H. Song and E.-S. Chung, "Comparison of quantile mapping methods for daily precipitation correction and future projections under the SSP5-8.5," J. Korea Water Resour. Assoc., vol. 57, no. 12, pp. 1069-1083, 2024, doi: 10.3741/JKWRA.2024.57.12.1069.',
    "[20] S. E. Perkins, A. J. Pitman, N. J. Holbrook, and J. McAneney, \"Evaluation of the AR4 climate models' simulated daily maximum temperature, minimum temperature, and precipitation over Australia using probability density functions,\" J. Clim., vol. 20, no. 17, pp. 4356-4376, 2007, doi: 10.1175/JCLI4253.1.",
    '[21] G. J. Székely and M. L. Rizzo, "Energy statistics: A class of statistics based on distances," J. Stat. Plan. Inference, vol. 143, no. 8, pp. 1249-1272, 2013, doi: 10.1016/j.jspi.2013.03.018.',
    '[22] S. Sharma, A. S. Raghuvanshi, and A. Agarwal, "Evaluating the performance of uni- and multivariate bias correction techniques: Challenges in preserving temporal and dependence structures," Water Resour. Res., vol. 62, no. 2, Art. no. e2025WR041526, 2026, doi: 10.1029/2025WR041526.',
    '[23] K. Tedprasith and W. Lohpaisankrit, "Development of intensity-duration-frequency relationships in Khon Kaen City, Thailand under changing climate using GCMs and a simple scaling method," J. Water Clim. Change, vol. 15, no. 3, pp. 1204-1217, 2024, doi: 10.2166/wcc.2024.533.',
]

FIGURE_ALTS = [
    "Map of the 12 rain gauges in Phetchaburi and Prachuap Khiri Khan with rainfall shading, elevation-scaled markers, and wet-day-frequency colour.",
    "Workflow for MBCn showing initial QDM, iterative rotation-based dependence adjustment, and final rank substitution.",
    "Boxplots of six single-site validation diagnostics for raw, QDM, and MBCn output.",
    "Observed, raw, QDM, and MBCn inter-station Spearman correlation matrices for MIROC6, with sorted pairwise errors.",
    "Paired QDM-to-MBCn changes in three dependence-error diagnostics for seven climate models.",
    "Binned inter-station Spearman correlation versus station separation for observations, raw output, QDM, and MBCn.",
    "Observed versus simulated pairwise probabilities of concurrent 95th-percentile wet-day exceedance.",
    "Pooled distribution of the daily number of wet gauges for observations, raw output, QDM, and MBCn.",
    "Three dependence scores versus MBCn iteration count for three representative climate models.",
    "Raw versus corrected projected changes in five wet-day quantiles for QDM and MBCn.",
]


def set_cell_margins(cell, top=40, start=55, bottom=40, end=55):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, size="8", color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Times New Roman", size=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def replace_text(paragraph, new_text):
    paragraph.clear()
    run = paragraph.add_run(new_text)
    set_run_font(run)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        r = p.add_run(text)
        set_run_font(r)
    return p


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def find_paragraph(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_document(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.0
    pf.space_after = Pt(3)
    pf.widow_control = True

    for name, size, level in (("Heading 1", 12, 1), ("Heading 2", 12, 2)):
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(8 if level == 1 else 6)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.keep_with_next = True

    if "Article Title" not in styles:
        st = styles.add_style("Article Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = styles["Article Title"]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(14)
    st.font.bold = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(8)

    for name, size, italic in (("Author Block", 12, False), ("Affiliation", 11, True), ("Caption EA", 10, False), ("Table Note", 9, False), ("Reference EA", 10, False), ("Gap", 4, False)):
        if name not in styles:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.italic = italic
        if name in ("Author Block", "Affiliation", "Caption EA"):
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "Reference EA":
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.left_indent = Inches(0.27)
            st.paragraph_format.first_line_indent = Inches(-0.27)
            st.paragraph_format.space_after = Pt(3)
            st.paragraph_format.line_spacing = 1.0
        elif name == "Table Note":
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.space_after = Pt(6)
        elif name == "Gap":
            st.paragraph_format.space_before = Pt(2)
            st.paragraph_format.space_after = Pt(2)
        else:
            st.paragraph_format.space_after = Pt(3)


def add_author_block(doc):
    title = doc.paragraphs[0]
    replace_text(title, TITLE)
    title.style = "Article Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    article_type = find_paragraph(doc, "Article type:")
    author = insert_paragraph_after(title, style="Author Block")
    name = author.add_run("Surasit Punyawansiri")
    set_run_font(name, size=12)
    name.underline = True
    marker = author.add_run("1,*")
    set_run_font(marker, size=9)
    marker.font.superscript = True
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER

    affiliation = insert_paragraph_after(author, style="Affiliation")
    m = affiliation.add_run("1")
    set_run_font(m, size=8, italic=True)
    m.font.superscript = True
    r = affiliation.add_run("Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok, Thailand")
    set_run_font(r, size=11, italic=True)
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER

    corresponding = insert_paragraph_after(affiliation, style="Affiliation")
    r = corresponding.add_run("*Corresponding author: Surasit.pum@thaimooc.ac.th")
    set_run_font(r, size=10, italic=True)
    corresponding.alignment = WD_ALIGN_PARAGRAPH.CENTER

    article_type._p.getparent().remove(article_type._p)
    corresponding._p.addnext(article_type._p)
    article_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    article_type.paragraph_format.space_after = Pt(8)
    for run in article_type.runs:
        set_run_font(run, size=10, italic=True)

    anon = find_paragraph(doc, "Author names and affiliations")
    delete_paragraph(anon)


def update_content(doc):
    for prefix, new_text in REPLACEMENTS.items():
        p = find_paragraph(doc, prefix)
        replace_text(p, new_text)

    for p in doc.paragraphs:
        s = p.text.strip()
        for prefix, new_text in CAPTION_REPLACEMENTS.items():
            if s.startswith(prefix):
                replace_text(p, new_text)
                break

    abstract_heading = find_paragraph(doc, "Abstract")
    abstract_heading.style = "Heading 1"
    abstract_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    keywords = find_paragraph(doc, "Keywords:")
    replace_text(keywords, "Keywords: multivariate bias correction; MBCn; inter-station dependence; multi-site precipitation; CMIP6")
    keywords.paragraph_format.space_after = Pt(8)

    for p in doc.paragraphs:
        text = p.text.strip()
        if re.match(r"^[1-4]\.\s", text):
            p.style = "Heading 1"
        elif re.match(r"^[1-4]\.\d+\s", text):
            p.style = "Heading 2"
        elif text in {"Acknowledgment", "Author Contributions", "Funding", "Conflicts of Interest", "Data and Code Availability", "References"}:
            p.style = "Heading 1"

    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith("Figure ") and re.match(r"^Figure \d+\.", text):
            p.style = "Caption EA"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.keep_together = True


def update_references(doc):
    first = find_paragraph(doc, "[1]")
    old_refs = [p for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
    anchor = first
    for i, ref in enumerate(REFERENCES):
        if i < len(old_refs):
            p = old_refs[i]
            replace_text(p, ref)
        else:
            p = insert_paragraph_after(anchor, ref, "Reference EA")
        p.style = "Reference EA"
        anchor = p
    for p in old_refs[len(REFERENCES):]:
        delete_paragraph(p)


def move_and_style_tables(doc):
    caption_starts = ["Table 1.", "Table 2.", "Table 3.", "Table 4.", "Table 5."]
    note_prefixes = ["Coordinates (WGS84)", "All seven models provide", "pp = percentage points", "Values are means of per-model scores", "Every model improves under MBCn"]
    widths = [
        [0.65, 0.85, 0.90, 0.75, 1.65, 1.70],
        [1.25, 2.65, 1.45, 1.15],
        [0.70, 0.85, 1.05, 0.75, 0.65, 0.85, 1.65],
        [0.80, 0.85, 0.80, 0.80, 1.15, 1.05, 0.90],
        [1.85, 1.10, 0.65, 0.65, 1.20, 1.05],
    ]

    for idx, table in enumerate(doc.tables):
        cap = find_paragraph(doc, caption_starts[idx])
        note = find_paragraph(doc, note_prefixes[idx])

        gap_xml = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), "Gap")
        ppr.append(pstyle)
        gap_xml.append(ppr)
        if idx == 3:
            run = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            run.append(br)
            gap_xml.append(run)
        table._tbl.addprevious(gap_xml)

        table._tbl.addnext(cap._p)
        cap._p.addnext(note._p)
        cap.style = "Caption EA"
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = True
        note.style = "Table Note"
        note.paragraph_format.keep_with_next = False

        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_repeat_table_header(table.rows[0])
        for r_idx, row in enumerate(table.rows):
            row.height = None
            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(widths[idx]):
                    cell.width = Inches(widths[idx][c_idx])
                    cell._tc.tcPr.tcW.set(qn("w:w"), str(int(widths[idx][c_idx] * 1440)))
                    cell._tc.tcPr.tcW.set(qn("w:type"), "dxa")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                set_cell_border(cell)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        set_run_font(run, size=8, bold=(True if r_idx == 0 else run.bold))


def apply_global_formatting(doc):
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            continue
        if p.style.name not in {"Article Title", "Author Block", "Affiliation", "Caption EA", "Table Note", "Reference EA", "Heading 1", "Heading 2", "Gap"}:
            p.style = "Normal"
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.widow_control = True
        for run in p.runs:
            size = None
            if p.style.name == "Normal":
                size = 12
            elif p.style.name == "Caption EA":
                size = 10
            elif p.style.name == "Table Note":
                size = 9
            elif p.style.name == "Reference EA":
                size = 10
            set_run_font(run, size=size)

    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith("Keywords:") and p.runs:
            p.runs[0].bold = False


def add_alt_text(doc):
    doc_prs = doc._element.xpath("//wp:docPr")
    for i, doc_pr in enumerate(doc_prs[: len(FIGURE_ALTS)]):
        doc_pr.set("title", f"Figure {i + 1}")
        doc_pr.set("descr", FIGURE_ALTS[i])


def resize_images(doc):
    # Preserve legibility while recovering enough vertical space to keep Table 2
    # with its caption and note on the same page.
    if doc.inline_shapes:
        doc.inline_shapes[0].width = Inches(3.70)
        doc.inline_shapes[0].height = Inches(4.87)


def replace_workflow_media(docx_path):
    temp_fd, temp_name = tempfile.mkstemp(suffix=".docx", dir=str(BASE))
    os.close(temp_fd)
    temp_path = Path(temp_name)
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/media/image2.png":
                data = WORKFLOW.read_bytes()
            zout.writestr(item, data)
    os.replace(temp_path, docx_path)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not WORKFLOW.exists():
        raise FileNotFoundError(WORKFLOW)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    style_document(doc)
    add_author_block(doc)
    update_content(doc)
    update_references(doc)
    move_and_style_tables(doc)
    apply_global_formatting(doc)
    resize_images(doc)
    add_alt_text(doc)

    props = doc.core_properties
    props.title = TITLE
    props.subject = "Evaluation of multivariate precipitation bias correction over a coastal gauge network in Thailand"
    props.author = "Surasit Punyawansiri"
    props.last_modified_by = "Surasit Punyawansiri"
    props.comments = ""
    props.keywords = "MBCn, quantile delta mapping, multi-site precipitation, CMIP6, Thailand"

    doc.save(OUTPUT)
    replace_workflow_media(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
