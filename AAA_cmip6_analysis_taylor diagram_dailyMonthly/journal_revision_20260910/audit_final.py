import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


BASE = Path(__file__).resolve().parent
DOCX = BASE / "Manuscript_MBCn_Engineering_Access_Ready.docx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def words(text):
    return re.findall(r"\b[\w-]+\b", text, flags=re.UNICODE)


doc = Document(DOCX)
paras = doc.paragraphs
texts = [p.text.strip() for p in paras]

abstract_index = texts.index("Abstract")
keywords_index = next(i for i, t in enumerate(texts) if t.startswith("Keywords:"))
abstract_text = " ".join(texts[abstract_index + 1 : keywords_index])
keywords = [x.strip() for x in texts[keywords_index].split(":", 1)[1].split(";") if x.strip()]

references_index = texts.index("References")
body_text = "\n".join(texts[:references_index])
reference_texts = [t for t in texts[references_index + 1 :] if re.match(r"^\[\d+\]", t)]

citations_in_sequence = []
for match in re.finditer(r"\[(\d+)\](?:\s*-\s*\[(\d+)\])?|\[(\d+)\s*-\s*(\d+)\]", body_text):
    start = int(match.group(1) or match.group(3))
    end = int(match.group(2) or match.group(4) or start)
    citations_in_sequence.extend(range(start, end + 1))
first_appearance = []
for n in citations_in_sequence:
    if n not in first_appearance:
        first_appearance.append(n)

figure_captions = [t for t in texts if re.match(r"^Figure \d+\.", t)]
table_captions = [t for t in texts if re.match(r"^Table \d+\.", t)]

with zipfile.ZipFile(DOCX) as zf:
    bad_crc = zf.testzip()
    document_xml = zf.read("word/document.xml")
    root = etree.fromstring(document_xml)
    ns = {"w": W, "wp": WP}
    insertions = len(root.xpath("//w:ins", namespaces=ns))
    deletions = len(root.xpath("//w:del", namespaces=ns))
    comments = [name for name in zf.namelist() if name.startswith("word/comments")]
    doc_prs = root.xpath("//wp:docPr", namespaces=ns)
    image_alt = [
        {"title": node.get("title"), "description": node.get("descr")}
        for node in doc_prs
    ]
    body = root.find(f"{{{W}}}body")
    children = list(body)
    table_positions = [i for i, node in enumerate(children) if node.tag == f"{{{W}}}tbl"]
    caption_positions = []
    for i, node in enumerate(children):
        if node.tag != f"{{{W}}}p":
            continue
        text = "".join(node.xpath(".//w:t/text()", namespaces=ns)).strip()
        if re.match(r"^Table \d+\.", text):
            caption_positions.append(i)

author_ok = (
    "Surasit Punyawansiri" in body_text
    and "Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok, Thailand" in body_text
    and "Surasit.pum@thaimooc.ac.th" in body_text
)

report = {
    "file": str(DOCX),
    "zip_crc_error": bad_crc,
    "page_layout": {
        "sections": len(doc.sections),
        "page_width_in": round(doc.sections[0].page_width / 914400, 2),
        "page_height_in": round(doc.sections[0].page_height / 914400, 2),
        "margins_in": {
            "top": round(doc.sections[0].top_margin / 914400, 2),
            "bottom": round(doc.sections[0].bottom_margin / 914400, 2),
            "left": round(doc.sections[0].left_margin / 914400, 2),
            "right": round(doc.sections[0].right_margin / 914400, 2),
        },
    },
    "abstract_word_count": len(words(abstract_text)),
    "keyword_count": len(keywords),
    "keywords": keywords,
    "author_block_present": author_ok,
    "table_count": len(doc.tables),
    "table_caption_count": len(table_captions),
    "table_captions_after_tables": len(table_positions) == len(caption_positions) == 5
    and all(cap > table for table, cap in zip(table_positions, caption_positions)),
    "figure_count": len(doc.inline_shapes),
    "figure_caption_count": len(figure_captions),
    "figure_caption_numbers": [int(re.match(r"^Figure (\d+)\.", t).group(1)) for t in figure_captions],
    "reference_count": len(reference_texts),
    "reference_numbers": [int(re.match(r"^\[(\d+)\]", t).group(1)) for t in reference_texts],
    "citation_first_appearance": first_appearance,
    "citation_order_complete": first_appearance == list(range(1, 24)),
    "tracked_insertions": insertions,
    "tracked_deletions": deletions,
    "comment_parts": comments,
    "images_with_alt_text": sum(bool(x["description"]) for x in image_alt),
    "image_alt": image_alt,
}

print(json.dumps(report, ensure_ascii=False, indent=2))
