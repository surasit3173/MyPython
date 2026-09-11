from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


def emu_to_inches(value: int | None) -> float | None:
    return None if value is None else round(value / 914400, 4)


def color_value(run) -> str | None:
    try:
        color = run.font.color.rgb
        return None if color is None else str(color)
    except Exception:
        return None


def run_record(run) -> dict[str, Any]:
    fonts = run._element.get_or_add_rPr().rFonts
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": None if run.underline is None else str(run.underline),
        "font_name": run.font.name,
        "font_ascii": fonts.get(qn("w:ascii")) if fonts is not None else None,
        "font_hansi": fonts.get(qn("w:hAnsi")) if fonts is not None else None,
        "font_eastasia": fonts.get(qn("w:eastAsia")) if fonts is not None else None,
        "font_size_pt": None if run.font.size is None else round(run.font.size.pt, 2),
        "color": color_value(run),
        "superscript": run.font.superscript,
        "subscript": run.font.subscript,
        "highlight": None if run.font.highlight_color is None else str(run.font.highlight_color),
    }


def paragraph_record(paragraph, index: int | None = None) -> dict[str, Any]:
    fmt = paragraph.paragraph_format
    return {
        "index": index,
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style is not None else None,
        "alignment": None if paragraph.alignment is None else str(paragraph.alignment),
        "left_indent_in": emu_to_inches(fmt.left_indent),
        "right_indent_in": emu_to_inches(fmt.right_indent),
        "first_line_indent_in": emu_to_inches(fmt.first_line_indent),
        "space_before_pt": None if fmt.space_before is None else round(fmt.space_before.pt, 2),
        "space_after_pt": None if fmt.space_after is None else round(fmt.space_after.pt, 2),
        "line_spacing": None if fmt.line_spacing is None else str(fmt.line_spacing),
        "keep_with_next": fmt.keep_with_next,
        "keep_together": fmt.keep_together,
        "page_break_before": fmt.page_break_before,
        "runs": [run_record(run) for run in paragraph.runs],
    }


def document_record(path: Path) -> dict[str, Any]:
    doc = Document(path)
    result: dict[str, Any] = {
        "path": str(path.resolve()),
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "keywords": doc.core_properties.keywords,
            "comments": doc.core_properties.comments,
        },
        "sections": [],
        "styles": [],
        "paragraphs": [],
        "tables": [],
        "headers": [],
        "footers": [],
        "package_parts": [],
        "body_blocks": [],
    }
    for section in doc.sections:
        result["sections"].append(
            {
                "start_type": str(section.start_type),
                "orientation": str(section.orientation),
                "page_width_in": emu_to_inches(section.page_width),
                "page_height_in": emu_to_inches(section.page_height),
                "top_margin_in": emu_to_inches(section.top_margin),
                "bottom_margin_in": emu_to_inches(section.bottom_margin),
                "left_margin_in": emu_to_inches(section.left_margin),
                "right_margin_in": emu_to_inches(section.right_margin),
                "header_distance_in": emu_to_inches(section.header_distance),
                "footer_distance_in": emu_to_inches(section.footer_distance),
            }
        )
        result["headers"].append([paragraph_record(p) for p in section.header.paragraphs])
        result["footers"].append([paragraph_record(p) for p in section.footer.paragraphs])
    for style in doc.styles:
        if style.type is None:
            continue
        font = getattr(style, "font", None)
        pf = getattr(style, "paragraph_format", None)
        base_style = getattr(style, "base_style", None)
        result["styles"].append(
            {
                "name": style.name,
                "type": str(style.type),
                "base_style": None if base_style is None else base_style.name,
                "font_name": None if font is None else font.name,
                "font_size_pt": None if font is None or font.size is None else round(font.size.pt, 2),
                "bold": None if font is None else font.bold,
                "italic": None if font is None else font.italic,
                "space_before_pt": None if pf is None or pf.space_before is None else round(pf.space_before.pt, 2),
                "space_after_pt": None if pf is None or pf.space_after is None else round(pf.space_after.pt, 2),
                "line_spacing": None if pf is None or pf.line_spacing is None else str(pf.line_spacing),
            }
        )
    for i, paragraph in enumerate(doc.paragraphs):
        result["paragraphs"].append(paragraph_record(paragraph, i))
    for ti, table in enumerate(doc.tables):
        tdata = {
            "index": ti,
            "style": None if table.style is None else table.style.name,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "cells": [],
        }
        for ri, row in enumerate(table.rows):
            row_data = []
            for ci, cell in enumerate(row.cells):
                row_data.append(
                    {
                        "row": ri,
                        "col": ci,
                        "width_in": emu_to_inches(cell.width),
                        "vertical_alignment": None if cell.vertical_alignment is None else str(cell.vertical_alignment),
                        "paragraphs": [paragraph_record(p) for p in cell.paragraphs],
                    }
                )
            tdata["cells"].append(row_data)
        result["tables"].append(tdata)
    paragraph_index = {id(paragraph._element): i for i, paragraph in enumerate(doc.paragraphs)}
    table_index = {id(table._element): i for i, table in enumerate(doc.tables)}
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            images = []
            for blip in block._element.xpath(".//a:blip"):
                rel_id = blip.get(qn("r:embed"))
                if rel_id and rel_id in doc.part.rels:
                    images.append({"rel_id": rel_id, "target": str(doc.part.rels[rel_id].target_ref)})
            result["body_blocks"].append(
                {
                    "kind": "paragraph",
                    "index": paragraph_index.get(id(block._element)),
                    "style": block.style.name if block.style is not None else None,
                    "text": block.text,
                    "images": images,
                }
            )
        elif isinstance(block, Table):
            result["body_blocks"].append(
                {
                    "kind": "table",
                    "index": table_index.get(id(block._element)),
                    "rows": len(block.rows),
                    "cols": len(block.columns),
                }
            )
    with ZipFile(path) as archive:
        result["package_parts"] = sorted(archive.namelist())
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    args.out.write_text(json.dumps(document_record(args.input), ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
