from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly")
OUT = ROOT / "journal_revision_q3"
FONT = "Palatino Linotype"


def paragraphs_and_cells(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def cite_numbers(text: str) -> list[int]:
    out: list[int] = []
    for content in re.findall(r"\[([^\]]+)\]", text):
        if not re.fullmatch(r"[0-9,\-\s]+", content):
            continue
        for token in content.split(","):
            token = token.strip()
            if "-" in token:
                a, b = [int(x.strip()) for x in token.split("-", 1)]
                out.extend(range(a, b + 1))
            elif token:
                out.append(int(token))
    return out


def verify_doc(path: Path, *, blinded: bool) -> dict:
    doc = Document(path)
    text = "\n".join(p.text for p in paragraphs_and_cells(doc))
    refs_index = next((i for i, p in enumerate(doc.paragraphs) if p.text == "References"), len(doc.paragraphs))
    references = [p for p in doc.paragraphs[refs_index + 1:] if re.match(r"^\d+\.\s", p.text)]
    captions = sorted(int(m.group(1)) for m in re.finditer(r"^Figure (\d+)\.", "\n".join(p.text for p in doc.paragraphs), re.MULTILINE))
    body_text = "\n".join(p.text for p in doc.paragraphs[:refs_index])
    citations = []
    for p in doc.paragraphs[:refs_index]:
        citations.extend(cite_numbers(p.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                citations.extend(cite_numbers(cell.text))
    citation_first = []
    for number in citations:
        if number not in citation_first:
            citation_first.append(number)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    visible_text = "\n".join(p.text for p in doc.paragraphs)
    equation_labels = sorted(set(re.findall(r"\((\d+)\)", visible_text)), key=int)
    all_runs = [run for p in paragraphs_and_cells(doc) for run in p.runs]
    bad_fonts = [run.text for run in all_runs if run.text and run.font.name != FONT]
    result = {
        "file": str(path),
        "pages_verified_by_word_render": None,
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "figure_captions": captions,
        "references": len(references),
        "citation_first_appearance": citation_first,
        "equation_labels": equation_labels,
        "bad_font_runs": len(bad_fonts),
        "contains_placeholders": any(x in text or x in xml for x in ("20XX", "xx-xx", "vxxix", "xxxxxx")),
        "contains_removed_multiple_testing_claim": "Benjamini" in text or "false-discovery" in text,
        "blinded_author_text_found": ("Punyawansiri" in text or "Surasit" in text) if blinded else False,
    }
    assert len(doc.tables) == 6 if blinded else len(doc.tables) == 0
    assert len(doc.inline_shapes) == 7 if blinded else True
    assert captions == list(range(1, 8)) if blinded else True
    assert len(references) == 27 if blinded else True
    assert citation_first == list(range(1, 28)) if blinded else True
    assert equation_labels == [str(i) for i in range(1, 15)] if blinded else True
    assert not bad_fonts
    assert not result["contains_placeholders"]
    assert not result["contains_removed_multiple_testing_claim"]
    assert not result["blinded_author_text_found"]
    return result


def main() -> None:
    main_doc = verify_doc(OUT / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN.docx", blinded=True)
    title_doc = verify_doc(OUT / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_TITLE_PAGE.docx", blinded=False)
    figure_files = sorted((OUT / "submission_figures").glob("Figure_*.tiff"))
    assert len(figure_files) == 7
    payload = {"main": main_doc, "title_page": title_doc, "separate_tiff_figures": [p.name for p in figure_files]}
    print(json.dumps(payload, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
