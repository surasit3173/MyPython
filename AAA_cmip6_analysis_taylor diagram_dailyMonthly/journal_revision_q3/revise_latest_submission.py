from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly")
SOURCE = Path(r"D:\วารสาร ASEAN J. Sci\New paper\AJSTR_QDM_CMIP6_Q2Q3_REVISED_FINAL (1).docx")
OUT_DIR = ROOT / "journal_revision_q3"
MAIN = OUT_DIR / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN.docx"
TITLE_PAGE = OUT_DIR / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_TITLE_PAGE.docx"
FIGURES = OUT_DIR / "submission_figures"
ASSETS = OUT_DIR / "source_assets"
FONT = "Palatino Linotype"

FIGURE_CAPTIONS = {
    "Figure 1. Location of Prachuap Khiri Khan Province and the 12 rainfall stations used in the analysis.":
        "Figure 1. Location of Prachuap Khiri Khan Province and the 12 rainfall stations used in the analysis.",
    "Figure 1. Fit-freeze-apply framework. Observation-dependent transfer functions were estimated in 1981-2000 and applied without re-estimation in 2001-2014; uncertainty used paired year-block bootstrap resampling.":
        "Figure 2. Fit-freeze-apply framework. Observation-dependent transfer functions were estimated in 1981-2000 and applied without re-estimation in 2001-2014; uncertainty used paired year-block bootstrap resampling.",
    "Figure 2. Observed rainfall in calibration and validation: (a) monthly climatology with interquartile ribbons across stations, (b) annual rainfall, and (c) paired station-level wet-day frequencies. Colored connecting lines show the direction of change; these summaries are descriptive and motivate the independent holdout test.":
        "Figure 3. Observed rainfall in calibration and validation: (a) monthly climatology with interquartile ribbons across stations, (b) annual rainfall, and (c) paired station-level wet-day frequencies. Colored connecting lines show the direction of change; these summaries are descriptive and motivate the independent holdout test.",
    "Figure 3. Validation distributions for 60 model-station combinations per variant. Boxes show the interquartile range, whiskers show the 10th-90th percentiles, diamonds show means, centre lines show medians, and dashed vertical lines show the raw mean.":
        "Figure 4. Validation distributions for 60 model-station combinations per variant. Boxes show the interquartile range, whiskers show the 10th-90th percentiles, diamonds show means, centre lines show medians, and dashed vertical lines show the raw mean.",
    "Figure 4. Percentage of 60 paired model-station evaluations with a favorable point change relative to raw CMIP6 simulations. Values above 50% indicate a favorable change for most pairs; percentages are descriptive and pairs are not independent.":
        "Figure 5. Percentage of 60 paired model-station evaluations with a favorable point change relative to raw CMIP6 simulations. Values above 50% indicate a favorable change for most pairs; percentages are descriptive and pairs are not independent.",
    "Figure 5. Station-level response of monthly QDM. (a) Heatmap of corrected-minus-raw KGE for NP monthly QDM by station and CMIP6 model. (b) Paired station means for non-parametric and parametric monthly QDM. Values are descriptive because gauge-model combinations are dependent.":
        "Figure 6. Station-level response of monthly QDM. (a) Heatmap of corrected-minus-raw KGE for NP monthly QDM by station and CMIP6 model. (b) Paired station means for non-parametric and parametric monthly QDM. Values are descriptive because gauge-model combinations are dependent.",
    "Figure 6. Sensitivity of monthly QDM validation metrics to the observed wet-day threshold. Points and lines are means across 60 paired model-station evaluations; translucent ribbons show the interquartile range across those pairs, and the dashed line shows the raw mean. Lower is favorable for RMSE and absolute q99 bias, whereas higher is favorable for KGE. Lines aid comparison and do not imply continuous interpolation.":
        "Figure 7. Sensitivity of monthly QDM validation metrics to the observed wet-day threshold. Points and lines are means across 60 paired model-station evaluations; translucent ribbons show the interquartile range across those pairs, and the dashed line shows the raw mean. Lower is favorable for RMSE and absolute q99 bias, whereas higher is favorable for KGE. Lines aid comparison and do not imply continuous interpolation.",
}

# New citation order follows the first appearance in the revised manuscript. The
# multiple-testing reference is intentionally omitted because no adjusted p-values
# are reported in the results.
CITATION_MAP = {
    1: 1, 2: 2, 3: 3, 4: 4, 6: 5, 7: 6, 5: 7, 8: 8, 9: 9, 10: 10,
    14: 11, 15: 12, 11: 13, 12: 14, 13: 15, 27: 16, 22: 17,
    16: 18, 17: 19, 18: 20, 19: 21, 20: 22, 28: 23, 24: 24,
    21: 25, 25: 26, 26: 27,
}
REFERENCE_ORDER = [1, 2, 3, 4, 6, 7, 5, 8, 9, 10, 14, 15, 11, 12, 13,
                   27, 22, 16, 17, 18, 19, 20, 28, 24, 21, 25, 26]


def set_run_font(run, size: float = 10.0, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)


def set_para_text(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)


def all_body_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def add_text(doc: Document, text: str, *, size: float = 10.0, bold: bool = False,
             italic: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, before: float = 0,
             after: float = 2):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic)
    return p


def add_labeled(doc: Document, label: str, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(label)
    set_run_font(r, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def renumber_citations(text: str) -> str:
    def repl(match: re.Match[str]) -> str:
        content = match.group(1)
        if not re.fullmatch(r"[0-9,\-\s]+", content):
            return match.group(0)
        numbers: list[int] = []
        for token in content.split(","):
            token = token.strip()
            if "-" in token:
                start, end = [int(x.strip()) for x in token.split("-", 1)]
                numbers.extend(range(start, end + 1))
            elif token:
                numbers.append(int(token))
        mapped = [CITATION_MAP[n] for n in numbers if n in CITATION_MAP]
        if not mapped:
            return ""
        return "[" + ",".join(str(n) for n in mapped) + "]"

    return re.sub(r"\[([^\]]+)\]", repl, text)


def add_equation_after(reference: Paragraph, expression: str, number: int) -> Paragraph:
    p = Paragraph(OxmlElement("w:p"), reference._parent)
    reference._p.addnext(p._p)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.10), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.10), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = expression
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)
    number_run = p.add_run(f"\t({number})")
    set_run_font(number_run)
    return p


def add_note_after(reference: Paragraph, text: str) -> Paragraph:
    p = Paragraph(OxmlElement("w:p"), reference._parent)
    reference._p.addnext(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=9.5, italic=True)
    return p


def add_picture_after(reference: Paragraph, path: Path, width_inches: float = 6.25) -> Paragraph:
    p = Paragraph(OxmlElement("w:p"), reference._parent)
    reference._p.addnext(p._p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    return p


def extract_map_and_figures() -> None:
    ASSETS.mkdir(exist_ok=True)
    FIGURES.mkdir(exist_ok=True)
    with zipfile.ZipFile(SOURCE) as archive:
        map_data = archive.read("word/media/image1.png")
    map_png = ASSETS / "Figure_1_location_map.png"
    map_png.write_bytes(map_data)
    try:
        from PIL import Image
        with Image.open(map_png) as image:
            image.convert("RGB").save(FIGURES / "Figure_1.tiff", dpi=(600, 600), compression="tiff_lzw")
    except Exception:
        shutil.copyfile(map_png, FIGURES / "Figure_1.png")

    source_names = [
        "Figure1_framework.tif", "Figure2_observed_periods.tif", "Figure3_method_comparison.tif",
        "Figure4_improvement_heatmap.tif", "Figure5_station_spatial_response.tif",
        "Figure6_wet_threshold_sensitivity.tif",
    ]
    for new_number, source_name in enumerate(source_names, 2):
        shutil.copyfile(OUT_DIR / "figures_q1" / source_name, FIGURES / f"Figure_{new_number}.tiff")


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def clear_template_metadata(doc: Document) -> None:
    for section in doc.sections:
        for container in (section.header, section.first_page_header, section.even_page_header,
                          section.footer, section.first_page_footer, section.even_page_footer):
            for p in container.paragraphs:
                if "20XX" in p.text or "doi.org/10.55164" in p.text:
                    set_para_text(p, "ASEAN J. Sci. Tech. Report.")
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""


def citation_and_prose_revision(doc: Document) -> None:
    replacements = {
        "Figure 1. Location of Prachuap Khiri Khan Province and the 12 rainfall stations used in the analysis.":
            "Figure 1. Location of Prachuap Khiri Khan Province and the 12 rainfall stations used in the analysis.",
        "Figure 1. Fit-freeze-apply framework.": "Figure 2. Fit-freeze-apply framework.",
        "Figure 2. Observed rainfall in calibration and validation:": "Figure 3. Observed rainfall in calibration and validation:",
        "Figure 3. Validation distributions for 60 model-station combinations per variant.": "Figure 4. Validation distributions for 60 model-station combinations per variant.",
        "Figure 4. Percentage of 60 paired model-station evaluations": "Figure 5. Percentage of 60 paired model-station evaluations",
        "Figure 5. Station-level response of monthly QDM.": "Figure 6. Station-level response of monthly QDM.",
        "Figure 6. Sensitivity of monthly QDM validation metrics": "Figure 7. Sensitivity of monthly QDM validation metrics",
        "Mean annual observed rainfall ranged from 944 to 1,376 mm among stations, and maximum daily rainfall ranged from 54.2 to 298.5 mm day-1. Calibration and validation retained the same broad bimodal seasonal cycle, but several stations changed wet-day frequency between periods (Figure 2).":
            "Mean annual observed rainfall ranged from 944 to 1,376 mm among stations, and maximum daily rainfall ranged from 54.2 to 298.5 mm day⁻¹. Calibration and validation retained the same broad bimodal seasonal cycle, but several stations changed wet-day frequency between periods (Figure 3).",
        "(Table 3; Figure 3)": "(Table 3; Figure 4)",
        "Station-aggregated responses were heterogeneous (Figure 5).": "Station-aggregated responses were heterogeneous (Figure 6).",
        "Wet-day threshold sensitivity was method-dependent (Figure 6).": "Wet-day threshold sensitivity was method-dependent (Figure 7).",
        "The station codes and coordinates in the Hydro-Informatics Institute/Royal Irrigation Department project archive match the official 12-station inventory reported for the province [27].":
            "The station codes and coordinates in the Hydro-Informatics Institute/Royal Irrigation Department project archive match the official 12-station inventory reported for the province [27] (Figure 1).",
        "The paired bootstrap did not support a broad parametric advantage over the empirical monthly method.":
            "The paired bootstrap did not support a broad parametric advantage over the empirical monthly method (Table 4).",
        "Upper-tail performance followed a different ordering.":
            "Upper-tail performance followed a different ordering (Figure 5).",
        "27.38% for NP monthly": "27.37% for NP monthly",
        "All computational procedures and maximum likelihood optimizations were executed using Python 3.12.13 and the SciPy 1.18.1 library [24]. ":
            "All computational procedures and maximum likelihood optimizations were executed using Python 3.12.3 and the SciPy 1.17.1 library [24].",
        "The source, experiment, member, grid, variable, frequency, and date range are encoded in the analysis-ready filenames and reported in Table 2; model descriptions follow the corresponding publications [16-20].":
            "The source, experiment, member, grid, variable, frequency, and date range are encoded in the analysis-ready filenames and summarized in Table 2; model descriptions follow the corresponding publications [16-20]. The EC-Earth3 gr identifier is retained exactly as recorded in the analysis-ready filename.",
        "Note: Cell area is a first-order latitude-longitude approximation at 12 degrees N (111.2 km per degree latitude; 108.8 km per degree longitude), not an exact curvilinear-grid area.":
            "Note: Grid identifiers (gn/gr) are retained as recorded in the analysis-ready files. Cell area is a first-order latitude-longitude approximation at 12 degrees N (111.2 km per degree latitude; 108.8 km per degree longitude), not an exact curvilinear-grid area.",
        "The apparent deficits of 5 observed and 33 model fits in the earlier table were omitted gamma and Weibull categories, not optimization failures.": "",
        "This rules out the uncorrected AIC penalty as the cause of the observed selection pattern for these samples, but it does not establish a physical rainfall law; Pearson type III retained greater shape and location flexibility than several constrained candidates.":
            "Because AICc reproduced the selections and validation metrics in this analysis, the selection pattern was not sensitive to the small-sample correction. This does not establish a physical rainfall law; Pearson type III retained greater shape and location flexibility than several constrained candidates.",
        "NP monthly is an assumption-lean default for seasonal reservoir inflow, irrigation scheduling, and monthly drought or water-balance screening.":
            "NP monthly is an assumption-lean default for seasonal reservoir-inflow and monthly drought or water-balance screening; it should not be interpreted as a substitute for event-scale irrigation scheduling.",
        "The reported facets and checksum manifest identify the exact derivatives used, while independent recreation requires repeating the spatial extraction from the corresponding ESGF datasets.":
            "The metadata summary and checksum manifest identify the exact derivative files used; independent recreation would additionally require repeating the spatial extraction from the corresponding ESGF datasets.",
        "CMIP6 source datasets are discoverable through ESGF using the complete facets reported in Table 2.":
            "CMIP6 source datasets are discoverable through ESGF using the metadata summarized in Table 2.",
    }
    revised_bootstrap = (
        "Uncertainty in corrected-minus-raw differences was estimated separately for each model-station combination using "
        "2,000 paired year-block bootstrap replicates with a fixed random seed (seed 42). Validation monthly totals were "
        "structured as a 14-year by 12-month matrix. Fourteen years were sampled with replacement in each replicate, and "
        "the exact same year indices were applied simultaneously to observations, raw simulations, and all corrected variants "
        "to preserve temporal correlation. The 95% interval was extracted from the 2.5th and 97.5th percentiles of the "
        "bootstrap distribution. These intervals and favorable-point counts are descriptive pair-specific summaries; no "
        "network-wide p-value or multiple-testing claim is made because the 60 combinations are spatially and structurally dependent. "
        "The fit-freeze-apply sequence is summarized in Figure 2."
    )
    revised_metric_intro = (
        "Secondary endpoints included monthly root mean square error (RMSE), mean absolute error (MAE), percent bias (PBIAS), "
        "the absolute relative bias of the 99th wet-day percentile (q99 bias), the Kolmogorov-Smirnov statistic (D) for wet-day "
        "distributions, monthly Pearson correlation (r), and correlation between the 12 monthly climatological means (r_clim). "
        "The equations for these metrics are expressed below."
    )
    revised_metric_note = (
        "Here, N is the number of paired monthly records; Sᵢ and Oᵢ are simulated and observed monthly totals; q₉₉ is the "
        "99th percentile of wet-day intensity; and F_S and F_O are empirical cumulative distribution functions of wet-day "
        "precipitation. The climatology correlation is calculated from the 12 month-specific means, not from independent observations. "
        "KGE and correlations are better at larger values, whereas the remaining endpoints are better at smaller values. "
        "Direction-based rankings are labelled descriptive and do not imply statistical superiority."
    )
    for p in list(doc.paragraphs):
        text = p.text
        if text == "Article":
            text = "Research Article"
        if text.startswith("Evaluating Parametric and Non-Parametric"):
            text = "Evaluating parametric and non-parametric quantile delta mapping for CMIP6 precipitation bias correction: independent validation in coastal Thailand"
        if text.startswith("Secondary endpoints included"):
            text = revised_metric_intro
        if text.startswith("For a model precipitation value"):
            text = "For a model precipitation value x at time t in target period p, the multiplicative Quantile Delta Mapping (QDM) procedure was applied following Cannon et al. [5]."
        if text.startswith("Uncertainty in the corrected-minus-raw"):
            text = revised_bootstrap
        if text.startswith("where ") and "number of paired monthly records" in text:
            text = revised_metric_note
        for old, new in replacements.items():
            text = text.replace(old, new)
        text = text.replace("mm day-1", "mm day⁻¹")
        text = text.replace("kg m-2 s-1", "kg m⁻² s⁻¹")
        text = text.replace("km2", "km²")
        text = renumber_citations(text)
        if text != p.text:
            set_para_text(p, text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    updated = renumber_citations(p.text)
                    if updated != p.text:
                        set_para_text(p, updated)

    # Compact header labels keep the wide station/model tables readable at the
    # journal's 10-point Palatino setting without changing any values.
    if len(doc.tables) >= 6:
        table_headers = {
            0: ["Station", "Lat. (°N)", "Long. (°E)", "Elev. (m)", "Miss. (%)", "Annual rain. (mm)", "SDII (mm day⁻¹)", "Max. daily rain. (mm day⁻¹)", "Lag-1 ACF"],
            1: ["Model", "Institution", "Variant", "Grid", "Resolution; area (10³ km²)", "Cal.", "Days", "Series", "Ref."],
            3: ["Method", "Met.", "Raw mean", "Corrected mean", "Median delta", "Median 95% interval", "Favorable point change", "CI F/A/O"],
            4: ["Method", "Eligible/structural", "Observed selection", "Model selection", "Annual fallback", "Fit failures", "Capped", "Clipped events"],
            5: ["Cap factor", "P annual RMSE", "P annual KGE", "P monthly RMSE", "P monthly KGE", "P monthly |q99| bias (%)", "P annual capped", "P monthly capped"],
        }
        for table_index, headers in table_headers.items():
            for cell, header in zip(doc.tables[table_index].rows[0].cells, headers):
                set_para_text(cell.paragraphs[0], header)
        for row in doc.tables[1].rows[1:]:
            if row.cells[5].text.strip().lower() == "standard":
                set_para_text(row.cells[5].paragraphs[0], "Std.")
        for row in doc.tables[3].rows[1:]:
            for cell in row.cells:
                if cell.text.strip() in {"mRMSE", "mKGE"}:
                    set_para_text(cell.paragraphs[0], cell.text.strip()[1:])

    # Add only definitions for metrics already reported in the article; no values
    # are recalculated and no new experiment is introduced.
    metric_note = next((p for p in doc.paragraphs if p.text.startswith("Here, N is the number of paired")), None)
    if metric_note is not None:
        equations = [
            ("PBIAS = 100 × (S̄ − Ō) / Ō", 11),
            ("MAE = (1/N) Σᵢ₌₁ᴺ |Sᵢ − Oᵢ|", 12),
            ("r = Σᵢ₌₁ᴺ(Sᵢ − S̄)(Oᵢ − Ō) / √[Σᵢ₌₁ᴺ(Sᵢ − S̄)² Σᵢ₌₁ᴺ(Oᵢ − Ō)²]", 13),
            ("r_clim = corr({S̄₁, …, S̄₁₂}, {Ō₁, …, Ō₁₂})", 14),
        ]
        previous = metric_note
        inserted = []
        for expression, number in equations:
            previous = add_equation_after(previous, expression, number)
            inserted.append(previous)
        add_note_after(
            inserted[-1],
            "The PBIAS and MAE equations use the same paired monthly records as RMSE; r is the Pearson correlation of monthly totals; and r_clim is the Pearson correlation of the 12 simulated and observed monthly climatological means.",
        )

    # Main-manuscript figure sequence: map, workflow, observed context, metrics,
    # favorable changes, station response, and threshold sensitivity.
    for p in doc.paragraphs:
        if p.text.startswith("Figure "):
            for old, new in FIGURE_CAPTIONS.items():
                if p.text == old:
                    set_para_text(p, new)
                    break

    # In the supplied latest file the new location map shares its XML paragraph
    # with the caption. Reinsert the extracted raster after the edited caption so
    # changing the caption cannot remove the figure object.
    map_caption = next((p for p in doc.paragraphs if p.text.startswith("Figure 1. Location of Prachuap")), None)
    if map_caption is not None and not any(s.width.inches > 6.8 for s in doc.inline_shapes):
        add_picture_after(map_caption, ASSETS / "Figure_1_location_map.png")


def remove_blinded_content(doc: Document) -> None:
    for p in list(doc.paragraphs):
        text = p.text.strip()
        if text in {
            "Surasit Punyawansiri1*",
            "1 Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok 10300, Thailand",
            "* Correspondence: Surasit.irri@gmail.com",
            "5. Acknowledgements",
            "The author acknowledges the institutions that generated and distributed the CMIP6 simulations and the custodians of the Thai rain-gauge data used in the supplied project archive.",
            "Author Contributions: S.P.: conceptualization, methodology, software, validation, formal analysis, data curation, visualization, writing-original draft, and writing-review and editing.",
        }:
            remove_paragraph(p)


def rebuild_references(doc: Document, reference_texts: dict[int, str]) -> None:
    refs_start = None
    for p in doc.paragraphs:
        if p.text.strip() == "References":
            refs_start = p
            break
    if refs_start is None:
        return
    started = False
    for p in list(doc.paragraphs):
        if p._p is refs_start._p:
            started = True
            continue
        if started and re.match(r"^\d+\.\s", p.text.strip()):
            remove_paragraph(p)
    for new_number, old_number in enumerate(REFERENCE_ORDER, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{new_number}. {reference_texts[old_number]}")
        set_run_font(r, size=9.5)


def style_document(doc: Document, *, title_page: bool = False) -> None:
    for style in doc.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = FONT
        style.font.size = Pt(10)
        if style._element.rPr is None:
            style._element.get_or_add_rPr()
        r_fonts = style._element.rPr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            style._element.rPr.insert(0, r_fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{attr}"), FONT)

    for p in all_body_paragraphs(doc):
        if p.text.startswith("Evaluating parametric"):
            size = 16
        elif p.text.startswith("References") or re.match(r"^\d+\.\s", p.text):
            size = 9.5 if p.text[:1].isdigit() and not p.text.startswith(("1. ", "2. ", "3. ", "4. ", "5. ")) else 10
        elif p.text.startswith(("Figure ", "Table ")):
            size = 9.5
        else:
            size = 10
        for run in p.runs:
            set_run_font(run, size=size)

    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if row_index == 0 and not tr_pr.xpath("./w:tblHeader"):
                tr_pr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_together = True
                    for run in p.runs:
                        set_run_font(run, size=9.5, bold=(row_index == 0) or run.bold)

    for shape, alt in zip(doc.inline_shapes, [
        "Map of Prachuap Khiri Khan Province and the 12 rainfall stations.",
        "Fit-freeze-apply workflow for independent temporal validation.",
        "Observed rainfall climatology, annual rainfall, and wet-day frequency by period.",
        "Validation distributions for four QDM variants across paired evaluations.",
        "Favorable point-change percentages for corrected versus raw rainfall metrics.",
        "Station by model heatmap and paired station means for monthly QDM response.",
        "Sensitivity of monthly QDM metrics to the wet-day threshold.",
    ]):
        shape._inline.docPr.set("descr", alt)
        shape._inline.docPr.set("title", alt)

    clear_template_metadata(doc)
    if title_page:
        doc.core_properties.author = "Surasit Punyawansiri"
        doc.core_properties.last_modified_by = "Surasit Punyawansiri"


def build_title_page(source_doc: Document, statements: dict[str, str]) -> Document:
    doc = Document(SOURCE)
    clear_body(doc)
    add_text(doc, "Title Page", size=12, bold=True, before=10, after=8)
    add_text(doc, "Evaluating parametric and non-parametric quantile delta mapping for CMIP6 precipitation bias correction: independent validation in coastal Thailand", size=16, bold=True, after=8)
    add_text(doc, "Surasit Punyawansiri1*", bold=True, after=2)
    add_text(doc, "1 Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok 10300, Thailand", after=2)
    add_text(doc, "* Correspondence: Surasit.irri@gmail.com", after=10)
    add_text(doc, "Statements", size=12, bold=True, before=6, after=4)
    add_labeled(doc, "Acknowledgements: ", statements["ack"])
    add_labeled(doc, "Author Contributions: ", statements["contrib"])
    add_labeled(doc, "Funding: ", statements["funding"])
    data_statement = (statements["data"].replace(" [27]", "").replace("in.", "in the project archive.")
                      .replace("complete facets reported in Table 2", "metadata summarized in Table 2"))
    add_labeled(doc, "Data Availability Statement: ", data_statement)
    add_labeled(doc, "Conflicts of Interest: ", statements["conflict"])
    add_labeled(doc, "Author Approval: ", "The author has read and approved the final version of the manuscript for submission.")
    style_document(doc, title_page=True)
    doc.save(TITLE_PAGE)
    return doc


def build() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    extract_map_and_figures()
    source_doc = Document(SOURCE)

    reference_texts: dict[int, str] = {}
    statement_texts = {"ack": "", "contrib": "", "funding": "", "data": "", "conflict": ""}
    in_refs = False
    for p in source_doc.paragraphs:
        text = p.text.strip()
        if text == "References":
            in_refs = True
            continue
        if in_refs:
            match = re.match(r"^(\d+)\.\s+(.*)$", text)
            if match:
                reference_texts[int(match.group(1))] = match.group(2)
        if text.startswith("The author acknowledges"):
            statement_texts["ack"] = text
        elif text.startswith("Author Contributions:"):
            statement_texts["contrib"] = text.split(": ", 1)[1]
        elif text.startswith("Funding:"):
            statement_texts["funding"] = text.split(": ", 1)[1]
        elif text.startswith("Data Availability Statement:"):
            statement_texts["data"] = text.split(": ", 1)[1]
        elif text.startswith("Conflicts of Interest:"):
            statement_texts["conflict"] = text.split(": ", 1)[1]

    main = Document(SOURCE)
    citation_and_prose_revision(main)
    remove_blinded_content(main)
    rebuild_references(main, reference_texts)
    style_document(main)
    main.save(MAIN)
    build_title_page(source_doc, statement_texts)
    print(MAIN)
    print(TITLE_PAGE)


if __name__ == "__main__":
    build()
