from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("C:/MyPython/AAA_cmip6_analysis_taylor diagram_dailyMonthly")
sys.path.insert(0, str(ROOT / "journal_template_work"))
import build_ajstr_manuscript as base  # noqa: E402


REFERENCE = ROOT / "journal_template_work" / "AJSTR_template_reference.docx"
OUTPUT = ROOT / "journal_revision_q3" / "AJSTR_QDM_CMIP6_Q2Q3_REVISED_complete.docx"
RESULTS = ROOT / "qdm_p_np_publication_out_q3_final"
STATS = ROOT / "journal_revision_q3" / "statistics"
FIGURES = ROOT / "journal_revision_q3" / "figures_q1"
REVISION2 = ROOT / "journal_revision_q3" / "revision2_analysis"
BODY_FONT = "Palatino Linotype"


def set_font_on_run(run, font: str = BODY_FONT) -> None:
    run.font.name = font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font)


def iter_container_paragraphs(container):
    yield from container.paragraphs
    for table_obj in container.tables:
        for row in table_obj.rows:
            for cell in row.cells:
                yield from iter_container_paragraphs(cell)


def apply_document_font(doc: Document) -> None:
    for style in doc.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = BODY_FONT
        if style._element.rPr is None:
            style._element.get_or_add_rPr()
        r_fonts = style._element.rPr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            style._element.rPr.insert(0, r_fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{attr}"), BODY_FONT)

    containers = [doc]
    for section in doc.sections:
        containers.extend([section.header, section.first_page_header, section.even_page_header,
                           section.footer, section.first_page_footer, section.even_page_footer])
    seen = set()
    for container in containers:
        marker = id(container._element)
        if marker in seen:
            continue
        seen.add(marker)
        for paragraph in iter_container_paragraphs(container):
            for run in paragraph.runs:
                set_font_on_run(run)


def set_article_type(doc: Document) -> None:
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if "Article type" in run.text:
                                    run.text = run.text.replace("Article type", "Research Article")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_figure(doc, path: Path, number: int, caption: str) -> None:
    caption_paragraph = base.add_caption(doc, "Figure", number, caption)
    caption_paragraph.paragraph_format.keep_with_next = True
    figure_paragraph = doc.add_paragraph()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.keep_together = True
    figure_paragraph.add_run().add_picture(str(path), width=Inches(6.25))


def polish_table(table) -> None:
    set_repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            mar = tc_pr.first_child_found_in("w:tcMar")
            if mar is None:
                mar = OxmlElement("w:tcMar")
                tc_pr.append(mar)
            for side, value in (("top", 70), ("left", 80), ("bottom", 70), ("right", 80)):
                node = mar.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    mar.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_together = True
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(doc, headers, rows, widths, font_size=8):
    result = base.add_table(doc, headers, rows, widths=widths, font_size=font_size)
    polish_table(result)
    return result


def add_equation(doc: Document, expression: str, number: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(3.10), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Inches(6.10), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = expression
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)
    number_run = p.add_run(f"\t({number})")
    base.set_run(number_run, font=BODY_FONT, size=12)


def station_rows():
    old_doc = Document(ROOT / "cmip6bc_q1_editable" / "ID1558_Manuscript_BSJ.docx")
    source = old_doc.tables[0]
    return (
        [cell.text.replace("\n", " ").strip() for cell in source.rows[0].cells],
        [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in source.rows[1:]],
    )


def model_rows():
    return [
        ["ACCESS-ESM1-5", "CSIRO", "r1i1p1f1", "gn", "1.875 x 1.25; 28.4", "standard", "12,418", "2", "[20]"],
        ["CESM2", "NCAR", "r11i1p1f1", "gn", "1.25 x 0.90; 13.6", "365-day", "12,410", "3", "[17]"],
        ["CanESM5", "CCCma", "r1i1p1f1", "gn", "2.8 x 2.8; 94.9", "365-day", "12,410", "2", "[16]"],
        ["EC-Earth3", "EC-Earth Consortium", "r1i1p1f1", "gr", "1.125 x 1.125; 15.3", "standard", "12,418", "3", "[18]"],
        ["MIROC6", "MIROC Consortium", "r1i1p1f1", "gn", "1.4 x 1.4; 23.7", "standard", "12,418", "2", "[19]"],
    ]


def validation_rows():
    summary = pd.read_csv(RESULTS / "summary.csv")
    labels = {
        "QDM_NP_annual": "NP annual",
        "QDM_NP_monthly": "NP monthly",
        "QDM_P_annual": "P annual",
        "QDM_P_monthly": "P monthly",
    }
    specs = [
        ("Mean absolute PBIAS (%)", "PBIAS", 2, "lower"),
        ("Kolmogorov-Smirnov D", "KS_D", 3, "lower"),
        ("Mean absolute q99 bias (%)", "q99_relbias_pct", 2, "lower"),
        ("Monthly RMSE (mm)", "mRMSE", 2, "lower"),
        ("Monthly MAE (mm)", "mMAE", 2, "lower"),
        ("Monthly KGE", "mKGE", 3, "higher"),
        ("Monthly correlation r", "mr", 3, "higher"),
        ("Monthly climatology r", "clim_r", 3, "higher"),
    ]
    rows = []
    validation = summary[summary.period == "validation"]
    for title, metric, decimals, direction in specs:
        values = validation[validation.metric == metric].set_index("method")
        corrected = {method: float(values.loc[method, "corrected_mean"]) for method in labels}
        raw = float(values.iloc[0].raw_mean)
        if direction == "lower":
            best = min(corrected, key=lambda key: abs(corrected[key]) if metric in {"PBIAS", "q99_relbias_pct"} else corrected[key])
        else:
            best = max(corrected, key=corrected.get)
        rows.append([
            title,
            f"{raw:.{decimals}f}",
            f"{corrected['QDM_NP_annual']:.{decimals}f}",
            f"{corrected['QDM_NP_monthly']:.{decimals}f}",
            f"{corrected['QDM_P_annual']:.{decimals}f}",
            f"{corrected['QDM_P_monthly']:.{decimals}f}",
            labels[best],
        ])
    return rows


def bootstrap_rows():
    data = pd.read_csv(STATS / "bootstrap_summary.csv")
    labels = {
        "QDM_NP_annual": "NP annual",
        "QDM_NP_monthly": "NP monthly",
        "QDM_P_annual": "P annual",
        "QDM_P_monthly": "P monthly",
    }
    rows = []
    for _, row in data.iterrows():
        rows.append([
            labels[row.method],
            row.metric,
            f"{row.mean_raw:.3f}" if row.metric == "mKGE" else f"{row.mean_raw:.2f}",
            f"{row.mean_corrected:.3f}" if row.metric == "mKGE" else f"{row.mean_corrected:.2f}",
            f"{row.median_delta:+.3f}" if row.metric == "mKGE" else f"{row.median_delta:+.2f}",
            (f"[{row.median_ci_low:+.3f}, {row.median_ci_high:+.3f}]" if row.metric == "mKGE"
             else f"[{row.median_ci_low:+.2f}, {row.median_ci_high:+.2f}]"),
            f"{int(row.point_improved_n)}/60",
            f"{int(row.ci_improved_n)}/{int(row.ci_worsened_n)}/{int(row.ci_overlaps_zero_n)}",
        ])
    return rows


def diagnostic_rows():
    diag = pd.read_csv(RESULTS / "fit_diagnostics.csv")
    validation = diag[(diag.period == "validation") & (diag.family == "parametric")]
    validation = validation.drop_duplicates(["method", "model", "station", "group"])
    rows = []
    abbreviations = {
        "pearson3": "P3",
        "genextreme": "GEV",
        "lognorm": "LN",
        "gamma": "Gamma",
        "weibull_min": "Weibull",
    }
    for method, label, structural in (
        ("QDM_P_annual", "P annual", 60),
        ("QDM_P_monthly", "P monthly", 720),
    ):
        d = validation[validation.method == method]
        eligible = int(d.usable.sum())
        obs = d.loc[d.usable, "obs_dist"].value_counts()
        mod = d.loc[d.usable, "mod_dist"].value_counts()
        all_rows = diag[(diag.method == method) & (diag.period == "validation")]
        clipped = int(all_rows.n_delta_clipped_low.sum() + all_rows.n_delta_clipped_high.sum())
        capped = int(all_rows.n_output_capped.sum())
        rows.append([
            label,
            f"{eligible}/{structural}",
            "; ".join(f"{abbreviations.get(name, name)} {count}" for name, count in obs.items()),
            "; ".join(f"{abbreviations.get(name, name)} {count}" for name, count in mod.items()),
            str(int(d.fallback_used.sum())),
            "0",
            str(capped),
            str(clipped),
        ])
    return rows


def sensitivity_rows():
    roots = [
        ("1.5", ROOT / "qdm_p_np_publication_out_q3_cap1p5"),
        ("5.0 (primary)", RESULTS),
        ("None", ROOT / "qdm_p_np_publication_out_q3_uncapped"),
    ]
    rows = []
    for label, root in roots:
        summary = pd.read_csv(root / "summary.csv")
        diag = pd.read_csv(root / "fit_diagnostics.csv")
        v = summary[summary.period == "validation"].set_index(["method", "metric"])
        rows.append([
            label,
            f"{v.loc[('QDM_P_annual', 'mRMSE'), 'corrected_mean']:.2f}",
            f"{v.loc[('QDM_P_annual', 'mKGE'), 'corrected_mean']:.3f}",
            f"{v.loc[('QDM_P_monthly', 'mRMSE'), 'corrected_mean']:.2f}",
            f"{v.loc[('QDM_P_monthly', 'mKGE'), 'corrected_mean']:.3f}",
            f"{v.loc[('QDM_P_monthly', 'q99_relbias_pct'), 'corrected_mean']:.2f}",
            str(int(diag[(diag.period == "validation") & (diag.method == "QDM_P_annual")].n_output_capped.sum())),
            str(int(diag[(diag.period == "validation") & (diag.method == "QDM_P_monthly")].n_output_capped.sum())),
        ])
    return rows


def add_body_paragraphs(doc, paragraphs):
    for text in paragraphs:
        base.add_para(doc, text, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def build() -> None:
    shutil.copyfile(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    base.clear_body(doc)
    set_article_type(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.79)
    section.right_margin = Inches(0.79)
    section.top_margin = Inches(0.98)
    section.bottom_margin = Inches(0.89)
    section.different_first_page_header_footer = True
    doc.styles["Normal"].font.name = BODY_FONT
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    doc.styles["Normal"].font.size = Pt(12)

    base.add_para(doc, "Article", size=12, italic=True, before=8, after=4)
    base.add_para(
        doc,
        "Evaluating Parametric and Non-Parametric Quantile Delta Mapping for CMIP6 Precipitation Bias Correction: Independent Validation in Coastal Thailand",
        font="Palatino Linotype", size=18, bold=True, after=6,
    )
    base.add_para(doc, "Surasit Punyawansiri1*", size=12, bold=True, after=2)
    base.add_para(doc, "1 Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok 10300, Thailand", size=12, after=2)
    base.add_para(doc, "* Correspondence: Surasit.irri@gmail.com", size=12, after=8)

    abstract = (
        "Bias correction of climate-model precipitation is often assessed in the calibration period, which can overstate transferability. "
        "This study compared annual and month-wise parametric and non-parametric quantile delta mapping (QDM) for daily rainfall from five CMIP6 models at 12 stations in Prachuap Khiri Khan, Thailand. Transfer functions were fitted in 1981-2000 and frozen before independent validation in 2001-2014. Parametric candidates were selected by maximum-likelihood Akaike information criterion (AIC); robustness tests used small-sample AIC (AICc), wet-day thresholds of 0.1, 0.5, and 1.0 mm day-1, and alternative upper-tail caps. At the primary 1.0-mm threshold and five-times cap, parametric monthly QDM reduced mean monthly RMSE from 116.20 to 105.06 mm and increased mean Kling-Gupta efficiency (KGE) from 0.229 to 0.337; non-parametric monthly QDM yielded 109.45 mm and 0.329. Their paired median KGE difference was 0.0002, and the ranking reversed without the cap. AICc reproduced every AIC selection. Non-parametric monthly KGE remained 0.329-0.332 across wet-day thresholds, whereas parametric monthly KGE decreased to 0.217 at 0.1 mm. Non-parametric annual QDM gave the lowest mean absolute 99th-percentile bias (24.79% versus 52.98% raw). The 60 model-station pairs share model grid cells and are descriptive paired evaluations, not independent spatial replicates. Month-wise calibration is supported for monthly water-balance applications, but parametric gains remain conditional on occurrence and tail safeguards; empirical monthly QDM is the more assumption-lean default."
    )
    base.add_mixed_para(doc, "Abstract: ", abstract, before=6, after=4)
    base.add_mixed_para(doc, "Keywords: ", "bias correction; quantile delta mapping; CMIP6; independent validation; Thailand", after=8)

    base.add_heading(doc, "1", "Introduction")
    add_body_paragraphs(doc, [
        "Daily rainfall controls water-resources planning, flood and drought risk, agricultural operations, and climate adaptation in Southeast Asia. CMIP6 simulations provide physically based climate information, but precipitation biases remain because of coarse resolution and imperfect representations of convection, topography, and regional circulation [1,2]. These limitations are especially consequential in Prachuap Khiri Khan, where monsoon seasonality, a narrow coastal plain, and the Tenasserim Range generate strongly intermittent and skewed rainfall.",
        "Quantile mapping is widely used to align simulated and observed distributions [3,4,6,7]. Quantile Delta Mapping (QDM) extends this family by preserving modelled relative changes at each quantile [5], while scaled distribution mapping provides a related trend-preserving approach [8]. Method performance nevertheless depends on reference uncertainty, spatial mismatch, dry-day treatment, and the statistical representation of extremes [9,10]. Recent global and multi-region studies likewise show that the preferred method varies with climate, terrain, variable, and impact metric [10,14,15]. Direct holdout comparisons of parametric and empirical QDM, including occurrence and tail sensitivity, remain limited for narrow tropical coastal provinces.",
        "Three unresolved issues motivate the present comparison. First, evaluation in the calibration period measures construction-dependent agreement and can overstate out-of-sample performance. Second, univariate quantile mapping adjusts marginal distributions but cannot reconstruct missing event timing or wet-spell persistence [11]. Third, finite samples make upper-tail behaviour sensitive to extrapolation and parameterization [12,13]. These issues are particularly relevant when monthly stratification improves seasonality but reduces the wet-day sample available for fitting.",
        "This study therefore compares four multiplicative QDM variants: non-parametric annual, non-parametric monthly, parametric annual, and parametric monthly. All observation-dependent quantities are fitted in 1981-2000, frozen, and applied to 2001-2014. The primary endpoint is validation-period monthly Kling-Gupta efficiency (KGE); secondary endpoints characterize distributional fidelity, upper-tail bias, seasonal structure, and temporal correspondence. The objective is to determine which conclusions are transferable, which depend on numerical safeguards, and which uses of the corrected rainfall remain defensible.",
    ])

    base.add_heading(doc, "2", "Materials and Methods")
    base.add_subheading(doc, "2.1", "Study area, observations, and CMIP6 simulations")
    add_body_paragraphs(doc, [
        "Prachuap Khiri Khan is a narrow province between the Gulf of Thailand and the Tenasserim Range. The analysis used daily rainfall at 12 Thai Meteorological Department gauges from 1 January 1981 through 31 December 2014 (12,418 days). The station codes and coordinates in the Hydro-Informatics Institute/Royal Irrigation Department project archive match the official 12-station inventory reported for the province [27]. The retained period is a common 1981-2014 subset of that longer gauge archive.",
        "All stations had complete daily records. Declared missing codes were masked before negative-value handling. Negative values between -0.1 and 0 mm day-1 were set to zero, and values below -0.1 mm day-1 were treated as missing; no such corrections were required in the observation file used here. Annual totals were screened with the Pettitt change-point test at alpha = 0.05 [22]. No station rejected the null hypothesis (p = 0.059-1.000), which is evidence of no detected change point rather than proof of homogeneity. No homogenisation was applied.",
        "The CMIP6 input comprised daily precipitation (pr) from the historical experiment for ACCESS-ESM1-5, CESM2, CanESM5, EC-Earth3, and MIROC6. The source, experiment, member, grid, variable, frequency, and date range are encoded in the analysis-ready filenames and reported in Table 2; model descriptions follow the corresponding publications [16-20]. These identifiers were cross-checked against the federated Earth System Grid Federation (ESGF) catalogue. Because ESGF datasets can be replicated across data nodes, the complete CMIP6 facets are used as the stable source description rather than a transient mirror endpoint. Values were already in mm day-1; the workflow checks units and would multiply flux units in kg m-2 s-1 by 86,400. The station-series files contain two or three distinct model series across the 12 gauges, consistent with nearest native-grid extraction and shared grid cells. A machine-readable manifest records the exact derivative filenames, byte sizes, and SHA-256 checksums used in the analysis.",
        "Standard-calendar models and observations contain 7,305 calibration and 5,113 validation days. CESM2 and CanESM5 use a 365-day calendar, with 7,300 and 5,110 days. Analysis used dates common to each model and the observations; 29 February was omitted for 365-day models and was not interpolated. The 60 model-station combinations are therefore paired evaluation units, not 60 independent spatial experiments, because stations can share model grid cells and the models share components and development lineages.",
    ])
    headers, rows = station_rows()
    base.add_caption(doc, "Table", 1, "Rain-gauge station characteristics for 1981-2014.")
    table(doc, headers, rows, [0.60, 0.52, 0.56, 0.52, 0.60, 0.82, 0.62, 0.82, 0.54], font_size=7.5)
    base.add_para(doc, "Note: n/a indicates that elevation metadata were absent from the supplied station file. SDII is mean rainfall on wet days (>=1.0 mm day-1), and ACF is lag-1 autocorrelation.", size=10, italic=True, after=4)

    base.add_caption(doc, "Table", 2, "CMIP6 historical simulations and metadata recovered from the analysis-ready files.")
    table(doc, ["Model", "Institution", "Variant", "Grid", "Resolution (degrees); approximate area (10^3 km2)", "Calendar", "Days", "Distinct series", "Reference"], model_rows(), [0.82, 0.84, 0.70, 0.40, 0.96, 0.60, 0.52, 0.58, 0.45], font_size=7.2)
    base.add_para(doc, "Note: Cell area is a first-order latitude-longitude approximation at 12 degrees N (111.2 km per degree latitude; 108.8 km per degree longitude), not an exact curvilinear-grid area. Distinct series is the number of unique daily station-series signatures across 12 gauges and indicates shared native grid cells; it is not a formal effective sample size.", size=10, italic=True, after=4)

    base.add_subheading(doc, "2.2", "Split-sample protocol and wet-day adaptation")
    add_body_paragraphs(doc, [
        "Calibration covered 1981-2000 and independent validation covered 2001-2014. Every observed threshold, empirical distribution, and parametric calibration distribution was estimated from calibration data only and then frozen. Validation observations were used solely to calculate evaluation metrics.",
        "Wet-day occurrence and positive intensity were corrected separately. A day was wet when observed rainfall was at least 1.0 mm day-1, consistent with the Expert Team on Climate Change Detection and Indices convention used for wet-day precipitation indices [28]. For each model, station, and calibration group, p_dry was the fraction of observed calibration days below 1.0 mm day-1. The model threshold was the empirical p_dry quantile of all model calibration values. Values at or below that threshold were set to zero; with ties, this rule ensures that the adapted model wet-day count does not exceed the observed count. The threshold was reused unchanged in validation. Annual variants used one threshold per model-station pair, whereas monthly variants used calendar-month-specific thresholds. The full workflow was also rerun at 0.1 and 0.5 mm day-1 to quantify threshold dependence.",
        "A minimum of 30 wet days was required for a monthly transfer function. Of 720 model-station-month blocks, 100 dry-season blocks did not meet this threshold. These blocks were not discarded or set to zero; they used the corresponding annual transfer function fitted from calibration data. This fallback was declared uniformly and recorded in the diagnostics.",
    ])

    base.add_subheading(doc, "2.3", "Multiplicative QDM")
    base.add_para(doc, "For a model value x at time t in target period p, multiplicative QDM followed Cannon et al. [5]:", first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_equation(doc, "tau(t) = F_m,p[x_m,p(t)]", 1)
    add_equation(doc, "Delta(t) = x_m,p(t) / F_m,h^(-1)[tau(t)]", 2)
    add_equation(doc, "x_hat(t) = F_o,h^(-1)[tau(t)] x Delta(t)", 3)
    add_body_paragraphs(doc, [
        "Here F_m,p is the model cumulative distribution in the target period; F_m,h^(-1) and F_o,h^(-1) are model and observed calibration-period quantile functions; tau is non-exceedance probability; Delta is the modelled relative-change ratio; and x_hat is corrected rainfall. The target-period probability uses model values only and therefore does not violate the observation holdout. Retaining Delta distinguishes QDM from ordinary quantile mapping.",
        "Delta was bounded a priori to [0.2, 5.0] to prevent division by very small historical quantiles. For the primary parametric analysis, corrected rainfall was capped at five times the maximum observed wet-day rainfall in the corresponding calibration block. Both bounds are numerical safeguards, not hydrological return-level thresholds. The cap was applied after transformation. Sensitivity runs used a factor of 1.5 and no cap; all clipping and capping events were counted.",
    ])

    base.add_subheading(doc, "2.4", "Parametric and non-parametric variants")
    add_body_paragraphs(doc, [
        "Non-parametric variants used empirical cumulative distribution functions with plotting positions (i - 0.5)/n and linear interpolation. Annual variants pooled all months, whereas monthly variants fitted each calendar month separately.",
        "Parametric variants fitted gamma, generalized extreme value, Weibull minimum, lognormal, and Pearson type III candidates to wet-day intensities by maximum likelihood. Gamma, Weibull, and lognormal location parameters were fixed at zero; the generalized extreme value and Pearson type III location parameters were estimated. AIC = 2k - 2 ln(L), where k is the number of freely estimated parameters and L is maximized likelihood. The minimum-AIC candidate was selected deterministically; non-finite parameters or likelihoods disqualified a candidate. As a small-sample sensitivity, selection and the complete correction workflow were repeated with AICc = AIC + 2k(k + 1)/(n - k - 1). The Kolmogorov-Smirnov statistic was stored as a fit diagnostic. The generalized extreme value distribution was used as a flexible whole-distribution candidate, not as a claim that daily wet-day rainfall is a block-maximum sample. Calculations used Python 3.12.13 and SciPy 1.18.1 [24].",
    ])

    base.add_subheading(doc, "2.5", "Evaluation endpoints and uncertainty")
    add_body_paragraphs(doc, [
        "The primary endpoint was monthly KGE [21], selected because it combines correlation, variability, and mean bias. Secondary endpoints were the Kolmogorov-Smirnov statistic for wet-day distributions, absolute relative bias of the 99th wet-day percentile, monthly RMSE, monthly correlation, and monthly-climatology correlation. KGE and correlations are better when larger; the remaining endpoints are better when smaller. Direction-based rankings of means are labelled descriptive and do not imply statistical superiority; discordant secondary endpoints are reported separately.",
        "Uncertainty in corrected-minus-raw differences was estimated separately for each model-station combination with 2,000 paired year-block bootstrap replicates and seed 42. Validation monthly totals were arranged as a 14-year by 12-month matrix. Fourteen years were sampled with replacement in each replicate, and the same year indices were applied to observations, raw simulations, and all corrected variants. The 95% interval was the 2.5th and 97.5th percentiles. Two-sided p values used a plus-one correction, and Benjamini-Hochberg adjustment controlled the false-discovery rate at 0.05 within each method-metric family [23]. Because combinations are spatially and structurally dependent, network summaries remain descriptive and are not treated as 60 independent experiments.",
    ])
    add_figure(doc, FIGURES / "Figure1_framework.png", 1, "Fit-freeze-apply framework. Observation-dependent transfer functions were estimated in 1981-2000 and applied without re-estimation in 2001-2014; uncertainty used paired year-block bootstrap resampling.")

    base.add_heading(doc, "3", "Results and Discussion")
    base.add_subheading(doc, "3.1", "Observed rainfall and model context")
    add_body_paragraphs(doc, [
        "Mean annual observed rainfall ranged from 944 to 1,376 mm among stations, and maximum daily rainfall ranged from 54.2 to 298.5 mm day-1. Calibration and validation retained the same broad bimodal seasonal cycle, but several stations changed wet-day frequency between periods (Figure 2). This non-stationarity makes the holdout test more demanding than same-period evaluation.",
        "The five model files represented only two or three distinct native-grid series across the 12 stations (Table 2). Consequently, averaging across all 60 combinations would overstate spatial replication. Results are therefore reported as paired model-station distributions and counts, with this dependence carried into interpretation.",
    ])
    add_figure(doc, FIGURES / "Figure2_observed_periods.png", 2, "Observed rainfall in calibration and validation: (a) monthly climatology with interquartile ribbons across stations, (b) annual rainfall, and (c) paired station-level wet-day frequencies. Colored connecting lines show the direction of change; these summaries are descriptive and motivate the independent holdout test.")

    base.add_subheading(doc, "3.2", "Independent validation performance")
    add_body_paragraphs(doc, [
        "Month-wise calibration produced the clearest descriptive gains (Table 3; Figure 3). P monthly had the most favorable mean primary endpoint, increasing monthly KGE from 0.229 to 0.337, and the lowest monthly RMSE, decreasing it from 116.20 to 105.06 mm. NP monthly was close, with KGE 0.329 and RMSE 109.45 mm. Both monthly methods increased monthly-climatology correlation to at least 0.864 and produced a favorable point change for that metric in 59 of 60 paired evaluations.",
        "The paired bootstrap did not support a broad parametric advantage over the empirical monthly method. For P monthly minus NP monthly, the median combination-level KGE difference was 0.0002, with 30 combinations favouring each method. The median RMSE difference was -1.45 mm, with P monthly lower in 36 combinations and NP monthly lower in 24. Thus, the aggregate ranking is driven by heterogeneous station-model responses rather than uniform dominance.",
        "Upper-tail performance followed a different ordering. NP annual reduced mean absolute q99 bias from 52.98% to 24.79%, compared with 27.38% for NP monthly and 32.54% for P monthly. Pooling months provides a larger empirical wet-day sample and therefore more stable upper quantiles, whereas month-wise fitting directly resolves seasonality at the cost of smaller samples. No method was best on every endpoint.",
        "Annual variants were weaker for monthly performance: NP annual increased mean RMSE to 124.74 mm and reduced mean KGE to 0.203, while P annual yielded 123.72 mm and 0.152. Distributional correction alone therefore does not guarantee better month-to-month agreement. This is consistent with the limitation that univariate mapping cannot reconstruct event sequencing absent from free-running climate simulations [11].",
        "Station-aggregated responses were heterogeneous (Figure 5). Mean corrected-minus-raw KGE was positive at 9 of 12 stations for NP monthly and 11 of 12 for P monthly. Stations 500008 and 500009 showed the largest favorable KGE changes (approximately 0.29-0.31), whereas station 500002 decreased under both methods. The station-by-model display is descriptive: the small network, missing elevation metadata, and shared model cells do not support a causal geographic interpretation.",
    ])
    base.add_caption(doc, "Table", 3, "Independent-validation means across paired model-station evaluations.")
    table(doc, ["Metric", "Raw", "NP annual", "NP monthly", "P annual", "P monthly", "Favorable mean (descriptive)"], validation_rows(), [1.36, 0.54, 0.69, 0.72, 0.63, 0.67, 1.02], font_size=7.5)
    base.add_para(doc, "Note: PBIAS and q99 bias are summarized as absolute values. The final column is a direction-based descriptive ranking and does not imply statistical significance; the 60 pairs are not independent spatial replicates.", size=10, italic=True, after=4)

    base.add_caption(doc, "Table", 4, "Paired year-block bootstrap summary for corrected-minus-raw monthly metrics.")
    table(doc, ["Method", "Metric", "Raw mean", "Corrected mean", "Median delta", "Median 95% interval", "Favorable point change", "CI F/A/O"], bootstrap_rows(), [0.72, 0.48, 0.58, 0.72, 0.68, 1.05, 0.78, 1.00], font_size=7.3)
    base.add_para(doc, "Note: Interval bounds are medians of 60 pair-specific bootstrap limits, not a confidence interval for a network-wide effect. For RMSE, negative delta is favorable; for KGE, positive delta is favorable. CI F/A/O denotes favorable, adverse, and overlapping-zero intervals; counts are descriptive because pairs are dependent.", size=10, italic=True, after=4)

    add_figure(doc, FIGURES / "Figure3_method_comparison.png", 3, "Validation distributions for 60 model-station combinations per variant. Boxes show the interquartile range, whiskers show the 10th-90th percentiles, diamonds show means, centre lines show medians, and dashed vertical lines show the raw mean in each panel; these are not bootstrap distributions.")
    add_figure(doc, FIGURES / "Figure4_improvement_heatmap.png", 4, "Percentage of 60 paired model-station evaluations with a favorable point change relative to raw CMIP6 simulations. Values above 50% indicate a favorable change for most pairs; percentages are descriptive and pairs are not independent spatial replicates.")
    add_figure(doc, FIGURES / "Figure5_station_spatial_response.png", 5, "Station-level response of monthly QDM. (a) Heatmap of corrected-minus-raw KGE for NP monthly QDM by station and CMIP6 model. (b) Paired station means for non-parametric and parametric monthly QDM. Values are descriptive because gauges share model grid cells; station patterns are not interpreted as causal geographic effects.")

    base.add_subheading(doc, "3.3", "Parametric selection, fallback, and tail sensitivity")
    add_body_paragraphs(doc, [
        "Maximum-likelihood fitting changed the distribution-selection pattern relative to the earlier method-of-moments implementation. Pearson type III was selected for 55 of 60 observed annual fits and 610 of 620 eligible observed monthly fits. Model-side fits were more diverse: Pearson type III accounted for 496 monthly fits, generalized extreme value for 106, lognormal for 12, gamma for 3, and Weibull minimum for 3 (Table 5). All 620 eligible monthly blocks fitted successfully. The apparent deficits of 5 observed and 33 model fits in the earlier table were omitted gamma and Weibull categories, not optimization failures.",
        "One hundred monthly blocks were ineligible because at least one calibration wet-day sample contained fewer than 30 values. These blocks clustered in dry-season months and used the annual fallback. This distinction is essential: the structural total remains 720, selection counts sum to the 620 eligible monthly fits, and fallback counts account for the remaining 100.",
        "Tail safeguards affected the parametric ranking (Table 6). With the primary factor of 5.0, P monthly KGE was 0.337 versus 0.329 for NP monthly; without a cap, P monthly KGE decreased to 0.327. A stricter 1.5 factor capped 171 P-monthly validation values and improved mean RMSE to 101.91 mm, illustrating that aggressive truncation can improve central error metrics while imposing a strong assumption on extremes. The cap sensitivity therefore precludes a universal claim that parametric monthly QDM is superior.",
        "The q99 metric changed little across the three parametric cap settings because most capped values were above the station-level 99th percentile. This does not make the cap inconsequential: a small number of altered maxima can materially affect return-level, IDF, or flood applications. Recent work similarly emphasizes that tail handling can behave differently inside and outside calibration and must be evaluated explicitly [12,13].",
        "AICc produced exactly the same selections and validation metrics as AIC: observed fits selected Pearson type III in 55 of 60 annual and 610 of 620 eligible monthly blocks, while model-side monthly counts remained 496 Pearson type III, 106 generalized extreme value, 12 lognormal, 3 gamma, and 3 Weibull. This rules out the uncorrected AIC penalty as the cause of the observed selection pattern for these samples, but it does not establish a physical rainfall law; Pearson type III retained greater shape and location flexibility than several constrained candidates.",
        "Wet-day threshold sensitivity was method-dependent (Figure 6). NP monthly remained stable across 0.1, 0.5, and 1.0 mm day-1 (KGE 0.332, 0.330, and 0.329; RMSE 109.01, 109.26, and 109.45 mm). P monthly changed more strongly: KGE was 0.217, 0.327, and 0.337, while absolute q99 bias was 49.36%, 45.87%, and 32.54%. Thus, the 1.0-mm convention is defensible and gives the most favorable tested parametric result, but the parametric ranking is conditional on how trace rainfall is classified.",
    ])
    base.add_caption(doc, "Table", 5, "Parametric fit accounting and numerical safeguards in validation.")
    table(doc, ["Method", "Eligible/structural", "Observed selection", "Model selection", "Annual fallback", "Fit failures", "Capped values", "Clipped delta events"], diagnostic_rows(), [0.62, 0.72, 1.28, 1.65, 0.68, 0.58, 0.56, 0.64], font_size=7.1)
    base.add_para(doc, "Note: P3 = Pearson type III; GEV = generalized extreme value; LN = lognormal. Monthly selection counts sum to 620 eligible blocks. The 100 ineligible blocks used annual calibration-only transfer functions. Capped values count validation outputs truncated above the selected cap. Clipped delta events are days on which the delta ratio was constrained at the lower or upper bound [0.2, 5.0]; event counts are not numbers of fitted blocks.", size=10, italic=True, after=4)

    base.add_caption(doc, "Table", 6, "Sensitivity of parametric validation metrics to the upper-tail cap factor.")
    table(doc, ["Cap factor", "P annual RMSE", "P annual KGE", "P monthly RMSE", "P monthly KGE", "P monthly |q99 bias| (%)", "P annual capped", "P monthly capped"], sensitivity_rows(), [0.70, 0.75, 0.66, 0.78, 0.68, 1.02, 0.70, 0.72], font_size=7.4)
    base.add_para(doc, "Note: The cap is factor times the maximum observed wet-day rainfall in the relevant calibration block and is a numerical safeguard rather than a return-level threshold. None denotes the uncapped sensitivity run. NP results are invariant to this parameter.", size=10, italic=True, after=4)
    add_figure(doc, FIGURES / "Figure6_wet_threshold_sensitivity.png", 6, "Sensitivity of monthly QDM validation metrics to the observed wet-day threshold. Points and lines are means across 60 paired model-station evaluations; translucent ribbons show the interquartile range across those pairs, and the dashed line shows the raw mean. Lower is favorable for RMSE and absolute q99 bias, whereas higher is favorable for KGE. Lines aid comparison and do not imply continuous interpolation.")

    base.add_subheading(doc, "3.4", "Implications and limitations")
    add_body_paragraphs(doc, [
        "The model-specific raw results illuminate why a single correction function was insufficient. ACCESS-ESM1-5 and MIROC6 overestimated wet-day frequency by 0.227 and 0.232, respectively, and had positive mean PBIAS of 18.7% and 25.3%. EC-Earth3 instead had a negative mean PBIAS of -15.9% but the lowest raw monthly RMSE (92.30 mm) among the five models. Thus, similar aggregate errors arose from different combinations of event frequency, intensity, and seasonal phase. This heterogeneity agrees with Southeast Asian CMIP6 evaluations showing that rainfall skill reflects both monsoon circulation and model-specific parameterization, and that good large-scale circulation does not guarantee local rainfall skill [25,26].",
        "A physically plausible interpretation follows from the province's narrow coastal geometry. Moist southwesterly flow encounters the Tenasserim Range along the western boundary, while Gulf moisture and the northeast monsoon contribute to late-year rainfall on the eastern coastal plain [27]. The native cells span approximately 13,600-94,900 km2 at 12 degrees N, so even the finest model cell is large relative to the province's mountain-coast transition. The especially large wet-frequency excess in ACCESS-ESM1-5 and MIROC6 is consistent with overly frequent resolved or parameterized rainfall, whereas EC-Earth3's negative volume bias suggests a different compensation between frequency and intensity. Grid size alone cannot explain this ordering because model physics, circulation, and convection schemes also differ. These are process-informed interpretations, not causal attribution, because circulation and moisture-flux diagnostics were not analysed.",
        "The most stable result is the advantage of month-wise calibration for seasonal and monthly metrics. Annual pooling combines dry-season occurrence, southwest-monsoon rainfall, and the late-year coastal maximum in one transfer function. Month-wise QDM instead estimates separate wet-day thresholds and intensity distributions for each phase of the seasonal cycle, which better respects the changing atmospheric regime without claiming to simulate its dynamics. Both monthly methods consequently improved RMSE, KGE, monthly correlation, and climatology more consistently than annual calibration. The countervailing cost is a roughly twelve-fold reduction in sample size per transfer function, which explains the dry-season fallback and contributes to upper-tail uncertainty.",
        "Method choice should follow application needs. NP monthly is an assumption-lean default for seasonal reservoir inflow, irrigation scheduling, and monthly drought or water-balance screening. NP annual is preferable when marginal q99 fidelity is the priority, although q99 is not a return-level estimator. P monthly can reduce monthly error further, but its advantage is conditional on distribution, wet-day threshold, and tail constraint. Fixed caps are unsuitable for design-storm, IDF, or flood-frequency estimation; such uses require uncapped sensitivity and dedicated extreme-value analysis. None of the variants restores observed daily event timing, spatial coherence, or wet-spell persistence.",
        "Four limitations remain. The study covers one province, 12 gauges, five historical runs, and a 14-year validation window. Shared native grid cells reduce spatial replication, and the analysis-ready CMIP6 products are station-matched derivatives rather than original NetCDF files with per-file tracking identifiers. The reported facets and checksum manifest identify the exact derivatives used, while independent recreation requires repeating the spatial extraction from the corresponding ESGF datasets. No future SSP experiment was analysed, so preservation of projected quantile change was implemented by the QDM formulation but not empirically evaluated. Future work should compare QDM directly with ordinary quantile mapping and scaled distribution mapping under the same holdout, test moving-window seasonal fits and explicit extreme-value tails, analyse circulation and spatial dependence, and evaluate regional or dynamically downscaled ensembles across future scenarios.",
    ])

    base.add_heading(doc, "4", "Conclusions")
    add_body_paragraphs(doc, [
        "Independent temporal validation indicates that month-wise QDM gives the most favorable descriptive pattern for monthly rainfall performance in Prachuap Khiri Khan. P monthly had the highest mean KGE and lowest mean RMSE under the five-times cap, but its KGE advantage over NP monthly was negligible in paired comparisons and reversed without the cap. AICc reproduced AIC selections, whereas lowering the wet-day threshold to 0.1 mm reduced P-monthly KGE to 0.217 while NP monthly remained stable. Seasonal stratification, rather than a universally superior distribution family, is therefore the principal transferable finding.",
        "NP annual gave the lowest q99 bias, whereas monthly methods best represented seasonal structure. For monthly water-balance applications, NP monthly is the conservative default; NP annual can support upper-quantile screening; and P monthly requires documented occurrence, fallback, clipping, and cap sensitivity. Fixed-cap outputs should not be used directly for design extremes. The 60 paired evaluations share grid cells and model lineages and are not independent spatial replicates, so network-wide inferential superiority is not claimed. The fit-freeze-apply protocol and explicit sensitivities provide a stronger basis for further regional and future-scenario evaluation than same-period fitting alone.",
    ])

    base.add_heading(doc, "5", "Acknowledgements")
    base.add_para(doc, "The author acknowledges the institutions that generated and distributed the CMIP6 simulations and the custodians of the Thai rain-gauge data used in the supplied project archive.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    base.add_mixed_para(doc, "Author Contributions: ", "S.P.: conceptualization, methodology, software, validation, formal analysis, data curation, visualization, writing-original draft, and writing-review and editing.")
    base.add_mixed_para(doc, "Funding: ", "This research received no external funding.")
    base.add_mixed_para(doc, "Data Availability Statement: ", "The rainfall observations are Thai Meteorological Department gauge records distributed through the Hydro-Informatics Institute/Royal Irrigation Department project archive; the official station inventory is documented in [27]. Access requests should be directed to the responsible Thai data custodians, and redistribution remains subject to their terms. CMIP6 source datasets are discoverable through ESGF using the complete facets reported in Table 2. The complete reproducibility package, including the analysis code, input-file checksum manifest, derived metrics, diagnostics, sensitivity outputs, and figure-generation scripts, is available from the corresponding author upon reasonable request. No public repository DOI had been assigned at the time of manuscript preparation.")
    base.add_mixed_para(doc, "Conflicts of Interest: ", "The author declares no conflict of interest.")

    base.add_para(doc, "References", size=12, bold=True, before=8, after=2)
    references = [
        "Eyring, V.; Bony, S.; Meehl, G.A.; Senior, C.A.; Stevens, B.; Stouffer, R.J.; Taylor, K.E. Overview of the Coupled Model Intercomparison Project Phase 6 (CMIP6) Experimental Design and Organization. Geosci. Model Dev. 2016, 9, 1937-1958. https://doi.org/10.5194/gmd-9-1937-2016.",
        "Tebaldi, C.; Knutti, R. The Use of the Multi-Model Ensemble in Probabilistic Climate Projections. Philos. Trans. R. Soc. A 2007, 365, 2053-2075. https://doi.org/10.1098/rsta.2007.2076.",
        "Maraun, D. Bias Correction, Quantile Mapping, and Downscaling: Revisiting the Inflation Issue. J. Clim. 2013, 26, 2137-2143. https://doi.org/10.1175/JCLI-D-12-00821.1.",
        "Themeßl, M.J.; Gobiet, A.; Leuprecht, A. Empirical-Statistical Downscaling and Error Correction of Daily Precipitation from Regional Climate Models. Int. J. Climatol. 2011, 31, 1530-1544. https://doi.org/10.1002/joc.2168.",
        "Cannon, A.J.; Sobie, S.R.; Murdock, T.Q. Bias Correction of GCM Precipitation by Quantile Mapping: How Well Do Methods Preserve Changes in Quantiles and Extremes? J. Clim. 2015, 28, 6938-6959. https://doi.org/10.1175/JCLI-D-14-00754.1.",
        "Gudmundsson, L.; Bremnes, J.B.; Haugen, J.E.; Engen-Skaugen, T. Technical Note: Downscaling RCM Precipitation to the Station Scale Using Statistical Transformations. Hydrol. Earth Syst. Sci. 2012, 16, 3383-3390. https://doi.org/10.5194/hess-16-3383-2012.",
        "Piani, C.; Haerter, J.O.; Coppola, E. Statistical Bias Correction for Daily Precipitation in Regional Climate Models over Europe. Theor. Appl. Climatol. 2010, 99, 187-192. https://doi.org/10.1007/s00704-009-0134-9.",
        "Switanek, M.B.; Troch, P.A.; Castro, C.L.; Leuprecht, A.; Chang, H.I.; Mukherjee, R.; Demaria, E.M.C. Scaled Distribution Mapping: A Bias Correction Method That Preserves Raw Climate Model Projected Changes. Hydrol. Earth Syst. Sci. 2017, 21, 2649-2666. https://doi.org/10.5194/hess-21-2649-2017.",
        "Casanueva, A.; Herrera, S.; Iturbide, M.; Lange, S.; Jury, M.; Dosio, A.; Maraun, D.; Gutiérrez, J.M. Testing Bias Adjustment Methods for Regional Climate Change Applications under Observational Uncertainty and Resolution Mismatch. Atmos. Sci. Lett. 2020, 21, e978. https://doi.org/10.1002/asl.978.",
        "Spuler, F.R.; Wessel, J.B.; Comyn-Platt, E.; Varndell, J.; Cagnazzo, C. ibicus: A New Open-Source Python Package and Comprehensive Interface for Statistical Bias Adjustment and Evaluation in Climate Modelling (v1.0.1). Geosci. Model Dev. 2024, 17, 1249-1269. https://doi.org/10.5194/gmd-17-1249-2024.",
        "Potter, N.J.; Chiew, F.H.S.; Charles, S.P.; Fu, G.; Zheng, H.; Zhang, L. Bias in Dynamically Downscaled Rainfall Characteristics for Hydroclimatic Projections. Hydrol. Earth Syst. Sci. 2020, 24, 2963-2979. https://doi.org/10.5194/hess-24-2963-2020.",
        "Berg, P.; Bosshard, T.; Bozhinova, D.; Bärring, L.; Löw, J.; Nilsson, C.; Strandberg, G.; Södling, J.; Thuresson, J.; Wilcke, R.; et al. Robust Handling of Extremes in Quantile Mapping - Murder Your Darlings. Geosci. Model Dev. 2024, 17, 8173-8179. https://doi.org/10.5194/gmd-17-8173-2024.",
        "Padulano, R.; Gomez-Mogollon, L.A.; Napolitano, L.; Rianna, G. Quantile-Based Bias-Correction of Extreme Rainfall: Pros and Cons of Popular Methods for Climate Signal Preservation. J. Hydrol. 2025, 653, 132814. https://doi.org/10.1016/j.jhydrol.2025.132814.",
        "Song, Y.H.; Chung, E.-S. Intercomparison of Bias Correction Methods for Precipitation of Multiple GCMs across Six Continents. Geosci. Model Dev. 2025, 18, 8017-8045. https://doi.org/10.5194/gmd-18-8017-2025.",
        "Gergel, D.R.; Malevich, S.B.; McCusker, K.E.; Tenezakis, E.; Delgado, M.T.; Fish, M.A.; Kopp, R.E. Global Downscaled Projections for Climate Impacts Research: Preserving Quantile Trends for Modeling Future Climate Impacts. Geosci. Model Dev. 2024, 17, 191-227. https://doi.org/10.5194/gmd-17-191-2024.",
        "Swart, N.C.; Cole, J.N.S.; Kharin, V.V.; Lazare, M.; Scinocca, J.F.; Gillett, N.P.; Anstey, J.; Arora, V.; Christian, J.R.; Jiao, Y.; et al. The Canadian Earth System Model Version 5 (CanESM5.0.3). Geosci. Model Dev. 2019, 12, 4823-4873. https://doi.org/10.5194/gmd-12-4823-2019.",
        "Danabasoglu, G.; Lamarque, J.-F.; Bacmeister, J.; Bailey, D.A.; DuVivier, A.K.; Edwards, J.; Emmons, L.K.; Fasullo, J.; Garcia, R.; Gettelman, A.; et al. The Community Earth System Model Version 2 (CESM2). J. Adv. Model Earth Syst. 2020, 12, e2019MS001916. https://doi.org/10.1029/2019MS001916.",
        "Döscher, R.; Acosta, M.; Alessandri, A.; Anthoni, P.; Arneth, A.; Arsouze, T.; Bergmann, T.; Bernadello, R.; Boussetta, S.; Caron, L.P.; et al. The EC-Earth3 Earth System Model for the Coupled Model Intercomparison Project 6. Geosci. Model Dev. 2022, 15, 2973-3020. https://doi.org/10.5194/gmd-15-2973-2022.",
        "Tatebe, H.; Ogura, T.; Nitta, T.; Komuro, Y.; Ogochi, K.; Takemura, T.; Sudo, K.; Sekiguchi, M.; Abe, M.; Saito, F.; et al. Description and Basic Evaluation of Simulated Mean State, Internal Variability, and Climate Sensitivity in MIROC6. Geosci. Model Dev. 2019, 12, 2727-2765. https://doi.org/10.5194/gmd-12-2727-2019.",
        "Ziehn, T.; Chamberlain, M.A.; Law, R.M.; Lenton, A.; Bodman, R.W.; Dix, M.; Stevens, L.; Wang, Y.-P.; Srbinovsky, J. The Australian Earth System Model: ACCESS-ESM1.5. J. South. Hemisph. Earth Syst. Sci. 2020, 70, 193-214. https://doi.org/10.1071/ES19035.",
        "Gupta, H.V.; Kling, H.; Yilmaz, K.K.; Martinez, G.F. Decomposition of the Mean Squared Error and NSE Performance Criteria: Implications for Improving Hydrological Modelling. J. Hydrol. 2009, 377, 80-91. https://doi.org/10.1016/j.jhydrol.2009.08.003.",
        "Pettitt, A.N. A Non-Parametric Approach to the Change-Point Problem. Appl. Stat. 1979, 28, 126-135. https://doi.org/10.2307/2346729.",
        "Benjamini, Y.; Hochberg, Y. Controlling the False Discovery Rate: A Practical and Powerful Approach to Multiple Testing. J. R. Stat. Soc. Ser. B 1995, 57, 289-300. https://doi.org/10.1111/j.2517-6161.1995.tb02031.x.",
        "Virtanen, P.; Gommers, R.; Oliphant, T.E.; Haberland, M.; Reddy, T.; Cournapeau, D.; Burovski, E.; Peterson, P.; Weckesser, W.; Bright, J.; et al. SciPy 1.0: Fundamental Algorithms for Scientific Computing in Python. Nat. Methods 2020, 17, 261-272. https://doi.org/10.1038/s41592-019-0686-2.",
        "Khadka, D.; Babel, M.S.; Abatan, A.A.; Collins, M. An Evaluation of CMIP5 and CMIP6 Climate Models in Simulating Summer Rainfall in the Southeast Asian Monsoon Domain. Int. J. Climatol. 2022, 42, 1181-1202. https://doi.org/10.1002/joc.7296.",
        "Liu, S.; Raghavan, S.V.; Ona, B.J.; Nguyen, N.S. Bias Evaluation in Rainfall over Southeast Asia in CMIP6 Models. J. Hydrol. 2023, 621, 129593. https://doi.org/10.1016/j.jhydrol.2023.129593.",
        "Pimsing, P.; Paengmun, S.; Thuengsaeng, A.; Wattanangkun, P.; Kasibut, S.; Sonphan, S. A Study on the Variability of Rainfall and Effective Rainfall under Climate Change of Prachuap Khiri Khan Province. In Proceedings of the 18th THAICID National Symposium, Thailand, 7 July 2025. Official event page: https://thaicid.rid.go.th/announcement/31/18th%20THAICID%20National%20Symposium; full paper: https://thaicid.rid.go.th/public/files/proposal/2025/fullpaper/1752739614_a49127fa0ffaf725ff44.pdf (accessed 29 August 2026).",
        "Zhang, X.; Alexander, L.; Hegerl, G.C.; Jones, P.D.; Tank, A.K.; Peterson, T.C.; Trewin, B.; Zwiers, F.W. Indices for Monitoring Changes in Extremes Based on Daily Temperature and Precipitation Data. Wiley Interdiscip. Rev. Clim. Change 2011, 2, 851-870. https://doi.org/10.1002/wcc.147.",
    ]
    for index, reference in enumerate(references, 1):
        base.add_para(doc, f"{index}. {reference}", size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=1)

    apply_document_font(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
