from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def text_of(node: etree._Element) -> str:
    return "".join(node.xpath(".//w:t/text() | .//w:delText/text()", namespaces=NS)).strip()


def inspect(path: Path) -> dict:
    doc = Document(path)
    result = {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "section_count": len(doc.sections),
        "paragraphs": [],
        "tables": [],
        "tracked_insertions": [],
        "tracked_deletions": [],
        "comments": [],
    }
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text = paragraph.text.replace("\n", " ").strip()
        if text:
            result["paragraphs"].append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else None,
                    "text": text,
                }
            )
    for table_index, table in enumerate(doc.tables, 1):
        result["tables"].append(
            {
                "index": table_index,
                "rows": [
                    [cell.text.replace("\n", " / ").strip() for cell in row.cells]
                    for row in table.rows
                ],
            }
        )

    with zipfile.ZipFile(path) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        result["tracked_insertions"] = [
            text_of(node) for node in document_xml.xpath(".//w:ins", namespaces=NS) if text_of(node)
        ]
        result["tracked_deletions"] = [
            text_of(node) for node in document_xml.xpath(".//w:del", namespaces=NS) if text_of(node)
        ]
        if "word/comments.xml" in archive.namelist():
            comments_xml = etree.fromstring(archive.read("word/comments.xml"))
            for node in comments_xml.xpath(".//w:comment", namespaces=NS):
                result["comments"].append(
                    {
                        "id": node.get(f"{{{NS['w']}}}id"),
                        "author": node.get(f"{{{NS['w']}}}author"),
                        "text": text_of(node),
                    }
                )
    return result


def main() -> None:
    if len(sys.argv) < 3:
        raise SystemExit("usage: inspect_revision_docs.py OUT.json INPUT.docx [INPUT.docx ...]")
    output = Path(sys.argv[1])
    payload = [inspect(Path(value)) for value in sys.argv[2:]]
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
