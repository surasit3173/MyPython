from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "journal_revision_q3" / "AJSTR_QDM_CMIP6_Q2Q3_REVISED_FINAL.docx"


def main() -> None:
    assert FINAL.is_file() and FINAL.stat().st_size > 1_000_000
    doc = Document(FINAL)
    assert len(doc.sections) == 1
    assert len(doc.tables) == 6
    text = "\n".join(p.text for p in doc.paragraphs)
    assert not any(token in text for token in ("___", "PLACEHOLDER", "ตามที่คุณ", "เอกสารแนะนำ", "118.68"))
    abstract = next(p.text for p in doc.paragraphs if p.text.startswith("Abstract:"))
    assert len(abstract.split()) <= 250
    references_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == "References")
    assert "Reviewers suggestion" not in text
    header_text = "\n".join(
        cell.text
        for section in doc.sections
        for header in (section.header, section.first_page_header, section.even_page_header)
        for table in header.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Research Article" in header_text
    assert "Article type" not in header_text
    references = [p.text for p in doc.paragraphs[references_index + 1:] if re.match(r"^\d+\. ", p.text)]
    assert len(references) == 28
    assert sum("https://doi.org/" in reference for reference in references) == 27
    assert "thaicid.rid.go.th" in references[-2]
    assert "10.1002/wcc.147" in references[-1]

    with zipfile.ZipFile(FINAL) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        assert xml.count("<m:oMath>") == 3
        assert xml.count("<wp:inline") == 6
        font_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
        assert "Palatino Linotype" in font_xml
        assert "Angsana" not in font_xml
        embedded_media = [archive.read(name) for name in archive.namelist() if name.startswith("word/media/")]

    figure_dir = ROOT / "journal_revision_q3" / "figures_q1"
    stems = [
        "Figure1_framework",
        "Figure2_observed_periods",
        "Figure3_method_comparison",
        "Figure4_improvement_heatmap",
        "Figure5_station_spatial_response",
        "Figure6_wet_threshold_sensitivity",
    ]
    for stem in stems:
        png = figure_dir / f"{stem}.png"
        tif = figure_dir / f"{stem}.tif"
        pdf = figure_dir / f"{stem}.pdf"
        assert png.read_bytes() in embedded_media
        with Image.open(png) as image:
            assert image.width >= 3700
            assert image.info.get("dpi", (0, 0))[0] >= 599
        with Image.open(tif) as image:
            assert image.width >= 3700
            assert image.info.get("dpi", (0, 0))[0] >= 599
            assert image.info.get("compression") == "tiff_lzw"
        assert pdf.read_bytes().startswith(b"%PDF")

    a11y = json.loads((ROOT / "journal_revision_q3" / "a11y_revised_final.json").read_text(encoding="utf-8"))
    assert a11y["counts"] == {"high": 0, "medium": 0, "low": 0}

    summary = pd.read_csv(ROOT / "qdm_p_np_publication_out_q3_final" / "summary.csv")
    validation = summary[summary.period.eq("validation")].set_index(["method", "metric"])
    checks = {
        ("QDM_P_monthly", "mRMSE"): 105.0599,
        ("QDM_P_monthly", "mKGE"): 0.3368,
        ("QDM_NP_monthly", "mKGE"): 0.3293,
        ("QDM_NP_annual", "q99_relbias_pct"): 24.785,
    }
    for key, expected in checks.items():
        actual = float(validation.loc[key, "corrected_mean"])
        assert abs(actual - expected) < 0.001

    threshold = pd.read_csv(ROOT / "journal_revision_q3" / "revision2_analysis" / "wet_threshold_sensitivity.csv")
    np_monthly = threshold[(threshold.method.eq("QDM_NP_monthly")) & (threshold.metric.eq("mKGE"))].set_index("scenario")
    p_monthly = threshold[(threshold.method.eq("QDM_P_monthly")) & (threshold.metric.eq("mKGE"))].set_index("scenario")
    assert abs(float(np_monthly.loc["0.1 mm", "corrected_mean"]) - 0.3322) < 0.001
    assert abs(float(p_monthly.loc["0.1 mm", "corrected_mean"]) - 0.2173) < 0.001

    aicc = pd.read_csv(ROOT / "journal_revision_q3" / "revision2_analysis" / "aicc_validation_summary.csv")
    aicc_p = aicc[(aicc.method.eq("QDM_P_monthly")) & (aicc.metric.eq("mKGE"))]
    assert abs(float(aicc_p.corrected_mean.iloc[0]) - 0.3368) < 0.001

    calendar = pd.read_csv(ROOT / "journal_revision_q3" / "revision2_analysis" / "calendar_audit.csv")
    assert calendar.duplicates.eq(0).all()
    assert calendar.monotonic.all()

    inspection = json.loads((ROOT / "journal_revision_q3" / "final_inspection_revised.json").read_text(encoding="utf-8"))[0]
    assert not inspection["tracked_insertions"]
    assert not inspection["tracked_deletions"]
    assert not inspection["comments"]
    print(
        "PASS: final DOCX uses Palatino Linotype; 6 tables; 6 verified 600-dpi body figures; 3 Word equations; "
        "28 references (27 DOI + official RID source); abstract <=250 words; a11y 0/0/0; no "
        "placeholders, comments, or tracked changes; primary, threshold, AICc, and calendar values verified."
    )


if __name__ == "__main__":
    main()
