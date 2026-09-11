from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


DOCX = Path(
    r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\journal_revision_q3"
    r"\final_revision\AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN_FINAL.docx"
)
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W, "m": M, "wp": WP}


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))


def expand_citation(content: str) -> list[int]:
    values: list[int] = []
    for token in content.split(","):
        token = token.strip()
        if not token:
            continue
        pieces = re.split(r"[-\u2013]", token)
        if len(pieces) == 2:
            values.extend(range(int(pieces[0]), int(pieces[1]) + 1))
        else:
            values.append(int(token))
    return values


def main() -> None:
    doc = Document(DOCX)
    with zipfile.ZipFile(DOCX) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))

    paragraphs = root.xpath("//w:body//w:p", namespaces=NS)
    texts = [paragraph_text(paragraph) for paragraph in paragraphs]
    references_index = texts.index("References")
    body_text = "\n".join(texts[:references_index])

    references = [
        text for text in texts[references_index + 1 :] if re.match(r"^\d+\.\s", text)
    ]
    reference_numbers = [int(text.split(".", 1)[0]) for text in references]
    dois = []
    for reference in references:
        match = re.search(r"https://doi\.org/([^\s.]+(?:\.[^\s.]+)*)\.?$", reference)
        if not match:
            raise AssertionError(f"missing DOI: {reference}")
        dois.append(match.group(1).rstrip("."))

    citations: list[int] = []
    for match in re.finditer(r"\[([0-9,\-\u2013\s]+)\]", body_text):
        citations.extend(expand_citation(match.group(1)))
    first_seen: list[int] = []
    for citation in citations:
        if citation not in first_seen:
            first_seen.append(citation)

    equation_data: dict[int, etree._Element] = {}
    for paragraph in paragraphs:
        labels = re.findall(r"\((1[0-4]|[1-9])\)", paragraph_text(paragraph))
        formulas = paragraph.xpath(".//m:oMath", namespaces=NS)
        if labels and formulas:
            assert len(labels) == len(formulas), (labels, len(formulas))
            for label, formula in zip(labels, formulas):
                equation_data[int(label)] = formula

    tables = root.xpath("//w:body//w:tbl", namespaces=NS)
    table_headers = [
        bool(table.xpath("./w:tr[1]/w:trPr/w:tblHeader", namespaces=NS))
        for table in tables
    ]
    figure_captions = [text for text in texts if re.match(r"^Figure\s+[1-7]\.\s", text)]
    figure_alt = [
        (node.get("descr") or node.get("title") or "").strip()
        for node in root.xpath("//wp:docPr", namespaces=NS)
    ]

    fonts = {
        value
        for attr in styles.xpath("//w:rFonts/@*", namespaces=NS)
        for value in [attr]
        if isinstance(value, str) and value
    }

    assert len(doc.tables) == 7
    assert len(references) == 26
    assert reference_numbers == list(range(1, 27))
    assert len(set(dois)) == 26
    assert first_seen == list(range(1, 27)), first_seen
    assert sorted(equation_data) == list(range(1, 15))
    assert equation_data[11].xpath(".//m:f", namespaces=NS)
    assert equation_data[12].xpath(".//m:nary", namespaces=NS)
    assert equation_data[13].xpath(".//m:f", namespaces=NS)
    assert equation_data[13].xpath(".//m:rad", namespaces=NS)
    assert equation_data[14].xpath(".//m:sSub", namespaces=NS)
    assert len(tables) == 7 and all(table_headers)
    assert len(figure_captions) == 7
    assert len(figure_alt) >= 7 and all(figure_alt[:7])
    assert "Palatino Linotype" in fonts
    assert "*653*" not in body_text and all("*653*" not in ref for ref in references)
    assert "[27]" not in body_text
    assert "original ESGF node" not in body_text
    assert "NP results are invariant" not in body_text
    assert "13,600\u201394,900" not in body_text

    print(f"Document: {DOCX}")
    print(f"Body paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(tables)}; repeating header rows: {sum(table_headers)}/7")
    print(f"Figures: {len(figure_captions)}; alt-text entries: {sum(bool(x) for x in figure_alt[:7])}/7")
    print(f"Equations: {sorted(equation_data)}; 11-14 use structured OMML")
    print(f"References: {len(references)}; unique DOI: {len(set(dois))}")
    print(f"First-appearance citation order: {first_seen}")
    print("Palatino Linotype present in styles: yes")
    print("Structural audit: PASS")


if __name__ == "__main__":
    main()
