from __future__ import annotations

import copy
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\journal_revision_q3")
WORK = ROOT / "final_revision"
SOURCE = WORK / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN_SOURCE_SNAPSHOT.docx"
INTERMEDIATE = WORK / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN_INTERMEDIATE.docx"
OUTPUT = WORK / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN_FINAL.docx"
TABLE_HELPER = (
    Path(r"C:\Users\PC\.codex\plugins\cache\openai-primary-runtime")
    / "documents"
    / "26.826.12353"
    / "skills"
    / "documents"
    / "scripts"
)
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W, "m": M, "wp": WP}


def paragraphs_and_cells(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def find_paragraph(doc: Document, startswith: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(startswith):
            return paragraph
    raise ValueError(f"paragraph not found: {startswith!r}")


def clear_paragraph(paragraph: Paragraph) -> None:
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)


def first_run_properties(paragraph: Paragraph):
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        return copy.deepcopy(paragraph.runs[0]._r.rPr)
    return None


def add_run_with_properties(paragraph: Paragraph, text: str, r_pr=None):
    run = paragraph.add_run(text)
    if r_pr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, copy.deepcopy(r_pr))
    return run


def set_text(paragraph: Paragraph, text: str) -> None:
    r_pr = first_run_properties(paragraph)
    clear_paragraph(paragraph)
    add_run_with_properties(paragraph, text, r_pr)


def set_labeled_text(paragraph: Paragraph, label: str, text: str) -> None:
    r_pr = first_run_properties(paragraph)
    clear_paragraph(paragraph)
    label_run = add_run_with_properties(paragraph, label, r_pr)
    label_run.bold = True
    body_run = add_run_with_properties(paragraph, text, r_pr)
    body_run.bold = False


def insert_after(template: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    if template._p.pPr is not None:
        new_p.append(copy.deepcopy(template._p.pPr))
    template._p.addnext(new_p)
    paragraph = Paragraph(new_p, template._parent)
    paragraph.add_run(text)
    return paragraph


def insert_before_with_label(
    anchor: Paragraph, label: str, text: str
) -> Paragraph:
    paragraph = anchor.insert_paragraph_before()
    if anchor._p.pPr is not None:
        paragraph._p.insert(0, copy.deepcopy(anchor._p.pPr))
    paragraph.add_run(label).bold = True
    paragraph.add_run(text)
    return paragraph


def remove_table_column(table, index: int) -> None:
    for row in table.rows:
        tc = row.cells[index]._tc
        tc.getparent().remove(tc)
    grid = table._tbl.tblGrid
    grid_cols = list(grid)
    grid.remove(grid_cols[index])


def mark_header_row(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def rebuild_reference(paragraph: Paragraph, new_number: int, text: str) -> None:
    journal_names = sorted(
        {
            "Wiley Interdiscip. Rev. Clim. Change",
            "J. South. Hemisph. Earth Syst. Sci.",
            "Hydrol. Earth Syst. Sci.",
            "Philos. Trans. R. Soc. A",
            "J. Adv. Model Earth Syst.",
            "Geosci. Model Dev.",
            "Theor. Appl. Climatol.",
            "Int. J. Climatol.",
            "Atmos. Sci. Lett.",
            "Nat. Methods",
            "Appl. Stat.",
            "J. Hydrol.",
            "J. Clim.",
            "Hydrology",
        },
        key=len,
        reverse=True,
    )
    clean = text.replace("\u00a0", " ").replace("*", " ")
    clean = re.sub(r"\s+", " ", clean).strip()
    clean = re.sub(r"^\d+\.\s*", "", clean)

    doi_match = re.search(r"\s+(https?://doi\.org/\S+)\.?$", clean)
    if not doi_match:
        raise ValueError(f"DOI not found in reference: {clean}")
    doi = doi_match.group(1).rstrip(".")
    bibliographic = clean[: doi_match.start()].rstrip().rstrip(".")

    year_match = re.search(r"\s((?:19|20)\d{2}),\s", bibliographic)
    if not year_match:
        raise ValueError(f"year not found in reference: {clean}")
    year = year_match.group(1)
    prefix = bibliographic[: year_match.start()].rstrip()
    suffix = bibliographic[year_match.end() :]

    journal = next((name for name in journal_names if prefix.endswith(name)), None)
    if not journal:
        raise ValueError(f"journal not recognized in reference: {clean}")
    before_journal = prefix[: -len(journal)].rstrip()
    volume_match = re.match(r"(\d+)(.*)", suffix)
    if not volume_match:
        raise ValueError(f"volume not found in reference: {clean}")
    volume, remainder = volume_match.groups()
    remainder = re.sub(r"\s*,\s*", ", ", remainder.strip())
    remainder = re.sub(r"(?<=\d)-(?=\d)", "\u2013", remainder)

    clear_paragraph(paragraph)
    paragraph.add_run(f"{new_number}. {before_journal} ")
    paragraph.add_run(journal).italic = True
    paragraph.add_run(f" {year}, ")
    paragraph.add_run(volume).italic = True
    paragraph.add_run(remainder)
    paragraph.add_run(f". {doi}.")
    for run in paragraph.runs:
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(9.5)
        run._r.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Palatino Linotype"
        )
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def high_level_edits() -> None:
    doc = Document(SOURCE)

    # Required labels and front/back matter.
    abstract = find_paragraph(doc, "Abstract:")
    set_labeled_text(abstract, "Abstract: ", abstract.text[len("Abstract: ") :])

    funding = find_paragraph(doc, "Funding:")
    insert_before_with_label(
        funding,
        "Acknowledgements: ",
        "The author acknowledges the World Climate Research Programme (WCRP), "
        "which coordinated and promoted CMIP6 through its Working Group on Coupled "
        "Modelling; the participating climate modelling groups for producing and "
        "making their model output available; and the Earth System Grid Federation "
        "(ESGF) for archiving and providing access to CMIP6 data. The author also "
        "acknowledges the responsible Thai data custodians who provided access to "
        "the rainfall records used in this study.",
    )
    insert_before_with_label(
        funding,
        "Author Contributions: ",
        "Conceptualization, Methodology, Software, Validation, Formal analysis, "
        "Investigation, Data curation, Visualization, Writing \u2013 original draft, "
        "and Writing \u2013 review and editing were performed by the author. "
        "The author has read and agreed to the published version of the manuscript.",
    )

    data_availability = find_paragraph(doc, "Data Availability Statement:")
    set_labeled_text(
        data_availability,
        "Data Availability Statement: ",
        "The rainfall observations used in this study were supplied through the "
        "Hydro-Informatics Institute/Royal Irrigation Department project archive "
        "and remain subject to the terms of the responsible Thai data custodians. "
        "Access requests should be directed to those custodians. CMIP6 datasets are "
        "discoverable through ESGF using the source, experiment, member, grid, "
        "variable, frequency, and date-range metadata summarized in Table 2. The "
        "complete reproducibility package, including the analysis code, input-file "
        "checksum manifest, derived metrics, diagnostics, sensitivity outputs, and "
        "figure-generation scripts, is available from the corresponding author "
        "upon reasonable request.",
    )

    conflicts = find_paragraph(doc, "Conflicts of Interest:")
    set_labeled_text(
        conflicts,
        "Conflicts of Interest: ",
        "The author declares no conflict of interest. The funders had no role in "
        "the design of the study; in the collection, analyses, or interpretation "
        "of data; in the writing of the manuscript; or in the decision to publish "
        "the results.",
    )

    # Methods: direct Table 1 citation, DEM/elevation scope, and exact equation text.
    study_area = find_paragraph(doc, "Located in coastal Thailand")
    set_text(
        study_area,
        study_area.text.replace(
            "from a 12-station inventory distributed throughout the province (Figure 1).",
            "from a 12-station inventory distributed throughout the province "
            "(Figure 1; Table 1).",
        ),
    )

    figure1_caption = find_paragraph(doc, "Figure 1.")
    elevation_note = insert_after(
        figure1_caption,
        "The elevation surface in Figure 1 is provided only as regional "
        "topographic context. Station elevations in Table 1 are reproduced from "
        "the supplied station metadata; no DEM-derived elevations were substituted "
        "where station metadata were unavailable.",
    )
    elevation_note.paragraph_format.first_line_indent = study_area.paragraph_format.first_line_indent
    elevation_note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    equation6_text = find_paragraph(doc, "where  is the sample size")
    set_text(
        equation6_text,
        equation6_text.text.replace(
            "where  is the sample size.", "where n is the sample size."
        ),
    )

    table3_definition = find_paragraph(doc, "To prevent generating precipitation")
    set_text(
        table3_definition,
        table3_definition.text.replace("defined in table 3", "defined in Table 3"),
    )

    # Table 2: retain only metadata recoverable from the station-series files.
    table2 = doc.tables[1]
    if len(table2.columns) != 9:
        raise ValueError(f"expected 9 columns in Table 2, found {len(table2.columns)}")
    remove_table_column(table2, 4)
    table2.rows[0].cells[0].text = "Model"
    table2.rows[0].cells[1].text = "Institution"
    table2.rows[0].cells[2].text = "Variant"
    table2.rows[0].cells[3].text = "Grid"
    table2.rows[0].cells[4].text = "Cal."
    table2.rows[0].cells[5].text = "Days"
    table2.rows[0].cells[6].text = "Series"
    table2.rows[0].cells[7].text = "Ref."
    for row in table2.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    table2_note = find_paragraph(doc, "Note: Grid identifiers")
    set_text(
        table2_note,
        "Note: Grid identifiers (gn/gr) are retained as recorded in the "
        "analysis-ready files. Distinct series is the number of unique daily "
        "station-series signatures across 12 gauges and indicates shared native "
        "grid cells; it is not a formal effective sample size.",
    )

    # Results: disclose observational structure without new computations.
    observed_context = find_paragraph(doc, "Mean annual observed rainfall ranged")
    station_note = insert_after(
        observed_context,
        "Table 1 and Figure 3 show two distinct observational structures. "
        "Stations 500005-500009 have lower SDII (5.1-6.5 mm day\u207b\u00b9) and "
        "substantially higher wet-day frequencies than the other stations, while "
        "500008 and 500009 also have lower recorded daily maxima (68.0 and "
        "54.2 mm, respectively). These contrasts may reflect genuine gauge-regime "
        "differences or data-handling practices, but the available metadata do not "
        "permit attribution. All stations were therefore retained, and station-level "
        "comparisons are interpreted descriptively rather than as independent "
        "spatial evidence.",
    )
    station_note.paragraph_format.first_line_indent = observed_context.paragraph_format.first_line_indent
    station_note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    paired = find_paragraph(doc, "The paired bootstrap did not support")
    summary_note = insert_after(
        paired,
        "Table 4 reports arithmetic means of metric levels, whereas Table 5 "
        "reports medians of pair-specific corrected-minus-raw changes. The two "
        "summaries therefore need not coincide when pairwise changes are skewed "
        "or contain influential values.",
    )
    summary_note.paragraph_format.first_line_indent = paired.paragraph_format.first_line_indent
    summary_note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    metric_definitions = find_paragraph(doc, "Here, N is the number of paired monthly records")
    set_text(
        metric_definitions,
        "Here, N is the number of paired monthly records; S\u1d62 and O\u1d62 are simulated "
        "and observed monthly totals; and q99 is the 99th percentile of wet-day intensity. "
        "In Equation (10), D is the supremum absolute difference between the simulated "
        "and observed empirical cumulative distribution functions of wet-day precipitation. "
        "The climatology correlation is calculated from the 12 month-specific means, not "
        "from independent observations. KGE and correlations are better at larger values, "
        "whereas the remaining endpoints are better at smaller values. Direction-based "
        "rankings are labeled descriptive and do not imply statistical superiority.",
    )

    paired_metric_note = find_paragraph(doc, "The PBIAS and MAE equations")
    set_text(
        paired_metric_note,
        "The PBIAS and MAE equations use the same paired monthly records as RMSE; r is "
        "the Pearson correlation of monthly totals; and the climatology correlation is "
        "calculated from the 12 simulated and observed monthly climatological means.",
    )

    secondary_endpoints = find_paragraph(doc, "Secondary endpoints included")
    set_text(
        secondary_endpoints,
        "Secondary endpoints included monthly root mean square error (RMSE), mean "
        "absolute error (MAE), percent bias (PBIAS), absolute relative bias of the "
        "99th wet-day percentile (q99 bias), the Kolmogorov-Smirnov statistic (D) "
        "for wet-day distributions, monthly Pearson correlation (r), and correlation "
        "between the 12 monthly climatological means. These metrics are defined in "
        "Equations (8)\u2013(14).",
    )

    aicc = find_paragraph(doc, "AICc produced exactly")
    set_text(
        aicc,
        aicc.text
        + " Because family-specific boundary and parameter-shape diagnostics were "
        "not prespecified, the selection frequencies are descriptive and should "
        "not be interpreted as evidence that Pearson type III uniquely represents "
        "regional rainfall physics.",
    )

    table7_note = find_paragraph(doc, "Note: The cap is factor times")
    set_text(
        table7_note,
        table7_note.text.replace(
            "NP results are invariant to this parameter.",
            "This cap-sensitivity experiment applies only to the parametric "
            "variants; non-parametric results are not included in Table 7.",
        ),
    )

    physical = find_paragraph(doc, "A physically plausible interpretation")
    old_grid_sentence = (
        "The native cells span approximately 13,600-94,900 km\u00b2 at 12 degrees N, "
        "so even the finest model cell is large relative to the province's "
        "mountain-coast transition."
    )
    new_grid_sentence = (
        "The station-series files show that several gauges share identical "
        "native-grid extractions, so the effective spatial resolution remains "
        "coarse relative to the province's mountain-coast transition."
    )
    if old_grid_sentence not in physical.text:
        raise ValueError("expected grid-size sentence was not found")
    set_text(physical, physical.text.replace(old_grid_sentence, new_grid_sentence))

    # Reference list: remove duplicate old [15], renumber old [16]-[27], and use
    # the AJSTR template's italic journal/volume styling.
    references_heading_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text == "References"
    )
    reference_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs[references_heading_index + 1 :]
        if re.match(r"^\d+\.\s", paragraph.text)
    ]
    if len(reference_paragraphs) != 27:
        raise ValueError(f"expected 27 references, found {len(reference_paragraphs)}")
    for paragraph in reference_paragraphs:
        old_number = int(paragraph.text.split(".", 1)[0])
        original_text = paragraph.text
        if old_number == 15:
            delete_paragraph(paragraph)
            continue
        new_number = old_number if old_number < 15 else old_number - 1
        rebuild_reference(paragraph, new_number, original_text)

    # Freeze table geometry in DXA so Word, LibreOffice, and production systems
    # use the same column proportions as the source manuscript.
    sys.path.insert(0, str(TABLE_HELPER))
    from table_geometry import (  # type: ignore[import-not-found]
        apply_table_geometry,
        column_widths_from_weights,
        section_content_width_dxa,
    )

    content_width = section_content_width_dxa(doc.sections[0])
    dense_table_widths = {
        4: [1250, 900, 900, 1150, 1050, 1550, 1550, 1283],
        5: [1100, 1700, 1550, 1700, 1050, 850, 800, 883],
    }
    for table_index, table in enumerate(doc.tables):
        grid_columns = list(table._tbl.tblGrid.gridCol_lst)
        weights = [int(column.get(qn("w:w"))) for column in grid_columns]
        if len(weights) != len(table.columns):
            raise ValueError(
                f"table grid mismatch: {len(weights)} grid columns versus "
                f"{len(table.columns)} table columns"
            )
        widths = dense_table_widths.get(
            table_index, column_widths_from_weights(weights, content_width)
        )
        if table_index in dense_table_widths:
            apply_table_geometry(
                table,
                widths,
                table_width_dxa=content_width,
                indent_dxa=60,
                cell_margins_dxa={"start": 60, "end": 60},
            )
        else:
            apply_table_geometry(table, widths, table_width_dxa=content_width)

        if table_index == 5:
            for cell in table.rows[0].cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)

        # All first rows are true column headers and should repeat if a table splits.
        mark_header_row(table)

    # Keep each inline figure with its following caption. Caption paragraphs should
    # stay intact but must not drag the next figure or section onto the same page.
    for caption in doc.paragraphs:
        if not re.match(r"^Figure [1-7]\.\s", caption.text):
            continue
        caption.paragraph_format.keep_together = True
        caption.paragraph_format.keep_with_next = False
        previous = caption._p.getprevious()
        while previous is not None and previous.tag == qn("w:p"):
            if previous.xpath(".//w:drawing"):
                image_paragraph = Paragraph(previous, caption._parent)
                image_paragraph.paragraph_format.keep_together = True
                image_paragraph.paragraph_format.keep_with_next = True
                break
            if "".join(previous.itertext()).strip():
                break
            previous = previous.getprevious()

    doc.save(INTERMEDIATE)


def mathml_to_omml(mathml: str, transform: etree.XSLT):
    math = etree.fromstring(mathml.encode("utf-8"))
    transformed = transform(math)
    root = transformed.getroot()
    if root.tag == f"{{{M}}}oMath":
        return copy.deepcopy(root)
    nodes = root.xpath(".//m:oMath", namespaces=NS)
    if not nodes:
        raise ValueError("MML2OMML transform did not produce m:oMath")
    return copy.deepcopy(nodes[0])


def bar(variable: str) -> str:
    return f"<mover><mi>{variable}</mi><mo>\u00af</mo></mover>"


def summation(term: str) -> str:
    return (
        "<munderover><mo>\u2211</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn>"
        "</mrow><mi>N</mi></munderover>" + term
    )


def new_equation_mathml() -> dict[int, str]:
    ns = 'xmlns="http://www.w3.org/1998/Math/MathML"'
    sbar = bar("S")
    obar = bar("O")
    si = "<msub><mi>S</mi><mi>i</mi></msub>"
    oi = "<msub><mi>O</mi><mi>i</mi></msub>"

    eq11 = (
        f"<math {ns}><mrow><mi>PBIAS</mi><mo>=</mo><mn>100</mn><mo>\u00d7</mo>"
        f"<mfrac><mrow>{sbar}<mo>\u2212</mo>{obar}</mrow>{obar}</mfrac>"
        "</mrow></math>"
    )
    eq12 = (
        f"<math {ns}><mrow><mi>MAE</mi><mo>=</mo><mfrac><mn>1</mn><mi>N</mi>"
        f"</mfrac>{summation(f'<mrow><mo>|</mo>{si}<mo>\u2212</mo>{oi}<mo>|</mo></mrow>')}"
        "</mrow></math>"
    )
    numerator = summation(
        f"<mrow><mo>(</mo>{si}<mo>\u2212</mo>{sbar}<mo>)</mo>"
        f"<mo>(</mo>{oi}<mo>\u2212</mo>{obar}<mo>)</mo></mrow>"
    )
    s_squared = summation(
        f"<msup><mrow><mo>(</mo>{si}<mo>\u2212</mo>{sbar}<mo>)</mo></mrow>"
        "<mn>2</mn></msup>"
    )
    o_squared = summation(
        f"<msup><mrow><mo>(</mo>{oi}<mo>\u2212</mo>{obar}<mo>)</mo></mrow>"
        "<mn>2</mn></msup>"
    )
    eq13 = (
        f"<math {ns}><mrow><mi>r</mi><mo>=</mo><mfrac><mrow>{numerator}</mrow>"
        f"<msqrt><mrow><mo>(</mo>{s_squared}<mo>)</mo><mo>(</mo>{o_squared}"
        "<mo>)</mo></mrow></msqrt></mfrac></mrow></math>"
    )
    s1 = f"<msub>{sbar}<mn>1</mn></msub>"
    s12 = f"<msub>{sbar}<mn>12</mn></msub>"
    o1 = f"<msub>{obar}<mn>1</mn></msub>"
    o12 = f"<msub>{obar}<mn>12</mn></msub>"
    eq14 = (
        f"<math {ns}><mrow><msub><mi>r</mi><mi>clim</mi></msub><mo>=</mo>"
        "<mi mathvariant=\"normal\">corr</mi><mo>(</mo><mo>{</mo>"
        f"{s1}<mo>,</mo><mo>\u2026</mo><mo>,</mo>{s12}<mo>}}</mo><mo>,</mo>"
        f"<mo>{{</mo>{o1}<mo>,</mo><mo>\u2026</mo><mo>,</mo>{o12}<mo>}}</mo>"
        "<mo>)</mo></mrow></math>"
    )
    return {11: eq11, 12: eq12, 13: eq13, 14: eq14}


def equation_paragraph(number: int, formula) -> etree._Element:
    paragraph = etree.Element(f"{{{W}}}p")
    p_pr = etree.SubElement(paragraph, f"{{{W}}}pPr")
    tabs = etree.SubElement(p_pr, f"{{{W}}}tabs")
    center = etree.SubElement(tabs, f"{{{W}}}tab")
    center.set(f"{{{W}}}val", "center")
    center.set(f"{{{W}}}pos", "4464")
    right = etree.SubElement(tabs, f"{{{W}}}tab")
    right.set(f"{{{W}}}val", "right")
    right.set(f"{{{W}}}pos", "8784")
    spacing = etree.SubElement(p_pr, f"{{{W}}}spacing")
    spacing.set(f"{{{W}}}before", "40")
    spacing.set(f"{{{W}}}after", "40")
    spacing.set(f"{{{W}}}line", "240")
    spacing.set(f"{{{W}}}lineRule", "auto")

    first_run = etree.SubElement(paragraph, f"{{{W}}}r")
    etree.SubElement(first_run, f"{{{W}}}tab")
    paragraph.append(copy.deepcopy(formula))
    second_run = etree.SubElement(paragraph, f"{{{W}}}r")
    etree.SubElement(second_run, f"{{{W}}}tab")
    label_run = etree.SubElement(paragraph, f"{{{W}}}r")
    label = etree.SubElement(label_run, f"{{{W}}}t")
    label.text = f"({number})"
    return paragraph


def format_citation(numbers: list[int]) -> str:
    mapped = []
    for number in numbers:
        if number == 15:
            number = 8
        elif number >= 16:
            number -= 1
        if number not in mapped:
            mapped.append(number)
    mapped.sort()

    tokens: list[str] = []
    i = 0
    while i < len(mapped):
        j = i
        while j + 1 < len(mapped) and mapped[j + 1] == mapped[j] + 1:
            j += 1
        if j - i >= 2:
            tokens.append(f"{mapped[i]}\u2013{mapped[j]}")
        else:
            tokens.extend(str(value) for value in mapped[i : j + 1])
        i = j + 1
    return "[" + ", ".join(tokens) + "]"


def update_body_text(root: etree._Element) -> None:
    references_seen = False
    numeric_citation = re.compile(r"\[([0-9,\-\s]+)\]")

    def replace_citation(match: re.Match[str]) -> str:
        content = match.group(1)
        numbers: list[int] = []
        for token in content.split(","):
            token = token.strip()
            if not token:
                continue
            if "-" in token:
                start, end = (int(value.strip()) for value in token.split("-", 1))
                numbers.extend(range(start, end + 1))
            else:
                numbers.append(int(token))
        return format_citation(numbers)

    spelling = {
        "modelled": "modeled",
        "recognised": "recognized",
        "behaviour": "behavior",
        "homogenisation": "homogenization",
        "analysed": "analyzed",
        "labelled": "labeled",
        "centre": "center",
        "favouring": "favoring",
    }

    for paragraph in root.xpath("//w:body//w:p", namespaces=NS):
        paragraph_text = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
        if paragraph_text.strip() == "References":
            references_seen = True
            continue
        if references_seen:
            continue
        for node in paragraph.xpath(".//w:t", namespaces=NS):
            text = node.text or ""
            text = numeric_citation.sub(replace_citation, text)
            text = re.sub(r"(?<!ESM1)(?<=\d)-(?=\d)", "\u2013", text)
            text = text.replace("10th-90th", "10th\u201390th")
            text = text.replace(">=1.0", "\u22651.0")
            text = text.replace("defined in table 3", "defined in Table 3")
            for british, american in spelling.items():
                text = re.sub(rf"\b{british}\b", american, text, flags=re.IGNORECASE)
            node.text = text


def patch_ooxml() -> None:
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    mathml = new_equation_mathml()

    with zipfile.ZipFile(INTERMEDIATE, "r") as source_zip:
        entries = {name: source_zip.read(name) for name in source_zip.namelist()}

    root = etree.fromstring(entries["word/document.xml"])

    equation_paragraphs: list[tuple[etree._Element, list[int]]] = []
    existing_formulas: dict[int, etree._Element] = {}
    for paragraph in root.xpath("//w:body/w:p", namespaces=NS):
        formulas = paragraph.xpath(".//m:oMath", namespaces=NS)
        if not formulas:
            continue
        labels = [
            int(value)
            for value in re.findall(
                r"\((1[0-4]|[1-9])\)",
                "".join(paragraph.xpath(".//w:t/text()", namespaces=NS)),
            )
        ]
        if not labels:
            continue
        if len(formulas) != len(labels):
            raise ValueError(
                f"equation paragraph labels/formulas mismatch: {labels}, {len(formulas)}"
            )
        equation_paragraphs.append((paragraph, labels))
        for number, formula in zip(labels, formulas):
            if number <= 10:
                existing_formulas[number] = copy.deepcopy(formula)

    if sorted(existing_formulas) != list(range(1, 11)):
        raise ValueError(f"missing existing equations: {sorted(existing_formulas)}")
    for number, source_mathml in mathml.items():
        existing_formulas[number] = mathml_to_omml(source_mathml, transform)

    for old_paragraph, labels in equation_paragraphs:
        parent = old_paragraph.getparent()
        index = parent.index(old_paragraph)
        for offset, number in enumerate(labels):
            parent.insert(
                index + offset,
                equation_paragraph(number, existing_formulas[number]),
            )
        parent.remove(old_paragraph)

    update_body_text(root)

    # Meaningful Figure 2 alt text.
    doc_props = root.xpath(".//wp:inline/wp:docPr", namespaces=NS)
    if len(doc_props) != 7:
        raise ValueError(f"expected 7 body figures, found {len(doc_props)}")
    figure2 = doc_props[1]
    figure2.set("title", "Fit-freeze-apply framework for independent validation")
    figure2.set(
        "descr",
        "Calibration uses observations and CMIP6 data from 1981-2000 to estimate "
        "wet-day adjustment and QDM transfer functions. Frozen functions are "
        "applied without re-estimation to CMIP6 simulations and withheld "
        "observations from 2001-2014, followed by paired evaluation and "
        "validation-year bootstrap and cap-sensitivity assessment.",
    )

    entries["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    with tempfile.NamedTemporaryFile(
        suffix=".docx", dir=WORK, delete=False
    ) as temp_handle:
        temp_path = Path(temp_handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as out_zip:
            for name, data in entries.items():
                out_zip.writestr(name, data)
        shutil.move(temp_path, OUTPUT)
    finally:
        temp_path.unlink(missing_ok=True)


def final_assertions() -> None:
    doc = Document(OUTPUT)
    text = "\n".join(paragraph.text for paragraph in paragraphs_and_cells(doc))
    references_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text == "References"
    )
    references = [
        paragraph
        for paragraph in doc.paragraphs[references_index + 1 :]
        if re.match(r"^\d+\.\s", paragraph.text)
    ]
    assert len(references) == 26
    assert len(doc.tables) == 7
    assert len(doc.inline_shapes) == 7
    assert "*653*" not in text
    assert "official station inventory is documented" not in text
    assert "where is the sample size" not in text
    assert "Author Contributions:" in text
    assert "Acknowledgements:" in text
    assert "NP results are invariant" not in text
    assert "13,600-94,900" not in text
    assert "13,600\u201394,900" not in text
    assert ">=1.0" not in text
    assert "defined in table 3" not in text
    assert "[27]" not in text

    with zipfile.ZipFile(OUTPUT) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        core = etree.fromstring(archive.read("docProps/core.xml"))
        assert "docProps/custom.xml" not in archive.namelist()
    core_ns = {
        "dc": "http://purl.org/dc/elements/1.1/",
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    }
    assert not "".join(core.xpath(".//dc:creator/text()", namespaces=core_ns)).strip()
    assert not "".join(core.xpath(".//cp:lastModifiedBy/text()", namespaces=core_ns)).strip()
    labels: list[int] = []
    for paragraph in root.xpath("//w:body/w:p", namespaces=NS):
        if not paragraph.xpath(".//m:oMath", namespaces=NS):
            continue
        labels.extend(
            int(value)
            for value in re.findall(
                r"\((1[0-4]|[1-9])\)",
                "".join(paragraph.xpath(".//w:t/text()", namespaces=NS)),
            )
        )
    assert labels == list(range(1, 15)), labels
    assert len(root.xpath("//w:tbl", namespaces=NS)) == 7
    assert len(root.xpath("//w:tbl/w:tr[1]/w:trPr/w:tblHeader", namespaces=NS)) == 7
    assert all(
        root.xpath(
            f"count(//w:body/w:p[.//w:t='({number})']//m:oMath) = 1",
            namespaces=NS,
        )
        for number in range(1, 15)
    )
    eq13 = next(
        paragraph
        for paragraph in root.xpath("//w:body/w:p", namespaces=NS)
        if "(13)" in "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
    )
    assert eq13.xpath(".//m:f", namespaces=NS)
    assert eq13.xpath(".//m:rad", namespaces=NS)
    print(OUTPUT)


def scrub_privacy_metadata() -> None:
    sys.path.insert(0, str(TABLE_HELPER))
    from privacy_scrub import scrub  # type: ignore[import-not-found]

    scrubbed = OUTPUT.with_name(f"{OUTPUT.stem}_SCRUBBED.docx")
    stats = scrub(str(OUTPUT), str(scrubbed))
    shutil.move(scrubbed, OUTPUT)
    if not stats["core_props_scrubbed"]:
        raise ValueError("privacy scrub did not clear core author metadata")


def main() -> None:
    high_level_edits()
    patch_ooxml()
    scrub_privacy_metadata()
    final_assertions()


if __name__ == "__main__":
    main()
