from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image


WORK = Path(
    r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly"
    r"\journal_revision_q3\final_revision"
)
DEST = Path(r"D:\วารสาร ASEAN J. Sci\New paper\New folder")
SOURCE = DEST / "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN.docx"
MAIN_NAME = "AJSTR_QDM_CMIP6_Q3_SUBMISSION_MAIN_FINAL.docx"
ZIP_NAME = "AJSTR_QDM_CMIP6_Q3_SUBMISSION_FINAL_PACKAGE.zip"
EXPECTED_SOURCE_SHA256 = (
    "B26DE9E17A1D714A97659EFAE1B9A6BCED3A90C0326D0F1ABA5E1789D05E7566"
)
EXPECTED_ENTRIES = {
    MAIN_NAME,
    "AJSTR_QDM_CMIP6_Q3_SUBMISSION_TITLE_PAGE.docx",
    *(f"Figure_{number}.tiff" for number in range(1, 8)),
}
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = {"w": W, "m": M}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def main() -> None:
    work_main = WORK / MAIN_NAME
    work_zip = WORK / ZIP_NAME
    dest_main = DEST / MAIN_NAME
    dest_zip = DEST / ZIP_NAME

    assert sha256(SOURCE) == EXPECTED_SOURCE_SHA256
    assert sha256(work_main) == sha256(dest_main)
    assert sha256(work_zip) == sha256(dest_zip)

    document = Document(dest_main)
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""
    assert len(document.tables) == 7
    assert len(document.inline_shapes) == 7

    with zipfile.ZipFile(dest_main) as archive:
        names = set(archive.namelist())
        assert "word/comments.xml" not in names
        assert "docProps/custom.xml" not in names
        root = etree.fromstring(archive.read("word/document.xml"))
    assert not root.xpath("//w:ins | //w:del", namespaces=NS)
    texts = [
        "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
        for paragraph in root.xpath("//w:body//w:p", namespaces=NS)
    ]
    refs_index = texts.index("References")
    references = [
        text for text in texts[refs_index + 1 :] if re.match(r"^\d+\.\s", text)
    ]
    assert len(references) == 26
    assert len({re.search(r"https://doi\.org/(\S+)\.?$", ref).group(1).rstrip(".") for ref in references}) == 26
    assert len([text for text in texts if re.match(r"^Figure [1-7]\.\s", text)]) == 7
    equation_labels = []
    for paragraph in root.xpath("//w:body/w:p[.//m:oMath]", namespaces=NS):
        equation_labels.extend(
            int(value)
            for value in re.findall(
                r"\((1[0-4]|[1-9])\)",
                "".join(paragraph.xpath(".//w:t/text()", namespaces=NS)),
            )
        )
    assert equation_labels == list(range(1, 15))

    with zipfile.ZipFile(dest_zip) as package:
        assert set(package.namelist()) == EXPECTED_ENTRIES
        assert package.testzip() is None
        Document(io.BytesIO(package.read(MAIN_NAME)))
        Document(io.BytesIO(package.read("AJSTR_QDM_CMIP6_Q3_SUBMISSION_TITLE_PAGE.docx")))
        for number in range(1, 8):
            name = f"Figure_{number}.tiff"
            image = Image.open(io.BytesIO(package.read(name)))
            assert image.mode == "RGB", (name, image.mode)
            assert tuple(round(value) for value in image.info["dpi"]) == (600, 600)
            assert image.info.get("compression") == "tiff_lzw"

    print("Source SHA256:", sha256(SOURCE))
    print("Delivered main SHA256:", sha256(dest_main))
    print("Delivered ZIP SHA256:", sha256(dest_zip))
    print("Main DOCX: 7 tables, 7 figures, 14 equations, 26 unique references")
    print("Privacy: core author metadata blank; no comments, custom properties, or revisions")
    print("ZIP: 9 expected files; Figure 1-7 are RGB, 600 dpi, LZW TIFF")
    print("DELIVERY VERIFICATION: PASS")


if __name__ == "__main__":
    main()
