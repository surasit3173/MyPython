from pathlib import Path
from docx import Document


PATH = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def find_heading(doc, prefix):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    raise ValueError(prefix)


doc = Document(PATH)

# Remove duplicated 3.5 result paragraphs that were accidentally inserted before the 3.5 heading.
i34 = find_heading(doc, "3.4 ")
i35 = find_heading(doc, "3.5 ")
for p in list(doc.paragraphs[i34 + 3 : i35]):
    delete_paragraph(p)

for p in doc.paragraphs:
    text = p.text
    if "ผลระดับคู่ทั้งหมดของทั้งสองรูปแบบพร้อมค่า p ก่อนปรับและค่า q หลังปรับ แสดงไว้ในเอกสารประกอบ" in text:
        set_text(
            p,
            text.replace(
                " ผลระดับคู่ทั้งหมดของทั้งสองรูปแบบพร้อมค่า p ก่อนปรับและค่า q หลังปรับ แสดงไว้ในเอกสารประกอบ",
                "",
            ),
        )
    if p.text.strip().startswith("4.1 การตรวจสอบข้ามช่วงเวลาแบบอิสระ"):
        set_text(p, p.text.replace("การตรวจสอบข้ามช่วงเวลาแบบอิสระ", "การตรวจสอบด้วยช่วงเวลาอิสระ"))
    elif "กรอบการตรวจสอบข้ามช่วงเวลาแบบอิสระ" in p.text:
        set_text(p, p.text.replace("กรอบการตรวจสอบข้ามช่วงเวลาแบบอิสระ", "กรอบการตรวจสอบด้วยช่วงเวลาอิสระ"))

doc.save(PATH)
print(PATH)
