from pathlib import Path
import re
from collections import Counter
from docx import Document

docx_path = Path(r"D:\วารสารบูรพา\Evaluation of Quantile Delta Mapping for Bias Correction of CMIP6 Daily Rainfall Using Independent Temporal Holdout.docx")
doc = Document(docx_path)
paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
full = "\n".join(paras)

heading_re = re.compile(r"^(\d+(\.\d+)?\s|บทคัดย่อ$|Abstract$|Introduction$|Methodology$|Results$|Discussion$|Conclusions$|Acknowledg|References$)")
headings = [(i, t) for i, t in enumerate(paras) if heading_re.match(t)]

print(f"FILE\t{docx_path}")
print(f"PARAGRAPHS\t{len(paras)}")
print(f"TABLES\t{len(doc.tables)}")
print(f"INLINE_SHAPES\t{len(doc.inline_shapes)}")

print("\nHEADINGS")
for i, h in headings:
    print(f"{i}\t{h}")

terms = [
    "Temporal Cross-Validation",
    "Temporal Holdout",
    "independent validation",
    "independent temporal holdout",
    "Benjamini and Yekutieli",
    "Benjamini-Yekutieli",
    "Benjamini and Hochberg",
    "Supplementary",
    "เอกสารประกอบ",
    "https:/doi.org",
    "doi.org10",
    "เขคดุสิต",
    "ตรวจสอบความถูกต้องข้ามช่วงเวลา",
    "ช่วงตรวจสอบแบบอิสระ",
    "ช่วงตรวจสอบอิสระ",
    "ตรวจสอบด้วยช่วงเวลาอิสระ",
]
print("\nTERM_COUNTS")
for term in terms:
    print(f"{term}\t{full.count(term)}")

print("\nEQUATION_LINES")
for i, t in enumerate(paras):
    if re.search(r"\(\d+[a-z]?\)\s*$", t) or "MSE =" in t or "CI" in t or "p =" in t or "c(m)" in t:
        print(f"{i}\t{t}")

print("\nTABLE_SUMMARY")
for ti, table in enumerate(doc.tables, 1):
    rows = len(table.rows)
    cols = len(table.columns) if rows else 0
    sample = " | ".join(cell.text.strip().replace("\n", " / ")[:80] for cell in table.rows[0].cells) if rows else ""
    print(f"Table{ti}\t{rows}x{cols}\t{sample}")

print("\nLONG_PARAS")
for i, t in enumerate(paras):
    if len(t) > 900:
        section = ""
        for hi, h in headings:
            if hi <= i:
                section = h
            else:
                break
        print(f"{i}\t{section}\tlen={len(t)}\t{t[:250]}")

print("\nREFERENCE_YEAR_CHECK")
ref_start = next((i for i, t in enumerate(paras) if t == "References"), None)
refs = paras[ref_start + 1:] if ref_start is not None else []
print(f"refs_count_guess\t{len(refs)}")
for r in refs[:80]:
    years = re.findall(r"\((\d{4})\)|\b(19\d{2}|20\d{2})\b", r)
    print(("YEAR_OK" if years else "NO_YEAR") + "\t" + r[:220])
