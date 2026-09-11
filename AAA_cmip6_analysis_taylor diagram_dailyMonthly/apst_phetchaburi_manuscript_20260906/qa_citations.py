from docx import Document
import re

d = Document("APST_Phetchaburi_Precipitation_Extremes_Manuscript.docx")
text = "\n".join(p.text for p in d.paragraphs)
body = text.split("References")[0]
nums = []
for match in re.findall(r"\[(.*?)\]", body):
    for token in re.split(r"[,–-]", match):
        if token.strip().isdigit():
            nums.append(int(token.strip()))
print("citation_numbers", sorted(set(nums)), "count", len(nums))
print("reference_tail")
for p in d.paragraphs[-15:]:
    print(p.text)
