from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
from docx import Document

p = Path("APST_Phetchaburi_Precipitation_Extremes_Manuscript.docx")
d = Document(str(p))
txt = "\n".join(x.text for x in d.paragraphs)
ref_paras = [x.text for x in d.paragraphs if re.match(r"^\d+\. ", x.text)]
abs_p = next(x for x in d.paragraphs if x.text.startswith("Phetchaburi Province is exposed"))
concl_i = next(i for i, x in enumerate(d.paragraphs) if x.text == "5. Conclusion")
concl_p = d.paragraphs[concl_i + 1]
with ZipFile(p) as z:
    xml = z.read("word/document.xml").decode("utf-8")
    media = [n for n in z.namelist() if n.startswith("word/media/")]
result = {
    "size_bytes": p.stat().st_size,
    "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
    "paragraphs": len(d.paragraphs),
    "tables": [(len(t.rows), len(t.columns)) for t in d.tables],
    "inline_shapes": len(d.inline_shapes),
    "media": media,
    "abstract_words": len(abs_p.text.split()),
    "conclusion_words": len(concl_p.text.split()),
    "references": len(ref_paras),
    "placeholders": [s for s in ["1st author name", "Please_fill_in", "This document is a Microsoft Word template", "microcapsules", "Gac aril"] if s in txt],
    "line_numbering": "<w:lnNumType" in xml,
    "vertical_borders_nil": "w:insideV w:val=\"nil\"" in xml,
    "page_size_inches": [d.sections[0].page_width.inches, d.sections[0].page_height.inches],
    "margins_inches": [d.sections[0].top_margin.inches, d.sections[0].bottom_margin.inches, d.sections[0].left_margin.inches, d.sections[0].right_margin.inches],
}
print(json.dumps(result, ensure_ascii=False, indent=2))
