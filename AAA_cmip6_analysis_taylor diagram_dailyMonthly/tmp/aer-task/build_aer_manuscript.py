import argparse
import os
import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


REPLACEMENTS = {
    5: (
        "Provincial drainage design may rely on a single design rainfall even when rainfall occurrence and intensity vary among gauges. "
        "We analyzed daily rainfall from 13 gauges in Uttaradit Province during 1981-2014. Each record was described by mean annual rainfall, "
        "wet-day count, and mean wet-day intensity. Annual maxima were fitted with Gumbel, generalized extreme value, and log-Pearson type III "
        "distributions; the lowest small-sample corrected Akaike information criterion determined the selected model. Return-level uncertainty was "
        "estimated with 2,000 bootstrap resamples under fixed and reselected models. Between-gauge return levels varied by factors of 3.83 at 2 years "
        "and 12.53 at 100 years, compared with median within-gauge distribution ratios of 1.09 and 1.56. Mean wet-day intensity correlated with return "
        "level at 0.73 for 2 years and 0.48 for 100 years. At 100 years, median relative 95% interval width was 145.9% with model selection fixed and "
        "217.7% with model reselection; the gauge-specific ratio of these widths ranged from 0.71 to 10.25. The provincial mean fell below the local "
        "100-year estimate at five gauges, with a maximum shortfall of 65.3%. Spatial heterogeneity therefore exceeded sensitivity to the three "
        "candidate distributions, while long-period uncertainty depended strongly on model-selection stability. Local gauge estimates and their "
        "uncertainty intervals should be preferred when records are available."
    ),
    7: "Keywords: design rainfall; extreme-value analysis; model uncertainty; rainfall intensity; spatial heterogeneity; Uttaradit",
    10: (
        "Design rainfall, the depth with annual exceedance probability 1/T for return period T, determines the capacity and failure risk of culverts, "
        "drains, farm ponds, and small irrigation structures. Projects without nearby gauges often apply one value across an administrative area, "
        "although its site-level representativeness is rarely tested."
    ),
    11: (
        "A provincial value can be unrepresentative for two reasons. Rainfall extremes vary spatially, and each gauge estimate is uncertain when the "
        "return period exceeds the available record. Design practice must therefore distinguish between-gauge heterogeneity from uncertainty caused "
        "by extrapolation [1]."
    ),
    12: (
        "Rainfall records also differ in occurrence and intensity. Similar annual totals may arise from few intense wet days or many light wet days, "
        "a structure commonly described with daily rainfall indices [2, 3]. Whether this pre-fit structure explains design rainfall determines whether "
        "a simple network descriptor can support ungauged locations."
    ),
    13: (
        "Thai studies report long-term variability and trends in total and extreme precipitation [3, 4], within regional [5] and global [6] evidence. "
        "Extreme-value theory [7-9], regional frequency analysis in Thailand [10], and Thai intensity-duration-frequency studies [11] provide methods "
        "for translating rainfall records into design quantities. However, no distribution is uniformly preferred for daily rainfall maxima [12], "
        "and long-return-period uncertainty may exceed differences among fitted models [1]. Few studies compare spatial variation, distributional "
        "sensitivity, and sampling uncertainty within one province on a common scale."
    ),
    14: (
        "This study therefore (i) quantifies how design rainfall differs among 13 gauges and across return periods, (ii) tests whether annual total, "
        "wet-day frequency, and wet-day intensity are associated with return levels and their uncertainty, and (iii) measures the site-level error from "
        "applying one provincial value. All gauges remain in the primary analysis; sensitivity scenarios examine how network composition affects the "
        "reported range."
    ),
    17: (
        "Uttaradit Province lies in northern Thailand between approximately 17.2-18.1 degrees N and 100.0-101.1 degrees E. Daily rainfall from 1 January "
        "1981 to 31 December 2014 was obtained from the Royal Irrigation Department and the Thai Meteorological Department for 13 gauges at elevations "
        "of 55-428 m above mean sea level. All gauges contain numeric daily totals for all 12,418 days, with no negative values or missing-value codes. "
        "Record length and completeness therefore do not explain between-gauge differences. The agencies' quality-control procedures were not documented "
        "in the supplied files; Section 2.11 describes targeted artifact checks. Table 1 summarizes the stations and records."
    ),
    20: (
        "A wet day was defined as a day with at least 1.0 mm at every gauge. The mean annual total (P), mean wet-day count (W), and mean wet-day intensity "
        "(S) were calculated as follows:"
    ),
    22: "where x(t) is daily depth, n is the number of days, and N(y) is the number of years. Mean wet-day count was calculated as",
    24: "where I is the indicator function. Mean wet-day intensity was calculated as",
    26: (
        "Because P includes days below 1.0 mm while W and S do not, W x S was 0.58-3.31% below P across the records. The definition of S follows the "
        "Expert Team on Climate Change Detection and Indices [2]."
    ),
    28: "For each gauge and calendar year, annual maximum daily depth was defined as",
    30: (
        "Years with at least 300 recorded days were retained. Every gauge-year met this criterion, yielding 442 station-years (13 gauges x 34 years)."
    ),
    32: "Three distributions commonly used in engineering [13] were fitted to each annual maximum series by maximum likelihood. The Gumbel distribution [7] is",
    34: "where μ and σ are location and scale. The generalized extreme value (GEV) distribution [8] adds shape parameter ξ:",
    36: (
        "The GEV reduces to the Gumbel distribution as ξ approaches zero. The log-Pearson type III (LP3) distribution was fitted by maximum likelihood "
        "to z = ln(x) using the SciPy skewness, location, and scale parameterization. Its log-likelihood on the original scale includes the transformation Jacobian:"
    ),
    38: "Return levels from the LP3 model were calculated on the logarithmic scale and exponentiated.",
    39: "Models were compared at each gauge with the small-sample corrected Akaike information criterion (AICc) [14, 15]:",
    41: (
        "where k is the number of fitted parameters, n = 34, and L is the maximized likelihood. Equation (7) places the LP3 likelihood on the same scale as "
        "the other models. The uncorrected criterion selected the same model at all gauges. We report the gap to the second-ranked model (Delta AICc) and "
        "normalized Akaike weights. The Kolmogorov-Smirnov statistic was calculated descriptively as"
    ),
    43: (
        "where F(n) is the empirical distribution function. No p-value was assigned because parameters were estimated from the same sample. The model with "
        "the lowest AICc was used for return-level estimation; the other models were retained for sensitivity analysis."
    ),
    45: "The design rainfall for return period T is the quantile with annual exceedance probability 1/T:",
    47: (
        "Equation (10) was evaluated for T = 2, 5, 10, 25, 50, and 100 years. In each of 2,000 bootstrap replicates, the 34 annual maxima were resampled with "
        "replacement [16]. The fixed-selection bootstrap refitted the model selected from the original record. The selection-inclusive bootstrap refitted all "
        "three models and reselected by AICc within each replicate. Both procedures assume independent annual maxima. The 2.5th and 97.5th percentiles formed "
        "the interval, whose relative width was"
    ),
    49: "A value of 100% means that the full interval width equals the point estimate. This is resampling variability, not prediction error.",
    51: "Three ratios placed spatial, distributional, and resampling spread on comparable relative scales. Between-gauge spread was",
    53: "where i indexes the 13 gauges. Within-gauge spread across candidate distributions d was",
    55: "The median of Equation (11) represented resampling spread. These ratios are comparisons, not additive components of variance.",
    57: (
        "Spearman rank correlations related P, W, and S to R(T), w(T), and the GEV shape parameter xi. GEV shape was estimated at every gauge, regardless "
        "of the selected model, because shape parameters from the three distributions are not comparable. For ranks a(i) and b(i),"
    ),
    59: "with n = 13. These descriptive rank associations do not imply a mechanism.",
    61: (
        "One provincial value R(p)(T) was calculated as the mean, median, 75th percentile, or 90th percentile of the 13 gauge estimates. Its relative error "
        "at gauge i was"
    ),
    63: "Negative values indicate a provincial estimate below the local estimate; positive values indicate the reverse. The two cases were counted separately.",
    65: (
        "Stationary frequency analysis assumes independent annual maxima from one distribution. Dependence was examined with lag-one autocorrelation and "
        "the Ljung-Box statistic over five lags. Temporal stability was assessed with the Mann-Kendall test and Theil-Sen slope. These diagnostics qualify "
        "the stationary results; they do not fit dependent or nonstationary models. Two-sided significance was assessed at 5%."
    ),
    67: (
        "Calculations used Python 3.12.3 with NumPy 2.4.4, SciPy 1.17.1, pandas 3.0.2, Matplotlib 3.10.8, statsmodels 0.15.0, and pyMannKendall 1.4.3. "
        "Distributions were fitted with scipy.stats.gumbel_r, scipy.stats.genextreme, and scipy.stats.pearson3. SciPy's GEV shape c was reported as ξ = -c. "
        "Bootstrap seeds were 0 for fixed selection and 2024 for selection-inclusive analysis. Scripts, configuration, and result tables are available from "
        "the corresponding author on request."
    ),
    69: (
        "Because two gauges had substantially lower S than the others, sensitivity analyses considered all 13 gauges, exclusion of the two lowest-S gauges, "
        "and exclusion of the five lowest-S gauges. These scenarios are not corrections and do not define a quality threshold."
    ),
    70: (
        "Three checks assessed whether unusual occurrence-intensity records showed recording artifacts: terminal-digit frequencies among wet-day values, a "
        "chi-square test of wet-day counts by weekday, and rainfall recorded at each gauge on the 50 days when the other 12 gauges were wettest on average."
    ),
    73: (
        "The 13 records differed more in rainfall occurrence and intensity than in annual total (Table 1; Figure 1). Annual totals ranged from 944 to 1,376 mm "
        "(factor 1.46), wet-day counts from 79 to 229 d yr-1 (factor 2.91), and S from 5.11 to 13.06 mm (factor 2.56). Wet-day count and S were strongly inversely "
        "correlated (ρ = -0.978), so their variation largely offset in annual totals."
    ),
    74: (
        "Record maxima ranged from 54.2 to 298.5 mm (factor 5.51). Gauges 351009 and 351010 recorded no day above 100 mm despite annual totals of 1,206 and "
        "1,177 mm, whereas four gauges recorded more than 20 such days."
    ),
    78: (
        "AICc selected LP3 at six gauges, Gumbel at five, and GEV at two (Table 2). The uncorrected criterion produced the same selections."
    ),
    79: (
        "Model separation was generally weak. Selected-model Akaike weights ranged from 0.43 to 0.99 (median 0.54), and 11 gauges had a second-ranked model "
        "within 2 AICc units. Gauge 351003 was the only case with clear support (weight 0.990; gap 9.65). Selection counts therefore identify the first-ranked "
        "model rather than an unequivocal distribution."
    ),
    82: (
        "Between-gauge differences increased with return period (Table 3; Figure 2). The ratio of maximum to minimum return level rose from 3.83 at 2 years "
        "to 5.56 at 10 years, 7.67 at 25 years, and 12.53 at 100 years; the 100-year estimates ranged from 57.8 to 724.7 mm."
    ),
    83: (
        "Median within-gauge spread across the three models increased from a factor of 1.09 at 2 years to 1.56 at 100 years. Thus, spatial spread exceeded model "
        "spread by approximately 3.5 times at 2 years and eight times at 100 years. This comparison applies only to these gauges and candidate distributions."
    ),
    88: (
        "Mean wet-day intensity was positively associated with return level, whereas wet-day count was negatively associated; both relationships weakened at "
        "longer return periods (Table 4). Correlation between S and return level declined from 0.73 at 2 years to 0.48 at 100 years; the corresponding values for "
        "wet-day count were -0.80 and -0.45. Mean annual total had weaker associations (-0.40 to -0.54)."
    ),
    89: (
        "None of the three descriptors was associated with bootstrap interval width (correlations -0.19 to 0.19) or GEV shape (-0.22 to 0.06; p = 0.47-0.86)."
    ),
    90: (
        "Gauge-level results confirm that low S did not imply low design rainfall. Gauge 351009 had the lowest S but a higher 100-year estimate than 351010, "
        "and gauge 351007 combined low S with a much higher estimate than 351003. These unadjusted correlations across 13 gauges are descriptive."
    ),
    93: (
        "Under fixed model selection, median relative interval width increased from 30.8% at 2 years to 44.5% at 10 years, 75.0% at 25 years, and 145.9% at "
        "100 years (Table 3; Figure 3). At 100 years, seven gauges exceeded 100%, with widths from 12.5% to 343.6%."
    ),
    94: (
        "With model reselection, median width reached 39.6%, 57.6%, 113.0%, and 217.7% at the same return periods, and 10 gauges exceeded 100% at 100 years. "
        "The selection-inclusive to fixed-width ratio ranged from 0.71 to 10.25 at 100 years, although its median was 1.11."
    ),
    95: (
        "The largest increases occurred where selection was least stable. At gauge 351012, near-equal support for Gumbel and LP3 widened the interval from "
        "44.3% to 454.0%; gauge 351004 widened from 39.2% to 318.1%. Reselection narrowed intervals at gauges 351006 and 351007. The modal bootstrap choice matched "
        "the full-record choice at 12 gauges; 351004 was the exception."
    ),
    96: (
        "These intervals describe variation produced by resampling 34 annual maxima. They exclude prediction error, bias, distributions outside the candidate "
        "set, nonstationarity, and regionalization. A 100-year value without an interval therefore conveys more precision than these records support."
    ),
    100: (
        "At 100 years, the provincial mean was 251.8 mm and fell below five local estimates by as much as 65.3%; it exceeded eight by as much as 335.3%. Median "
        "absolute error was 25.8% (Table 5). At 25 years, the mean was 162.2 mm, with four gauges below by up to 55.7% and nine above by up to 240.2%."
    ),
    101: (
        "Alternative provincial statistics shifted rather than removed the imbalance (Figure 4). The median reduced maximum overestimation but left six gauges "
        "below their local values. The 90th percentile left only two below, but maximum overestimation rose to 486.1% and median absolute error to 53.2%. At "
        "100 years, median representation error (25.8%) remained smaller than median fixed-selection interval width (145.9%); these measures are not additive."
    ),
    106: (
        "Excluding the two lowest-S gauges reduced the 100-year between-gauge ratio from 12.53 to 4.70 and raised the network mean from 251.8 to 285.3 mm. "
        "Excluding the five lowest-S gauges left the ratio unchanged and raised the mean to 290.3 mm (Table 3)."
    ),
    107: (
        "The unchanged ratio shows that S did not identify the gauges defining the range. The scenarios demonstrate sensitivity to network composition but do "
        "not establish that any excluded record is erroneous."
    ),
    109: (
        "Diagnostics supported independence and stationarity at most gauges but flagged several departures. Lag-one autocorrelation ranged from -0.344 to "
        "+0.492. Ljung-Box tests were significant at gauges 351006, 351007, and 351201. Mann-Kendall tests identified a decrease of 2.226 mm yr-1 at 351003 and "
        "increases of 1.307, 1.111, and 0.670 mm yr-1 at 351004, 351007, and 351006, respectively."
    ),
    110: (
        "Gauges 351006 and 351007 were flagged by both diagnostics, suggesting that temporal change may contribute to serial association. Gauge 351003 had the "
        "strongest trend but no significant serial dependence (Ljung-Box p = 0.24), so the tests are not interchangeable."
    ),
    111: (
        "No gauge was excluded on this basis. Instead, the diagnostics qualify the stationary return levels, especially because gauges 351003 and 351007 help "
        "define the low and high ends of the design range. Nonstationary models provide an alternative when temporal stability is untenable [17], but were not "
        "fitted here."
    ),
    113: (
        "Terminal-zero frequencies ranged from 10.1% to 18.0%, and terminal-five frequencies from 9.4% to 14.1%; the two lowest-S gauges lay within both ranges. "
        "Weekday wet-day tests gave p = 0.29-0.998 at all gauges, including 0.997 and 0.998 at the two lowest-S gauges."
    ),
    114: (
        "On the 50 days when the other gauges were wettest on average, gauges 351009 and 351010 recorded medians of 21.2 and 20.4 mm and exceeded 20 mm on 58% "
        "and 54% of days. Four gauges with higher S had lower medians, including 5.3 mm at gauge 351002."
    ),
    115: (
        "The checks found no evidence of the examined artifacts: coarse rounding, weekday-linked observation, or systematic failure to register network rainfall. "
        "Rainfall records alone cannot identify the cause of the unusual structure; instrument, exposure, and siting histories were unavailable."
    ),
    117: (
        "Three implications follow. First, local gauge estimates should be preferred because a provincial value produced a median 100-year error of about 25% "
        "and site shortfalls above 50%. Second, intervals should accompany long-return-period estimates: seven gauges had fixed-selection widths above 100%, and "
        "changing among three distributions did not remove the dominant sampling spread. Longer records or regional frequency analysis [18] may increase effective "
        "information if regional homogeneity is established. Third, unusual occurrence-intensity structure alone is insufficient grounds for exclusion because the "
        "most unusual records passed all three artifact checks."
    ),
    119: (
        "The analysis used 13 gauges across a topographically varied province and did not interpolate between them. Thirty-four years provide limited support for "
        "100-year quantiles. Results apply only to daily annual maxima and the three fitted distributions [9, 12]. Independent resampling does not preserve serial "
        "dependence, so intervals at three flagged gauges require additional caution; block resampling could address this assumption. Nonparametric bootstrap performance "
        "for extremes can also be sensitive to sample size [19]. Fixed selection omits model-selection uncertainty, whereas the second bootstrap includes reselection "
        "within the candidate set. The 39 rank correlations were unadjusted for multiple comparisons and are descriptive. All return levels remain stationary-model estimates "
        "conditional on 1981-2014."
    ),
    121: "Analysis of 13 gauges in Uttaradit Province produced four main conclusions.",
    122: (
        "Spatial heterogeneity exceeded sensitivity to distribution choice and increased with return period: between-gauge ratios rose from 3.83 at 2 years to "
        "12.53 at 100 years, compared with a median within-gauge ratio of 1.56 at 100 years."
    ),
    123: (
        "Occurrence-intensity structure explained shorter-period estimates better than the upper tail. Correlation between wet-day intensity and return level declined "
        "from 0.73 at 2 years to 0.48 at 100 years, and record descriptors did not explain interval width or GEV shape."
    ),
    124: (
        "Long-period estimates were weakly constrained. Median 100-year interval width was 145.9% under fixed selection and 217.7% with reselection; the gauge-specific "
        "ratio between them ranged from 0.71 to 10.25."
    ),
    125: (
        "The provincial mean fell below five local 100-year estimates by as much as 65.3%, while no alternative provincial statistic removed both underestimation and "
        "overestimation."
    ),
    126: (
        "These stationary results are conditional on 1981-2014, three candidate distributions, and independent resampling. Design practice should therefore use local "
        "records where available, report uncertainty intervals, and avoid excluding gauges solely because their occurrence-intensity structure is unusual."
    ),
    128: (
        "The author thanks the Royal Irrigation Department and the Thai Meteorological Department for providing the rain-gauge records. No external funding was received."
    ),
    130: (
        "S.P.: conceptualization, methodology, software, formal analysis, data curation, visualization, writing-original draft, and writing-review and editing. The author "
        "approved the submitted manuscript."
    ),
    132: "The author declares no conflict of interest.",
    134: (
        "During preparation of this manuscript, the author used OpenAI Codex (OpenAI, accessed September 2026) for language editing and document formatting. All "
        "AI-assisted revisions were critically reviewed, verified, and revised by the author. The author accepts full responsibility for the published work."
    ),
}


CAPTIONS = {
    18: ("Table 1", "Characteristics of 13 rain gauges in Uttaradit Province and occurrence-intensity descriptors for 1981-2014. P, W, and S are defined in Equations (1)-(3)."),
    76: ("Figure 1", "Occurrence-intensity structure of the 13 records: (a) annual total versus wet-day count, (b) wet-day intensity versus wet-day count, and (c) record maximum versus wet-day intensity."),
    80: ("Table 2", "AICc values, separation from the second-ranked model, selected-model Akaike weight, and selected distribution at each gauge."),
    84: ("Table 3", "Return-level spread and bootstrap interval width by return period, followed by sensitivity scenarios for network composition."),
    86: ("Figure 2", "Gauge-specific design rainfall and 95% fixed-selection bootstrap intervals. Color identifies the selected distribution; both axes are logarithmic."),
    91: ("Table 4", "Spearman correlations between occurrence-intensity descriptors and return level, relative interval width, and GEV shape across 13 gauges."),
    98: ("Figure 3", "Spatial, distributional, and bootstrap spread across return periods: (a) between-gauge and within-gauge ratios, (b) median relative interval widths, and (c) selection-inclusive to fixed-width ratios."),
    102: ("Table 5", "Provincial representation error at 25- and 100-year return periods for four provincial statistics."),
    104: ("Figure 4", "Gauge-level representation error at a 100-year return period for four provincial statistics. Negative values indicate underestimation relative to the local estimate."),
}


HEADINGS_1 = {
    9: "1 Introduction",
    15: "2 Materials and methods",
    71: "3 Results and discussion",
    120: "4 Conclusions",
    127: "5 Acknowledgements",
    129: "6 Author contributions",
    131: "7 Conflict of interest",
    133: "8 Declaration of the use of generative AI and AI-assisted technologies",
    135: "References",
}


HEADINGS_2 = {
    16: "2.1 Study area and data",
    19: "2.2 Description of each record",
    27: "2.3 Annual maximum series",
    31: "2.4 Distributions fitted",
    44: "2.5 Return levels and uncertainty",
    50: "2.6 Comparison of spread",
    56: "2.7 Association between record structure and design estimate",
    60: "2.8 Provincial representation error",
    64: "2.9 Diagnostics of annual maxima",
    66: "2.10 Software and reproducibility",
    68: "2.11 Sensitivity scenarios and artifact checks",
    72: "3.1 Occurrence-intensity structure",
    77: "3.2 Distribution selection",
    81: "3.3 Spatial and distributional spread",
    87: "3.4 Associations with record structure",
    92: "3.5 Bootstrap spread of return levels",
    99: "3.6 Error from a provincial design value",
    105: "3.7 Sensitivity to gauge exclusion",
    108: "3.8 Independence and stationarity",
    112: "3.9 Checks on unusual records",
    116: "3.10 Implications for design practice",
    118: "3.11 Limitations",
}


REFERENCES = [
    "[1] Serinaldi, F., Kilsby, C.G. Stationarity is undead: Uncertainty dominates the distribution of extremes. Advances in Water Resources, 2015, 77, 17-36. https://doi.org/10.1016/j.advwatres.2014.12.013",
    "[2] Zhang, X., Alexander, L., Hegerl, G.C., Jones, P., Klein Tank, A., Peterson, T.C., ..., Zwiers, F.W. Indices for monitoring changes in extremes based on daily temperature and precipitation data. WIREs Climate Change, 2011, 2(6), 851-870. https://doi.org/10.1002/wcc.147",
    "[3] Limjirakan, S., Limsakul, A., Sriburi, T. Trends in temperature and rainfall extreme changes in Bangkok Metropolitan Area. Applied Environmental Research, 2010, 32(1), 31-48.",
    "[4] Limsakul, A., Singhruck, P. Long-term trends and variability of total and extreme precipitation in Thailand. Atmospheric Research, 2016, 169, 301-317. https://doi.org/10.1016/j.atmosres.2015.10.015",
    "[5] Manton, M.J., Della-Marta, P.M., Haylock, M.R., Hennessy, K.J., Nicholls, N., Chambers, L.E., ..., Yee, D. Trends in extreme daily rainfall and temperature in Southeast Asia and the South Pacific: 1961-1998. International Journal of Climatology, 2001, 21(3), 269-284. https://doi.org/10.1002/joc.610",
    "[6] Alexander, L.V., Zhang, X., Peterson, T.C., Caesar, J., Gleason, B., Klein Tank, A.M.G., ..., Vazquez-Aguirre, J.L. Global observed changes in daily climate extremes of temperature and precipitation. Journal of Geophysical Research: Atmospheres, 2006, 111(D5), D05109. https://doi.org/10.1029/2005JD006290",
    "[7] Gumbel, E.J. Statistics of extremes. New York, NY, USA: Columbia University Press, 1958.",
    "[8] Coles, S. An introduction to statistical modeling of extreme values. London, UK: Springer, 2001. https://doi.org/10.1007/978-1-4471-3675-0",
    "[9] Katz, R.W., Parlange, M.B., Naveau, P. Statistics of extremes in hydrology. Advances in Water Resources, 2002, 25(8-12), 1287-1304. https://doi.org/10.1016/S0309-1708(02)00056-8",
    "[10] Prahadchai, T., Busababodhin, P., Park, J.-S. Regional flood frequency analysis of extreme rainfall in Thailand, based on L-moments. Communications for Statistical Applications and Methods, 2024, 31(1), 37-53. https://doi.org/10.29220/CSAM.2024.31.1.037",
    "[11] Yamoat, N., Hanchoowong, R., Yamoad, O., Chaimoon, N., Kangrang, A. Estimation of regional intensity-duration-frequency relationships of extreme rainfall by simple scaling in Thailand. Journal of Water and Climate Change, 2023, 14(3), 796-810. https://doi.org/10.2166/wcc.2023.430",
    "[12] Papalexiou, S.M., Koutsoyiannis, D. Battle of extreme value distributions: A global survey on extreme daily rainfall. Water Resources Research, 2013, 49(1), 187-201. https://doi.org/10.1029/2012WR012557",
    "[13] Kite, G.W. Frequency and risk analyses in hydrology. Littleton, CO, USA: Water Resources Publications, 1977.",
    "[14] Akaike, H. A new look at the statistical model identification. IEEE Transactions on Automatic Control, 1974, 19(6), 716-723. https://doi.org/10.1109/TAC.1974.1100705",
    "[15] Burnham, K.P., Anderson, D.R. Model selection and multimodel inference: A practical information-theoretic approach. 2nd ed. New York, NY, USA: Springer, 2002. https://doi.org/10.1007/b97636",
    "[16] Efron, B., Tibshirani, R.J. An introduction to the bootstrap. New York, NY, USA: Chapman and Hall, 1993. https://doi.org/10.1007/978-1-4899-4541-9",
    "[17] Cheng, L., AghaKouchak, A. Nonstationary precipitation intensity-duration-frequency curves for infrastructure design in a changing climate. Scientific Reports, 2014, 4, 7093. https://doi.org/10.1038/srep07093",
    "[18] Hosking, J.R.M., Wallis, J.R. Regional frequency analysis: An approach based on L-moments. Cambridge, UK: Cambridge University Press, 1997. https://doi.org/10.1017/CBO9780511529443",
    "[19] Kyselý, J. A cautionary note on the use of nonparametric bootstrap for estimating uncertainties in extreme-value models. Journal of Applied Meteorology and Climatology, 2008, 47(12), 3236-3251. https://doi.org/10.1175/2008JAMC1763.1",
]


def set_run_font(run, size=12, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_text(paragraph, text, bold=False, italic=False):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, 12, bold=bold, italic=italic)
    return run


def set_caption(paragraph, label, text):
    clear_paragraph(paragraph)
    r1 = paragraph.add_run(label)
    set_run_font(r1, 12, bold=True, italic=False)
    r2 = paragraph.add_run(" " + text)
    set_run_font(r2, 12, bold=False, italic=False)


def set_affiliation_block(paragraphs):
    set_text(paragraphs[0], "Rainfall Occurrence-Intensity Heterogeneity and Its Implications for Provincial Extreme Rainfall Design", bold=True)
    clear_paragraph(paragraphs[1])
    r = paragraphs[1].add_run("Surasit Punyawansiri")
    set_run_font(r, 12, bold=True)
    r = paragraphs[1].add_run("1,*")
    set_run_font(r, 12, bold=True)
    r.font.superscript = True
    clear_paragraph(paragraphs[2])
    r = paragraphs[2].add_run("1")
    set_run_font(r, 12)
    r.font.superscript = True
    r = paragraphs[2].add_run(" Office of Water Management and Hydrology, Royal Irrigation Department, 811 Samsen Road, Dusit, Bangkok 10300, Thailand")
    set_run_font(r, 12)
    set_text(paragraphs[3], "*Corresponding author. E-mail: Surasit.pu@ku.th")


def set_style_font(style, size, bold=None, italic=None):
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def set_cell_margins(cell, top=40, start=50, bottom=40, end=50):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = trPr.find(qn("w:tblHeader"))
    if tblHeader is None:
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)
    tblHeader.set(qn("w:val"), "true")


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("left", "right", "insideV", "insideH"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")
    for edge in ("top", "bottom"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "000000")
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn("w:shd"))
        if shd is not None:
            tcPr.remove(shd)
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            tcPr.append(tcBorders)
        bottom = tcBorders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            tcBorders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")


def format_table(table, widths, font_size=9, no_wrap_body=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    grid_cols = table._tbl.tblGrid.gridCol_lst
    for ci, width in enumerate(widths):
        if ci < len(grid_cols):
            grid_cols[ci].set(qn("w:w"), str(int(width * 1440)))
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    set_repeat_header(table.rows[0])
    set_table_borders(table)
    for ri, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cant_split = trPr.find(qn("w:cantSplit"))
        if cant_split is None:
            trPr.append(OxmlElement("w:cantSplit"))
        for ci, cell in enumerate(row.cells):
            if ci < len(widths):
                cell.width = Inches(widths[ci])
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(int(widths[ci] * 1440)))
                tcW.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if no_wrap_body and ri > 0:
                tcPr = cell._tc.get_or_add_tcPr()
                no_wrap = tcPr.find(qn("w:noWrap"))
                if no_wrap is None:
                    tcPr.append(OxmlElement("w:noWrap"))
            shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            if shd is not None:
                cell._tc.get_or_add_tcPr().remove(shd)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci > 0) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, font_size, bold=(True if ri == 0 else None))


def add_page_number(paragraph):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 10)


def add_line_numbering(section):
    sectPr = section._sectPr
    old = sectPr.find(qn("w:lnNumType"))
    if old is not None:
        sectPr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:restart"), "continuous")
    sectPr.append(ln)


def contains_math(paragraph):
    return bool(paragraph._p.xpath(".//m:oMath | .//m:oMathPara"))


def contains_drawing(paragraph):
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def remove_paragraph_borders(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is not None:
        pPr.remove(pBdr)


def build(source, output):
    source = Path(source)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(source))
    if len(doc.paragraphs) != 155 or len(doc.tables) != 5:
        raise RuntimeError("Unexpected source structure; refusing index-based edit")

    # Page and section system.
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        section.different_first_page_header_footer = False
        add_line_numbering(section)
        add_page_number(section.footer.paragraphs[0])

    # Styles.
    normal = doc.styles["Normal"]
    set_style_font(normal, 12, bold=False, italic=False)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    title = doc.styles["Title"]
    set_style_font(title, 14, bold=True, italic=False)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    remove_paragraph_borders(doc.paragraphs[0])

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 12, bold=True, italic=False)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(0)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 12, bold=True, italic=True)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(0)
    h2.paragraph_format.keep_with_next = True

    pars = doc.paragraphs
    set_affiliation_block(pars)
    set_text(pars[4], "Abstract", bold=True)

    for idx, text in REPLACEMENTS.items():
        set_text(pars[idx], text)
    for idx, text in HEADINGS_1.items():
        set_text(pars[idx], text, bold=True)
        pars[idx].style = h1
    for idx, text in HEADINGS_2.items():
        set_text(pars[idx], text, bold=True, italic=True)
        pars[idx].style = h2
    for idx, (label, text) in CAPTIONS.items():
        set_caption(pars[idx], label, text)

    for i, reference in enumerate(REFERENCES, start=136):
        set_text(pars[i], reference)

    # Paragraph formatting by semantic role.
    centered = {0, 1, 2, 3}
    no_indent = centered | {4, 5, 7} | set(HEADINGS_1) | set(HEADINGS_2) | set(CAPTIONS)
    for idx, p in enumerate(pars):
        pf = p.paragraph_format
        if idx not in set(HEADINGS_1) | set(HEADINGS_2):
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
        if idx in centered:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
        elif idx in no_indent:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx != 0 else WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
        elif idx >= 136:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Inches(0.25)
            pf.first_line_indent = Inches(-0.25)
        elif contains_math(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
            pf.keep_together = True
        elif contains_drawing(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
            pf.line_spacing = 1.0
            pf.keep_together = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.first_line_indent = Inches(0.5)
        if idx in CAPTIONS:
            pf.keep_with_next = True if CAPTIONS[idx][0].startswith("Table") else False
        remove_paragraph_borders(p)
        for run in p.runs:
            if not contains_math(p):
                set_run_font(run, 12)

    # Title-page refinements.
    pars[0].style = title
    pars[0].paragraph_format.keep_with_next = True
    for idx in (1, 2, 3):
        pars[idx].paragraph_format.keep_with_next = True
    pars[4].paragraph_format.space_before = Pt(12)
    pars[4].paragraph_format.keep_with_next = True
    pars[5].paragraph_format.first_line_indent = Pt(0)
    pars[7].paragraph_format.space_before = Pt(6)
    pars[7].paragraph_format.space_after = Pt(6)
    if pars[7].runs:
        prefix, rest = pars[7].text.split(":", 1)
        clear_paragraph(pars[7])
        rr = pars[7].add_run(prefix + ":")
        set_run_font(rr, 12, bold=True)
        rr = pars[7].add_run(rest)
        set_run_font(rr, 12)

    # Tables: 15 cm maximum width, no vertical rules or shading.
    widths = [
        [0.38, 0.68, 0.68, 0.71, 0.60, 0.78, 0.54, 0.66, 0.87],
        [0.72, 0.85, 0.78, 0.78, 0.65, 1.08, 1.04],
        [0.45, 0.63, 0.63, 0.68, 0.62, 0.60, 0.72, 0.78, 0.55, 0.24],
        [1.65, 0.65, 1.20, 1.20, 1.20],
        [0.48, 1.00, 0.74, 0.74, 0.74, 0.76, 0.76, 0.68],
    ]
    for ti, (table, w) in enumerate(zip(doc.tables, widths)):
        format_table(table, w, 9, no_wrap_body=(ti == 0))

    # Add thousands separators in Table 1 annual totals.
    for row in doc.tables[0].rows[1:]:
        text = row.cells[5].text.strip()
        if text.isdigit() and int(text) >= 1000:
            row.cells[5].text = f"{int(text):,}"
            for p in row.cells[5].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Pt(0)
                for run in p.runs:
                    set_run_font(run, 9)

    # Keep source images embedded and scale only if they exceed the text width.
    for shape in doc.inline_shapes:
        max_width = Inches(6.45)
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)

    # Ask Word to refresh PAGE fields on open.
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.core_properties.title = pars[0].text
    doc.core_properties.subject = "Original research manuscript prepared for Applied Environmental Research"
    doc.core_properties.author = "Surasit Punyawansiri"
    doc.core_properties.last_modified_by = "Surasit Punyawansiri"
    doc.core_properties.keywords = "design rainfall; extreme-value analysis; model uncertainty; rainfall intensity; spatial heterogeneity; Uttaradit"

    doc.save(str(output))
    print(str(output))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("output")
    args = ap.parse_args()
    build(args.source, args.output)
