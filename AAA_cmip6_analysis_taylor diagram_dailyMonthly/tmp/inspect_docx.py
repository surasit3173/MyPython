import argparse
import hashlib
import json
import os
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W_NS, "r": R_NS}


def q(tag):
    prefix, local = tag.split(":", 1)
    return "{%s}%s" % ({"w": W_NS, "r": R_NS}[prefix], local)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def twip_in(value):
    return None if value is None else round(int(value) / 1440, 4)


def safe_text(el):
    return "".join(el.xpath(".//w:t/text()", namespaces=NS)).strip()


def para_record(p, idx, scope):
    fmt = p.paragraph_format
    runs = []
    for r in p.runs:
        rpr = r._element.rPr
        fonts = {}
        if rpr is not None and rpr.rFonts is not None:
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                val = rpr.rFonts.get(qn(f"w:{key}"))
                if val:
                    fonts[key] = val
        runs.append({
            "text": r.text,
            "bold": r.bold,
            "italic": r.italic,
            "underline": None if r.underline is None else str(r.underline),
            "size_pt": None if r.font.size is None else round(r.font.size.pt, 2),
            "font": r.font.name,
            "fonts": fonts,
            "color": None if r.font.color.rgb is None else str(r.font.color.rgb),
            "superscript": r.font.superscript,
            "subscript": r.font.subscript,
        })
    return {
        "scope": scope,
        "index": idx,
        "style": p.style.name if p.style is not None else None,
        "text": p.text,
        "alignment": None if p.alignment is None else str(p.alignment),
        "keep_with_next": fmt.keep_with_next,
        "keep_together": fmt.keep_together,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
        "space_before_pt": None if fmt.space_before is None else round(fmt.space_before.pt, 2),
        "space_after_pt": None if fmt.space_after is None else round(fmt.space_after.pt, 2),
        "line_spacing": None if fmt.line_spacing is None else str(fmt.line_spacing),
        "left_indent_in": None if fmt.left_indent is None else round(fmt.left_indent.inches, 4),
        "right_indent_in": None if fmt.right_indent is None else round(fmt.right_indent.inches, 4),
        "first_line_indent_in": None if fmt.first_line_indent is None else round(fmt.first_line_indent.inches, 4),
        "runs": runs,
    }


def style_record(style):
    p = getattr(style, "paragraph_format", None)
    f = getattr(style, "font", None)
    el = style._element
    based = el.find(q("w:basedOn"))
    nxt = el.find(q("w:next"))
    return {
        "style_id": style.style_id,
        "name": style.name,
        "type": str(style.type),
        "builtin": style.builtin,
        "based_on": based.get(q("w:val")) if based is not None else None,
        "next": nxt.get(q("w:val")) if nxt is not None else None,
        "font": None if f is None else f.name,
        "size_pt": None if f is None or f.size is None else round(f.size.pt, 2),
        "bold": None if f is None else f.bold,
        "italic": None if f is None else f.italic,
        "color": None if f is None or f.color.rgb is None else str(f.color.rgb),
        "alignment": None if p is None or p.alignment is None else str(p.alignment),
        "space_before_pt": None if p is None or p.space_before is None else round(p.space_before.pt, 2),
        "space_after_pt": None if p is None or p.space_after is None else round(p.space_after.pt, 2),
        "line_spacing": None if p is None or p.line_spacing is None else str(p.line_spacing),
        "left_indent_in": None if p is None or p.left_indent is None else round(p.left_indent.inches, 4),
        "right_indent_in": None if p is None or p.right_indent is None else round(p.right_indent.inches, 4),
        "first_line_indent_in": None if p is None or p.first_line_indent is None else round(p.first_line_indent.inches, 4),
        "keep_with_next": None if p is None else p.keep_with_next,
        "keep_together": None if p is None else p.keep_together,
        "page_break_before": None if p is None else p.page_break_before,
    }


def inspect_docx(path, outdir):
    path = Path(path)
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    doc = Document(str(path))

    paragraphs = [para_record(p, i, "body") for i, p in enumerate(doc.paragraphs)]
    for si, sec in enumerate(doc.sections):
        for i, p in enumerate(sec.header.paragraphs):
            paragraphs.append(para_record(p, i, f"header-{si}"))
        for i, p in enumerate(sec.first_page_header.paragraphs):
            paragraphs.append(para_record(p, i, f"first-header-{si}"))
        for i, p in enumerate(sec.even_page_header.paragraphs):
            paragraphs.append(para_record(p, i, f"even-header-{si}"))
        for i, p in enumerate(sec.footer.paragraphs):
            paragraphs.append(para_record(p, i, f"footer-{si}"))
        for i, p in enumerate(sec.first_page_footer.paragraphs):
            paragraphs.append(para_record(p, i, f"first-footer-{si}"))
        for i, p in enumerate(sec.even_page_footer.paragraphs):
            paragraphs.append(para_record(p, i, f"even-footer-{si}"))

    tables = []
    for ti, table in enumerate(doc.tables):
        rec = {
            "index": ti,
            "style": table.style.name if table.style is not None else None,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "autofit": table.autofit,
            "cells": [],
        }
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                rec["cells"].append({
                    "row": ri,
                    "col": ci,
                    "width_in": None if cell.width is None else round(cell.width.inches, 4),
                    "vertical_alignment": None if cell.vertical_alignment is None else str(cell.vertical_alignment),
                    "text": cell.text,
                    "paragraphs": [para_record(p, pi, f"table-{ti}-r{ri}c{ci}") for pi, p in enumerate(cell.paragraphs)],
                })
        tables.append(rec)

    sections = []
    for si, s in enumerate(doc.sections):
        sectPr = s._sectPr
        cols = sectPr.find(q("w:cols"))
        sections.append({
            "index": si,
            "start_type": str(s.start_type),
            "page_width_in": round(s.page_width.inches, 4),
            "page_height_in": round(s.page_height.inches, 4),
            "orientation": str(s.orientation),
            "top_margin_in": round(s.top_margin.inches, 4),
            "bottom_margin_in": round(s.bottom_margin.inches, 4),
            "left_margin_in": round(s.left_margin.inches, 4),
            "right_margin_in": round(s.right_margin.inches, 4),
            "header_distance_in": round(s.header_distance.inches, 4),
            "footer_distance_in": round(s.footer_distance.inches, 4),
            "gutter_in": round(s.gutter.inches, 4),
            "different_first_page": s.different_first_page_header_footer,
            "columns": None if cols is None else {
                "num": cols.get(q("w:num"), "1"),
                "space_twips": cols.get(q("w:space")),
                "equal_width": cols.get(q("w:equalWidth")),
                "separator": cols.get(q("w:sep")),
                "items": [{k.split('}')[-1]: v for k, v in c.attrib.items()} for c in cols.findall(q("w:col"))],
            },
        })

    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        pkg = [{"path": n, "size": z.getinfo(n).file_size, "sha256": hashlib.sha256(z.read(n)).hexdigest()} for n in names]
        xml_counts = Counter()
        all_text = []
        xml_parts = [n for n in names if n.startswith("word/") and n.endswith(".xml")]
        for name in xml_parts:
            try:
                root = etree.fromstring(z.read(name))
            except Exception:
                continue
            for tag in ["sdt", "fldSimple", "instrText", "ins", "del", "commentRangeStart", "footnoteReference", "endnoteReference", "drawing", "pict", "tbl", "sectPr"]:
                xml_counts[f"{name}:{tag}"] = len(root.xpath(f".//w:{tag}", namespaces=NS))
            text = "\n".join(t.strip() for t in root.xpath(".//w:t/text()", namespaces=NS) if t.strip())
            if text:
                all_text.append({"part": name, "text": text})

        media_dir = outdir / "media"
        media_dir.mkdir(exist_ok=True)
        media = []
        for n in names:
            if n.startswith("word/media/"):
                dest = media_dir / Path(n).name
                dest.write_bytes(z.read(n))
                media.append({"part": n, "size": z.getinfo(n).file_size, "out": str(dest)})

    styles = [style_record(s) for s in doc.styles]
    core = doc.core_properties
    report = {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "core_properties": {
            "title": core.title,
            "subject": core.subject,
            "author": core.author,
            "last_modified_by": core.last_modified_by,
            "created": None if core.created is None else core.created.isoformat(),
            "modified": None if core.modified is None else core.modified.isoformat(),
            "keywords": core.keywords,
            "comments": core.comments,
        },
        "sections": sections,
        "paragraph_count": len(doc.paragraphs),
        "paragraphs": paragraphs,
        "tables": tables,
        "styles": styles,
        "style_usage": Counter(p["style"] for p in paragraphs),
        "package_parts": pkg,
        "xml_feature_counts": dict(xml_counts),
        "media": media,
        "all_part_text": all_text,
    }
    report["style_usage"] = dict(report["style_usage"])
    (outdir / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = []
    lines.append(f"PATH: {path}")
    lines.append(f"SHA256: {report['sha256']}")
    lines.append(f"SECTIONS: {json.dumps(sections, ensure_ascii=False)}")
    lines.append(f"PARAGRAPHS: {len(doc.paragraphs)} TABLES: {len(tables)} MEDIA: {len(media)}")
    lines.append("\n=== BODY PARAGRAPHS ===")
    for p in paragraphs:
        if p["scope"] == "body":
            lines.append(f"P{p['index']:04d} [{p['style']}] {p['text']}")
    lines.append("\n=== TABLES ===")
    for t in tables:
        lines.append(f"TABLE {t['index']} style={t['style']} rows={t['rows']} cols={t['cols']}")
        for c in t["cells"]:
            lines.append(f"  R{c['row']}C{c['col']}: {c['text'].replace(chr(10), ' | ')}")
    lines.append("\n=== HEADERS FOOTERS ===")
    for p in paragraphs:
        if p["scope"] != "body" and p["text"]:
            lines.append(f"{p['scope']} P{p['index']:04d} [{p['style']}] {p['text']}")
    lines.append("\n=== STYLE USAGE ===")
    for k, v in Counter(p["style"] for p in paragraphs).most_common():
        lines.append(f"{k}: {v}")
    (outdir / "summary.txt").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({
        "path": str(path),
        "sha256": report["sha256"],
        "sections": sections,
        "paragraphs": len(doc.paragraphs),
        "tables": len(tables),
        "media": media,
        "style_usage": report["style_usage"],
        "report": str(outdir / "report.json"),
        "summary": str(outdir / "summary.txt"),
    }, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("outdir")
    args = ap.parse_args()
    inspect_docx(args.docx, args.outdir)
