"""Cross-artifact reconciliation for one completed run."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys

import openpyxl
import pandas as pd
from docx import Document
from PIL import Image


EXPECTED_SHEETS = [
    "Summary",
    "Data Quality",
    "Network Results",
    "Station Results",
    "Bootstrap CI",
    "Method Simulation",
    "FDR Simulation",
    "Sources",
]


def docx_text(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def require(condition: bool, message: str, checks: list[dict]) -> None:
    checks.append({"check": message, "pass": bool(condition)})
    if not condition:
        raise AssertionError(message)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("run_directory", type=Path)
    args = parser.parse_args()
    run_dir = args.run_directory.resolve()
    checks: list[dict] = []
    expected_files = [
        "run_manifest.json",
        "Prachuap_Rainfall_Trend_Results.xlsx",
        "Article_1_Method_Comparison.docx",
        "Article_2_Prachuap_Rainfall_Trends.docx",
        "tables/network_results.csv",
        "tables/station_results.csv",
        "tables/bootstrap_ci.csv",
        "tables/method_simulation.csv",
        "tables/fdr_simulation.csv",
    ]
    for relative in expected_files:
        require((run_dir / relative).is_file(), f"required file exists: {relative}", checks)

    network = pd.read_csv(run_dir / "tables" / "network_results.csv", encoding="utf-8-sig")
    station = pd.read_csv(run_dir / "tables" / "station_results.csv", encoding="utf-8-sig")
    bootstrap = pd.read_csv(run_dir / "tables" / "bootstrap_ci.csv", encoding="utf-8-sig")
    simulation = pd.read_csv(run_dir / "tables" / "method_simulation.csv", encoding="utf-8-sig")
    fdr = pd.read_csv(run_dir / "tables" / "fdr_simulation.csv", encoding="utf-8-sig")
    require(len(network) == 12, "network table contains 3 periods x 4 methods", checks)
    require(len(station) == 144, "station table contains 12 stations x 3 periods x 4 methods", checks)
    require(int(station["reject_bh"].sum()) == 7, "station table has 7 method-specific BH decisions", checks)
    require(int(network["reject_bh"].sum()) == 0, "network table has no BH decision", checks)
    require(len(bootstrap) == 39, "bootstrap table contains 36 station and 3 network intervals", checks)
    require((simulation["reps"] == 10000).all(), "all method simulation cells use 10,000 replicates", checks)
    require((fdr["reps"] == 5000).all(), "all FDR rows use 5,000 replicate families", checks)
    require(((fdr["fdr"] - fdr["fwer"]).abs() < 1e-15).all(), "complete-null FDR equals FWER", checks)

    workbook_path = run_dir / "Prachuap_Rainfall_Trend_Results.xlsx"
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    require(workbook.sheetnames == EXPECTED_SHEETS, "workbook sheet topology matches specification", checks)
    summary = workbook["Summary"]
    require(summary["B4"].value == 12, "workbook summary station count is 12", checks)
    require(summary["B5"].value == 7, "workbook summary station BH count is 7", checks)
    require(summary["B6"].value == 0, "workbook summary network BH count is 0", checks)
    error_tokens = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    workbook_errors = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in error_tokens:
                    workbook_errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    require(not workbook_errors, "workbook contains no formula error tokens", checks)
    workbook.close()

    method_text = docx_text(run_dir / "Article_1_Method_Comparison.docx")
    applied_text = docx_text(run_dir / "Article_2_Prachuap_Rainfall_Trends.docx")
    for token in ("10,000", "5,000", "HR-MMK-3", "FDP", "0.379", "0.399", "0.010", "0.466"):
        require(token in method_text, f"method manuscript contains controlling token {token}", checks)
    for token in ("149,016", "79.0%", "0.83", "-1.38", "3.46", "7 method-specific", "500002", "500006"):
        require(token in applied_text, f"applied manuscript contains controlling token {token}", checks)

    for folder, pages in (("docx_1_rendered", 4), ("docx_2_rendered", 5)):
        rendered = sorted((run_dir / "validation" / folder).glob("page-*.png"))
        require(len(rendered) == pages, f"{folder} contains {pages} rendered pages", checks)
        for image_path in rendered:
            with Image.open(image_path) as image:
                require(image.width > 500 and image.height > 700, f"rendered page has usable dimensions: {image_path.name}", checks)
    previews = sorted((run_dir / "validation" / "workbook_previews").glob("*.png"))
    require(len(previews) == 8, "all eight workbook sheets were rendered", checks)
    formula_report = json.loads((run_dir / "validation" / "workbook_formula_errors.json").read_text(encoding="utf-8"))
    require(formula_report == [], "artifact-tool formula-error scan is empty", checks)

    report = {"verdict": "confirmed", "checks_passed": len(checks), "checks": checks}
    output = run_dir / "validation" / "artifact_validation.json"
    output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"verdict": report["verdict"], "checks_passed": len(checks), "output": str(output)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
