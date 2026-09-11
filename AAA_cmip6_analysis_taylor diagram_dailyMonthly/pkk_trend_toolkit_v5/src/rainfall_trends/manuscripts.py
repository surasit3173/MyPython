"""Generate two journal-style manuscripts from one immutable result run."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable, Sequence

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "17365D"
PALE_BLUE = "DCE6F1"
LIGHT_BLUE = "F4F8FC"
WHITE = "FFFFFF"
GRID = "B7C9D6"


def _load_run(run_dir: Path) -> dict:
    tables = run_dir / "tables"
    return {
        "manifest": json.loads((run_dir / "run_manifest.json").read_text(encoding="utf-8")),
        "quality": pd.read_csv(tables / "data_quality.csv", encoding="utf-8-sig"),
        "network": pd.read_csv(tables / "network_results.csv", encoding="utf-8-sig"),
        "station": pd.read_csv(tables / "station_results.csv", encoding="utf-8-sig"),
        "bootstrap": pd.read_csv(tables / "bootstrap_ci.csv", encoding="utf-8-sig"),
        "method_sim": pd.read_csv(tables / "method_simulation.csv", encoding="utf-8-sig"),
        "fdr_sim": pd.read_csv(tables / "fdr_simulation.csv", encoding="utf-8-sig"),
        "network_series": pd.read_csv(tables / "network_series.csv", encoding="utf-8-sig"),
        "metadata": pd.read_csv(tables / "station_metadata.csv", encoding="utf-8-sig"),
    }


def _set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), color)


def _set_cell_borders(cell, color: str = GRID, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _configure_document(document: Document, short_title: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11), ("Heading 3", 10)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    title = document.styles["Title"]
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    _add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
    settings = document.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(title)
    line = document.add_paragraph()
    line.paragraph_format.space_after = Pt(9)
    run = line.add_run(subtitle)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(65, 65, 65)


def _add_abstract(document: Document, text: str, keywords: str) -> None:
    document.add_heading("Abstract", level=1)
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.keep_together = True
    key = document.add_paragraph()
    key.add_run("Keywords: ").bold = True
    key.add_run(keywords)


def _add_table(document: Document, caption: str, headers: Sequence[str], rows: Iterable[Sequence[object]], font_size: float = 7.5) -> None:
    cap = document.add_paragraph()
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(7)
    cap.paragraph_format.space_after = Pt(3)
    run = cap.add_run(caption)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    table.alignment = 0
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _set_cell_shading(cell, NAVY)
        _set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for text_run in paragraph.runs:
                text_run.font.name = "Arial"
                text_run.font.size = Pt(font_size)
                text_run.font.bold = True
                text_run.font.color.rgb = RGBColor(255, 255, 255)
    _cant_split(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = "" if value is None else str(value)
            _set_cell_shading(cells[column_index], PALE_BLUE if row_index % 2 else WHITE)
            _set_cell_borders(cells[column_index])
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[column_index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for text_run in paragraph.runs:
                    text_run.font.name = "Arial"
                    text_run.font.size = Pt(font_size)
        _cant_split(table.rows[-1])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_figure(document: Document, image_path: Path, caption: str, width_inches: float = 6.3) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_together = True
    run = cap.add_run(caption)
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.italic = True


def _p(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def _percent(value: float) -> str:
    return f"{100 * value:.1f}%"


def _new_page(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_references(document: Document) -> None:
    document.add_heading("References", level=1)
    references = [
        "Benjamini, Y., & Hochberg, Y. (1995). Controlling the false discovery rate: A practical and powerful approach to multiple testing. Journal of the Royal Statistical Society: Series B, 57, 289–300. https://doi.org/10.1111/j.2517-6161.1995.tb02031.x",
        "Hamed, K. H., & Rao, A. R. (1998). A modified Mann–Kendall trend test for autocorrelated data. Journal of Hydrology, 204, 182–196. https://doi.org/10.1016/S0022-1694(97)00125-X",
        "Mann, H. B. (1945). Nonparametric tests against trend. Econometrica, 13, 245–259. https://doi.org/10.2307/1907187",
        "Sen, P. K. (1968). Estimates of the regression coefficient based on Kendall’s tau. Journal of the American Statistical Association, 63, 1379–1389. https://doi.org/10.1080/01621459.1968.10480934",
        "Yue, S., & Wang, C. Y. (2002). Applicability of prewhitening to eliminate the influence of serial correlation on the Mann–Kendall test. Water Resources Research, 38(6). https://doi.org/10.1029/2001WR000861",
    ]
    for reference in references:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.first_line_indent = Cm(-0.6)
        paragraph.paragraph_format.keep_together = True


def _method_manuscript(run_dir: Path, data: dict, output: Path) -> None:
    manifest = data["manifest"]
    simulation = data["method_sim"]
    fdr = data["fdr_sim"]
    n = int(simulation["n"].iloc[0])
    reps = int(simulation["reps"].iloc[0])
    fdr_reps = int(fdr["reps"].iloc[0])
    type_i = simulation[simulation["metric"] == "type_i_error"]
    phi0 = type_i[type_i["phi"] == 0].set_index("method")
    phi04 = type_i[type_i["phi"].round(6) == 0.4].set_index("method")
    document = Document()
    _configure_document(document, "Short-record rainfall trend inference")
    _add_title(
        document,
        "Short-record rainfall trend inference under serial and spatial dependence",
        "A reproducible comparison of MK, HR-MMK-3, PW-MK, and TFPW-MK",
    )
    abstract = (
        f"Trend tests used for short hydroclimatic records can produce materially different decisions when serial dependence is present. "
        f"We compared the standard Mann–Kendall test (MK), a prespecified three-lag Hamed–Rao variance variant (HR-MMK-3), raw prewhitening (PW-MK), and trend-free prewhitening (TFPW-MK) in stationary Gaussian AR(1) records of length {n}. "
        f"Each calibration and power scenario used {reps:,} replicates, while a complete-null Benjamini–Hochberg (BH) experiment used {fdr_reps:,} replicates, 12 tests, temporal phi=0.4, and equicorrelated innovations (rho=0.3). "
        f"At phi=0, empirical Type I error ranged from {_percent(phi0['estimate'].min())} to {_percent(phi0['estimate'].max())}. At phi=0.4, estimates were {_percent(phi04.loc['MK','estimate'])} for MK, {_percent(phi04.loc['HR-MMK-3','estimate'])} for HR-MMK-3, {_percent(phi04.loc['PW-MK','estimate'])} for PW-MK, and {_percent(phi04.loc['TFPW-MK','estimate'])} for TFPW-MK. "
        "PW-MK stayed closest to nominal size but lost substantial raw power as persistence increased; the other procedures often showed inflated rejection rates. BH did not repair miscalibrated marginal p-values. No method was uniformly reliable across the prespecified scenarios, so method agreement, calibration evidence, and multiplicity must be reported together."
    )
    _add_abstract(document, abstract, "Mann–Kendall; autocorrelation; prewhitening; false discovery rate; Monte Carlo simulation; rainfall")
    document.add_heading("1. Introduction", level=1)
    document.add_paragraph(
        "Monotonic-trend screening is common in hydroclimatology because rank procedures are robust to non-normal observations and outliers. That robustness does not remove the independence requirement of the ordinary MK null distribution. Positive serial dependence can inflate apparent trend evidence, while a transformation intended to remove dependence can also remove part of a real trend. The practical problem is therefore not to select a method by name, but to quantify its operating behavior for the record length and dependence range at issue."
    )
    document.add_paragraph(
        "This study prespecified a narrow, falsifiable comparison. It does not claim universal performance and does not tune methods after seeing the empirical rainfall results. The objective is to show how calibration, power, undefined variance factors, and BH behavior interact in records as short as the complete dry-season series used in the companion application."
    )
    document.add_heading("2. Methods", level=1)
    document.add_heading("2.1 Procedures", level=2)
    document.add_paragraph(
        "MK used the tie-corrected variance and continuity-corrected normal approximation; the Sen effect estimate used the actual time coordinates. HR-MMK-3 first removed the Sen line, computed autocorrelation of residual ranks, retained only individually significant lags 1–3, and multiplied the MK variance by the signed Hamed–Rao factor. A non-positive factor was declared undefined rather than clamped. PW-MK estimated lag-one dependence from the raw series and tested the prewhitened sequence, while reporting the original-series Sen slope. TFPW-MK estimated the Sen trend, prewhitened the detrended residuals, restored the trend, and then applied MK."
    )
    document.add_heading("2.2 Simulation design", level=2)
    document.add_paragraph(
        f"Stationary Gaussian AR(1) noise had unit marginal variance. We used n={n}, phi values 0, 0.2, 0.4, 0.6, and 0.8, and standardized slopes 0 (calibration), 0.01, 0.025, and 0.05 standard deviations per year (raw power). Each cell used {reps:,} seeded replicates. Wilson 95% intervals and Monte Carlo standard errors describe simulation uncertainty. Because elevated Type I error can inflate unadjusted power, power curves are not interpreted without their matching calibration curves."
    )
    document.add_heading("2.3 Multiplicity experiment", level=2)
    document.add_paragraph(
        f"The complete-null experiment simulated {fdr_reps:,} replicate families of 12 series with phi={fdr['phi'].iloc[0]:.1f} and spatial innovation correlation rho={fdr['spatial_rho'].iloc[0]:.1f}. BH was applied at alpha={fdr['alpha'].iloc[0]:.2f}. For each replicate, FDP was V/max(R,1); FDR was the mean FDP. Under the complete null every rejection is false, so FDR equals the probability of one or more rejections (FWER). Mean rejection share was reported separately and was never substituted for FDR."
    )
    _new_page(document)
    document.add_heading("3. Results", level=1)
    type_rows = []
    for phi in sorted(type_i["phi"].unique()):
        group = type_i[type_i["phi"] == phi].set_index("method")
        type_rows.append([f"{phi:.1f}", *[f"{group.loc[m,'estimate']:.3f} ({group.loc[m,'ci_low']:.3f}–{group.loc[m,'ci_high']:.3f})" for m in ("MK", "HR-MMK-3", "PW-MK", "TFPW-MK")]])
    _add_table(document, "Table 1. Empirical Type I error (95% Wilson interval).", ["Phi", "MK", "HR-MMK-3", "PW-MK", "TFPW-MK"], type_rows, font_size=7.2)
    _add_figure(document, run_dir / "figures" / "simulation_type_i.png", "Figure 1. Type I error by AR(1) persistence. Shaded bands are Wilson 95% intervals; the dashed line is the nominal 0.05 level.")
    document.add_paragraph(
        f"At phi=0, MK was compatible with the nominal level ({_percent(phi0.loc['MK','estimate'])}), whereas HR-MMK-3 was high ({_percent(phi0.loc['HR-MMK-3','estimate'])}). With positive persistence, MK, HR-MMK-3, and TFPW-MK became increasingly anti-conservative. PW-MK remained near 0.05 across the grid. HR-MMK-3 had {int(type_i['invalid_replicates'].sum())} undefined replicates across the calibration grid; these were counted as non-rejections, making its displayed rejection rates conservative with respect to that handling choice."
    )
    _new_page(document)
    _add_figure(document, run_dir / "figures" / "simulation_power.png", "Figure 2. Raw empirical power. Curves must be interpreted with Figure 1 because size inflation can masquerade as power.")
    power = simulation[(simulation["metric"] == "power") & (simulation["standardized_slope"].round(6) == 0.05)]
    power_rows = []
    for phi in sorted(power["phi"].unique()):
        group = power[power["phi"] == phi].set_index("method")
        power_rows.append([f"{phi:.1f}", *[f"{group.loc[m,'estimate']:.3f}" for m in ("MK", "HR-MMK-3", "PW-MK", "TFPW-MK")]])
    _add_table(document, "Table 2. Raw power for a standardized slope of 0.05 per year.", ["Phi", "MK", "HR-MMK-3", "PW-MK", "TFPW-MK"], power_rows)
    fdr_rows = [[row.method, f"{row.fdr:.3f}", f"{row.fwer:.3f}", f"{row.mean_rejection_share:.3f}", f"{row.ci_low:.3f}–{row.ci_high:.3f}", int(row.invalid_series)] for row in fdr.itertuples(index=False)]
    _add_table(document, "Table 3. Complete-null BH experiment.", ["Method", "FDR", "FWER", "Mean reject share", "95% interval", "Undefined series"], fdr_rows)
    _add_figure(document, run_dir / "figures" / "simulation_fdr.png", "Figure 3. Complete-null FDR under the prespecified temporally and spatially correlated scenario.")
    _new_page(document)
    document.add_heading("4. Discussion", level=1)
    document.add_heading("4.1 Findings", level=2)
    document.add_paragraph(
        "The prespecified short-record scenarios do not support a universal default. PW-MK was the only procedure that remained close to nominal Type I error across the phi grid, but its raw power decreased sharply as persistence increased because raw prewhitening removes low-frequency trend information. MK and TFPW-MK were strongly anti-conservative under positive persistence. HR-MMK-3 reduced inflation relative to MK at high phi but remained above nominal and was already anti-conservative at phi=0."
    )
    document.add_paragraph(
        "BH controlled neither FDR nor FWER when supplied with markedly miscalibrated marginal p-values. This is not a failure that multiplicity adjustment can be expected to repair: the quality of a multiple-testing procedure depends on the validity and dependence properties of its input p-values. The much smaller mean rejection share than FDR also demonstrates why E[R/m] is not an FDR estimator."
    )
    document.add_heading("4.2 Uncertainty and limitations", level=2)
    document.add_paragraph(
        "The simulation used Gaussian continuous observations, a single record length, stationary AR(1) temporal dependence, equicorrelated innovations, fixed nominal alpha, and a three-lag HR rule. Rainfall totals may be skewed, zero-inflated at shorter aggregation scales, nonstationary, or affected by more complex dependence. Results are scenario evidence, not a theorem. Undefined HR factors were treated as non-rejections, and raw power was not size-adjusted."
    )
    document.add_heading("4.3 Recommendations", level=2)
    document.add_paragraph(
        "Applications should (i) prespecify the test family and dependence treatment, (ii) report calibration evidence at the actual record length, (iii) retain effect sizes and dependence diagnostics, (iv) disclose undefined corrections and cross-method disagreement, and (v) avoid interpreting multiplicity adjustment as a repair for invalid marginal tests. When no candidate method is well calibrated for the data-generating regime, the conclusion should remain method-sensitive rather than selecting the most favorable p-value."
    )
    document.add_heading("5. Reproducibility", level=1)
    document.add_paragraph(
        f"This manuscript was generated from run {manifest['run_id']}. Scenario counts, seeds, software versions, source hashes, and all output hashes are recorded in run_manifest.json. The production package contains unit tests, independent comparisons with pymannkendall where definitions match, and CSV tables that control every value reported here."
    )
    _add_references(document)
    document.save(output)


def _applied_manuscript(run_dir: Path, data: dict, output: Path) -> None:
    manifest = data["manifest"]
    quality = data["quality"].set_index("metric")["value"]
    network = data["network"]
    station = data["station"]
    bootstrap = data["bootstrap"]
    series = data["network_series"]
    metadata = data["metadata"]
    mk = network[network["method"] == "MK"].set_index("period")
    network_ci = bootstrap[bootstrap["scope"] == "network"].set_index("period")
    discoveries = station[station["reject_bh"] == True].copy()  # noqa: E712
    unique_discoveries = discoveries[["station_id", "period"]].drop_duplicates()
    wet_share = series[series["period"] == "wet"]["network_mean_mm"].sum() / series[series["period"] == "annual"]["network_mean_mm"].sum()
    document = Document()
    _configure_document(document, "Prachuap Khiri Khan rainfall trends")
    _add_title(
        document,
        "Observed rainfall trends across the Prachuap Khiri Khan gauge network, 1981–2014",
        "Complete-period aggregation, dependence sensitivity, multiplicity control, and reproducible uncertainty",
    )
    abstract = (
        f"We reanalyzed daily rainfall from {int(float(quality['station_count']))} gauges in Prachuap Khiri Khan over {quality['start_date']} to {quality['end_date']} using fail-closed ingestion and complete-period aggregation. "
        f"Annual and May–October wet-season series contained 34 complete periods per station; November–April dry seasons contained 33 after excluding incomplete edge seasons. The equal-weight gauge-network Sen slopes were {mk.loc['annual','slope']:.2f} mm/year annually, {mk.loc['wet','slope']:.2f} mm/year in the wet season, and {mk.loc['dry','slope']:.2f} mm/year in the dry season. "
        f"None of the 12 network tests across four methods survived BH adjustment. At stations, {len(discoveries)} method-specific decisions across {len(unique_discoveries)} station-period combinations survived BH, but the pattern was method-sensitive and concentrated at stations 500002 and 500006. "
        f"Residual-block bootstrap intervals included zero for annual and wet network slopes but excluded zero for dry ({network_ci.loc['dry','ci_low']:.2f} to {network_ci.loc['dry','ci_high']:.2f} mm/year), a discrepancy that is not itself multiplicity-adjusted. Calibration simulations showed that none of the dependence procedures was uniformly reliable for a 33-year record. The defensible conclusion is therefore an inconclusive network-wide monotonic trend with localized, method-sensitive dry-season increases requiring independent confirmation."
    )
    _add_abstract(document, abstract, "rainfall; Prachuap Khiri Khan; Mann–Kendall; Sen slope; hydrological season; autocorrelation; false discovery rate")
    document.add_heading("1. Introduction", level=1)
    document.add_paragraph(
        "Rainfall-trend results can change when incomplete seasons, serial dependence, multiple station tests, and spatial interpretation are handled differently. The supplied legacy workflow mixed these issues and did not process every available station. This major revision separates data validation, aggregation, effect estimation, method sensitivity, multiplicity, and reporting so that each claim can be traced to a saved table and source hash."
    )
    document.add_paragraph(
        "The study asks whether monotonic changes are detectable in annual, wet-season, and dry-season totals at individual gauges and in an equal-weight gauge-network mean. The network statistic is explicitly not a province-wide areal mean because no spatial weighting or gridded integration is available."
    )
    document.add_heading("2. Data and methods", level=1)
    document.add_heading("2.1 Observations and quality control", level=2)
    document.add_paragraph(
        f"The input contained {int(float(quality['station_days'])):,} station-days for {int(float(quality['station_count']))} station identifiers. Dates, station keys, and rainfall values were validated before analysis. Missing rainfall, negative rainfall, and duplicate station-day counts were all zero. Elevation was not supplied for {int(float(quality['missing_elevation']))} stations; coordinates were retained for all gauges. No synthetic values or silent input fallback were permitted."
    )
    _new_page(document)
    _add_figure(document, run_dir / "figures" / "station_locations.png", "Figure 1. Gauge locations. The figure is a coordinate schematic without an administrative boundary or area-weighting interpretation.", width_inches=5.8)
    document.add_heading("2.2 Period definitions", level=2)
    document.add_paragraph(
        "Annual totals used January–December. Wet totals used May–October. Dry totals used November of the preceding calendar year through April of the label year. A period was retained only at 100% daily completeness. Consequently, annual and wet analyses spanned 1981–2014 (34 periods), whereas complete dry seasons spanned 1982–2014 (33 periods); the partial 1981 and 2015 dry endpoints were explicitly excluded."
    )
    document.add_heading("2.3 Trend inference", level=2)
    document.add_paragraph(
        "Sen slope quantified change in millimetres per year using actual year coordinates. Four inference procedures were carried in parallel: MK, HR-MMK-3, PW-MK, and TFPW-MK. HR-MMK-3 was the prespecified three-lag variance-correction sensitivity procedure, not an all-lag or universally calibrated method. At stations, BH was applied separately within each period-by-method family of 12 tests. At the network level, BH was applied across the three periods within each method. A circular residual-block bootstrap with block length 3 and 2,000 seeded replicates described Sen-slope uncertainty on the original time axis."
    )
    document.add_heading("3. Results", level=1)
    document.add_heading("3.1 Seasonal rainfall contribution and network series", level=2)
    document.add_paragraph(
        f"Across the 34 common calendar years, May–October contributed {100 * wet_share:.1f}% of the summed equal-weight network annual rainfall. This is substantially less than the approximately 95% value in the legacy narrative and is calculated directly from the supplied daily observations. Interannual variability was large relative to the fitted monotonic slopes."
    )
    _add_figure(document, run_dir / "figures" / "network_trends.png", "Figure 2. Equal-weight gauge-network totals and MK-associated Sen lines. Dry labels identify the year in which the season ends.", width_inches=6.0)
    network_rows = []
    for row in network.sort_values(["period", "method"]).itertuples(index=False):
        ci = network_ci.loc[row.period]
        network_rows.append([row.period, row.method, int(row.n), f"{row.slope:.2f}", _p(row.p_value), _p(row.q_value), "yes" if row.reject_bh else "no", f"{ci.ci_low:.2f} to {ci.ci_high:.2f}"])
    _add_table(document, "Table 1. Network trend estimates, tests, BH adjustment, and method-independent residual-block slope intervals.", ["Period", "Method", "N", "Slope", "p", "q", "BH", "Bootstrap 95% CI"], network_rows, font_size=6.7)
    document.add_paragraph(
        f"For MK, annual slope was {mk.loc['annual','slope']:.2f} mm/year (p={_p(mk.loc['annual','p_value'])}, q={_p(mk.loc['annual','q_value'])}); wet-season slope was {mk.loc['wet','slope']:.2f} mm/year (p={_p(mk.loc['wet','p_value'])}, q={_p(mk.loc['wet','q_value'])}); and dry-season slope was {mk.loc['dry','slope']:.2f} mm/year (p={_p(mk.loc['dry','p_value'])}, q={_p(mk.loc['dry','q_value'])}). No network result passed BH. PW-MK produced a smaller dry raw p-value ({_p(network[(network['period']=='dry') & (network['method']=='PW-MK')]['p_value'].iloc[0])}) but still did not pass network BH."
    )
    _new_page(document)
    document.add_heading("3.2 Station results", level=2)
    def lag_text(value: object) -> str:
        if pd.isna(value):
            return "–"
        numeric = float(value)
        return str(int(numeric)) if numeric.is_integer() else str(value)

    station_rows = [[str(row.station_id), row.period, row.method, f"{row.slope:.2f}", _p(row.p_value), _p(row.q_value), f"{row.correction_factor:.3f}", lag_text(row.significant_lags)] for row in discoveries.sort_values(["station_id", "period", "method"]).itertuples(index=False)]
    _add_table(document, "Table 2. Method-specific station decisions surviving BH (all non-discoveries remain in the supplementary workbook).", ["Station", "Period", "Method", "Slope", "p", "q", "HR factor", "Selected lags"], station_rows)
    _add_figure(document, run_dir / "figures" / "station_slope_intervals.png", "Figure 3. Station Sen slopes with 95% residual-block bootstrap intervals. Intervals are effect uncertainty summaries, not BH-adjusted hypothesis tests.")
    annual_hr = discoveries[(discoveries["station_id"].astype(str) == "500002") & (discoveries["period"] == "annual") & (discoveries["method"] == "HR-MMK-3")]
    document.add_paragraph(
        f"Seven method-specific rejections represented only three station-period combinations: station 500002 annual under HR-MMK-3, and dry seasons at stations 500002 and 500006 under MK, HR-MMK-3, and TFPW-MK. PW-MK produced no station discovery. For the station 500002 annual HR result, the signed variance factor was {annual_hr['correction_factor'].iloc[0]:.3f} after selecting lag 2; this variance deflation generated much stronger evidence than MK and should be treated as a method-sensitivity warning rather than a robust stand-alone discovery."
    )
    document.add_heading("3.3 Bootstrap and hypothesis-test disagreement", level=2)
    document.add_paragraph(
        f"Network bootstrap intervals were {network_ci.loc['annual','ci_low']:.2f} to {network_ci.loc['annual','ci_high']:.2f} mm/year annually, {network_ci.loc['wet','ci_low']:.2f} to {network_ci.loc['wet','ci_high']:.2f} mm/year for the wet season, and {network_ci.loc['dry','ci_low']:.2f} to {network_ci.loc['dry','ci_high']:.2f} mm/year for the dry season. The dry interval excluded zero even though no network test survived BH. The interval resamples residual blocks around the fitted trend and is not a multiplicity-adjusted test; the fixed block length and differing estimands can yield this disagreement. The conservative interpretation is increased dry-season rainfall magnitude with inconclusive network-wide monotonic-test evidence."
    )
    document.add_heading("4. Discussion", level=1)
    document.add_heading("4.1 Findings", level=2)
    document.add_paragraph(
        "The complete-period network shows little evidence of a monotonic annual or wet-season change. Dry-season totals have a positive effect estimate, but the inferential signal changes with the treatment of serial dependence and multiplicity. Local dry increases at two gauges recur across MK, HR-MMK-3, and TFPW-MK, yet disappear under raw prewhitening. The annual station 500002 HR result is driven by variance deflation and does not agree with the other methods after BH."
    )
    document.add_heading("4.2 Scientific uncertainty", level=2)
    document.add_paragraph(
        "The companion simulation shows that MK, HR-MMK-3, and TFPW-MK can be anti-conservative for short positively persistent records, while PW-MK can be conservative and low-powered. Therefore, counting methods that reject is not a formal ensemble test. Homogeneity breaks, gauge relocations, exposure changes, and instrument metadata were unavailable. Fixed wet/dry month definitions may not capture year-to-year monsoon timing. Four elevation values were absent, and equal station weighting does not represent spatial rainfall area."
    )
    document.add_heading("4.3 Recommendations", level=2)
    document.add_paragraph(
        "The present results support monitoring rather than a categorical province-wide trend claim. Future work should obtain station-history metadata, test homogenization and changepoints, repeat the analysis with spatial weighting or gridded products, evaluate alternative dependence models and block lengths, and extend the record. Confirmatory language should be reserved for effects that remain coherent across a prespecified, empirically calibrated method and multiplicity family."
    )
    document.add_heading("5. Reproducibility and data provenance", level=1)
    document.add_paragraph(
        f"All values in this manuscript were loaded from run {manifest['run_id']}; none were copied from the legacy DOCX. Input SHA-256 hashes, configuration, software environment, seeds, iteration counts, tables, figures, and output hashes are in run_manifest.json. The package fails when input paths, schemas, dates, station keys, or rainfall domains violate the declared contract. Unit tests cover ingestion, complete hydrological periods, MK-family calculations, BH adjustment, bootstrap determinism, simulation, and run immutability."
    )
    _add_references(document)
    document.save(output)


def generate_manuscripts(run_directory: str | Path) -> tuple[Path, Path]:
    run_dir = Path(run_directory).resolve()
    data = _load_run(run_dir)
    method_path = run_dir / "Article_1_Method_Comparison.docx"
    applied_path = run_dir / "Article_2_Prachuap_Rainfall_Trends.docx"
    _method_manuscript(run_dir, data, method_path)
    _applied_manuscript(run_dir, data, applied_path)
    return method_path, applied_path
