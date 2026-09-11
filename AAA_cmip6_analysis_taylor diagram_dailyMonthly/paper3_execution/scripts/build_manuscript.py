from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript" / "Paper3_CMJS_manuscript.md"
OUT = ROOT / "manuscript" / "Paper3_CMJS_manuscript.docx"
FIGURES = {
    1: ROOT / "output" / "figures" / "Paper3_Figure_01_ENSO_classification.png",
    2: ROOT / "output" / "figures" / "Paper3_Figure_02_observed_responses.png",
    3: ROOT / "output" / "figures" / "Paper3_Figure_03_raw_qdm_comparison.png",
    4: ROOT / "output" / "figures" / "Paper3_Figure_04_preservation_asymmetry.png",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_line_numbers(section) -> None:
    sect_pr = section._sectPr
    ln = sect_pr.find(qn("w:lnNumType"))
    if ln is None:
        ln = OxmlElement("w:lnNumType")
        sect_pr.append(ln)
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")


def style_run(run, size=11, bold=False, italic=False, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_inline(paragraph, text: str, size=11) -> None:
    text = text.replace("`", "")
    pattern = re.compile(r"(<sup>.*?</sup>|\*\*.*?\*\*|\*.*?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            style_run(paragraph.add_run(text[pos:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("<sup>"):
            run = paragraph.add_run(re.sub(r"</?sup>", "", token))
            style_run(run, size=max(size - 2, 7))
            run.font.superscript = True
        elif token.startswith("**"):
            style_run(paragraph.add_run(token[2:-2]), size=size, bold=True)
        else:
            style_run(paragraph.add_run(token[1:-1]), size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        style_run(paragraph.add_run(text[pos:]), size=size)


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.75)
    sec.right_margin = Cm(1.55)
    sec.header_distance = Cm(0.65)
    sec.footer_distance = Cm(0.65)
    add_line_numbers(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    pf.widow_control = True

    for name, size in (("Title", 15), ("Heading 1", 12), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_table(doc: Document, rows: list[list[str]], table_number: int) -> None:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, value in enumerate(rows[0]):
        cell = hdr.cells[j]
        set_cell_shading(cell, "D9EAF2")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, value, size=7.2 if table_number == 2 else 7.8)
        for run in p.runs:
            run.bold = True
    for row_values in rows[1:]:
        cells = table.add_row().cells
        for j, value in enumerate(row_values):
            cell = cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=28, bottom=28)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j < 3 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, size=6.8 if table_number == 2 else 7.5)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 1.0


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_document(doc)

    i = 0
    table_number = 0
    in_tables = False
    in_figures = False
    while i < len(lines):
        raw = lines[i].rstrip()
        if not raw or raw == "---":
            i += 1
            continue

        if raw.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, raw[2:], size=15)
            i += 1
            continue
        if raw == "## Tables":
            doc.add_page_break()
            in_tables = True
            p = doc.add_paragraph("Tables", style="Heading 1")
            i += 1
            continue
        if raw == "## Figure Captions":
            doc.add_page_break()
            in_figures = True
            p = doc.add_paragraph("Figures", style="Heading 1")
            i += 1
            continue
        if raw.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, raw[3:], size=12)
            i += 1
            continue
        if raw.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, raw[4:], size=11)
            i += 1
            continue

        if in_tables and raw.startswith("**Table "):
            table_number += 1
            if table_number > 1:
                doc.add_page_break()
            cap = doc.add_paragraph()
            cap.paragraph_format.line_spacing = 1.0
            cap.paragraph_format.space_after = Pt(3)
            add_inline(cap, raw, size=9)
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            rows: list[list[str]] = []
            if i < len(lines) and lines[i].lstrip().startswith("|"):
                while i < len(lines) and lines[i].lstrip().startswith("|"):
                    values = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r"-+:?", v.replace(":", "")) for v in values):
                        if not all(set(v) <= {"-", ":"} for v in values):
                            rows.append(values)
                    i += 1
                add_table(doc, rows, table_number)
            continue

        if in_figures and raw.startswith("**Figure "):
            match = re.match(r"\*\*Figure (\d+)\.\*\*\s*(.*)", raw)
            assert match
            num = int(match.group(1))
            if num == 3:
                doc.add_page_break()
            pic_p = doc.add_paragraph()
            pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic_p.paragraph_format.space_after = Pt(1)
            pic_p.paragraph_format.line_spacing = 1.0
            width = Inches(6.75)
            pic_p.add_run().add_picture(str(FIGURES[num]), width=width)
            cap = doc.add_paragraph()
            cap.paragraph_format.line_spacing = 1.0
            cap.paragraph_format.space_after = Pt(3)
            cap.paragraph_format.keep_with_next = num in (1, 3)
            add_inline(cap, raw, size=8.2)
            i += 1
            continue

        if raw.startswith("|"):
            i += 1
            continue

        p = doc.add_paragraph()
        if len(doc.paragraphs) <= 4:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
        if raw.startswith("**Keywords:**"):
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(4)
        if raw.startswith("\\["):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, raw.replace("\\[", "").replace("\\]", ""), size=11)
        i += 1

    props = doc.core_properties
    props.title = "ENSO-Conditioned Seasonal Rainfall Extremes and Signal Preservation over Uttaradit"
    props.author = "Surasit Punyawansiri"
    props.subject = "Chiang Mai Journal of Science research paper"
    props.keywords = "ENSO, CMIP6, quantile delta mapping, rainfall extremes, Thailand"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"MANUSCRIPT_DOCX_COMPLETE {OUT}")


if __name__ == "__main__":
    main()
