from __future__ import annotations

import calendar
import json
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "manuscript" / "Paper3_CMJS_manuscript.docx"
OUT = ROOT / "manuscript" / "Paper3_CMJS_Supplementary_Information.docx"
SECONDARY_CSV = ROOT / "output" / "tables" / "Paper3_Supplementary_Secondary_Responses.csv"
TOTAL_DXA = 10020
PHASE_LABEL = {
    "EL_NINO": "El Niño",
    "LA_NINA": "La Niña",
    "NEUTRAL": "Neutral",
    "TRANSITION_UNCLASSIFIED": "Transition/unclassified",
}
SEASON_LABEL = {"RAINY": "Rainy", "HOT_DRY": "Hot/dry"}


def set_font(run, size: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_table_geometry(table, widths: list[int]) -> None:
    assert sum(widths) == TOTAL_DXA
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout"):
        old = tbl_pr.find(qn(f"w:{tag}"))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(TOTAL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            set_cell_width(cell, width)


def caption(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(f"Supplementary Table {label}. "), 9, bold=True)
    set_font(p.add_run(text), 9)


def note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run("Note. "), 8, italic=True)
    set_font(p.add_run(text), 8)


def page_break(doc: Document) -> None:
    if doc.paragraphs and doc.paragraphs[-1].text == "":
        doc.paragraphs[-1].add_run().add_break()
    else:
        doc.add_page_break()


def add_table(
    doc: Document,
    rows: list[list[str]],
    widths: list[int],
    *,
    font_size: float = 7.4,
    left_columns: tuple[int, ...] = (0,),
) -> None:
    assert rows and len(widths) == len(rows[0])
    table = doc.add_table(rows=1, cols=len(widths))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_header(header)
    keep_row_together(header)
    for j, value in enumerate(rows[0]):
        cell = header.cells[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9EAF2")
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(value)), font_size, bold=True)
    for row_values in rows[1:]:
        row = table.add_row()
        keep_row_together(row)
        for j, value in enumerate(row_values):
            cell = row.cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=34, bottom=34)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in left_columns else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), font_size)
    set_table_geometry(table, widths)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def remove_line_numbering(doc: Document) -> None:
    """Supplementary table files do not require manuscript line numbers.

    Removing them also prevents Word 2024's line-number drawing layer from
    masking caption glyphs next to wide fixed-layout tables.
    """
    for section in doc.sections:
        sect_pr = section._sectPr
        line_numbers = sect_pr.find(qn("w:lnNumType"))
        if line_numbers is not None:
            sect_pr.remove(line_numbers)


def fmt(value, digits=1, missing="NE") -> str:
    if value is None or not np.isfinite(float(value)):
        return missing
    return f"{float(value):.{digits}f}"


def secondary_responses(metrics: pd.DataFrame, config: dict) -> pd.DataFrame:
    sys.path.insert(0, str(ROOT / "src"))
    from paper3core.statistics import model_response, observed_response

    rows = []
    for season in ("RAINY", "HOT_DRY"):
        for metric in config["metrics"]["secondary"]:
            for phase in ("EL_NINO", "LA_NINA"):
                record = {"season_type": season, "metric": metric, "phase": phase}
                for source in ("OBSERVED", "RAW", "QDM"):
                    d = metrics[(metrics.source_type == source) & (metrics.season_type == season) & metrics.complete].copy()
                    d["value"] = pd.to_numeric(d[metric], errors="coerce")
                    if source == "OBSERVED":
                        estimate, _ = observed_response(d, phase=phase, minimum_seasons=3)
                        record["observed_response_pct"] = estimate["response_pct"]
                        record["n_observed_stations"] = estimate["n_entities"]
                    else:
                        estimate, _, _ = model_response(d, phase=phase, minimum_seasons=3)
                        record[f"{source.lower()}_response_pct"] = estimate["response_pct"]
                        record[f"n_{source.lower()}_models"] = estimate["n_entities"]
                rows.append(record)
    out = pd.DataFrame(rows)
    out.to_csv(SECONDARY_CSV, index=False)
    return out


def main() -> None:
    assert REFERENCE.exists()
    shutil.copy2(REFERENCE, OUT)
    doc = Document(OUT)
    clear_body(doc)
    remove_line_numbering(doc)

    config = yaml.safe_load((ROOT / "config" / "paper3_enso.yaml").read_text(encoding="utf-8"))
    output = ROOT / "output"
    manifest = json.loads((ROOT / "inputs" / "enso" / "model_nino34_source_manifest.json").read_text(encoding="utf-8"))
    metrics = pd.read_csv(output / "seasonal_metrics_all.csv.gz")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("ENSO-Conditioned Seasonal Rainfall Extremes and Signal Preservation after\nCross-Fitted Quantile Delta Mapping over Uttaradit, Thailand"), 15, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.line_spacing = 1.0
    subtitle.paragraph_format.space_after = Pt(5)
    set_font(subtitle.add_run("SUPPLEMENTARY INFORMATION"), 12, bold=True)
    for text, size in (
        ("Surasit Punyawansiri¹,*", 11),
        ("¹Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok 10300, Thailand", 10),
        ("*Corresponding author: Surasit.ku@ku.th", 10),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(1)
        set_font(p.add_run(text), size)

    p = doc.add_paragraph(style="Heading 1")
    set_font(p.add_run("Scope"), 12, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    set_font(p.add_run(
        "These tables provide the model, station, event-classification, secondary-index, inference, "
        "bias-adjustment, quality-control, and provenance details supporting the main article. Values "
        "are generated from the frozen 1981-2014 analysis outputs. Secondary-index responses are "
        "descriptive and were not included in the pre-specified primary multiple-testing family. Full "
        "row-level QDM and quality-control diagnostics remain in the machine-readable reproducibility package."
    ), 11)

    # S1
    page_break(doc)
    caption(doc, "S1", "CMIP6 precipitation and exact-member Niño-3.4 source configuration.")
    grids = pd.read_csv(output / "raw_grid_signatures.csv").groupby("model").grid_signature.nunique().to_dict()
    s1 = [["Model", "Member", "pr grid", "tos grid", "tos version", "Niño-3.4 cells", "Unique pr grids"]]
    for item in manifest["datasets"]:
        s1.append([
            item["model"], item["member_id"], item["pr_grid_label"], item["tos_grid_label"],
            item["tos_version"], str(item["nino34_selected_grid_cells"]), str(grids[item["model"]]),
        ])
    add_table(doc, s1, [2100, 1350, 900, 900, 1350, 1650, 1770], font_size=7.8, left_columns=(0,))
    note(doc, "All model Niño-3.4 indices use area-weighted exact-member tos, a 1981-2010 monthly climatology, centred linear detrending, and a centred three-month mean. EC-Earth3 intentionally uses pr on gr and tos on gn.")

    # S2
    page_break(doc)
    caption(doc, "S2", "Uttaradit station metadata, probable-missing-month corrections, and complete management seasons.")
    obs = pd.read_csv(ROOT / "inputs" / "dataUttaradit" / "Observed_Rain_daily_198101_201412_Uttaradit.csv", nrows=1)
    stations = [int(c) for c in obs.columns if c not in {"YEAR", "MONTH", "DAY"}]
    coords = pd.read_excel(ROOT / "inputs" / "dataUttaradit" / "station_coordinates.xlsx")
    coords = coords[coords.station.astype(int).isin(stations)].copy()
    qc = pd.read_csv(output / "observed_zero_qc_diagnostics.csv")
    probable = qc[qc.classification == "probable_missing"].copy()
    probable["days_removed"] = [calendar.monthrange(int(y), int(m))[1] for y, m in zip(probable.year, probable.month)]
    qsum = probable.groupby("station").agg(probable_months=("month", "size"), days_removed=("days_removed", "sum"))
    complete = metrics[(metrics.source_type == "OBSERVED") & metrics.complete].groupby(["station", "season_type"]).size().unstack(fill_value=0)
    s2 = [["Station", "Latitude", "Longitude", "Elevation (m)", "Probable-missing months", "Days removed", "Complete rainy", "Complete hot/dry"]]
    for row in coords.sort_values("station").itertuples(index=False):
        st = int(row.station)
        s2.append([str(st), f"{row.latitude:.2f}", f"{row.longitude:.2f}", f"{getattr(row, '_4'):.1f}",
                   str(int(qsum.loc[st, "probable_months"])) if st in qsum.index else "0",
                   str(int(qsum.loc[st, "days_removed"])) if st in qsum.index else "0",
                   str(int(complete.loc[st, "RAINY"])) if st in complete.index else "0",
                   str(int(complete.loc[st, "HOT_DRY"])) if st in complete.index else "0"])
    add_table(doc, s2, [1100, 1050, 1100, 1200, 1700, 1100, 1385, 1385], font_size=7.5, left_columns=(0,))
    note(doc, "Observed zero screening was performed on the native Gregorian calendar. In total, 43 station-months (1,308 daily values) were classified as probable missing; confirmed dry and uncertain zeros were retained.")

    # S3
    page_break(doc)
    caption(doc, "S3", "Observed management-season ENSO classification from frozen NOAA CPC ERSSTv6 ONI.")
    classes = pd.read_csv(output / "season_classification_observed.csv")
    s3 = [["Climate year", "Rainy phase", "Rainy mean ONI (°C)", "Rainy windows", "Hot/dry phase", "Hot/dry mean ONI (°C)", "Hot/dry windows"]]
    for year in range(1981, 2015):
        rainy = classes[(classes.season_type == "RAINY") & (classes.climate_year == year)].iloc[0]
        hot = classes[(classes.season_type == "HOT_DRY") & (classes.climate_year == year)].iloc[0]
        hot_eligible = year <= 2013
        s3.append([
            str(year), PHASE_LABEL[rainy.enso_phase], fmt(rainy.index_mean_c, 2), str(int(rainy.n_overlapping_windows)),
            PHASE_LABEL[hot.enso_phase] if hot_eligible else "Not analysed",
            fmt(hot.index_mean_c, 2) if hot_eligible else "NE",
            str(int(hot.n_overlapping_windows)) if hot_eligible else "NE",
        ])
    add_table(doc, s3, [1050, 1550, 1450, 1150, 1550, 1450, 1820], font_size=7.0, left_columns=())
    note(doc, "Rainy is May-October. Hot/dry is November-April and is labelled by its November start year. Hot/dry climate year 2014 ends outside the 1981-2014 daily analysis window and is not analysed.")

    # S4
    page_break(doc)
    caption(doc, "S4", "Model-specific management-season ENSO sample sizes.")
    sizes = pd.read_csv(output / "enso_phase_sample_sizes.csv")
    sizes = sizes[sizes.source_type != "OBSERVED"]
    pivot = sizes.pivot_table(index=["model", "season_type"], columns="enso_phase", values="n_seasons", fill_value=0).reset_index()
    s4 = [["Model", "Season", "El Niño", "Neutral", "La Niña", "Transition/unclassified"]]
    for row in pivot.sort_values(["model", "season_type"]).itertuples(index=False):
        s4.append([row.model, SEASON_LABEL[row.season_type], str(int(row.EL_NINO)), str(int(row.NEUTRAL)), str(int(row.LA_NINA)), str(int(row.TRANSITION_UNCLASSIFIED))])
    add_table(doc, s4, [2600, 1400, 1200, 1200, 1200, 2420], font_size=7.8, left_columns=(0,))
    note(doc, "Each GCM is classified from its own exact-member Niño-3.4 index. Persistent episodes require at least five overlapping three-month windows at |anomaly| >= 0.5°C; season assignment uses a strict majority.")

    # S5
    page_break(doc)
    caption(doc, "S5", "Model-level raw and blocked cross-fitted-QDM primary rainfall responses relative to Neutral.")
    model = pd.read_csv(output / "model_primary_responses.csv")
    raw = model[model.source_type == "RAW"].rename(columns={"response_pct": "raw_pct"})
    qdm = model[model.source_type == "QDM"].rename(columns={"response_pct": "qdm_pct"})
    merged = raw.merge(qdm, on=["season_type", "metric", "phase", "model"], suffixes=("_raw", "_qdm"))
    s5 = [["Season", "Metric", "Phase", "Model", "Raw response (%)", "QDM response (%)", "QDM-raw (pp)", "Raw grids", "Phase/Neutral n"]]
    for row in merged.sort_values(["season_type", "metric", "phase", "model"]).itertuples(index=False):
        s5.append([SEASON_LABEL[row.season_type], row.metric.replace("wet_day_frequency_pct", "Wet-day frequency"), PHASE_LABEL[row.phase], row.model,
                   fmt(row.raw_pct), fmt(row.qdm_pct), fmt(row.qdm_pct - row.raw_pct), str(int(row.n_grid_signatures_raw)), f"{int(row.n_phase_raw)}/{int(row.n_neutral_raw)}"])
    add_table(doc, s5, [850, 1250, 1050, 1500, 1180, 1180, 1100, 850, 1060], font_size=6.7, left_columns=(1, 3))
    note(doc, "Responses are model medians after collapsing stations sharing an identical raw-grid signature. A minimum of three phase and three Neutral seasons is required; unavailable model-phase combinations are omitted. pp = percentage points.")

    # S6
    page_break(doc)
    caption(doc, "S6", "Descriptive secondary-index responses relative to Neutral for observations, raw CMIP6, and cross-fitted QDM.")
    secondary = secondary_responses(metrics, config)
    s6 = [["Season", "Metric", "Phase", "Observed response (%)", "Observed stations", "Raw ensemble (%)", "QDM ensemble (%)", "Raw/QDM models"]]
    for row in secondary.sort_values(["season_type", "metric", "phase"]).itertuples(index=False):
        s6.append([SEASON_LABEL[row.season_type], row.metric, PHASE_LABEL[row.phase], fmt(row.observed_response_pct), str(int(row.n_observed_stations)),
                   fmt(row.raw_response_pct), fmt(row.qdm_response_pct), f"{int(row.n_raw_models)}/{int(row.n_qdm_models)}"])
    add_table(doc, s6, [900, 1150, 1050, 1450, 1150, 1300, 1300, 1720], font_size=7.0, left_columns=(1,))
    note(doc, "Secondary indices were not included in the primary Benjamini-Hochberg family and are presented without inferential claims. NE = not estimable because the Neutral reference was zero or too few eligible entities were available. Definitions: Rx5day, maximum five-day total; CWD, maximum wet spell; R10/20/50mm, counts above thresholds; R95p/R99p, totals above observed wet-day percentiles; q50/q90/q95/q99, wet-day quantiles.")

    # S7
    page_break(doc)
    caption(doc, "S7", "Phase contrast and true neutral-centred asymmetry with event-level inference.")
    asym = pd.read_csv(output / "enso_asymmetry_primary.csv")
    s7 = [["Source", "Season", "Metric", "Phase contrast % [95% CI]", "p", "True asymmetry % [95% CI]", "p", "BH q"]]
    for row in asym.sort_values(["source_type", "season_type", "metric"]).itertuples(index=False):
        contrast = f"{fmt(row.phase_contrast_pct)} [{fmt(row.phase_contrast_pct_ci_low)}, {fmt(row.phase_contrast_pct_ci_high)}]"
        true_asym = f"{fmt(row.neutral_centered_asymmetry_pct)} [{fmt(row.neutral_centered_asymmetry_pct_ci_low)}, {fmt(row.neutral_centered_asymmetry_pct_ci_high)}]"
        s7.append([row.source_type.title(), SEASON_LABEL[row.season_type], row.metric.replace("wet_day_frequency_pct", "Wet-day frequency"), contrast,
                   f"{row.phase_contrast_pct_permutation_p:.3f}", true_asym, f"{row.neutral_centered_asymmetry_pct_permutation_p:.3f}", f"{row.true_asymmetry_p_bh_primary:.3f}"])
    add_table(doc, s7, [900, 850, 1200, 1900, 650, 1900, 650, 1970], font_size=6.7, left_columns=(2,))
    note(doc, "Phase contrast = ALa - AEl. True neutral-centred asymmetry = ALa + AEl. Bootstrap repetitions = 5,000; permutation repetitions = 4,999; resampling unit = management-season climate year.")

    # S8
    page_break(doc)
    caption(doc, "S8", "Cross-fitted-QDM ENSO signal-preservation diagnostics.")
    preserve = pd.read_csv(output / "qdm_enso_signal_preservation.csv")
    s8 = [["Season", "Metric", "Phase", "Observed (%)", "Raw (%)", "QDM (%)", "Category", "Ratio", "Shift (pp)", "Direction agreement raw/QDM", "Paired sign retention"]]
    for row in preserve.sort_values(["season_type", "metric", "phase"]).itertuples(index=False):
        s8.append([SEASON_LABEL[row.season_type], row.metric.replace("wet_day_frequency_pct", "Wet-day frequency"), PHASE_LABEL[row.phase], fmt(row.observed_response_pct), fmt(row.raw_response_pct), fmt(row.qdm_response_pct),
                   row.category.replace("INDETERMINATE_RAW_NEAR_ZERO", "Indeterminate"), fmt(row.magnitude_ratio, 2), fmt(row.shift_pct_points),
                   f"{fmt(row.raw_model_direction_agreement_with_observed, 2)}/{fmt(row.qdm_model_direction_agreement_with_observed, 2)}",
                   f"{fmt(row.paired_model_sign_preservation_fraction, 2)} ({int(row.n_paired_models)} models)"])
    add_table(doc, s8, [750, 1050, 900, 760, 700, 700, 1200, 650, 750, 1400, 1160], font_size=6.3, left_columns=(1, 6))
    note(doc, "A raw response with magnitude <5 percentage points is indeterminate. Otherwise: reversed = sign changed; attenuated = |QDM/raw| <0.8; preserved = 0.8-1.2; amplified = >1.2.")

    # S9
    page_break(doc)
    caption(doc, "S9", "Blocked cross-fitted-QDM calibration and leakage diagnostic summary by model.")
    diag = pd.read_csv(output / "crossfit_qdm_diagnostics.csv")
    diag["sparse_fallback_used"] = diag.sparse_fallback_used.astype(str).str.lower().eq("true")
    grouped = diag.groupby(["model", "member_id"], as_index=False).agg(
        fits=("month", "size"), blocks=("block_start", "nunique"), fallback_fits=("sparse_fallback_used", "sum"),
        target_obs_used=("n_target_observations_used", "sum"), min_obs_wet=("n_observed_wet", "min"),
        min_model_wet=("n_model_wet", "min"), min_target=("n_target", "min"), max_target=("n_target", "max"),
    )
    s9 = [["Model", "Member", "Fits", "Blocks", "Fallback fits", "Target observations used", "Min observed wet", "Min model wet", "Target n range"]]
    for row in grouped.itertuples(index=False):
        s9.append([row.model, row.member_id, str(int(row.fits)), str(int(row.blocks)), str(int(row.fallback_fits)), str(int(row.target_obs_used)),
                   str(int(row.min_obs_wet)), str(int(row.min_model_wet)), f"{int(row.min_target)}-{int(row.max_target)}"])
    add_table(doc, s9, [1900, 1250, 850, 800, 1100, 1500, 950, 900, 770], font_size=7.2, left_columns=(0,))
    note(doc, "All 7,644 fits used zero observations from their target fold. The adjacent-three-month calibration pool was activated only when a single calendar month failed the >=30 wet-day calibration requirement; the target remained in its original month.")

    # S10a and S10b
    page_break(doc)
    caption(doc, "S10a", "Pre-specified analysis acceptance gates.")
    gates = pd.read_csv(output / "PAPER3_ACCEPTANCE_GATES.csv")
    s10a = [["Gate", "Description", "Status"]] + [[row.gate, row.description, row.status] for row in gates.itertuples(index=False)]
    add_table(doc, s10a, [900, 7600, 1520], font_size=7.8, left_columns=(1,))

    caption(doc, "S10b", "Machine-readable supplementary artifacts and their analytical roles.")
    artifacts = [
        ("enso_episode_catalog_observed.csv", 45, "Observed persistent ONI episodes and intensity summaries"),
        ("enso_episode_catalog_models.csv", 128, "Exact-member persistent Niño-3.4 episodes"),
        ("season_classification_observed.csv", 68, "Observed rainy and hot/dry season labels"),
        ("season_classification_models.csv", 476, "Model-specific management-season labels"),
        ("observed_station_primary_responses.csv", 208, "Gauge-level primary phase responses"),
        ("model_primary_responses.csv", len(model), "Raw and QDM model-level primary responses"),
        ("primary_response_summary.csv", 48, "Observed/raw/QDM ensemble inference"),
        ("enso_asymmetry_primary.csv", len(asym), "Phase contrast and true asymmetry inference"),
        ("qdm_enso_signal_preservation.csv", len(preserve), "Preservation classification and directional agreement"),
        ("crossfit_qdm_diagnostics.csv", len(diag), "Fold/month calibration and leakage audit"),
        ("Paper3_Supplementary_Secondary_Responses.csv", len(secondary), "Descriptive secondary-index responses generated for Table S6"),
        ("PAPER3_ACCEPTANCE_GATES.csv", len(gates), "Pre-specified reproducibility gates"),
    ]
    s10b = [["File", "Rows", "Role"]] + [[name, str(rows), role] for name, rows, role in artifacts]
    add_table(doc, s10b, [3200, 900, 5920], font_size=7.3, left_columns=(0, 2))
    note(doc, "Paths, checksums, source capture metadata, the executable notebook, and the complete daily/seasonal intermediate outputs are retained in the reproducibility package. Station observations remain subject to provider terms.")

    doc.core_properties.title = "Supplementary Information - ENSO-Conditioned Seasonal Rainfall Extremes over Uttaradit"
    doc.core_properties.author = "Surasit Punyawansiri"
    doc.core_properties.subject = "Supplementary tables for Chiang Mai Journal of Science submission"
    doc.core_properties.keywords = "ENSO, CMIP6, QDM, supplementary information, Thailand"
    doc.save(OUT)
    print(f"SUPPLEMENTARY_DOCX_COMPLETE {OUT}")
    print(f"SECONDARY_RESPONSES_COMPLETE {SECONDARY_CSV} rows={len(secondary)}")


if __name__ == "__main__":
    main()
