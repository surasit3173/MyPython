import argparse
import json
import re
import zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image


def clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def para_text_with_tabs(p):
    chunks = []
    for run in p.runs:
        chunks.append(run.text)
    text = "".join(chunks)
    return clean(text)


def paragraph_record(p, index):
    text = clean(p.text)
    bold_chars = 0
    total_chars = 0
    breaks = 0
    drawings = 0
    for r in p.runs:
        n = len(r.text or "")
        total_chars += n
        if r.bold:
            bold_chars += n
        breaks += len(r._r.xpath(".//w:br"))
        drawings += len(r._r.xpath(".//w:drawing")) + len(r._r.xpath(".//w:pict"))
    ppr = p._p.pPr
    page_break_before = False
    if ppr is not None and ppr.pageBreakBefore is not None:
        page_break_before = bool(ppr.pageBreakBefore.val)
    return {
        "index": index,
        "style": p.style.name if p.style is not None else "",
        "text": text,
        "total_chars": total_chars,
        "bold_fraction": (bold_chars / total_chars) if total_chars else 0.0,
        "breaks": breaks,
        "drawings": drawings,
        "page_break_before": page_break_before,
    }


def cell_text(cell):
    return clean(" | ".join(clean(p.text) for p in cell.paragraphs if clean(p.text)))


def extract_docx(path):
    path = Path(path)
    doc = Document(path)
    paragraphs = [paragraph_record(p, i) for i, p in enumerate(doc.paragraphs)]
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            rows.append({"row": ri, "cells": [cell_text(c) for c in row.cells]})
        tables.append({"index": ti, "rows": rows})

    # Preserve body order so paragraphs and tables can be reconstructed section-by-section.
    body_order = []
    p_lookup = {id(p._p): rec for p, rec in zip(doc.paragraphs, paragraphs)}
    table_lookup = {id(t._tbl): i for i, t in enumerate(doc.tables)}
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            rec = p_lookup.get(id(child))
            if rec is not None:
                body_order.append({"type": "paragraph", "index": rec["index"], "style": rec["style"], "text": rec["text"]})
        elif child.tag == qn("w:tbl"):
            ti = table_lookup.get(id(child))
            if ti is not None:
                body_order.append({"type": "table", "index": ti, "text": " || ".join(" | ".join(r["cells"]) for r in tables[ti]["rows"])})

    sections = []
    for i, sec in enumerate(doc.sections):
        sections.append({
            "index": i,
            "page_width_in": sec.page_width.inches if sec.page_width else None,
            "page_height_in": sec.page_height.inches if sec.page_height else None,
            "top_margin_in": sec.top_margin.inches if sec.top_margin else None,
            "bottom_margin_in": sec.bottom_margin.inches if sec.bottom_margin else None,
            "left_margin_in": sec.left_margin.inches if sec.left_margin else None,
            "right_margin_in": sec.right_margin.inches if sec.right_margin else None,
        })

    core = doc.core_properties
    metadata = {
        "title": core.title,
        "subject": core.subject,
        "author": core.author,
        "last_modified_by": core.last_modified_by,
        "created": core.created.isoformat() if core.created else None,
        "modified": core.modified.isoformat() if core.modified else None,
    }

    package = {
        "omml_math": 0,
        "comments": 0,
        "tracked_insertions": 0,
        "tracked_deletions": 0,
        "footnotes": 0,
        "endnotes": 0,
        "images": [],
        "all_xml_text": "",
    }
    all_text_chunks = []
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        for name in names:
            if name.startswith("word/") and name.endswith(".xml"):
                raw = zf.read(name)
                try:
                    s = raw.decode("utf-8")
                except UnicodeDecodeError:
                    continue
                package["omml_math"] += len(re.findall(r"<(?:m:)?oMath(?:Para)?(?:\s|>)", s))
                package["tracked_insertions"] += len(re.findall(r"<w:ins(?:\s|>)", s))
                package["tracked_deletions"] += len(re.findall(r"<w:del(?:\s|>)", s))
                if name == "word/comments.xml":
                    package["comments"] = len(re.findall(r"<w:comment(?:\s|>)", s))
                if name == "word/footnotes.xml":
                    package["footnotes"] = max(0, len(re.findall(r"<w:footnote(?:\s|>)", s)) - 2)
                if name == "word/endnotes.xml":
                    package["endnotes"] = max(0, len(re.findall(r"<w:endnote(?:\s|>)", s)) - 2)
                texts = re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", s, re.S)
                all_text_chunks.extend(texts)
        for name in names:
            if name.startswith("word/media/"):
                data = zf.read(name)
                rec = {"name": name, "bytes": len(data)}
                try:
                    im = Image.open(BytesIO(data))
                    rec.update({
                        "format": im.format,
                        "width_px": im.width,
                        "height_px": im.height,
                        "dpi": im.info.get("dpi"),
                    })
                except Exception as exc:
                    rec["error"] = str(exc)
                package["images"].append(rec)
    package["all_xml_text"] = clean(" ".join(all_text_chunks))

    # Visible text counts. Tables are counted once; merged cells may appear duplicated in python-docx.
    paragraph_words = sum(len(re.findall(r"\b[\w'’-]+\b", p["text"], re.UNICODE)) for p in paragraphs)
    table_words = sum(len(re.findall(r"\b[\w'’-]+\b", c, re.UNICODE)) for t in tables for r in t["rows"] for c in r["cells"])

    return {
        "path": str(path),
        "metadata": metadata,
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "inline_shape_count": len(doc.inline_shapes),
        "word_counts": {"paragraphs": paragraph_words, "tables": table_words, "total": paragraph_words + table_words},
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": tables,
        "body_order": body_order,
        "package": package,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    result = extract_docx(args.docx)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "path": result["path"],
        "paragraph_count": result["paragraph_count"],
        "table_count": result["table_count"],
        "inline_shape_count": result["inline_shape_count"],
        "word_counts": result["word_counts"],
        "package": {k: v for k, v in result["package"].items() if k not in {"all_xml_text", "images"}},
        "image_count": len(result["package"]["images"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
