from docx import Document
from pathlib import Path

docx_path = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")
doc = Document(docx_path)
paras = [p.text.strip() for p in doc.paragraphs]

def find_idx(prefix):
    for i, text in enumerate(paras):
        if text.startswith(prefix):
            return i
    return None

idx_27 = find_idx("2.7")
idx_results = find_idx("Results")
idx_35 = find_idx("3.5")
idx_36 = find_idx("3.6")
idx_fig5 = next((i for i, t in enumerate(paras) if t.startswith("Figure 5")), None)

full_text = "\n".join(paras)
bad_terms = [
    "Supplementary Material",
    "https:/doi.org",
    "doi.org10",
    "เขคดุสิต",
    "การกันข้อมูลช่วงเวลา",
    "ตรวจสอบความถูกต้องข้ามช่วงเวลา",
]

method_27 = "\n".join(paras[idx_27:idx_results]) if idx_27 is not None and idx_results is not None else ""
results_35 = "\n".join(paras[idx_35:idx_36]) if idx_35 is not None and idx_36 is not None else ""

print(f"docx={docx_path}")
print(f"paragraphs={len(paras)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"indices: 2.7={idx_27} Results={idx_results} 3.5={idx_35} 3.6={idx_36} Figure5={idx_fig5}")
print("bad_terms=" + ", ".join(f"{term}:{full_text.count(term)}" for term in bad_terms))
print(f"method_2_7_has_eq5={'MSE = b2' in method_27}")
print(f"method_2_7_has_eq6={'MSE = MSEclim + MSEanom' in method_27}")
print(f"results_3_5_has_display_equation={'MSE =' in results_35}")
print(f"results_3_5_mentions_eq5={'สมการที่ (5)' in results_35}")
print(f"results_3_5_mentions_eq6={'สมการที่ (6)' in results_35}")
print("section_3_5_preview=")
for t in paras[idx_35:idx_36]:
    if t:
        print("  " + t[:220])
print("figure5_caption=" + (paras[idx_fig5] if idx_fig5 is not None else "MISSING"))
