from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ALIGN = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    None: None,
}


def pt(value):
    return None if value is None else round(value.pt, 3)


def inches(value):
    return None if value is None else round(value.inches, 4)


def paragraph_info(p, index):
    fmt = p.paragraph_format
    return {
        "index": index,
        "text": p.text,
        "style": p.style.name if p.style else None,
        "alignment": ALIGN.get(p.alignment, str(p.alignment)),
        "keep_with_next": fmt.keep_with_next,
        "keep_together": fmt.keep_together,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
        "space_before_pt": pt(fmt.space_before),
        "space_after_pt": pt(fmt.space_after),
        "line_spacing": fmt.line_spacing,
        "first_line_indent_in": inches(fmt.first_line_indent),
        "left_indent_in": inches(fmt.left_indent),
        "right_indent_in": inches(fmt.right_indent),
        "runs": [
            {
                "text": r.text,
                "font": r.font.name,
                "size_pt": pt(r.font.size),
                "bold": r.bold,
                "italic": r.italic,
                "underline": r.underline,
            }
            for r in p.runs
        ],
    }


def style_info(style):
    font = style.font
    fmt = getattr(style, "paragraph_format", None)
    return {
        "name": style.name,
        "type": str(style.type),
        "base_style": style.base_style.name if style.base_style else None,
        "font": font.name,
        "size_pt": pt(font.size),
        "bold": font.bold,
        "italic": font.italic,
        "alignment": ALIGN.get(fmt.alignment, str(fmt.alignment)) if fmt else None,
        "space_before_pt": pt(fmt.space_before) if fmt else None,
        "space_after_pt": pt(fmt.space_after) if fmt else None,
        "line_spacing": fmt.line_spacing if fmt else None,
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: inspect_apst_template.py input.docx output.json")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    doc = Document(source)

    with zipfile.ZipFile(source) as archive:
        package_parts = []
        for info in archive.infolist():
            payload = archive.read(info.filename)
            package_parts.append({
                "name": info.filename,
                "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
            })

    result = {
        "source": str(source),
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "sections": [
            {
                "index": i,
                "page_width_in": inches(s.page_width),
                "page_height_in": inches(s.page_height),
                "top_margin_in": inches(s.top_margin),
                "bottom_margin_in": inches(s.bottom_margin),
                "left_margin_in": inches(s.left_margin),
                "right_margin_in": inches(s.right_margin),
                "header_distance_in": inches(s.header_distance),
                "footer_distance_in": inches(s.footer_distance),
                "different_first_page": s.different_first_page_header_footer,
                "header_text": " | ".join(p.text for p in s.header.paragraphs),
                "footer_text": " | ".join(p.text for p in s.footer.paragraphs),
            }
            for i, s in enumerate(doc.sections)
        ],
        "paragraphs": [paragraph_info(p, i) for i, p in enumerate(doc.paragraphs)],
        "styles": [style_info(s) for s in doc.styles if s.type in (1, 2, 3)],
        "tables": [
            {
                "index": ti,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "style": table.style.name if table.style else None,
                "cells": [[cell.text for cell in row.cells] for row in table.rows],
            }
            for ti, table in enumerate(doc.tables)
        ],
        "inline_shapes": [
            {
                "index": i,
                "width_in": inches(shape.width),
                "height_in": inches(shape.height),
                "type": str(shape.type),
            }
            for i, shape in enumerate(doc.inline_shapes)
        ],
        "package_parts": package_parts,
    }
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
