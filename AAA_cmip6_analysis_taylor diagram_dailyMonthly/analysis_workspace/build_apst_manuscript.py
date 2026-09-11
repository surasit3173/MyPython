from __future__ import annotations

import csv
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "analysis_workspace"
SOURCE_TEMPLATE = WORK / "template" / "APST_format_reference.docx"
DRAFT = WORK / "manuscript_draft.md"
AUDIT = WORK / "audit_corrected"
OUTPUT_DIR = WORK / "output"
OUTPUT = OUTPUT_DIR / "APST_Uttaradit_CMIP6_Manuscript.docx"
REPORT = OUTPUT_DIR / "APST_Uttaradit_CMIP6_Manuscript.build.json"

FIGURES = {
    1: WORK / "future_zip" / "output" / "future_q1_uttaradit" / "figures" / "FUTURE_Q1_FIGURE_01_study_area.png",
    2: AUDIT / "figures" / "FIGURE_02_baseline_sensitivity.png",
    3: AUDIT / "figures" / "FIGURE_03_change_heatmap.png",
    4: AUDIT / "figures" / "FIGURE_04_selected_profiles.png",
}

FIGURE_META = {
    1: (
        4.25,
        "Study area and 13-gauge network in Uttaradit Province. Gauge coordinates and elevations were supplied with the daily rainfall archive.",
        "Map of Uttaradit Province showing the 13 rain-gauge locations and station elevations.",
    ),
    2: (
        5.70,
        "Sensitivity of regional median relative changes to baseline source. The x-axis uses observations for 1995–2014; the y-axis uses each model's supplied bias-corrected historical series for 1995–2014. Each point is an index–period combination; R50mm is excluded because percentage changes are unstable at small baselines. Crosses mark sign changes and the dashed line is 1:1. Results are descriptive rather than significance tests.",
        "Scatter plot comparing future-change estimates from observed and model-consistent baseline sources.",
    ),
    3: (
        5.70,
        "Audited multimodel regional median relative changes from the model-consistent baseline. Values are percentages; an asterisk indicates that at least six of seven models agreed on direction. R50mm percentages are shown only for completeness and are denominator-sensitive.",
        "Heatmap of audited precipitation-index changes by scenario and future assessment period.",
    ),
    4: (
        5.70,
        "Selected model-consistent projection profiles. Points show seven-model medians and bands show interquartile ranges. Filled markers indicate agreement by at least six of seven models; open markers are ambiguous. R50mm is reported in absolute days per year.",
        "Six panels showing selected future precipitation-index changes and multimodel spread.",
    ),
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def set_run_font(run, size: float = 10.0, bold=None, italic=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(
    paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    line_spacing: float = 2.0,
    space_before: float = 0.0,
    space_after: float = 0.0,
    keep_with_next: bool = False,
    keep_together: bool = False,
) -> None:
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_with_next
    fmt.keep_together = keep_together
    fmt.widow_control = True


def add_markdown_runs(paragraph, text: str, size: float = 10.0) -> None:
    """Render the small bold/italic subset used in the draft."""
    pos = 0
    token_re = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*))")
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size)


def clear_body_keep_section(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.2688)
    section.page_height = Inches(11.6944)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.4924)
    section.footer_distance = Inches(1.0)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for name, italic in (("Heading 1", False), ("Heading 2", True)):
        style = document.styles[name]
        style_ppr = style._element.get_or_add_pPr()
        for tag in ("w:numPr", "w:ind"):
            inherited = style_ppr.find(qn(tag))
            if inherited is not None:
                style_ppr.remove(inherited)
        style_rpr = style._element.get_or_add_rPr()
        caps = style_rpr.find(qn("w:caps"))
        if caps is not None:
            style_rpr.remove(caps)
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(10)
        style.font.bold = True
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True

    for header in (section.header, section.first_page_header, section.even_page_header):
        for paragraph in header.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                set_run_font(run, 10)

    settings = document.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def add_title_block(document: Document, title: str, authors: str, affiliation: str, corresponding: str) -> None:
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=2.0, keep_with_next=True)
    run = p.add_run(title)
    set_run_font(run, 12, bold=True)

    for text, italic in ((authors, False), (affiliation, False), (corresponding, True)):
        p = document.add_paragraph()
        set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=2.0, keep_with_next=True)
        run = p.add_run(text)
        set_run_font(run, 10, italic=italic)


def add_heading(document: Document, text: str, level: int) -> None:
    p = document.add_paragraph()
    set_paragraph_format(
        p,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=2.0,
        keep_with_next=True,
        keep_together=True,
    )
    run = p.add_run(text)
    set_run_font(run, 10, bold=True, italic=(level == 3))
    p.style = document.styles["Heading 1" if level == 2 else "Heading 2"]


def add_body_paragraph(document: Document, text: str, *, reference: bool = False) -> None:
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=2.0)
    if reference:
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    add_markdown_runs(p, text, 10)


def set_cell_margins(cell, top=35, start=55, bottom=35, end=55) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "bottom"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    for name in ("left", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    tblPr.append(borders)


def set_cell_bottom_border(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)


def set_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def fill_cell(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_format(p, alignment=align, line_spacing=1.0, keep_together=True)
    run = p.add_run(text)
    set_run_font(run, 10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table_caption(document: Document, number: int, caption: str) -> None:
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0, keep_with_next=True)
    run = p.add_run(f"Table {number}. ")
    set_run_font(run, 10, bold=True)
    run = p.add_run(caption)
    set_run_font(run, 10)


def add_native_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = "Table Normal"
    except KeyError:
        pass
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_header(header_row)
    for i, (header, width) in enumerate(zip(headers, widths)):
        header_row.cells[i].width = Inches(width)
        fill_cell(header_row.cells[i], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bottom_border(header_row.cells[i])
    for row_values in rows:
        row = table.add_row()
        for i, (value, width) in enumerate(zip(row_values, widths)):
            row.cells[i].width = Inches(width)
            align = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER
            fill_cell(row.cells[i], value, align=align)
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)


def build_table_1(document: Document) -> None:
    data = read_csv(AUDIT / "index_registry.csv")
    assert len(data) == 11
    rows = [[d["index"], d["unit"].replace("-1", "⁻¹"), d["definition"]] for d in data]
    add_table_caption(document, 1, "Definitions of the 11 annual precipitation-extreme indices.")
    add_native_table(document, ["Index", "Unit", "Definition"], rows, [0.75, 1.15, 4.35])
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
    add_markdown_runs(p, "Note: R50mm is study-specific; the remaining measures follow ETCCDI or ETCCDI-style definitions. Wet days have precipitation ≥1 mm.", 10)


def build_table_2(document: Document) -> None:
    data = read_csv(AUDIT / "historical_skill_summary.csv")
    order = [("Daily", "Raw"), ("Daily", "Bias-corrected"), ("Monthly", "Raw"), ("Monthly", "Bias-corrected")]
    lookup = {(d["scale"], d["variant"]): d for d in data}
    rows = []
    for scale, variant in order:
        d = lookup[(scale, variant)]
        rows.append(
            [
                scale,
                variant,
                f'{float(d["median_r"]):.3f}',
                (
                    f'{float(d["median_rmse"]):.2f} mm d⁻¹'
                    if scale == "Daily"
                    else f'{float(d["median_rmse"]):.2f} mm month⁻¹'
                ),
                f'{float(d["median_pbias_pct"]):+.1f}',
                f'{float(d["median_nse"]):.3f}',
                f'{float(d["median_kge"]):.3f}',
            ]
        )
    add_table_caption(document, 2, "Median historical performance across 91 model–gauge pairs.")
    add_native_table(
        document,
        ["Scale", "Series", "r", "RMSE", "PBIAS (%)", "NSE", "KGE"],
        rows,
        [0.70, 1.25, 0.50, 1.25, 0.85, 0.75, 0.75],
    )
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
    add_markdown_runs(p, "Note: monthly RMSE is for monthly rainfall totals and daily RMSE is for daily rainfall. Lower RMSE and |PBIAS| are better; higher r, NSE, and KGE are better.", 10)


def build_table_3(document: Document) -> None:
    data = read_csv(AUDIT / "ensemble_change_sensitivity.csv")
    lookup = {(d["scenario"], d["window"], d["index"]): d for d in data}
    selected = [
        ("ssp245", "Mid-term", "Rx5day", "percent"),
        ("ssp245", "Mid-term", "R99p", "percent"),
        ("ssp245", "Mid-term", "R50mm", "absolute"),
        ("ssp585", "Mid-term", "PRCPTOT", "percent"),
        ("ssp585", "Mid-term", "SDII", "percent"),
        ("ssp585", "Long-term", "PRCPTOT", "percent"),
        ("ssp585", "Long-term", "SDII", "percent"),
        ("ssp585", "Long-term", "Rx5day", "percent"),
        ("ssp585", "Long-term", "R99p", "percent"),
        ("ssp585", "Long-term", "CWD", "percent"),
        ("ssp585", "Long-term", "R50mm", "absolute"),
    ]
    rows = []
    for scenario, period, index, mode in selected:
        d = lookup[(scenario, period, index)]
        agreement = "7/7" if float(d["model_agreement"]) > 0.99 else "6/7"
        if mode == "absolute":
            median = float(d["model_median_abs"])
            q25 = float(d["model_q25_abs"])
            q75 = float(d["model_q75_abs"])
            change = f"{median:+.2f} ({q25:+.2f} to {q75:+.2f}) d yr⁻¹"
        else:
            median = float(d["model_median_pct"])
            q25 = float(d["model_q25_pct"])
            q75 = float(d["model_q75_pct"])
            change = f"{median:+.1f} ({q25:+.1f} to {q75:+.1f})%"
        rows.append(
            [
                "SSP2-4.5" if scenario == "ssp245" else "SSP5-8.5",
                "2041–2060" if period == "Mid-term" else "2081–2100",
                index,
                change,
                agreement,
            ]
        )
    add_table_caption(document, 3, "Selected robust multimodel regional changes from the model-consistent baseline.")
    add_native_table(
        document,
        ["Scenario", "Period", "Index", "Median change (IQR)", "Models"],
        rows,
        [1.00, 1.05, 0.85, 2.35, 0.70],
    )
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
    add_markdown_runs(p, "Note: IQR is the interquartile range across seven model-level regional changes. Robustness requires agreement by ≥6/7 models and is descriptive, not a significance test. R50mm is reported in absolute units.", 10)


def add_figure(document: Document, number: int) -> None:
    path = FIGURES[number]
    width, caption, alt = FIGURE_META[number]
    if not path.exists():
        raise FileNotFoundError(path)
    p = document.add_paragraph()
    set_paragraph_format(
        p,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.0,
        keep_with_next=True,
        keep_together=True,
    )
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", alt)
        doc_pr[0].set("title", f"Figure {number}")

    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.0, keep_together=True)
    r = p.add_run(f"Figure {number}. ")
    set_run_font(r, 10, bold=True)
    r = p.add_run(caption)
    set_run_font(r, 10)


def parse_front_matter(text: str) -> tuple[str, str, str, str, str, str, list[str]]:
    lines = text.splitlines()
    title = lines[0].removeprefix("# ").strip()
    author_line = next(line for line in lines[1:] if line.startswith("**[") and "Author" in line)
    authors = author_line.replace("**", "").replace("*", "").strip()
    affiliation = next(line for line in lines[1:] if line.startswith("1 [Department"))
    corresponding = next(line for line in lines[1:] if line.startswith("*Corresponding author"))
    corresponding = corresponding.removeprefix("*")
    abstract = text.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0].strip()
    keywords = text.split("**Keywords:**", 1)[1].splitlines()[0].strip()
    body_lines = text.split("## 1. Introduction", 1)[1].splitlines()
    body_lines.insert(0, "## 1. Introduction")
    return title, authors, affiliation, corresponding, abstract, keywords, body_lines


def build_document() -> dict[str, object]:
    for required in [SOURCE_TEMPLATE, DRAFT, *FIGURES.values()]:
        if not required.exists():
            raise FileNotFoundError(required)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_TEMPLATE, OUTPUT)
    document = Document(OUTPUT)
    clear_body_keep_section(document)
    configure_document(document)

    text = DRAFT.read_text(encoding="utf-8")
    title, authors, affiliation, corresponding, abstract, keywords, body_lines = parse_front_matter(text)
    abstract_words = len(re.findall(r"\b[\w−–'-]+\b", abstract))
    if abstract_words > 250:
        raise ValueError(f"Abstract exceeds 250 words: {abstract_words}")

    add_title_block(document, title, authors, affiliation, corresponding)
    add_heading(document, "Abstract", 2)
    add_body_paragraph(document, abstract)
    p = document.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=2.0)
    run = p.add_run("Keywords: ")
    set_run_font(run, 10, bold=True)
    run = p.add_run(keywords)
    set_run_font(run, 10)

    in_references = False
    for raw in body_lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            add_heading(document, line[4:], 3)
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_references = heading == "12. References"
            add_heading(document, heading, 2)
            continue
        placeholder = re.fullmatch(r"\*\*\[Insert (Figure|Table) (\d+) near here\]\*\*", line)
        if placeholder:
            kind, number_text = placeholder.groups()
            number = int(number_text)
            if kind == "Figure":
                add_figure(document, number)
            elif number == 1:
                build_table_1(document)
            elif number == 2:
                build_table_2(document)
            elif number == 3:
                build_table_3(document)
            continue
        add_body_paragraph(document, line, reference=in_references)

    props = document.core_properties
    props.title = title
    props.subject = "Audited CMIP6 precipitation-extreme projections for Uttaradit Province"
    props.author = "[Authors to be supplied]"
    props.last_modified_by = ""
    props.keywords = keywords
    props.comments = "Prepared in the supplied APST template; author metadata and declarations require confirmation."
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    text_all = "\n".join(p.text for p in reopened.paragraphs)
    with zipfile.ZipFile(OUTPUT) as archive:
        header_xml = "".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        )
    page_field_count = header_xml.count("PAGE")
    placeholder_markers = sorted(set(re.findall(r"\[Insert (?:Figure|Table).*?\]", text_all)))
    report = {
        "output": str(OUTPUT),
        "abstract_words": abstract_words,
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "inline_shapes": len(reopened.inline_shapes),
        "page_field_tokens_in_headers": page_field_count,
        "unresolved_insert_markers": placeholder_markers,
        "expected_tables": 3,
        "expected_figures": 4,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(build_document(), ensure_ascii=False, indent=2))
