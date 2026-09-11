from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_revised_submission import (
    add_body_paragraph,
    add_heading,
    add_title_block,
    clear_body_keep_section,
    configure_document,
    set_cell_bottom_border,
    set_cell_margins,
    set_paragraph_format,
    set_repeat_header,
    set_run_font,
    set_table_borders,
)


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "analysis_workspace"
AUDIT = WORK / "audit_corrected"
TEMPLATE = WORK / "template" / "APST_format_reference.docx"
OUTPUT_DIR = WORK / "revision_round2" / "output"
OUTPUT = OUTPUT_DIR / "APST_Uttaradit_CMIP6_Supplementary_Information.docx"
REPORT = OUTPUT_DIR / "APST_Uttaradit_CMIP6_Supplementary_Information.build.json"

INDEX_ORDER = [
    "PRCPTOT",
    "SDII",
    "Rx1day",
    "Rx5day",
    "CDD",
    "CWD",
    "R10mm",
    "R20mm",
    "R50mm",
    "R95p",
    "R99p",
]
WINDOW_ORDER = ["Near-term", "Mid-term", "Long-term"]
SCENARIO_ORDER = ["ssp245", "ssp585"]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is None:
        existing = OxmlElement("w:shd")
        tc_pr.append(existing)
    existing.set(qn("w:fill"), fill)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def fill_cell(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, size: float = 7.4) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, alignment=align, line_spacing=1.0, keep_together=True)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=24, start=34, bottom=24, end=34)


def add_caption(document: Document, label: str, caption: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0, keep_with_next=True)
    run = paragraph.add_run(f"Table {label}. ")
    set_run_font(run, 9, bold=True)
    run = paragraph.add_run(caption)
    set_run_font(run, 9)


def add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_format(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.0)
    run = paragraph.add_run("Note: ")
    set_run_font(run, 8, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, 8)


def add_compact_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 7.4,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = "Table Normal"
    except KeyError:
        pass
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_header(header)
    prevent_row_split(header)
    for i, (text, width) in enumerate(zip(headers, widths)):
        header.cells[i].width = Inches(width)
        fill_cell(header.cells[i], text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size)
        shade_cell(header.cells[i], "D9EAF7")
        set_cell_bottom_border(header.cells[i])
    for row_number, values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for i, (value, width) in enumerate(zip(values, widths)):
            row.cells[i].width = Inches(width)
            align = WD_ALIGN_PARAGRAPH.LEFT if i in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
            fill_cell(row.cells[i], value, align=align, size=font_size)
            if row_number % 2 == 0:
                shade_cell(row.cells[i], "F5F5F5")
    spacer = document.add_paragraph()
    set_paragraph_format(spacer, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)


def set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.6944)
    section.page_height = Inches(8.2688)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.35)


def scenario_period(row: dict[str, str]) -> str:
    scenario = "SSP2-4.5" if row["scenario"] == "ssp245" else "SSP5-8.5"
    years = {"Near-term": "2021–2040", "Mid-term": "2041–2060", "Long-term": "2081–2100"}[row["window"]]
    return f"{scenario}\n{years}"


def signed_interval(row: dict[str, str], prefix: str, measure: str, suffix: str = "") -> str:
    median = float(row[f"{prefix}_median_{measure}"])
    q25 = float(row[f"{prefix}_q25_{measure}"])
    q75 = float(row[f"{prefix}_q75_{measure}"])
    return f"{median:+.1f} [{q25:+.1f}, {q75:+.1f}]{suffix}"


def agreement_class(row: dict[str, str], prefix: str) -> str:
    count = int(round(float(row[f"{prefix}_agreement"]) * 7))
    return f"{count}/7; {row[f'{prefix}_classification']}"


def bool_text(value: str) -> str:
    return "yes" if value.strip().lower() == "true" else "no"


def build_s1(document: Document) -> None:
    data = read_csv(AUDIT / "ensemble_change_sensitivity.csv")
    if len(data) != 66:
        raise ValueError(f"Expected 66 baseline-sensitivity rows, found {len(data)}")
    scenario_rank = {value: i for i, value in enumerate(SCENARIO_ORDER)}
    window_rank = {value: i for i, value in enumerate(WINDOW_ORDER)}
    index_rank = {value: i for i, value in enumerate(INDEX_ORDER)}
    data.sort(key=lambda row: (scenario_rank[row["scenario"]], window_rank[row["window"]], index_rank[row["index"]]))

    headers = [
        "Scenario–period",
        "Index",
        "Model-consistent signed native-unit change, median [Q1, Q3]",
        "Model-consistent change, median [Q1, Q3] (%)",
        "Model agreement; class",
        "Observed-reference signed native-unit change, median [Q1, Q3]",
        "Observed-reference change, median [Q1, Q3] (%)",
        "Observed agreement; class",
        "Baseline effect (obs − model)",
    ]
    widths = [0.86, 0.58, 1.34, 1.14, 1.02, 1.34, 1.14, 1.02, 1.26]
    for scenario, label in (("ssp245", "S1a"), ("ssp585", "S1b")):
        rows: list[list[str]] = []
        for row in (item for item in data if item["scenario"] == scenario):
            rows.append(
                [
                    scenario_period(row),
                    row["index"] + ("†" if row["index"] == "R50mm" else ""),
                    signed_interval(row, "model", "abs"),
                    signed_interval(row, "model", "pct", "%"),
                    agreement_class(row, "model"),
                    signed_interval(row, "obs", "abs"),
                    signed_interval(row, "obs", "pct", "%"),
                    agreement_class(row, "obs"),
                    (
                        f'{float(row["baseline_sensitivity_pp"]):+.1f} pp; '
                        f'sign reversal: {bool_text(row["sign_reversal"])}; '
                        f'class changed: {bool_text(row["classification_changed"])}'
                    ),
                ]
            )
        scenario_name = "SSP2-4.5" if scenario == "ssp245" else "SSP5-8.5"
        add_caption(document, label, f"Complete baseline-source sensitivity results for {scenario_name} (33 combinations).")
        add_compact_table(document, headers, rows, widths, font_size=7.1)
        add_note(
            document,
            "Model-consistent = future supplied bias-corrected minus supplied bias-corrected historical 1995–2014; observed-reference = future supplied bias-corrected minus observed 1995–2014. Values are medians and Q1–Q3 across seven model-specific regional contrasts after equal weighting of 13 stations. “Signed native-unit change” is not an absolute magnitude. Agreement is the number of models sharing the ensemble direction; ≥6/7 is a descriptive robustness screen, not a significance test. Baseline effect is the signed observed-reference median minus the model-consistent median in percentage points. †R50mm percentages are denominator-sensitive; signed days yr⁻¹ should be prioritized.",
        )


def format_metric(row: dict[str, str], metric: str, decimals: int) -> str:
    return (
        f'{float(row[f"median_{metric}"]):.{decimals}f} '
        f'[{float(row[f"q25_{metric}"]):.{decimals}f}, {float(row[f"q75_{metric}"]):.{decimals}f}]'
    )


def build_s2(document: Document) -> None:
    data = read_csv(AUDIT / "historical_skill_summary.csv")
    lookup = {(row["scale"], row["variant"]): row for row in data}
    order = [
        ("Daily", "Raw"),
        ("Daily", "Bias-corrected"),
        ("Monthly", "Raw"),
        ("Monthly", "Bias-corrected"),
    ]
    rows = []
    for scale, variant in order:
        row = lookup[(scale, variant)]
        rows.append(
            [
                scale,
                "Supplied bias-corrected" if variant == "Bias-corrected" else "Raw",
                row["n_model_station_pairs"],
                format_metric(row, "r", 3),
                format_metric(row, "rmse", 2),
                format_metric(row, "mae", 2),
                format_metric(row, "pbias_pct", 1),
                format_metric(row, "nse", 3),
                format_metric(row, "kge", 3),
            ]
        )
    add_caption(document, "S2a", "Historical performance across 91 model–station pairs, reported as median [Q1, Q3].")
    add_compact_table(
        document,
        ["Scale", "Series", "N", "r", "RMSE", "MAE", "PBIAS (%)", "NSE", "KGE"],
        rows,
        [0.65, 1.30, 0.45, 1.00, 1.20, 1.20, 1.20, 1.00, 1.00],
        font_size=7.8,
    )
    add_note(
        document,
        "Daily RMSE and MAE are in mm d⁻¹; monthly RMSE and MAE are in mm month⁻¹. Daily records contain 12,410 or 12,418 model-specific days; monthly records contain 408 months. Daily and monthly PBIAS derive from the same accumulated rainfall and are not independent evidence.",
    )

    rows_b = []
    for scale in ("Daily", "Monthly"):
        row = lookup[(scale, "Bias-corrected")]
        n = int(row["n_model_station_pairs"])
        fields = [
            "fraction_lower_rmse",
            "fraction_lower_abs_pbias",
            "fraction_higher_r",
            "fraction_higher_kge",
        ]
        values = []
        for field in fields:
            fraction = float(row[field])
            values.append(f"{int(round(fraction * n))}/{n} ({100 * fraction:.1f}%)")
        rows_b.append([scale, *values])
    add_caption(document, "S2b", "Paired model–station changes after correction relative to raw series.")
    add_compact_table(
        document,
        ["Scale", "Lower RMSE", "Lower |PBIAS|", "Higher r", "Higher KGE"],
        rows_b,
        [1.10, 2.00, 2.00, 2.00, 2.00],
        font_size=8.2,
    )
    add_note(
        document,
        "These fractions are paired correction-versus-raw comparisons. Improvement in one metric does not imply overall improvement; RMSE increased for most daily pairs, and median NSE remained negative.",
    )


def build_s3(document: Document) -> None:
    data = read_csv(AUDIT / "residual_bias_summary.csv")
    registry = {row["index"]: row["unit"] for row in read_csv(AUDIT / "index_registry.csv")}
    lookup = {row["index"]: row for row in data}
    rows = []
    for index in INDEX_ORDER:
        row = lookup[index]
        rows.append(
            [
                index,
                registry[index].replace("-1", "⁻¹"),
                f'{float(row["median_regional_bias_abs"]):+.2f}',
                f'{float(row["median_regional_bias_pct"]):+.1f}',
                f'{float(row["q25_regional_bias_pct"]):+.1f} to {float(row["q75_regional_bias_pct"]):+.1f}',
                f'{float(row["max_abs_regional_bias_pct"]):.1f}',
            ]
        )
    add_caption(document, "S3", "Residual historical bias of supplied corrected series relative to observations, 1995–2014.")
    add_compact_table(
        document,
        ["Index", "Unit", "Median signed regional difference", "Median regional bias (%)", "Q1–Q3 regional bias (%)", "Maximum absolute regional bias (%)"],
        rows,
        [0.85, 1.20, 1.70, 1.45, 1.70, 1.80],
        font_size=8.0,
    )
    add_note(
        document,
        "For each model, 1995–2014 annual index means were calculated by station, the 13 stations were equally averaged to a regional value, and corrected-minus-observed bias was then calculated. The table summarizes these seven model-level regional biases. The signed regional difference is expressed in the index's native unit; it is not an absolute magnitude. The maximum absolute percentage is across seven regional model values, not across stations.",
    )


def build_s4(document: Document) -> None:
    audit = json.loads((AUDIT / "trend_type1_audit.json").read_text(encoding="utf-8"))
    rows = [
        ["Supplied self-test", "TFPW-MK", "not recorded", "0.6", "not recorded", "not recorded", "0.05", "0.3830"],
    ]
    for test in ("MK", "Hamed-Rao", "TFPW-MK"):
        rows.append(
            [
                "Independent audit",
                test,
                str(audit["n"]),
                f'{float(audit["ar1_phi"]):.1f}',
                f'{int(audit["replicates"]):,}',
                str(audit["seed"]),
                f'{float(audit["nominal_alpha"]):.2f}',
                f'{float(audit["false_positive_rate"][test]):.4f}',
            ]
        )
    add_heading(document, "Supplementary Note S1. Trend-test null diagnostic", 2)
    add_body_paragraph(
        document,
        "The supplied trend module included a diagnostic self-test under AR(1) dependence. Its TFPW-MK false-positive fraction was 0.383 at nominal α = 0.05. An independent audit of the previously supplied implementation used 2,000 AR(1) null series of length 20 with lag-1 coefficient 0.6 and seed 20260829; all three empirical false-positive fractions remained well above 0.05 (Table S4). These results falsify confirmatory type-I error control for the archived trend p-values, which were therefore excluded from the manuscript.",
    )
    add_caption(document, "S4", "Empirical false-positive fractions from supplied and independent null diagnostics.")
    add_compact_table(
        document,
        ["Audit source", "Test", "n", "AR(1) φ", "Replicates", "Seed", "Nominal α", "False-positive fraction"],
        rows,
        [1.50, 1.25, 0.70, 0.80, 1.15, 1.30, 1.00, 1.40],
        font_size=8.0,
    )
    add_note(
        document,
        "The archived JSON records n, φ, replicates, seed, α, and empirical fractions but not the innovation distribution, initialization/burn-in, sidedness, or all implementation details. This is therefore a diagnostic falsification/quality-assurance exercise rather than a fully reproducible Monte Carlo study.",
    )


def build_document() -> dict[str, object]:
    for required in (
        TEMPLATE,
        AUDIT / "ensemble_change_sensitivity.csv",
        AUDIT / "historical_skill_summary.csv",
        AUDIT / "residual_bias_summary.csv",
        AUDIT / "index_registry.csv",
        AUDIT / "trend_type1_audit.json",
    ):
        if not required.exists():
            raise FileNotFoundError(required)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, OUTPUT)
    document = Document(OUTPUT)
    clear_body_keep_section(document)
    configure_document(document)
    add_title_block(
        document,
        "Supplementary Information",
        "Baseline-source sensitivity and robustness of supplied bias-corrected CMIP6 daily precipitation-extreme projections for Uttaradit Province, Thailand",
        "[Author list identical to the main manuscript]",
        "[Corresponding author and email]",
    )
    add_heading(document, "Scope and evidence boundary", 2)
    add_body_paragraph(
        document,
        "This Supplementary Information reports only tables and diagnostics derived from the previously audited outputs used in the main manuscript. No new climate-model simulation, bias-correction run, projection window, bootstrap analysis, model weighting, or spatial-independence analysis was added during revision.",
    )
    add_body_paragraph(
        document,
        "Table S1 provides all 66 scenario–period–index combinations; Tables S2–S3 expand the historical performance and residual-bias summaries; Supplementary Note S1 records why archived trend p-values were excluded. Missing bias-correction provenance, source-grid metadata, model genealogy, and author-owned declarations are not inferred.",
    )

    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(landscape)
    add_heading(document, "Complete baseline-source sensitivity", 2)
    build_s1(document)
    add_heading(document, "Extended historical evaluation", 2)
    build_s2(document)
    build_s3(document)
    build_s4(document)

    properties = document.core_properties
    properties.title = "Supplementary Information: Uttaradit CMIP6 baseline-source sensitivity"
    properties.subject = "Audited supplementary tables and trend-test diagnostic"
    properties.author = "[Authors to be supplied]"
    properties.last_modified_by = ""
    properties.comments = "Uses only previously audited analytical outputs; author metadata require confirmation."
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    full_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    report = {
        "output": str(OUTPUT),
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "expected_tables": 6,
        "inline_shapes": len(reopened.inline_shapes),
        "contains_table_s1a": "Table S1a." in full_text,
        "contains_table_s1b": "Table S1b." in full_text,
        "contains_supplementary_note_s1": "Supplementary Note S1" in full_text,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(build_document(), ensure_ascii=False, indent=2))
