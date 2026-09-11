from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "analysis_workspace"
DOCX = WORK / "output" / "APST_Uttaradit_CMIP6_Manuscript.docx"
PDF = WORK / "output" / "APST_Uttaradit_CMIP6_Manuscript.qa-preview.pdf"
DRAFT = WORK / "manuscript_draft.md"
TEMPLATE = WORK / "template" / "APST_format_reference.docx"
AUDIT = WORK / "audit_corrected"
FIGURES = [
    WORK / "future_zip" / "output" / "future_q1_uttaradit" / "figures" / "FUTURE_Q1_FIGURE_01_study_area.png",
    AUDIT / "figures" / "FIGURE_02_baseline_sensitivity.png",
    AUDIT / "figures" / "FIGURE_03_change_heatmap.png",
    AUDIT / "figures" / "FIGURE_04_selected_profiles.png",
]


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_path(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def close(actual: float, expected: float, tolerance: float = 0.051) -> bool:
    return math.isclose(actual, expected, rel_tol=0.0, abs_tol=tolerance)


checks: list[tuple[str, bool, str]] = []


def check(name: str, condition: bool, detail: str) -> None:
    checks.append((name, bool(condition), detail))


for required in [DOCX, PDF, DRAFT, TEMPLATE, *FIGURES]:
    check(f"exists:{required.name}", required.exists(), str(required))

document = Document(DOCX)
paragraph_text = "\n".join(p.text for p in document.paragraphs)
table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
all_text = paragraph_text + "\n" + table_text

check("docx_tables", len(document.tables) == 3, f"actual={len(document.tables)} expected=3")
check("docx_figures", len(document.inline_shapes) == 4, f"actual={len(document.inline_shapes)} expected=4")
check("insert_markers_absent", "[Insert Figure" not in all_text and "[Insert Table" not in all_text, "no insertion markers")
check("title_present", "Baseline sensitivity and robustness of bias-corrected CMIP6" in all_text, "title text")
check("rmse_units", "mm day−1" in all_text and "mm month−1" in all_text and "mm d⁻¹" in all_text, "daily/monthly units in prose and table")
check("ordered_station_range", "−51.6% to −27.3%" in all_text, "late-century station range")
check("trend_values", all(x in all_text for x in ("0.225", "0.227", "0.267", "20260829")), "Monte Carlo rates and seed")
check("no_duplicated_heading_numbers", re.search(r"\b\d+\.\d+\.\s+\d+\.\d+", paragraph_text) is None, "heading numbering")
check("author_placeholders_retained", "[Author 1 full name]" in all_text and "[email address]" in all_text, "metadata awaits author input")

heading_1 = sum(p.style.name == "Heading 1" for p in document.paragraphs)
heading_2 = sum(p.style.name == "Heading 2" for p in document.paragraphs)
check("semantic_headings", heading_1 == 13 and heading_2 == 11, f"Heading1={heading_1}, Heading2={heading_2}")

section = document.sections[0]
inch = 914400
check("one_section", len(document.sections) == 1, f"actual={len(document.sections)}")
check("a4_width", close(section.page_width / inch, 8.2688, 0.01), f"{section.page_width / inch:.4f} in")
check("a4_height", close(section.page_height / inch, 11.6944, 0.01), f"{section.page_height / inch:.4f} in")
for label, value in (
    ("top", section.top_margin),
    ("bottom", section.bottom_margin),
    ("left", section.left_margin),
    ("right", section.right_margin),
):
    check(f"margin_{label}", close(value / inch, 1.0, 0.01), f"{value / inch:.3f} in")

draft = DRAFT.read_text(encoding="utf-8")
abstract = draft.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0]
abstract_words = len(re.findall(r"\b[\w−–'-]+\b", abstract))
check("abstract_limit", abstract_words <= 250, f"words={abstract_words}")

body = draft.split("## 12. References", 1)[0]
cited: set[int] = set()
for match in re.finditer(r"\[([0-9,–-]+)\]", body):
    for part in match.group(1).split(","):
        range_match = re.fullmatch(r"(\d+)[–-](\d+)", part)
        if range_match:
            cited.update(range(int(range_match.group(1)), int(range_match.group(2)) + 1))
        else:
            cited.add(int(part))
reference_count = len(re.findall(r"(?m)^\d+\. ", draft.split("## 12. References", 1)[1]))
check("references_cited", cited == set(range(1, reference_count + 1)), f"cited={sorted(cited)}, references={reference_count}")
check("reference_limit", reference_count <= 35, f"references={reference_count}")

pdf = PdfReader(str(PDF))
check("page_limit", len(pdf.pages) == 14 and len(pdf.pages) <= 15, f"pages={len(pdf.pages)}")
media_box = pdf.pages[0].mediabox
check("pdf_a4", close(float(media_box.width), 595.32, 0.2) and close(float(media_box.height), 842.04, 0.2), f"{float(media_box.width):.2f}x{float(media_box.height):.2f} pt")

with zipfile.ZipFile(DOCX) as archive:
    names = archive.namelist()
    header_xml = "".join(
        archive.read(name).decode("utf-8", errors="ignore")
        for name in names
        if name.startswith("word/header") and name.endswith(".xml")
    )
    document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    embedded = [archive.read(f"word/media/image{i}.png") for i in range(2, 6)]

check("page_fields", header_xml.count("PAGE") == 3, f"PAGE tokens={header_xml.count('PAGE')}")
check("figure_alt_text", document_xml.count("descr=") >= 4, f"descr attributes={document_xml.count('descr=')}")
for i, (embedded_bytes, source_path) in enumerate(zip(embedded, FIGURES), start=1):
    check(f"figure_{i}_current", sha256_bytes(embedded_bytes) == sha256_path(source_path), source_path.name)

template_hash = sha256_path(TEMPLATE).upper()
check("template_unchanged", template_hash == "587388A7C2B6B35978D14CAEBDDDCDB11F6CEB3FFDD56F17C8E90E4A596FF1BC", template_hash)

summary = json.loads((AUDIT / "audit_summary.json").read_text(encoding="utf-8"))
check("audit_dimensions", summary["station_count"] == 13 and summary["model_count"] == 7 and summary["annual_rows"] == 417274, json.dumps({k: summary[k] for k in ("station_count", "model_count", "annual_rows")}))
check("baseline_sensitivity", summary["sign_reversals"] == 25 and summary["classification_changes"] == 37, f"sign={summary['sign_reversals']}, class={summary['classification_changes']}")
check("change_preservation", close(summary["change_signal_sign_preserved_fraction"] * 100, 66.9, 0.06) and close(summary["median_abs_change_difference_pp"], 12.1, 0.06), f"{summary['change_signal_sign_preserved_fraction'] * 100:.1f}%, {summary['median_abs_change_difference_pp']:.1f} pp")

skill_rows = read_csv(AUDIT / "historical_skill_summary.csv")
skill = {(d["scale"], d["variant"]): d for d in skill_rows}
check("daily_rmse_values", close(float(skill[("Daily", "Raw")]["median_rmse"]), 10.18) and close(float(skill[("Daily", "Bias-corrected")]["median_rmse"]), 11.84), "10.18 -> 11.84 mm d^-1")
check("monthly_rmse_values", close(float(skill[("Monthly", "Raw")]["median_rmse"]), 103.98) and close(float(skill[("Monthly", "Bias-corrected")]["median_rmse"]), 108.73), "103.98 -> 108.73 mm month^-1")

ensemble_rows = read_csv(AUDIT / "ensemble_change_sensitivity.csv")
ensemble = {(d["scenario"], d["window"], d["index"]): d for d in ensemble_rows}
targets = {
    ("ssp245", "Mid-term", "Rx5day"): (12.5, 6 / 7),
    ("ssp245", "Mid-term", "R99p"): (26.5, 6 / 7),
    ("ssp585", "Long-term", "PRCPTOT"): (-46.7, 1.0),
    ("ssp585", "Long-term", "SDII"): (-27.7, 1.0),
    ("ssp585", "Long-term", "Rx5day"): (-25.2, 6 / 7),
    ("ssp585", "Long-term", "R99p"): (-62.6, 6 / 7),
    ("ssp585", "Long-term", "CWD"): (-14.0, 6 / 7),
}
for key, (expected_change, expected_agreement) in targets.items():
    row = ensemble[key]
    check(
        "ensemble:" + ":".join(key),
        close(float(row["model_median_pct"]), expected_change) and close(float(row["model_agreement"]), expected_agreement, 0.001),
        f"change={float(row['model_median_pct']):.1f}%, agreement={float(row['model_agreement']):.3f}",
    )
check("cdd_ambiguous", ensemble[("ssp585", "Long-term", "CDD")]["model_classification"] == "ambiguous", ensemble[("ssp585", "Long-term", "CDD")]["model_classification"])

trend = json.loads((AUDIT / "trend_type1_audit.json").read_text(encoding="utf-8"))
check("trend_reproducible", trend["seed"] == 20260829 and trend["replicates"] == 2000 and trend["false_positive_rate"] == {"MK": 0.2245, "Hamed-Rao": 0.227, "TFPW-MK": 0.2665}, json.dumps(trend["false_positive_rate"], sort_keys=True))

failures = [item for item in checks if not item[1]]
for name, passed, detail in checks:
    print(f"{'PASS' if passed else 'FAIL'} | {name} | {detail}")
print(f"SUMMARY | passed={len(checks) - len(failures)} failed={len(failures)} total={len(checks)}")
print(f"DOCX_SHA256 | {sha256_path(DOCX).upper()}")
if failures:
    raise SystemExit(1)
