from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from PIL import Image
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent / "future_zip"
OUT = Path(__file__).resolve().parent / "future_zip_profile.json"


def scalar(value):
    if value is None:
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if np.isnan(value) else float(value)
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except Exception:
            pass
    return value


def read_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")


def profile_frame(frame: pd.DataFrame) -> dict:
    frame = frame.copy()
    frame.columns = [str(c) for c in frame.columns]
    numeric = frame.select_dtypes(include=[np.number])
    summaries = {}
    for col in numeric.columns:
        s = pd.to_numeric(numeric[col], errors="coerce")
        if s.notna().any():
            summaries[str(col)] = {
                "count": int(s.count()),
                "min": scalar(s.min()),
                "median": scalar(s.median()),
                "mean": scalar(s.mean()),
                "max": scalar(s.max()),
            }
    distinct = {}
    for col in frame.columns:
        n = int(frame[col].nunique(dropna=True))
        if n <= 30:
            vals = [scalar(v) for v in frame[col].dropna().drop_duplicates().head(30).tolist()]
            distinct[str(col)] = {"n": n, "values": vals}
    return {
        "rows": int(len(frame)),
        "columns": [str(c) for c in frame.columns],
        "dtypes": {str(k): str(v) for k, v in frame.dtypes.items()},
        "nulls": {str(k): int(v) for k, v in frame.isna().sum().items() if int(v)},
        "exact_duplicate_rows": int(frame.duplicated().sum()),
        "distinct_low_cardinality": distinct,
        "numeric_summary": summaries,
        "head": [
            {str(k): scalar(v) for k, v in row.items()}
            for row in frame.head(5).to_dict(orient="records")
        ],
        "tail": [
            {str(k): scalar(v) for k, v in row.items()}
            for row in frame.tail(3).to_dict(orient="records")
        ],
    }


def profile_workbook(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet_dims = {
        ws.title: {"max_row": int(ws.max_row), "max_column": int(ws.max_column)}
        for ws in wb.worksheets
    }
    wb.close()
    result = {"sheets": sheet_dims, "profiles": {}}
    for sheet_name in sheet_dims:
        try:
            result["profiles"][sheet_name] = profile_frame(read_sheet(path, sheet_name))
        except Exception as exc:
            result["profiles"][sheet_name] = {"error": repr(exc)}
    return result


def main():
    output = {
        "root": str(ROOT),
        "xlsx": {},
        "png": {},
        "pdf": {},
    }
    for path in sorted(ROOT.rglob("*.xlsx")):
        output["xlsx"][str(path.relative_to(ROOT))] = profile_workbook(path)
    for path in sorted(ROOT.rglob("*.png")):
        with Image.open(path) as image:
            output["png"][str(path.relative_to(ROOT))] = {
                "width": image.width,
                "height": image.height,
                "mode": image.mode,
                "dpi": image.info.get("dpi"),
            }
    for path in sorted(ROOT.rglob("*.pdf")):
        reader = PdfReader(path)
        output["pdf"][str(path.relative_to(ROOT))] = {
            "pages": len(reader.pages),
            "metadata": {str(k): str(v) for k, v in (reader.metadata or {}).items()},
        }
    OUT.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
