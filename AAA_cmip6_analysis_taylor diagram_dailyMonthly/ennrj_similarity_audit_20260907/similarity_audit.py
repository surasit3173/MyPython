from __future__ import annotations

import json
import math
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from PIL import Image


PHET = Path(r"D:\วารสาร APST\New เพชรบุรี\sent\Observed and Projected Precipitation Extremes over Phetchaburi Province under CMIP6 SSP Scenarios.docx")
UTT = Path(r"C:\Users\PC\Downloads\EnNRJ_Uttaradit_CMIP6_Precipitation_Extremes_Revised.docx")


def words(text: str) -> list[str]:
    return re.findall(r"[A-Za-z0-9]+(?:[–-][A-Za-z0-9]+)*", text.lower())


def normalize(text: str, abstract_values: bool = False) -> str:
    text = text.lower().replace("phetchaburi", "<province>").replace("uttaradit", "<province>")
    text = text.replace("เพชรบุรี", "<province>").replace("อุตรดิตถ์", "<province>")
    if abstract_values:
        text = re.sub(r"\b(?:35|46|45)\d{4}\b", "<station>", text)
        text = re.sub(r"(?<![a-z])[-+−]?\d+(?:[.,]\d+)*(?:%|°[ne]?|\s*(?:mm|days?|m\s*msl|w/m²))?", "<num>", text)
    return " ".join(words(text))


def ngrams(tokens: list[str], n: int) -> set[tuple[str, ...]]:
    return {tuple(tokens[i : i + n]) for i in range(max(0, len(tokens) - n + 1))}


def ngram_metrics(a: str, b: str, n: int, abstract_values: bool = False) -> dict[str, float | int]:
    aa = ngrams(words(normalize(a, abstract_values)), n)
    bb = ngrams(words(normalize(b, abstract_values)), n)
    common = aa & bb
    union = aa | bb
    return {
        "n": n,
        "common": len(common),
        "en_containment_pct": round(100 * len(common) / len(bb), 2) if bb else 0,
        "ph_containment_pct": round(100 * len(common) / len(aa), 2) if aa else 0,
        "jaccard_pct": round(100 * len(common) / len(union), 2) if union else 0,
    }


def table_text(table) -> str:
    return "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)


def image_inventory(path: Path) -> list[dict[str, object]]:
    out = []
    with zipfile.ZipFile(path) as zf:
        for name in sorted(n for n in zf.namelist() if n.startswith("word/media/") and not n.endswith("/")):
            with zf.open(name) as fh:
                try:
                    im = Image.open(fh)
                    out.append({"name": name, "format": im.format, "pixels": im.size, "dpi": im.info.get("dpi")})
                except Exception:
                    out.append({"name": name, "format": "unknown"})
    return out


def heading_like(text: str) -> bool:
    return bool(re.match(r"^\d+(?:\.\d+)*\.?\s+\S", text)) or text in {
        "Abstract", "Keywords", "Highlights", "Acknowledgments", "References"
    }


@dataclass
class Manuscript:
    path: Path
    document: object
    paragraphs: list[str]
    tables: list[str]
    headings: list[str]
    captions: list[str]
    ref_index: int
    pre_refs_paragraph_text: str
    methods_text: str
    results_text: str


def load(path: Path) -> Manuscript:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    tables = [table_text(t) for t in doc.tables]
    headings = [p for p in paragraphs if p and heading_like(p)]
    captions = [p for p in paragraphs if re.match(r"^(?:Table|Figure)\s+\d+\.", p)]
    ref_index = next((i for i, p in enumerate(paragraphs) if re.fullmatch(r"(?:\d+\.\s*)?References", p, re.I)), len(paragraphs))
    pre_refs = "\n".join(p for p in paragraphs[:ref_index] if p)
    method_start = next(i for i, p in enumerate(paragraphs) if re.match(r"^2\.\s|^2\s", p))
    results_start = next(i for i, p in enumerate(paragraphs) if re.match(r"^3\.\s|^3\s", p))
    conclusion_start = next((i for i, p in enumerate(paragraphs) if re.match(r"^4\.\s|^4\s", p)), ref_index)
    return Manuscript(
        path=path,
        document=doc,
        paragraphs=paragraphs,
        tables=tables,
        headings=headings,
        captions=captions,
        ref_index=ref_index,
        pre_refs_paragraph_text=pre_refs,
        methods_text="\n".join(p for p in paragraphs[method_start:results_start] if p),
        results_text="\n".join(p for p in paragraphs[results_start:conclusion_start] if p),
    )


def count_nonref(ms: Manuscript) -> dict[str, int]:
    paragraph_words = len(words(ms.pre_refs_paragraph_text))
    table_words = sum(len(words(t)) for t in ms.tables)
    return {"paragraph_words": paragraph_words, "table_words": table_words, "total_including_tables": paragraph_words + table_words}


def reference_entries(ms: Manuscript) -> list[str]:
    return [p for p in ms.paragraphs[ms.ref_index + 1 :] if p]


def near_paragraph_pairs(a: Manuscript, b: Manuscript) -> list[dict[str, object]]:
    pa = [(i, p) for i, p in enumerate(a.paragraphs[: a.ref_index]) if len(words(p)) >= 12]
    pb = [(i, p) for i, p in enumerate(b.paragraphs[: b.ref_index]) if len(words(p)) >= 12]
    candidates = []
    for ia, x in pa:
        nx = normalize(x, abstract_values=True)
        for ib, y in pb:
            ny = normalize(y, abstract_values=True)
            ratio = SequenceMatcher(None, nx, ny, autojunk=False).ratio()
            if ratio >= 0.68:
                candidates.append({
                    "ratio": round(ratio, 3),
                    "phet_paragraph": ia,
                    "utt_paragraph": ib,
                    "phet": x,
                    "utt": y,
                })
    candidates.sort(key=lambda x: x["ratio"], reverse=True)
    return candidates[:30]


def table_pair_scores(a: Manuscript, b: Manuscript) -> list[dict[str, object]]:
    scores = []
    for i, x in enumerate(a.tables, 1):
        for j, y in enumerate(b.tables, 1):
            raw = SequenceMatcher(None, normalize(x), normalize(y), autojunk=False).ratio()
            abstract = SequenceMatcher(None, normalize(x, True), normalize(y, True), autojunk=False).ratio()
            scores.append({"phet_table": i, "utt_table": j, "raw_ratio": round(raw, 3), "abstracted_ratio": round(abstract, 3)})
    return sorted(scores, key=lambda x: x["abstracted_ratio"], reverse=True)


def sentence_pairs(a: Manuscript, b: Manuscript) -> list[dict[str, object]]:
    splitter = re.compile(r"(?<=[.!?])\s+")
    sa = [s.strip() for s in splitter.split(a.pre_refs_paragraph_text) if len(words(s)) >= 10]
    sb = [s.strip() for s in splitter.split(b.pre_refs_paragraph_text) if len(words(s)) >= 10]
    out = []
    for x in sa:
        nx = normalize(x, True)
        for y in sb:
            ny = normalize(y, True)
            ratio = SequenceMatcher(None, nx, ny, autojunk=False).ratio()
            if ratio >= 0.78:
                out.append({"ratio": round(ratio, 3), "phet": x, "utt": y})
    out.sort(key=lambda x: x["ratio"], reverse=True)
    return out[:40]


phet = load(PHET)
utt = load(UTT)

report = {
    "files": {"phetchaburi": str(PHET), "uttaradit": str(UTT)},
    "word_counts_excluding_references": {
        "phetchaburi": count_nonref(phet),
        "uttaradit": count_nonref(utt),
    },
    "structure": {
        "phetchaburi": {
            "paragraphs": len(phet.paragraphs),
            "tables": [(len(t.rows), len(t.columns)) for t in phet.document.tables],
            "inline_shapes": len(phet.document.inline_shapes),
            "omml": len(phet.document.element.body.xpath(".//m:oMath")),
            "references": len(reference_entries(phet)),
            "headings": phet.headings,
            "captions": phet.captions,
            "images": image_inventory(PHET),
        },
        "uttaradit": {
            "paragraphs": len(utt.paragraphs),
            "tables": [(len(t.rows), len(t.columns)) for t in utt.document.tables],
            "inline_shapes": len(utt.document.inline_shapes),
            "omml": len(utt.document.element.body.xpath(".//m:oMath")),
            "references": len(reference_entries(utt)),
            "headings": utt.headings,
            "captions": utt.captions,
            "images": image_inventory(UTT),
        },
    },
    "whole_manuscript_ngram_overlap": {
        "raw_5gram": ngram_metrics(phet.pre_refs_paragraph_text, utt.pre_refs_paragraph_text, 5),
        "raw_8gram": ngram_metrics(phet.pre_refs_paragraph_text, utt.pre_refs_paragraph_text, 8),
        "abstracted_5gram": ngram_metrics(phet.pre_refs_paragraph_text, utt.pre_refs_paragraph_text, 5, True),
        "abstracted_8gram": ngram_metrics(phet.pre_refs_paragraph_text, utt.pre_refs_paragraph_text, 8, True),
    },
    "methods_ngram_overlap": {
        "raw_5gram": ngram_metrics(phet.methods_text, utt.methods_text, 5),
        "raw_8gram": ngram_metrics(phet.methods_text, utt.methods_text, 8),
        "abstracted_5gram": ngram_metrics(phet.methods_text, utt.methods_text, 5, True),
        "abstracted_8gram": ngram_metrics(phet.methods_text, utt.methods_text, 8, True),
    },
    "results_ngram_overlap": {
        "raw_5gram": ngram_metrics(phet.results_text, utt.results_text, 5),
        "raw_8gram": ngram_metrics(phet.results_text, utt.results_text, 8),
        "abstracted_5gram": ngram_metrics(phet.results_text, utt.results_text, 5, True),
        "abstracted_8gram": ngram_metrics(phet.results_text, utt.results_text, 8, True),
    },
    "table_pair_scores": table_pair_scores(phet, utt)[:12],
    "near_paragraph_pairs": near_paragraph_pairs(phet, utt),
    "near_sentence_pairs": sentence_pairs(phet, utt),
    "targeted_checks": {
        "utt_R95p_4_5_hits": [p for p in utt.paragraphs if "R95p" in p and "4.5%" in p],
        "utt_R95p_4_4_hits": [p for p in utt.paragraphs if "R95p" in p and "4.4" in p],
        "utt_Tebaldi_body_hits": [p for p in utt.paragraphs[: utt.ref_index] if "Tebaldi" in p],
        "utt_record_count_hits": [p for p in utt.paragraphs if "4,862" in p or "60,060" in p],
        "phet_record_count_hits": [p for p in phet.paragraphs if "4,862" in p or "60,060" in p],
    },
}

print(json.dumps(report, ensure_ascii=False, indent=2))
