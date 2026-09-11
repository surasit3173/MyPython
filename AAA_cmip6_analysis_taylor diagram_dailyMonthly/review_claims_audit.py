import io
import json
import re
import sys
import zipfile
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def words(text):
    return re.findall(r"[A-Za-z0-9]+(?:[\-\u2013\u2014'][A-Za-z0-9]+)*", text)


def norm(text):
    text = text.casefold().replace("phetchaburi", "province").replace("uttaradit", "province")
    text = re.sub(r"\d+(?:\.\d+)?", "#", text)
    return " ".join(words(text))


def iter_blocks(doc):
    body = doc.element.body
    p_by_el = {p._p: p for p in doc.paragraphs}
    t_by_el = {t._tbl: t for t in doc.tables}
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = p_by_el[child]
            yield "p", p.text, p.style.name if p.style else ""
        elif tag == "tbl":
            t = t_by_el[child]
            rows = [[c.text for c in row.cells] for row in t.rows]
            yield "t", rows, ""


def reference_split(blocks):
    start = None
    for i, (kind, data, style) in enumerate(blocks):
        if kind == "p" and data.strip().casefold() in {"references", "8. references", "8 references"}:
            start = i
            break
    return (blocks[:start], blocks[start + 1 :]) if start is not None else (blocks, [])


def block_text(blocks):
    out = []
    for kind, data, _ in blocks:
        if kind == "p":
            out.append(data)
        else:
            for row in data:
                out.extend(row)
    return "\n".join(out)


def paragraphs_by_section(doc):
    sections = {}
    current = "front"
    for p in doc.paragraphs:
        s = p.text.strip()
        if not s:
            continue
        if re.match(r"^(?:\d+(?:\.\d+)*\.?\s+|Abstract$|Keywords?:|Highlights?:|References$)", s, re.I):
            # Treat numbered lines that look like headings as headings only if short.
            if len(words(s)) <= 10 or s.lower().startswith(("keywords", "highlights")):
                current = s
        sections.setdefault(current, []).append(s)
    return sections


def image_info(path):
    infos = []
    with zipfile.ZipFile(path) as z:
        for name in sorted(n for n in z.namelist() if n.startswith("word/media/")):
            data = z.read(name)
            try:
                im = Image.open(io.BytesIO(data))
                dpi = im.info.get("dpi")
                infos.append({"name": name, "format": im.format, "size": im.size, "dpi": dpi})
            except Exception as exc:
                infos.append({"name": name, "error": str(exc), "bytes": len(data)})
    return infos


def inspect(path):
    p = Path(path)
    doc = Document(str(p))
    blocks = list(iter_blocks(doc))
    pre, refs = reference_split(blocks)
    all_text = block_text(blocks)
    body_text = block_text(pre)
    ref_text = block_text(refs)
    p_text = "\n".join(p.text for p in doc.paragraphs)
    t_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    with zipfile.ZipFile(p) as z:
        xml = z.read("word/document.xml")
        omath = xml.count(b"<m:oMath") - xml.count(b"<m:oMathPara")
        omath_para = xml.count(b"<m:oMathPara")
        drawings = xml.count(b"<w:drawing")
        inlines = xml.count(b"<wp:inline")
        anchors = xml.count(b"<wp:anchor")
    headings = [(i, p.text.strip(), p.style.name if p.style else "") for i, p in enumerate(doc.paragraphs)
                if p.text.strip() and (p.style and "heading" in p.style.name.lower() or re.match(r"^\d+(?:\.\d+)*\.?\s", p.text.strip()))]
    refs_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # Reference paragraphs after the References heading only.
    ref_pars = []
    in_refs = False
    for p0 in doc.paragraphs:
        s = p0.text.strip()
        if s.casefold() in {"references", "8. references", "8 references"}:
            in_refs = True
            continue
        if in_refs and s:
            ref_pars.append(s)
    table_shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
    table_first_rows = [[[c.text for c in row.cells] for row in t.rows[:2]] for t in doc.tables]
    result = {
        "path": str(p),
        "size": p.stat().st_size,
        "paragraph_words_all": len(words(p_text)),
        "table_words_all": len(words(t_text)),
        "all_words": len(words(all_text)),
        "body_words_including_tables": len(words(body_text)),
        "reference_words": len(words(ref_text)),
        "reference_paragraphs": len(ref_pars),
        "tables": len(doc.tables),
        "table_shapes": table_shapes,
        "table_first_rows": table_first_rows,
        "inline_shapes_api": len(doc.inline_shapes),
        "drawings_xml": drawings,
        "inlines_xml": inlines,
        "anchors_xml": anchors,
        "omath": omath,
        "omath_para": omath_para,
        "images": image_info(p),
        "headings": headings,
        "hits": {},
        "text": all_text,
        "body_text": body_text,
        "ref_pars": ref_pars,
    }
    for pattern in [
        r"R95p[^\n]{0,100}", r"4\.4%", r"4\.5%", r"Tebaldi[^\n]{0,200}",
        r"4,862", r"60,060", r"30/143", r"14/143", r"93/143", r"88/143",
        r"elevation[^\n]{0,180}", r"did not test[^\n]{0,180}", r"not used[^\n]{0,180}",
        r"IDW[^\n]{0,180}", r"convex hull[^\n]{0,180}", r"not[^\n]{0,60}validated[^\n]{0,100}",
    ]:
        result["hits"][pattern] = re.findall(pattern, all_text, flags=re.I)
    return result


def ngrams(text, n=8):
    ws = norm(text).split()
    return {tuple(ws[i:i+n]) for i in range(len(ws)-n+1)}


def compare(a, b):
    na = norm(a["body_text"]).split()
    nb = norm(b["body_text"]).split()
    ng_a, ng_b = ngrams(a["body_text"]), ngrams(b["body_text"])
    # Exact paragraphs >= 12 words (with province/numbers normalized).
    pa = Counter(norm(x) for x in a["body_text"].splitlines() if len(words(x)) >= 12)
    pb = Counter(norm(x) for x in b["body_text"].splitlines() if len(words(x)) >= 12)
    common_pars = sorted(((min(pa[k], pb[k]), k) for k in pa.keys() & pb.keys()), reverse=True)
    return {
        "sequence_ratio_normalized_words": SequenceMatcher(None, na, nb, autojunk=True).ratio(),
        "8gram_a": len(ng_a), "8gram_b": len(ng_b), "8gram_intersection": len(ng_a & ng_b),
        "8gram_jaccard": len(ng_a & ng_b) / max(1, len(ng_a | ng_b)),
        "8gram_recall_a": len(ng_a & ng_b) / max(1, len(ng_a)),
        "8gram_recall_b": len(ng_a & ng_b) / max(1, len(ng_b)),
        "common_exact_normalized_paragraph_instances": sum(c for c, _ in common_pars),
        "common_paragraph_samples": [x for c, x in common_pars[:20]],
    }


if __name__ == "__main__":
    a = inspect(sys.argv[1])
    b = inspect(sys.argv[2])
    stripped = lambda d: {k:v for k,v in d.items() if k not in {"text", "body_text", "ref_pars"}}
    rendered = json.dumps({"a": stripped(a), "b": stripped(b), "compare": compare(a,b),
                           "a_ref_pars": a["ref_pars"], "b_ref_pars": b["ref_pars"]},
                          ensure_ascii=False, indent=2, default=str)
    if len(sys.argv) > 3:
        Path(sys.argv[3]).write_text(rendered, encoding="utf-8")
        print(sys.argv[3])
    else:
        print(rendered)
