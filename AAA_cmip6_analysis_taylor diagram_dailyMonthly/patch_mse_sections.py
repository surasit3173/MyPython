from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


PATH = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")


METHODOLOGY_MSE = [
    "2.7 การแยกองค์ประกอบของค่าความคลาดเคลื่อนกำลังสองเฉลี่ย",
    "เพื่ออธิบายกลไกที่ทำให้ตัวชี้วัดเชิงการแจกแจงและตัวชี้วัดที่อาศัยการจับคู่ตามเวลาให้ผลต่างกัน การศึกษานี้แยกองค์ประกอบของค่าความคลาดเคลื่อนกำลังสองเฉลี่ย (mean square error: MSE) ของปริมาณฝนรายเดือนเป็นสองระดับ ระดับแรกใช้การแยกส่วนตาม Murphy (1988) ซึ่งเขียน MSE เป็นผลรวมขององค์ประกอบจากความลำเอียงของค่าเฉลี่ย ความแตกต่างของการกระจาย และการขาดสหสัมพันธ์ ดังสมการที่ (5)",
    "MSE = b2 + (σs − rσo)2 + σo2(1 − r2)          (5)",
    "โดยที่ b คือผลต่างของค่าเฉลี่ยระหว่างข้อมูลจำลองและข้อมูลสังเกตการณ์ σs และ σo คือส่วนเบี่ยงเบนมาตรฐานของข้อมูลจำลองและข้อมูลสังเกตการณ์ตามลำดับ และ r คือสัมประสิทธิ์สหสัมพันธ์เพียร์สันของอนุกรมรายเดือน สมการนี้ทำให้เห็นว่า MSE มิได้ขึ้นกับความลำเอียงเพียงอย่างเดียว แต่ขึ้นกับการกระจายและความสอดคล้องของจังหวะเวลาไปพร้อมกัน สำหรับค่า r ที่กำหนด MSE จะต่ำที่สุดเมื่อ σs = rσo หรือเมื่ออัตราส่วนการกระจาย σs/σo เท่ากับ r มิใช่เท่ากับหนึ่งเสมอไป",
    "การแยกส่วนระดับแรกคำนวณสำหรับแต่ละคู่สถานี-แบบจำลอง แล้วนำองค์ประกอบของข้อมูลหลังปรับแก้ลบด้วยองค์ประกอบของข้อมูลดิบเพื่อประเมินว่า MSE เปลี่ยนไปจากองค์ประกอบใด การสรุปข้ามคู่การประเมิน 60 คู่ใช้ค่าเฉลี่ยเลขคณิต เนื่องจากค่าเฉลี่ยรักษาสมบัติการบวกขององค์ประกอบ กล่าวคือ ผลรวมของค่าเฉลี่ยองค์ประกอบเท่ากับค่าเฉลี่ยของผลรวม ส่วนค่ามัธยฐานไม่ได้ใช้เป็นค่าสรุปหลักของการแยกส่วน เพราะมัธยฐานของแต่ละองค์ประกอบไม่จำเป็นต้องบวกกันเท่ากับมัธยฐานของ MSE รวม",
    "ระดับที่สองใช้เพื่อตรวจว่าความแตกต่างระหว่างการประมาณแบบรวมทุกเดือนกับการประมาณแยกตามเดือนปฏิทินเกิดจากการปรับวัฏจักรฤดูกาลหรือจากการปรับค่าเบี่ยงภายในฤดูกาล การศึกษานี้เขียนผลต่างรายเดือนระหว่างข้อมูลจำลองและข้อมูลสังเกตการณ์เป็นผลรวมของผลต่างค่าเฉลี่ยประจำเดือนปฏิทินและผลต่างของค่าเบี่ยงจากค่าเฉลี่ยประจำเดือนนั้น จากนั้นแยก MSE เป็นองค์ประกอบวัฏจักรฤดูกาลและองค์ประกอบค่าเบี่ยง ดังสมการที่ (6)",
    "MSE = MSEclim + MSEanom          (6)",
    "โดย MSEclim คือส่วนของความคลาดเคลื่อนที่เกิดจากความแตกต่างของค่าเฉลี่ยประจำเดือนปฏิทินระหว่างข้อมูลจำลองและข้อมูลสังเกตการณ์ และ MSEanom คือส่วนของความคลาดเคลื่อนที่เกิดจากค่าเบี่ยงรายเดือนหลังหักค่าเฉลี่ยประจำเดือนปฏิทินออกแล้ว พจน์ไขว้มีค่าเป็นศูนย์ตามนิยาม เนื่องจากค่าเบี่ยงมีค่าเฉลี่ยเป็นศูนย์ภายในแต่ละเดือนปฏิทิน จึงทำให้สององค์ประกอบรวมกันได้เท่ากับ MSE รวมอย่างพอดี การตรวจสอบเชิงตัวเลขยืนยันว่าผลรวมของสององค์ประกอบเท่ากับค่ารวมโดยมีความคลาดเคลื่อนไม่เกิน 1 x 10-11",
    "การแยกส่วนทั้งสองระดับใช้เป็นเครื่องมือวินิจฉัย มิใช่การทดสอบเชิงสาเหตุโดยลำพัง ผลจากสมการที่ (5) ใช้อธิบายว่าองค์ประกอบใดสัมพันธ์กับการเปลี่ยนแปลงของ RMSE ภายหลัง QDM ส่วนสมการที่ (6) ใช้แยกว่าการเปลี่ยนแปลงของ RMSE และ NSE ระหว่างรูปแบบการจัดกลุ่มเชิงเวลาเกิดจากวัฏจักรฤดูกาลหรือจากความสอดคล้องของค่าเบี่ยงรายเดือน ซึ่งเป็นประเด็นสำคัญต่อการตีความสมรรถนะเชิงเวลาในช่วงตรวจสอบแบบอิสระ",
]


RESULTS_35 = [
    "3.5 การแยกส่วนค่าความคลาดเคลื่อนกำลังสองเฉลี่ย",
    "การแยกองค์ประกอบของ MSE ช่วยอธิบายผลที่ดูเหมือนขัดกันระหว่างหัวข้อ 3.2 ถึง 3.4 กล่าวคือ QDM ทำให้การแจกแจงของปริมาณฝนสมจริงขึ้น แต่ดัชนีที่อาศัยการจับคู่ตามเวลาไม่ได้ดีขึ้นเสมอไป ผลตามรูปที่ 5(a) แสดงว่าสัมประสิทธิ์สหสัมพันธ์รายเดือนของข้อมูลดิบมีค่าเฉลี่ยเพียง 0.342 ขณะที่อัตราส่วนการกระจายของข้อมูลดิบมีค่าเฉลี่ย 1.169 ซึ่งสูงกว่าค่าที่ทำให้ MSE ต่ำสุดในทุกคู่การประเมิน ภายหลังการปรับแก้ด้วย QDM อัตราส่วนดังกล่าวเพิ่มขึ้นเป็น 1.286 และยังคงสูงกว่าค่าที่เหมาะสมในทุกคู่",
    "ความสัมพันธ์ดังกล่าวปรากฏชัดในรูปที่ 5(b) โดยการเปลี่ยนแปลงของอัตราส่วนการกระจายมีความสัมพันธ์เชิงบวกสูงกับการเปลี่ยนแปลงของ RMSE ค่าสัมประสิทธิ์สหสัมพันธ์อันดับของสเปียร์แมนเท่ากับ 0.95 และอยู่ระหว่าง 0.73 ถึง 1.00 เมื่อคำนวณแยกรายแบบจำลอง ทั้งนี้ ค่าดังกล่าวรายงานเชิงพรรณนาโดยไม่ระบุค่า p เพราะคู่การประเมินจำนวนมากไม่เป็นอิสระต่อกันตามข้อจำกัดของข้อมูลกริดที่ระบุไว้ในหัวข้อ 2.2",
    "เมื่อแยกการเปลี่ยนแปลงของ MSE ตามองค์ประกอบในสมการที่ (5) พบว่าองค์ประกอบด้านการกระจายเพิ่มขึ้นเฉลี่ย +2,891 (มิลลิเมตรต่อเดือน) ยกกำลังสอง ขณะที่องค์ประกอบด้านความลำเอียงลดลงเฉลี่ย -237 และองค์ประกอบจากการขาดสหสัมพันธ์เพิ่มขึ้นเฉลี่ย +77 ผลรวมของสามองค์ประกอบจึงเท่ากับการเปลี่ยนแปลงสุทธิ +2,731 ดังแสดงในรูปที่ 5(c) เมื่อพิจารณารายคู่ องค์ประกอบด้านการกระจายเป็นองค์ประกอบที่มีขนาดใหญ่ที่สุดใน 56 จาก 60 คู่ ผลนี้ชี้ว่า ในชุดข้อมูลนี้ การคืนความเข้มและความแปรปรวนของฝนให้ใกล้ข้อมูลตรวจวัดมากขึ้นเกิดขึ้นภายใต้สหสัมพันธ์เชิงเวลาที่ต่ำ จึงสัมพันธ์กับการเพิ่มขึ้นของ RMSE ในหลายคู่การประเมิน",
    "ผลการแยกส่วนระดับที่สองตามสมการที่ (6) แสดงในรูปที่ 5(d) และให้คำอธิบายต่อความแตกต่างระหว่างรูปแบบการจัดกลุ่มเชิงเวลา การประมาณแยกตามเดือนปฏิทินลดองค์ประกอบวัฏจักรฤดูกาลลง -3,202 (มิลลิเมตรต่อเดือน) ยกกำลังสอง ขณะที่การประมาณแบบรวมทุกเดือนเพิ่มองค์ประกอบนี้ขึ้น +412 ในทางตรงกันข้าม องค์ประกอบค่าเบี่ยงรายเดือนด้อยลงในขนาดใกล้เคียงกันทั้งสองรูปแบบ โดยเพิ่มขึ้น +2,132 สำหรับการประมาณแยกตามเดือนปฏิทิน และ +2,319 สำหรับการประมาณแบบรวมทุกเดือน",
    "เมื่อเปรียบเทียบผลต่างระหว่างสองรูปแบบ องค์ประกอบวัฏจักรฤดูกาลต่างกัน 3,614 ขณะที่องค์ประกอบค่าเบี่ยงต่างกันเพียง 187 จากผลต่างรวม 3,801 หรือคิดเป็นร้อยละ 95.1 และ 4.9 ตามลำดับ ดังนั้น การที่ RMSE, NSE, KGE และสัมประสิทธิ์สหสัมพันธ์รายเดือนดีขึ้นภายใต้การประมาณแยกตามเดือนปฏิทินจึงเกิดจากการแก้ความคลาดเคลื่อนของวัฏจักรฤดูกาลเป็นส่วนใหญ่ มิใช่จากการฟื้นคืนความสอดคล้องของค่าเบี่ยงรายเดือน ผลนี้สอดคล้องกับหัวข้อ 3.4 ซึ่งพบว่าสัมประสิทธิ์สหสัมพันธ์ของค่าเบี่ยงรายเดือนไม่ดีขึ้นในทั้งสองรูปแบบ",
    "โดยสรุป การแยกส่วน MSE สนับสนุนข้อค้นพบหลักว่า ความสมจริงเชิงการแจกแจง ความถูกต้องของวัฏจักรฤดูกาล และความสอดคล้องของลำดับเวลาเป็นสมรรถนะคนละมิติ QDM สามารถปรับการแจกแจงและวัฏจักรฤดูกาลได้ดีขึ้นภายใต้บางรูปแบบการประมาณ แต่ไม่ควรถูกตีความว่าสร้างสารสนเทศด้านจังหวะเวลาของเหตุการณ์ฝนที่ไม่มีอยู่ในแบบจำลองเดิม",
]


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def copy_paragraph_format(src, dst):
    dst.style = src.style
    dst.alignment = src.alignment
    dst.paragraph_format.left_indent = src.paragraph_format.left_indent
    dst.paragraph_format.right_indent = src.paragraph_format.right_indent
    dst.paragraph_format.first_line_indent = src.paragraph_format.first_line_indent
    dst.paragraph_format.space_before = src.paragraph_format.space_before
    dst.paragraph_format.space_after = src.paragraph_format.space_after
    dst.paragraph_format.line_spacing = src.paragraph_format.line_spacing


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_section(doc, start_idx, end_idx, new_texts):
    anchor = doc.paragraphs[end_idx]
    template = doc.paragraphs[start_idx]
    for p in list(doc.paragraphs[start_idx:end_idx]):
        delete_paragraph(p)
    for text in new_texts:
        new_p = anchor.insert_paragraph_before("")
        copy_paragraph_format(template, new_p)
        set_text(new_p, text)
        if text.startswith("MSE ="):
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            new_p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def insert_section_before(doc, before_idx, new_texts):
    anchor = doc.paragraphs[before_idx]
    inserted = []
    for text in new_texts:
        p = anchor.insert_paragraph_before("")
        copy_paragraph_format(anchor, p)
        set_text(p, text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if text.startswith("MSE =") else WD_ALIGN_PARAGRAPH.LEFT
        inserted.append(p)
    return list(reversed(inserted))


def find_heading(doc, prefix):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    raise ValueError(prefix)


def fix_figure5_caption(doc):
    for p in doc.paragraphs:
        if p.text.strip().startswith("Figure 5  "):
            set_text(
                p,
                "Figure 5  Factors associated with the change in error metrics after QDM correction, independent validation period. (a) Dispersion ratio against the monthly correlation, with the dashed line marking the RMSE-optimal ratio; (b) change in dispersion ratio against change in RMSE for the annual-pooled configuration; (c) decomposition of the change in mean square error into bias, dispersion and lack-of-correlation components following Equation (5); and (d) split of the change in mean square error into an annual-cycle component and an anomaly component following Equation (6), shown for both temporal-grouping configurations. Bars in (c) and (d) are means across the 60 evaluation combinations and are additive, so the components sum exactly to the total; whiskers in (c) show the interquartile range across combinations, not the uncertainty of the mean",
            )
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            break


def main():
    doc = Document(PATH)
    i_results = find_heading(doc, "Results")
    try:
        i27 = find_heading(doc, "2.7 การแยกองค์ประกอบ")
        insert_section(doc, i27, i_results, METHODOLOGY_MSE)
    except ValueError:
        insert_section_before(doc, i_results, METHODOLOGY_MSE)

    i35 = find_heading(doc, "3.5 ")
    i36 = find_heading(doc, "3.6 ")
    insert_section(doc, i35, i36, RESULTS_35)
    fix_figure5_caption(doc)
    doc.save(PATH)
    print(PATH)


if __name__ == "__main__":
    main()
