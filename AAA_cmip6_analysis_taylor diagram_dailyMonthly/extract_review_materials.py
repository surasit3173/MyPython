from pathlib import Path
from docx import Document

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

files = {
    "revised_article": Path(r"D:\วารสารบูรพา\Evaluation of Quantile Delta Mapping for Bias Correction of CMIP6 Daily Rainfall Using Independent Temporal Holdout.docx"),
    "response_doc": Path(r"C:\Users\PC\Downloads\ID1558_การวิเคราะห์คำชี้แจงผู้ทรงคุณวุฒิ_ฉบับภาษาไทย.docx"),
    "old_pdf": Path(r"C:\Users\PC\Downloads\Research_Article_ID1558_Reviewer.pdf"),
}

out_dir = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\review_response_extracts")
out_dir.mkdir(exist_ok=True)

def extract_docx(path: Path, out_name: str):
    doc = Document(path)
    lines = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            style = p.style.name if p.style else ""
            lines.append(f"{i:04d}\t{style}\t{text}")
    for ti, table in enumerate(doc.tables, 1):
        lines.append(f"\n[TABLE {ti}]")
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            lines.append(f"{ri:03d}\t" + "\t".join(cells))
    out = out_dir / out_name
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {out}")

def extract_pdf(path: Path, out_name: str):
    if PdfReader is None:
        print("pypdf unavailable")
        return
    reader = PdfReader(str(path))
    lines = []
    for pi, page in enumerate(reader.pages, 1):
        lines.append(f"\n===== PAGE {pi} =====")
        try:
            lines.append(page.extract_text() or "")
        except Exception as e:
            lines.append(f"[extract error: {e}]")
    out = out_dir / out_name
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {out}")

extract_docx(files["revised_article"], "revised_article.txt")
extract_docx(files["response_doc"], "response_doc.txt")
extract_pdf(files["old_pdf"], "old_pdf.txt")
