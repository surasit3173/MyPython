from pathlib import Path
from docx import Document


PATH = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")
TERMS = [
    "Supplementary Material",
    "cross-validation",
    "Cross-Validation",
    "ตรวจสอบความถูกต้องข้ามช่วงเวลา",
    "การกันข้อมูลช่วงเวลา",
    "https:/doi.org",
    "doi.org10",
    "เขคดุสิต",
]

doc = Document(PATH)
txt = "\n".join(p.text for p in doc.paragraphs)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
for term in TERMS:
    print(f"{term}: {txt.count(term)}")
print("FIRST 28 PARAGRAPHS")
for i, p in enumerate(doc.paragraphs[:28], 1):
    print(f"{i:03d}: {p.text}")
print("TAIL")
for i, p in enumerate(doc.paragraphs[-36:], len(doc.paragraphs) - 35):
    print(f"{i:03d}: {p.text}")
