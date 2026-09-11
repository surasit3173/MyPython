import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


BASE = Path(__file__).resolve().parent
DOCX = BASE / "Manuscript_MBCn_Q3_Revised_6Figures.docx"
OUT = BASE / "final_structural_audit.json"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def words(text):
    return re.findall(r"\b[\w-]+\b", text, flags=re.UNICODE)


doc = Document(DOCX)
paras = doc.paragraphs
texts = [p.text.strip() for p in paras]

abstract_index = texts.index("Abstract")
keywords_index = next(i for i, t in enumerate(texts) if t.startswith("Keywords:"))
abstract_text = " ".join(texts[abstract_index + 1 : keywords_index])
keywords = [x.strip() for x in texts[keywords_index].split(":", 1)[1].split(";") if x.strip()]

references_index = texts.index("References")
body_paragraphs = texts[:references_index]
body_text = "\n".join(body_paragraphs)
reference_texts = [t for t in texts[references_index + 1 :] if re.match(r"^\[\d+\]", t)]

citations_in_sequence = []
for match in re.finditer(r"\[(\d+)\](?:\s*[-\u2013]\s*\[(\d+)\])?|\[(\d+)\s*[-\u2013]\s*(\d+)\]", body_text):
    start = int(match.group(1) or match.group(3))
    end = int(match.group(2) or match.group(4) or start)
    citations_in_sequence.extend(range(start, end + 1))
first_appearance = []
for number in citations_in_sequence:
    if number not in first_appearance:
        first_appearance.append(number)

figure_captions = [t for t in texts if re.match(r"^Figure \d+\.", t)]
table_captions = [t for t in texts if re.match(r"^Table \d+\.", t)]
noncaption_body = "\n".join(
    t for t in body_paragraphs if not re.match(r"^(Figure|Table) \d+\.", t)
)
figure_refs = [int(n) for n in re.findall(r"\bFigure\s+(\d+)\b", noncaption_body)]
table_refs = [int(n) for n in re.findall(r"\bTable\s+(\d+)\b", noncaption_body)]

with zipfile.ZipFile(DOCX) as zf:
    bad_crc = zf.testzip()
    document_xml = zf.read("word/document.xml")
    root = etree.fromstring(document_xml)
    ns = {"w": W, "wp": WP}
    insertions = len(root.xpath("//w:ins", namespaces=ns))
    deletions = len(root.xpath("//w:del", namespaces=ns))
    comments = [name for name in zf.namelist() if name.startswith("word/comments")]
    media_parts = sorted(name for name in zf.namelist() if name.startswith("word/media/"))
    doc_prs = root.xpath("//wp:docPr", namespaces=ns)
    image_alt = [
        {"title": node.get("title"), "description": node.get("descr")}
        for node in doc_prs
    ]
    body = root.find(f"{{{W}}}body")
    children = list(body)
    table_positions = [i for i, node in enumerate(children) if node.tag == f"{{{W}}}tbl"]
    table_caption_positions = []
    figure_positions = []
    figure_caption_positions = []
    for i, node in enumerate(children):
        if node.tag != f"{{{W}}}p":
            continue
        text = "".join(node.xpath(".//w:t/text()", namespaces=ns)).strip()
        if node.xpath(".//wp:inline", namespaces=ns):
            figure_positions.append(i)
        if re.match(r"^Table \d+\.", text):
            table_caption_positions.append(i)
        if re.match(r"^Figure \d+\.", text):
            figure_caption_positions.append(i)

author_ok = (
    "Surasit Punyawansiri" in body_text
    and "Office of Water Management and Hydrology, Royal Irrigation Department, Dusit, Bangkok, Thailand" in body_text
    and "Surasit.pum@thaimooc.ac.th" in body_text
)

report = {
    "file": str(DOCX),
    "zip_crc_error": bad_crc,
    "sections": len(doc.sections),
    "abstract_word_count": len(words(abstract_text)),
    "keyword_count": len(keywords),
    "keywords": keywords,
    "author_block_present": author_ok,
    "table_count": len(doc.tables),
    "table_caption_count": len(table_captions),
    "table_caption_numbers": [int(re.match(r"^Table (\d+)\.", t).group(1)) for t in table_captions],
    "table_captions_after_tables": len(table_positions) == len(table_caption_positions) == 5
    and all(caption > table for table, caption in zip(table_positions, table_caption_positions)),
    "table_references": table_refs,
    "table_reference_range_ok": bool(table_refs) and min(table_refs) >= 1 and max(table_refs) <= 5,
    "figure_count": len(doc.inline_shapes),
    "figure_caption_count": len(figure_captions),
    "figure_caption_numbers": [int(re.match(r"^Figure (\d+)\.", t).group(1)) for t in figure_captions],
    "figure_captions_after_figures": len(figure_positions) == len(figure_caption_positions) == 6
    and all(caption > figure for figure, caption in zip(figure_positions, figure_caption_positions)),
    "figure_references": figure_refs,
    "figure_reference_range_ok": bool(figure_refs) and min(figure_refs) >= 1 and max(figure_refs) <= 6,
    "reference_count": len(reference_texts),
    "reference_numbers": [int(re.match(r"^\[(\d+)\]", t).group(1)) for t in reference_texts],
    "citation_first_appearance": first_appearance,
    "citation_order_complete": first_appearance == list(range(1, 25)),
    "tracked_insertions": insertions,
    "tracked_deletions": deletions,
    "comment_parts": comments,
    "media_part_count": len(media_parts),
    "media_parts": media_parts,
    "images_with_alt_text": sum(bool(x["description"] and x["description"].strip()) for x in image_alt),
    "image_alt": image_alt,
}

report["all_checks_pass"] = all(
    [
        report["zip_crc_error"] is None,
        report["author_block_present"],
        report["keyword_count"] == 5,
        report["table_count"] == 5,
        report["table_caption_numbers"] == list(range(1, 6)),
        report["table_captions_after_tables"],
        report["table_reference_range_ok"],
        report["figure_count"] == 6,
        report["figure_caption_numbers"] == list(range(1, 7)),
        report["figure_captions_after_figures"],
        report["figure_reference_range_ok"],
        report["reference_count"] == 24,
        report["reference_numbers"] == list(range(1, 25)),
        report["citation_order_complete"],
        report["tracked_insertions"] == 0,
        report["tracked_deletions"] == 0,
        not report["comment_parts"],
        report["media_part_count"] == 6,
        report["images_with_alt_text"] == 6,
    ]
)

OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
