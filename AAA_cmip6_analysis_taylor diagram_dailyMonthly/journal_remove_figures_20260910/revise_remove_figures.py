from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "source_manuscript.docx"
OUTPUT = BASE / "Manuscript_MBCn_Q3_Revised_6Figures.docx"


BODY_REPLACEMENTS = {
    "Twelve daily rain gauges in Prachuap Khiri Khan provinces": (
        "Twelve daily rain gauges in Prachuap Khiri Khan Province, on the western shore "
        "of the Gulf of Thailand, provide the observations (Table 1 and Figure 1). Each record is complete "
        "from 1 January 1981 to 31 December 2014. Mean annual precipitation ranges from approximately 944 "
        "to 1,376 mm, and wet-day frequency varies markedly across the network. This spatial contrast provides "
        "a demanding test of whether correction preserves local distributions while improving network dependence."
    ),
    "Coordinates and elevations were obtained from the project metadata archive": (
        "Coordinates and elevations were obtained from the project metadata archive. Figure 1 locates the gauges "
        "within Prachuap Khiri Khan Province; the elevation background and national inset provide geographic context "
        "only and are not used in the analysis."
    ),
    "The consistency across models is empirical rather than guaranteed": (
        "The model-level results show that the ensemble mean was not driven by a small number of influential "
        "models. The seven-of-seven improvement in Table 5 occurred for energy distance, three correlation "
        "measures, joint-exceedance error, and wet-extent distance. Agreement across these related diagnostics "
        "indicates that MBCn improved several features of the joint precipitation structure rather than a single "
        "summary statistic."
    ),
    "Figure 6 relates pairwise Spearman correlation to gauge separation": (
        "Spatially, pairwise Spearman correlation declined modestly with gauge separation in the observations "
        "(Figure 4). Raw output remained near 0.96 across most distances, indicating pronounced over-coherence. "
        "QDM reduced this bias, whereas MBCn most closely followed the observed distance relationship. Some raw "
        "over-coherence is expected when gauges share or occupy adjacent coarse grid cells. The binned relationship "
        "and bootstrap bands are descriptive because the 66 gauge pairs are not independent."
    ),
    "Two event-oriented diagnostics complement the correlation analysis": (
        "Concurrent-extreme diagnostics led to the same ordering. Figure 5 compares observed and simulated "
        "probabilities that both gauges in a pair exceeded their own 95th wet-day percentiles on the same date. "
        "Raw output was widely dispersed, QDM retained a positive bias, and MBCn lay closest to the 1:1 line without "
        "a systematic tilt."
    ),
    "Figure 8 shows the pooled distribution of the number of gauges wet on each day": (
        "Wet extent provides a network-scale complement to the pairwise exceedance diagnostic. The mean per-model "
        "Wasserstein-1 distance decreased from 0.107 for QDM to 0.079 for MBCn (Table 4), and all seven models improved "
        "(Table 5). MBCn therefore reproduced the daily spatial coverage of wet conditions more closely, although "
        "this metric does not identify where within the network the wet gauges occurred."
    ),
    "Figure 9 shows the convergence analysis for three representative climate models": (
        "Iteration and seed experiments showed that the multivariate adjustment was numerically stable. Most "
        "improvement occurred within 10-15 passes, after which the diagnostics oscillated within a narrow range; "
        "the 25-pass default therefore lay on the convergence plateau. Across 20 random seeds, correlation-based "
        "scores changed little, with a coefficient of variation of approximately 4%. Energy distance varied more "
        "because it used fixed-size random subsamples, whereas the full-sample correlation diagnostics remained stable."
    ),
    "Projected wet-day quantile changes were similarly preserved by both methods (Figure 10)": (
        "Projected wet-day quantile changes were similarly preserved by both methods (Figure 6). Across all gauges, "
        "models, pathways, and future periods, mean absolute deviation from the raw change signal was 9.4 percentage "
        "points for QDM and 9.2 for MBCn; the corresponding R-squared values were 0.485 and 0.465. MBCn therefore "
        "improved dependence without a material additional loss of QDM's change-preservation property [4]. This check "
        "does not cover sequence-dependent statistics, such as wet- and dry-spell lengths, which MBCn may alter."
    ),
}


CAPTION_REPLACEMENTS = {
    "Figure 1.": (
        "Figure 1. Location of the 12 rain gauges in Prachuap Khiri Khan Province, Thailand. Red circles mark gauge "
        "locations and labels identify the stations. Background shading represents elevation classes, and the inset "
        "shows the province's location within Thailand."
    ),
    "Figure 6.": (
        "Figure 4. Inter-station Spearman correlation versus gauge separation. MBCn follows the observed distance "
        "relationship more closely than QDM, whereas raw output remains uniformly high. Bootstrap bands are "
        "exploratory because gauge pairs are not independent."
    ),
    "Figure 7.": (
        "Figure 5. Pairwise probability that both gauges exceed their 95th wet-day percentile on the same day, "
        "shown as the seven-model mean against observations. MBCn lies closest to the 1:1 line and is the only "
        "method without a systematic tilt."
    ),
    "Figure 10.": (
        "Figure 6. Raw versus corrected projected change in wet-day quantiles for every gauge, model, pathway, "
        "and future period. QDM and MBCn preserve the model-projected quantile-change signal to a similar degree."
    ),
}


REMOVE_CAPTIONS = (
    "Figure 4. Inter-station Spearman correlation for the median-performing model",
    "Figure 5. Paired QDM",
    "Figure 8. Pooled distribution",
    "Figure 9. Dependence diagnostics versus MBCn iteration count",
)


ALT_TEXT = [
    "Map of the 12 rain gauges in Prachuap Khiri Khan Province with station labels, elevation shading, and a Thailand location inset.",
    "Workflow for MBCn showing initial QDM, iterative rotation-based dependence adjustment, and final rank substitution.",
    "Boxplots of six single-site validation diagnostics for raw, QDM, and MBCn output.",
    "Binned inter-station Spearman correlation versus station separation for observations, raw output, QDM, and MBCn.",
    "Observed versus simulated pairwise probabilities of concurrent 95th-percentile wet-day exceedance.",
    "Raw versus corrected projected changes in five wet-day quantiles for QDM and MBCn.",
]


def find_paragraph(doc, prefix):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise KeyError(prefix)


def set_run_font(run, size=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)


def replace_body(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, 12)


def replace_caption(paragraph, text):
    paragraph.clear()
    label, rest = text.split(" ", 2)[0:2], None
    # Split after the numbered label so only "Figure N." is bold.
    parts = text.split(". ", 1)
    lead = paragraph.add_run(parts[0] + ".")
    lead.bold = True
    set_run_font(lead, 10)
    if len(parts) == 2:
        tail = paragraph.add_run(" " + parts[1])
        set_run_font(tail, 10)


def delete_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def remove_figure(doc, caption_prefix):
    caption = find_paragraph(doc, caption_prefix)
    node = caption._p.getprevious()
    while node is not None and node.tag != qn("w:p"):
        node = node.getprevious()
    if node is None or not node.xpath(".//w:drawing | .//w:pict"):
        raise RuntimeError(f"No image paragraph immediately before {caption_prefix}")
    node.getparent().remove(node)
    delete_paragraph(caption)


def clean_unused_image_relationships(doc):
    used = {
        blip.get(qn("r:embed"))
        for blip in doc._element.xpath(".//a:blip")
        if blip.get(qn("r:embed"))
    }
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.IMAGE and rel_id not in used:
            del doc.part.rels[rel_id]


def update_alt_text(doc):
    doc_prs = doc._element.xpath("//wp:docPr")
    if len(doc_prs) != len(ALT_TEXT):
        raise RuntimeError(f"Expected {len(ALT_TEXT)} remaining figures, found {len(doc_prs)}")
    for index, (doc_pr, description) in enumerate(zip(doc_prs, ALT_TEXT), start=1):
        doc_pr.set("title", f"Figure {index}")
        doc_pr.set("descr", description)


def remove_obsolete_page_break_before_table4(doc):
    preceding = doc.tables[3]._tbl.getprevious()
    if preceding is None or preceding.tag != qn("w:p"):
        return
    for line_break in preceding.xpath(".//w:br[@w:type='page']"):
        line_break.getparent().remove(line_break)


def add_page_break_before_table5(doc):
    preceding = doc.tables[4]._tbl.getprevious()
    if preceding is None or preceding.tag != qn("w:p"):
        raise RuntimeError("Expected a spacer paragraph before Table 5")
    run = preceding.makeelement(qn("w:r"))
    line_break = preceding.makeelement(qn("w:br"))
    line_break.set(qn("w:type"), "page")
    run.append(line_break)
    preceding.append(run)


def main():
    doc = Document(SOURCE)

    for caption_prefix in REMOVE_CAPTIONS:
        remove_figure(doc, caption_prefix)

    for prefix, replacement in BODY_REPLACEMENTS.items():
        replace_body(find_paragraph(doc, prefix), replacement)

    for old_prefix, replacement in CAPTION_REPLACEMENTS.items():
        replace_caption(find_paragraph(doc, old_prefix), replacement)

    remove_obsolete_page_break_before_table4(doc)
    add_page_break_before_table5(doc)
    clean_unused_image_relationships(doc)
    update_alt_text(doc)

    doc.core_properties.title = (
        "Multivariate bias correction improves inter-station dependence relative to QDM in multi-site daily "
        "precipitation across a coastal gauge network in Thailand"
    )
    doc.core_properties.author = "Surasit Punyawansiri"
    doc.core_properties.last_modified_by = "Surasit Punyawansiri"
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
