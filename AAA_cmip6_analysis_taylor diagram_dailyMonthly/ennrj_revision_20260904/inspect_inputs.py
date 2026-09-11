import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from PIL import Image

SRC = Path(r'D:\วารสาร EnNRJ\sent')
OUT = Path(__file__).resolve().parent / 'inspection'
OUT.mkdir(parents=True, exist_ok=True)
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'a':'http://schemas.openxmlformats.org/drawingml/2006/main', 'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
for src in SRC.glob('*.docx'):
    label = 'manuscript' if src.name.startswith('Contrasting') else 'template' if src.name.startswith('Template') else src.stem
    doc = Document(src)
    rows = []
    for i, p in enumerate(doc.paragraphs):
        rows.append({'p':i,'text':p.text,'style':p.style.name,'xml':p._p.xml,'runs':[{'text':r.text,'font':r.font.name,'size':r.font.size.pt if r.font.size else None,'bold':r.bold,'italic':r.italic,'highlight':str(r.font.highlight_color)} for r in p.runs]})
    data = {'source':str(src),'sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'paragraphs':rows,'tables':[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables], 'sections':[{'page_width':s.page_width.inches,'page_height':s.page_height.inches,'left':s.left_margin.inches,'right':s.right_margin.inches,'top':s.top_margin.inches,'bottom':s.bottom_margin.inches,'xml':s._sectPr.xml} for s in doc.sections]}
    with ZipFile(src) as z:
        data['parts'] = [{'path':n,'bytes':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()]
        data['all_text'] = {n:etree.fromstring(z.read(n)).xpath('//w:t/text()',namespaces=NS) for n in z.namelist() if n.startswith('word/') and n.endswith('.xml')}
        imgs=[]
        for n in z.namelist():
            if n.startswith('word/media/'):
                blob = z.read(n)
                try:
                    im=Image.open(io.BytesIO(blob))
                    info={'part':n,'width':im.width,'height':im.height,'format':im.format,'dpi':im.info.get('dpi')}
                except Exception:
                    info={'part':n,'bytes':len(blob)}
                imgs.append(info)
                if label in ('manuscript','template'):
                    media=OUT/label/'media'; media.mkdir(parents=True,exist_ok=True)
                    (media/Path(n).name).write_bytes(blob)
        data['images']=imgs
    (OUT/f'{label}.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/f'{label}.txt').write_text('\n'.join(f'[{r["p"]}] {r["text"]}' for r in rows)+'\n\nTABLES\n'+json.dumps(data['tables'],ensure_ascii=False,indent=2),encoding='utf-8')
    print(label, len(rows), 'paragraphs',len(doc.tables),'tables',len(imgs),'images')
    print(json.dumps(data['sections'],ensure_ascii=False)[:450])
