from __future__ import annotations

import argparse
import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def inspect_xlsx(path: Path) -> None:
    print(f"WORKBOOK: {path}")
    wb_formula = load_workbook(path, read_only=True, data_only=False)
    wb_value = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws_formula in wb_formula.worksheets:
            ws_value = wb_value[ws_formula.title]
            formula_count = 0
            error_count = 0
            nonempty = 0
            preview = []
            for row_f, row_v in zip(ws_formula.iter_rows(), ws_value.iter_rows()):
                rendered = []
                for cf, cv in zip(row_f, row_v):
                    if cf.value is not None:
                        nonempty += 1
                    if isinstance(cf.value, str) and cf.value.startswith("="):
                        formula_count += 1
                    if isinstance(cv.value, str) and cv.value.startswith("#"):
                        error_count += 1
                    rendered.append(cv.value if cv.value is not None else cf.value)
                if any(v is not None for v in rendered) and len(preview) < 8:
                    preview.append(rendered[:12])
            print(
                f"  SHEET {ws_formula.title!r}: rows={ws_formula.max_row}, "
                f"cols={ws_formula.max_column}, nonempty={nonempty}, "
                f"formulas={formula_count}, formula_errors={error_count}"
            )
            for row in preview:
                print(f"    {row}")
    finally:
        wb_formula.close()
        wb_value.close()


def inspect_docx(path: Path) -> None:
    print(f"DOCX: {path}")
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    nonempty = [p for p in paragraphs if p.strip()]
    text = "\n".join(paragraphs)
    heading_counts: dict[str, int] = {}
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            heading_counts[style] = heading_counts.get(style, 0) + 1
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        media = sorted(n for n in names if n.startswith("word/media/"))
        comments = "word/comments.xml" in names
        main_xml = zf.read("word/document.xml")
        tracked_insertions = len(re.findall(rb"<w:ins(?:\s|>)", main_xml))
        tracked_deletions = len(re.findall(rb"<w:del(?:\s|>)", main_xml))
        page_breaks = main_xml.count(b'w:type="page"')
    print(
        f"  paragraphs={len(paragraphs)}, nonempty={len(nonempty)}, "
        f"tables={len(doc.tables)}, inline_shapes={len(doc.inline_shapes)}, "
        f"media={len(media)}, page_breaks={page_breaks}"
    )
    print(
        f"  headings={heading_counts}, comments={comments}, "
        f"tracked_insertions={tracked_insertions}, "
        f"tracked_deletions={tracked_deletions}"
    )
    print(f"  text_sha256={hashlib.sha256(text.encode('utf-8')).hexdigest()}")
    print("  first_nonempty:")
    for p in nonempty[:12]:
        print(f"    {p[:240]}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", action="append", default=[])
    parser.add_argument("--docx", action="append", default=[])
    args = parser.parse_args()
    for item in args.xlsx:
        inspect_xlsx(Path(item))
    for item in args.docx:
        inspect_docx(Path(item))


if __name__ == "__main__":
    main()
