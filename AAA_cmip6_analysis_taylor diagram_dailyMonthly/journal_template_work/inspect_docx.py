from docx import Document
from pathlib import Path

for p in [
    Path("journal_template_work/AJSTR_template_reference.docx"),
    Path("cmip6bc_q1_editable/ID1558_Manuscript_BSJ.docx"),
]:
    doc = Document(p)
    print(f"\n### {p}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
    for i, para in enumerate(doc.paragraphs[:120], 1):
        txt = para.text.replace("\n", " ").strip()
        if txt:
            style = para.style.name if para.style is not None else "None"
            print(f"{i:03d} [{style}] {txt[:240]}")
    for j, table in enumerate(doc.tables[:5], 1):
        print(f"\nTABLE {j}: rows={len(table.rows)} cols={len(table.columns)}")
        for row in table.rows[:5]:
            print(" | ".join(c.text.replace("\n", " / ")[:100] for c in row.cells))
