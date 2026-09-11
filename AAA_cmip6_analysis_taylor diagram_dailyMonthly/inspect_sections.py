from pathlib import Path
from docx import Document

PATH = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")

doc = Document(PATH)
for start, end in [(50, 76), (95, 108), (120, 134)]:
    print(f"\n--- paragraphs {start}-{end}")
    for j in range(start, min(end, len(doc.paragraphs))):
        print(f"{j:03d}: {doc.paragraphs[j].text}")
