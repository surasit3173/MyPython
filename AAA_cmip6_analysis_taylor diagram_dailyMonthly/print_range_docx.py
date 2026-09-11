import sys
from docx import Document

doc = Document(sys.argv[1])
start = int(sys.argv[2])
end = int(sys.argv[3])
for i in range(start, min(end, len(doc.paragraphs))):
    print(f"{i:03d}: {doc.paragraphs[i].text}")
