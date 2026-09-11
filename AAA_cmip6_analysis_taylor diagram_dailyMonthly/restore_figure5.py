from pathlib import Path
from zipfile import ZipFile
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


PATH = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")
SOURCE = Path(r"D:\วารสารบูรพา\ID1558_Manuscript_BSJ (11).docx")
CAPTION = (
    "Figure 5  Factors associated with the change in error metrics after QDM correction, independent validation period. "
    "(a) Dispersion ratio against the monthly correlation, with the dashed line marking the RMSE-optimal ratio; "
    "(b) change in dispersion ratio against change in RMSE for the annual-pooled configuration; "
    "(c) decomposition of the change in mean square error into bias, dispersion and lack-of-correlation components following Equation (5); "
    "and (d) split of the change in mean square error into an annual-cycle component and an anomaly component following Equation (6), "
    "shown for both temporal-grouping configurations. Bars in (c) and (d) are means across the 60 evaluation combinations and are additive, "
    "so the components sum exactly to the total; whiskers in (c) show the interquartile range across combinations, not the uncertainty of the mean"
)


def find_heading(doc, prefix):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i, p
    raise ValueError(prefix)


def main():
    doc = Document(PATH)
    if any(p.text.strip().startswith("Figure 5  ") for p in doc.paragraphs):
        print("Figure 5 already present")
        return

    with tempfile.TemporaryDirectory() as tmp:
        image = Path(tmp) / "figure5.png"
        with ZipFile(SOURCE) as zf:
            image.write_bytes(zf.read("word/media/image4.png"))

        _, anchor = find_heading(doc, "3.6 ")
        image_para = anchor.insert_paragraph_before("")
        image_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_para.add_run().add_picture(str(image), width=Inches(6.3))

        cap = anchor.insert_paragraph_before(CAPTION)
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.save(PATH)
        print(PATH)


if __name__ == "__main__":
    main()
