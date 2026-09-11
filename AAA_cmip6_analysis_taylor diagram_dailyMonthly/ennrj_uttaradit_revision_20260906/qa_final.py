from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader


ROOT = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ennrj_uttaradit_revision_20260906")
DOCX = ROOT / "EnNRJ_Uttaradit_CMIP6_Precipitation_Extremes_Revised.docx"
PDF = ROOT / "final_word_v3.pdf"
PAGES = ROOT / "final_pages_v3"
FIGURES = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ennrj_uttaradit_manuscript_20260906\figures_submission")
TITLE = "Projected Precipitation Extremes over Uttaradit Province under CMIP6 SSP Scenarios"


checks: list[dict[str, object]] = []


def check(name: str, ok: bool, evidence: object) -> None:
    checks.append({"check": name, "pass": bool(ok), "evidence": evidence})


check("DOCX exists and is non-empty", DOCX.is_file() and DOCX.stat().st_size > 0, DOCX.stat().st_size if DOCX.exists() else None)
check("Word-rendered PDF exists and is non-empty", PDF.is_file() and PDF.stat().st_size > 0, PDF.stat().st_size if PDF.exists() else None)

doc = Document(DOCX)
paragraphs = [p.text.strip() for p in doc.paragraphs]
all_text = "\n".join(paragraphs)

check("Exact approved title", paragraphs[0] == TITLE, paragraphs[0])
check("Single author and affiliation retained", "Surasit Punyawansiri1,*" in all_text and "Office of Water Management and Hydrology" in all_text, paragraphs[1:4])
check("Corresponding email retained", "Surasit.pu@ku.th" in all_text, "Surasit.pu@ku.th" in all_text)

section = doc.sections[0]
page_cm = (section.page_width.cm, section.page_height.cm)
margins_cm = (section.left_margin.cm, section.right_margin.cm, section.top_margin.cm, section.bottom_margin.cm)
check("A4 page size", abs(page_cm[0] - 21.0) < 0.08 and abs(page_cm[1] - 29.7) < 0.08, tuple(round(v, 2) for v in page_cm))
check(
    "EnNRJ margins",
    abs(margins_cm[0] - 2.54) < 0.08
    and abs(margins_cm[1] - 2.54) < 0.08
    and abs(margins_cm[2] - 1.9) < 0.08
    and abs(margins_cm[3] - 1.9) < 0.08,
    tuple(round(v, 2) for v in margins_cm),
)

abstract_heading = paragraphs.index("Abstract")
abstract = paragraphs[abstract_heading + 1]
abstract_words = re.findall(r"\b[\w%–-]+\b", abstract)
check("Abstract does not exceed 250 words", len(abstract_words) <= 250, len(abstract_words))

keyword_line = next(p for p in paragraphs if p.startswith("Keywords:"))
keywords = [x.strip() for x in keyword_line.split(":", 1)[1].split(";") if x.strip()]
check("Exactly six keywords", len(keywords) == 6, keywords)

highlights_heading = paragraphs.index("Highlights")
highlights = []
for p in paragraphs[highlights_heading + 1 :]:
    if not p:
        break
    highlights.append(re.sub(r"^[•\-]\s*", "", p))
highlight_lengths = [len(x) for x in highlights]
check("Five highlights", len(highlights) == 5, highlights)
check("Each highlight is at most 85 characters", all(n <= 85 for n in highlight_lengths), highlight_lengths)

refs_heading = next(i for i, p in enumerate(paragraphs) if re.fullmatch(r"(?:\d+\.\s*)?References", p))
reference_entries = [p for p in paragraphs[refs_heading + 1 :] if p]
check("Exactly 28 reference entries", len(reference_entries) == 28, len(reference_entries))

main_text = " ".join(p for p in paragraphs[15:refs_heading] if p)
main_words = re.findall(r"\b[\w%–-]+\b", main_text)
check("Main text does not exceed 4,000 words", len(main_words) <= 4000, len(main_words))

expected_table_shapes = [(14, 4), (8, 5), (12, 4), (12, 8), (23, 6)]
actual_table_shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
check("Five tables with expected dimensions", actual_table_shapes == expected_table_shapes, actual_table_shapes)
check("Three embedded figures", len(doc.inline_shapes) == 3, len(doc.inline_shapes))

omml_count = len(doc.element.body.xpath(".//m:oMath"))
check("Ten numbered equations are native OMML math", omml_count == 10, omml_count)

for kind, total in (("Table", 5), ("Figure", 3)):
    for number in range(1, total + 1):
        token = f"{kind} {number}"
        hits = len(re.findall(rf"\b{re.escape(token)}\b", all_text))
        check(f"{token} is captioned and cited", hits >= 2, hits)

required_phrases = {
    "AR6-consistent baseline rationale": "ends with the final year of the CMIP6 historical experiment",
    "Scenario interpretation": "conditional pathways, not forecasts of a single future",
    "Separate limitations section": "3.6 Limitations and future research",
    "Single-author conflict statement": "The author declares no conflict of interest.",
    "IDW interpretation caution": "not dynamically downscaled or independently validated continuous-resolution climate projections",
}
for label, phrase in required_phrases.items():
    check(label, phrase in all_text, phrase)

for forbidden in ("Phetchaburi", "coordinate sensitivity", "primary-coordinate", "maximum 3 consecutive", "TODO", "PLACEHOLDER"):
    check(f"Forbidden text absent: {forbidden}", forbidden.lower() not in all_text.lower(), len(re.findall(re.escape(forbidden), all_text, flags=re.I)))

pdf = PdfReader(str(PDF))
check("Final Word rendering is exactly 14 pages", len(pdf.pages) == 14, len(pdf.pages))
pdf_sizes = []
for page in pdf.pages:
    width = float(page.mediabox.width)
    height = float(page.mediabox.height)
    pdf_sizes.append((round(width, 1), round(height, 1)))
check("All rendered PDF pages are A4", all(abs(w - 595.3) < 2 and abs(h - 841.9) < 2 for w, h in pdf_sizes), sorted(set(pdf_sizes)))

page_pngs = sorted(PAGES.glob("page-*.png"))
check("All 14 pages have QA raster images", len(page_pngs) == 14 and all(p.stat().st_size > 0 for p in page_pngs), len(page_pngs))

figure_info = {}
if FIGURES.is_dir():
    for p in sorted(FIGURES.glob("*.tif*")):
        with Image.open(p) as im:
            dpi = im.info.get("dpi", (0, 0))
            figure_info[p.name] = {"pixels": im.size, "dpi": tuple(round(float(x)) for x in dpi)}
check(
    "Submission figures are high-resolution TIFFs",
    len(figure_info) == 3 and all(info["dpi"][0] >= 600 and info["dpi"][1] >= 600 for info in figure_info.values()),
    figure_info,
)

with zipfile.ZipFile(DOCX) as zf:
    media = [n for n in zf.namelist() if n.startswith("word/media/") and not n.endswith("/")]
check("DOCX package contains exactly three figure media files", len(media) == 3, media)

passed = sum(1 for item in checks if item["pass"])
failed = [item for item in checks if not item["pass"]]
report = {"passed": passed, "total": len(checks), "failed": failed, "checks": checks}
print(json.dumps(report, ensure_ascii=False, indent=2))
sys.exit(1 if failed else 0)
