from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "source_manuscript.docx"
OUTPUT = ROOT / "EnNRJ_Uttaradit_CMIP6_Precipitation_Extremes_Revised.docx"

TITLE = "Projected Precipitation Extremes over Uttaradit Province under CMIP6 SSP Scenarios"


def set_run_font(run, size: float = 12.0, bold=None, italic=None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def replace_paragraph(paragraph, text: str, *, bold=False, italic=False, align=None) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, 12.0, bold=bold, italic=italic)
    if align is not None:
        paragraph.alignment = align


def find_unique(document: Document, prefix: str):
    hits = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(hits) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}; found {len(hits)}")
    return hits[0]


doc = Document(SOURCE)

# Retain the user's explicitly selected, projection-focused title.
replace_paragraph(doc.paragraphs[0], TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# Explain the scenarios in the Introduction while preserving the cautious ensemble language.
p = find_unique(doc, "The Coupled Model Intercomparison Project Phase 6")
replace_paragraph(
    p,
    "The Coupled Model Intercomparison Project Phase 6 (CMIP6) and Scenario Model "
    "Intercomparison Project provide a coordinated framework for examining future climate "
    "under alternative forcing pathways (Eyring et al., 2016; O’Neill et al., 2016). "
    "SSP2-4.5 represents an intermediate pathway approaching 4.5 W/m² of radiative forcing "
    "by 2100, whereas SSP5-8.5 represents a very high-forcing, fossil-fueled pathway "
    "approaching 8.5 W/m² (Riahi et al., 2017). These scenarios are conditional pathways, "
    "not forecasts of a single future. Multi-model ensembles reveal structural differences "
    "among models, but their members are neither independent nor equally credible at every "
    "location. Ensemble summaries are therefore interpreted here as descriptive evidence "
    "conditioned on the selected models, scenarios, and processing chain rather than as "
    "calibrated probabilities (Knutti et al., 2017; Sillmann et al., 2013a; Sillmann et al., "
    "2013b). This distinction is especially important for precipitation tails, for which model "
    "spread and small baseline values can produce large relative changes."
)

# Expand the period rationale requested by the user and make the CMIP6 experiment boundary explicit.
p = find_unique(doc, "The observed input comprised daily precipitation")
replace_paragraph(
    p,
    "The observed input comprised daily precipitation from 1 January 1981 to 31 December "
    "2014. Seven supplied bias-corrected CMIP6 daily precipitation series were analyzed "
    "(Table 2). For each model, 1995–2014 was used as the model-consistent historical "
    "baseline, and 2021–2050 was used as the future period. The 1995–2014 interval was "
    "selected as a recent-past reference because it is widely used in the Intergovernmental "
    "Panel on Climate Change Sixth Assessment Report, ends with the final year of the CMIP6 "
    "historical experiment, and facilitates comparison with assessments framed against the "
    "AR6 reference period (Intergovernmental Panel on Climate Change, 2021). CMIP6 scenario "
    "simulations begin in 2015; 2015–2020 was excluded to retain a discrete 30-year future "
    "window. The file inventory included one ensemble member and a native or regridded grid "
    "label for each model. Because the supplied daily files did not contain complete metadata "
    "for the preceding bias-correction implementation, that procedure was neither reconstructed "
    "nor independently validated; all conclusions are conditional on the supplied corrected series."
)

p = find_unique(doc, "Note: Historical data for 1995–2014")
replace_paragraph(
    p,
    "Note: Historical files span 1981–2014; 1995–2014 was the model baseline. Future "
    "files cover 2021–2050 under SSP2-4.5 and SSP5-8.5. Bias-correction metadata were incomplete."
)

# Follow the example paper's rhetorical separation of interpretation and limitations.
p = find_unique(doc, "3.5 Implications and limitations")
replace_paragraph(p, "3.5 Discussion", bold=True, italic=True)

limitations = find_unique(doc, "Four limitations constrain interpretation")
heading = limitations.insert_paragraph_before("3.6 Limitations and future research")
heading.style = limitations.style
heading.paragraph_format.keep_with_next = True
heading.paragraph_format.space_before = Pt(6)
heading.paragraph_format.space_after = Pt(0)
for run in heading.runs:
    set_run_font(run, 12.0, bold=True, italic=True)

replace_paragraph(
    limitations,
    "Four limitations constrain interpretation. First, the seven models form a selected, "
    "non-independent ensemble, so the median and sign counts do not define probabilities. "
    "Second, the supplied bias-corrected series could not be independently audited against "
    "complete correction metadata; precipitation extremes may be sensitive to wet-day "
    "adjustment and temporal sequencing (Cannon et al., 2015; Maraun, 2016; Maraun et al., "
    "2017). Third, relative changes in rare-event indices are unstable when baseline values "
    "are small or zero, as illustrated by the undefined R99p changes at one station and the "
    "broad R50mm and R99p ranges. Fourth, IDW represents geographic proximity only and "
    "cannot resolve orographic gradients, convective organization, or model physics absent "
    "from the station values. Future work should compare fully documented bias-correction "
    "methods, evaluate additional observational references, and test terrain-aware or "
    "dynamical downscaling before the fields are used for site-specific design."
)

p = find_unique(doc, "The authors declare no conflict of interest")
replace_paragraph(p, "The author declares no conflict of interest.")

# Keep the take-home message projection-focused and concise enough for the
# journal-length constraint without dropping uncertainty qualifications.
p = find_unique(doc, "Fixed-threshold count indices")
replace_paragraph(
    p,
    "Fixed-threshold indices displayed broad inter-model ranges, underscoring the "
    "sensitivity of rare-event percentages to small baselines. The absence of significant "
    "historical PRCPTOT trends does not contradict the projected changes, and within-window "
    "trend diagnostics should remain separate from baseline-to-future climatological "
    "differences. For planning, the results support adaptation across a range of plausible "
    "futures because flood and erosion hazards may intensify even where annual rainfall "
    "changes little and station support is sparse."
)

# EnNRJ specifies bold main headings and bold-italic second-level headings.
for p in doc.paragraphs:
    text = p.text.strip()
    if re.match(r"^[23]\.\d+\s", text):
        for run in p.runs:
            set_run_font(run, 12.0, bold=True, italic=True)
        p.paragraph_format.keep_with_next = True
    elif re.match(r"^[1-8]\.\s", text):
        for run in p.runs:
            set_run_font(run, 12.0, bold=True, italic=False)
        p.paragraph_format.keep_with_next = True

# Keep the body at the required 12 pt/1.5 spacing. References remain 12 pt but are
# single-spaced, a compact bibliography treatment used to meet the user's 14-page limit.
ref_heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "8. References")
for i, p in enumerate(doc.paragraphs):
    if i <= ref_heading_index:
        continue
    for run in p.runs:
        set_run_font(run, 12.0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = False
    p.paragraph_format.widow_control = False

# Ensure single-author contribution text uses consistent editorial punctuation.
p = find_unique(doc, "S.P.: Conceptualization")
replace_paragraph(
    p,
    "S.P.: Conceptualization, methodology, software, validation, formal analysis, "
    "investigation, data curation, visualization, writing—original draft, and "
    "writing—review and editing."
)

doc.save(OUTPUT)
print(OUTPUT)
