from pathlib import Path
from docx import Document


PATH = Path(r"C:\MyPython\AAA_cmip6_analysis_taylor diagram_dailyMonthly\ID1558_revised_for_publication_final.docx")


def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


doc = Document(PATH)
for p in doc.paragraphs:
    text = p.text
    if text.startswith("วิธีดำเนินการวิจัย :"):
        text = text.replace(
            "พร้อมประเมินทั้งรายแบบจำลองและกลุ่มแบบจำลองเฉลี่ยแบบไม่ถ่วงน้ำหนัก",
            "พร้อมประเมินทั้งรายแบบจำลอง กลุ่มแบบจำลองเฉลี่ยแบบไม่ถ่วงน้ำหนัก และแยกองค์ประกอบของค่าความคลาดเคลื่อนกำลังสองเฉลี่ยเพื่ออธิบายที่มาของการเปลี่ยนแปลง RMSE",
        )
        set_text(p, text)
    elif text.startswith("Methodology :"):
        text = text.replace(
            "Individual model performance and an unweighted multi-model ensemble were assessed separately.",
            "Individual model performance and an unweighted multi-model ensemble were assessed separately, and mean square error was decomposed to diagnose the components associated with RMSE changes.",
        )
        set_text(p, text)

doc.save(PATH)
print(PATH)
