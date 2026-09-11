#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Uttaradit Extreme Precipitation Trend Analysis
================================================

Publication-oriented analysis for:

    Spatiotemporal Trends of ETCCDI-Based Extreme Precipitation Indices
    in Uttaradit Province, Thailand

Core scope ONLY
----------------
1. 11 extreme-precipitation indices:
   PRCPTOT, SDII, Rx1day, Rx5day, CDD, CWD,
   R10mm, R20mm, R50mm, R95p, R99p

2. Two periods:
   - Observed: 1981-2014
   - Future:   2021-2050

3. Trend methods:
   - Standard Mann-Kendall (MK)
   - Yue & Wang (2004) effective-sample-size modified MK (MMK-2004)
   - Sen's slope + 95% CI

4. Future spatial change:
   - model-consistent baseline: 1995-2014
   - SSP2-4.5 and SSP5-8.5
   - 7-model median at each station
   - IDW interpolation for visualization

Explicitly excluded
--------------------
Pettitt, FDR, Bonferroni, pre-whitening, TFPW, seasonal trends,
monthly trends, Taylor diagrams, and non-target indices.

IMPORTANT TERMINOLOGY
---------------------
The requested "MMK 2004" is not a 2004 Hamed & Rao method.
The appropriate citation is:

Yue, S., & Wang, C. Y. (2004). The Mann-Kendall test modified by
effective sample size to detect trend in serially correlated
hydrological series. Water Resources Management, 18, 201-218.
doi:10.1023/B:WARM.0000043140.61082.60

The implementation below follows the Yue & Wang (2004) effective
sample size / variance-correction approach:
- estimate Sen slope;
- detrend the series using that slope;
- estimate serial correlation from the detrended series;
- use the full lag structure in the effective-sample-size correction;
- adjust Var(S) by n/n*;
- calculate the corrected Z and two-sided p-value.

The standard MK result is retained as a transparent reference.
The Sen slope itself is NOT changed by MMK; MMK changes inference
about significance through the variance correction.

Data assumptions
----------------
Observed CSV:
    YEAR, MONTH, DAY, <station columns>

CMIP6 CSV:
    YEAR, MONTH, DAY, <station columns>

The archive supplied with this project contains 63 station columns
for the CMIP6 files, but the analysis automatically intersects those
columns with the observed rainfall file and station_coordinates.xlsx.
The intended Uttaradit network is therefore data-driven rather than
hard-coded.

Expected archive layout after extraction
-----------------------------------------
dataUttaradit/
    Observed_Rain_daily_198101_201412_Uttaradit.csv
    station_coordinates.xlsx
    GCM raindata/
        ACCESS-ESM1-5/
            bc_pr_day_...historical....csv
            bc_pr_day_...ssp245....csv
            bc_pr_day_...ssp585....csv
        CanESM5/
            ...
        CESM2/
            ...
        EC-Earth3_CSV_FILE/
            ...
        FGOALS-g3/
            ...
        MIROC6/
            ...
        MRI-ESM2-0/
            ...

Outputs
-------
output_utt_trend/
    results/
        observed_trend_1981_2014.csv
        future_trend_2021_2050.csv
        future_change_2021_2050.csv
        observed_indices_1981_2014.csv
        future_indices_2021_2050.csv
        station_coordinates_used.csv
        qc_observed.csv
        qc_models.csv
        summary_by_index.csv
        summary_by_model_scenario.csv
        README_results.txt
        Uttaradit_Trend_Analysis.xlsx
    figures/
        observed_trends_1981_2014_<index>.png
        future_trends_2021_2050_<scenario>_<index>.png
        future_change_IDW_2021_2050_<scenario>_<index>.png
        observed_trend_heatmap.png
        future_trend_heatmap_<scenario>.png

Dependencies
------------
numpy
pandas
scipy
matplotlib
openpyxl
(optional) geopandas for province boundary plotting
(optional) rarfile + UnRAR/WinRAR/7-Zip if passing the .rar archive
"""

from __future__ import annotations

import argparse
import math
import os
import re
import shutil
import sys
import warnings
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import TwoSlopeNorm
from matplotlib.lines import Line2D


# ---------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------

VERSION = "1.0.0"
ALPHA = 0.05
MIN_YEARS = 10
MIN_ANNUAL_COMPLETENESS = 0.90

OBS_START = 1981
OBS_END = 2014

FUTURE_START = 2021
FUTURE_END = 2050

BASELINE_START = 1995
BASELINE_END = 2014

WET_THRESHOLD = 1.0
R10_THRESHOLD = 10.0
R20_THRESHOLD = 20.0
R50_THRESHOLD = 50.0

P95 = 95.0
P99 = 99.0

INDICES = [
    "PRCPTOT",
    "SDII",
    "Rx1day",
    "Rx5day",
    "CDD",
    "CWD",
    "R10mm",
    "R20mm",
    "R50mm",
    "R95p",
    "R99p",
]

INDEX_UNITS = {
    "PRCPTOT": "mm yr-1",
    "SDII": "mm wet-day-1",
    "Rx1day": "mm",
    "Rx5day": "mm",
    "CDD": "days",
    "CWD": "days",
    "R10mm": "days yr-1",
    "R20mm": "days yr-1",
    "R50mm": "days yr-1",
    "R95p": "mm yr-1",
    "R99p": "mm yr-1",
}

SCENARIOS = {
    "ssp245": "SSP2-4.5",
    "ssp585": "SSP5-8.5",
}

EXPECTED_MODELS = [
    "ACCESS-ESM1-5",
    "CanESM5",
    "CESM2",
    "EC-Earth3",
    "FGOALS-g3",
    "MIROC6",
    "MRI-ESM2-0",
]

MISS_FLAGS = [-99.9, -99.99, -999.0, -999.9, -9999.0]


# ---------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------

def log(msg: str) -> None:
    print(msg, flush=True)


def safe_float(x, default=np.nan) -> float:
    try:
        v = float(x)
        return v if np.isfinite(v) else default
    except (TypeError, ValueError):
        return default


def trend_label(z: float, p: float, alpha: float = ALPHA) -> str:
    if not np.isfinite(z) or not np.isfinite(p):
        return "Insufficient data"
    if p < alpha:
        if z > 0:
            return "Increasing"
        if z < 0:
            return "Decreasing"
    return "No significant trend"


def significance_code(p: float, alpha: float = ALPHA) -> str:
    if not np.isfinite(p):
        return "NA"
    return "*" if p < alpha else "ns"


def normalize_station_id(value) -> str:
    s = str(value).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


def canonical_model_name(name: str) -> str:
    s = name.strip()
    aliases = {
        "EC-Earth3_CSV_FILE": "EC-Earth3",
        "EC-Earth3_CSV": "EC-Earth3",
        "EC-Earth3": "EC-Earth3",
    }
    return aliases.get(s, s)


def parse_model_from_filename(path: Path) -> Optional[str]:
    stem = path.stem
    for model in EXPECTED_MODELS:
        if model.lower() in stem.lower():
            return model
    if "EC-Earth3" in stem:
        return "EC-Earth3"
    return None


def parse_scenario_from_filename(path: Path) -> Optional[str]:
    stem = path.stem.lower()
    if "ssp245" in stem or "ssp2-4.5" in stem:
        return "ssp245"
    if "ssp585" in stem or "ssp5-8.5" in stem:
        return "ssp585"
    return None


# ---------------------------------------------------------------------
# Archive handling
# ---------------------------------------------------------------------

def prepare_input_root(input_path: Path) -> Path:
    """
    Accept either an extracted data directory or dataUttaradit.rar.

    RAR extraction requires an installed external extractor supported by
    the rarfile package. If unavailable, the error is explicit.
    """
    input_path = input_path.expanduser().resolve()

    if input_path.is_dir():
        return input_path

    if input_path.suffix.lower() != ".rar":
        raise FileNotFoundError(
            f"Input must be a directory or .rar archive: {input_path}"
        )

    extracted = input_path.with_name(input_path.stem + "_extracted")
    marker = extracted / ".extraction_complete"

    if marker.exists():
        return extracted

    try:
        import rarfile
    except ImportError as exc:
        raise RuntimeError(
            "RAR input detected, but package 'rarfile' is not installed. "
            "Install it with: pip install rarfile"
        ) from exc

    extracted.mkdir(parents=True, exist_ok=True)

    try:
        with rarfile.RarFile(str(input_path)) as rf:
            rf.extractall(str(extracted))
    except Exception as exc:
        raise RuntimeError(
            "RAR extraction failed. The Python rarfile package also needs "
            "a working external extractor (UnRAR, WinRAR, 7-Zip, etc.). "
            f"Original error: {exc}"
        ) from exc

    marker.write_text("complete\n", encoding="utf-8")
    return extracted


def locate_data_root(root: Path) -> Path:
    candidates = [root, root / "dataUttaradit"]
    for p in candidates:
        if (p / "Observed_Rain_daily_198101_201412_Uttaradit.csv").exists():
            return p

    matches = list(root.rglob("Observed_Rain_daily_198101_201412_Uttaradit.csv"))
    if matches:
        return matches[0].parent

    raise FileNotFoundError(
        "Could not locate Observed_Rain_daily_198101_201412_Uttaradit.csv"
    )


# ---------------------------------------------------------------------
# Data loading and QC
# ---------------------------------------------------------------------

def load_daily_csv(path: Path) -> pd.DataFrame:
    """
    Load a daily rainfall CSV with YEAR/MONTH/DAY columns.

    No meteorological extreme is removed as an outlier. Negative values
    and configured missing-value sentinels become NaN.
    """
    df = pd.read_csv(path)
    df.columns = [str(c).strip() for c in df.columns]

    required = {"YEAR", "MONTH", "DAY"}
    if not required.issubset(df.columns):
        raise ValueError(
            f"{path.name}: required columns YEAR, MONTH, DAY not found."
        )

    for flag in MISS_FLAGS:
        df.replace(flag, np.nan, inplace=True)

    station_cols = [c for c in df.columns if c not in required]

    for c in station_cols:
        df[c] = pd.to_numeric(df[c], errors="coerce")
        df.loc[df[c] < 0, c] = np.nan

    dt = pd.to_datetime(
        {
            "year": pd.to_numeric(df["YEAR"], errors="coerce"),
            "month": pd.to_numeric(df["MONTH"], errors="coerce"),
            "day": pd.to_numeric(df["DAY"], errors="coerce"),
        },
        errors="coerce",
    )

    good = dt.notna()
    out = df.loc[good, station_cols].copy()
    out.index = pd.DatetimeIndex(dt.loc[good])
    out.index.name = "date"

    # Duplicate dates are retained only once, explicitly documented.
    out = out[~out.index.duplicated(keep="first")]
    out = out.sort_index()

    return out


def qc_daily(
    df: pd.DataFrame,
    expected_start: Optional[str] = None,
    expected_end: Optional[str] = None,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    QC without deleting valid extreme rainfall.

    Short-gap interpolation is NOT used by default. For extreme indices,
    artificial rainfall can alter counts, spell lengths, percentile
    exceedances, and maxima. Missing values therefore remain NaN.

    Completeness is reported at station-year level.
    """
    out = df.copy()

    if expected_start:
        out = out.loc[out.index >= pd.Timestamp(expected_start)]
    if expected_end:
        out = out.loc[out.index <= pd.Timestamp(expected_end)]

    rows = []
    for station in out.columns:
        s = out[station]
        n = len(s)
        missing = int(s.isna().sum())
        valid = int(s.notna().sum())

        years = []
        for year, g in s.groupby(s.index.year):
            n_days = len(g)
            n_valid = int(g.notna().sum())
            completeness = n_valid / max(n_days, 1)
            years.append(completeness)

        rows.append(
            {
                "Station": normalize_station_id(station),
                "N_records": n,
                "N_valid": valid,
                "N_missing": missing,
                "Missing_pct": 100.0 * missing / max(n, 1),
                "Min_valid": float(s.min()) if valid else np.nan,
                "Max_valid": float(s.max()) if valid else np.nan,
                "Median_valid": float(s.median()) if valid else np.nan,
                "Mean_valid": float(s.mean()) if valid else np.nan,
                "Years_with_data": len(years),
                "Median_annual_completeness": (
                    100.0 * float(np.median(years)) if years else np.nan
                ),
            }
        )

    return out, pd.DataFrame(rows)


def read_coordinates(path: Path) -> pd.DataFrame:
    """
    Read station_coordinates.xlsx and normalize expected fields:
    station, latitude, longitude, optional elevation.
    """
    xl = pd.ExcelFile(path)
    chosen = None
    for sheet in xl.sheet_names:
        tmp = pd.read_excel(path, sheet_name=sheet)
        cols = {str(c).strip().lower(): c for c in tmp.columns}
        if {"station", "latitude", "longitude"}.issubset(cols):
            chosen = tmp.rename(
                columns={
                    cols["station"]: "station",
                    cols["latitude"]: "latitude",
                    cols["longitude"]: "longitude",
                }
            )
            if "elevation" in cols:
                chosen = chosen.rename(
                    columns={cols["elevation"]: "elevation"}
                )
            break

    if chosen is None:
        raise ValueError(
            "station_coordinates.xlsx must contain station, latitude, longitude."
        )

    chosen["station"] = chosen["station"].map(normalize_station_id)
    chosen["latitude"] = pd.to_numeric(chosen["latitude"], errors="coerce")
    chosen["longitude"] = pd.to_numeric(chosen["longitude"], errors="coerce")
    if "elevation" in chosen.columns:
        chosen["elevation"] = pd.to_numeric(
            chosen["elevation"], errors="coerce"
        )

    chosen = chosen.dropna(subset=["station", "latitude", "longitude"])
    chosen = chosen.drop_duplicates("station", keep="first")
    return chosen.reset_index(drop=True)


# ---------------------------------------------------------------------
# ETCCDI-based indices
# ---------------------------------------------------------------------

def max_consecutive(mask: np.ndarray) -> float:
    if len(mask) == 0:
        return np.nan

    best = 0
    run = 0
    for value in mask:
        if bool(value):
            run += 1
            best = max(best, run)
        else:
            run = 0

    return float(best)


def rolling_rx5(values: np.ndarray) -> float:
    if len(values) < 5:
        return np.nan

    # NaN in any 5-day window invalidates that window.
    valid = np.isfinite(values)
    if not valid.all():
        arr = pd.Series(values)
        roll = arr.rolling(5, min_periods=5).sum().to_numpy()
        return float(np.nanmax(roll)) if np.isfinite(roll).any() else np.nan

    return float(np.max(np.convolve(values, np.ones(5), mode="valid")))


def station_thresholds_from_observed(
    observed_daily: pd.DataFrame,
) -> Dict[str, Tuple[float, float]]:
    """
    Station-specific P95/P99 thresholds from all observed wet days
    (1981-2014), then held fixed for both observed and future calculations.
    """
    thresholds = {}

    for station in observed_daily.columns:
        s = observed_daily[station].dropna()
        wet = s[s >= WET_THRESHOLD]

        if len(wet) < 10:
            thresholds[normalize_station_id(station)] = (np.nan, np.nan)
        else:
            thresholds[normalize_station_id(station)] = (
                float(np.percentile(wet, P95)),
                float(np.percentile(wet, P99)),
            )

    return thresholds


def annual_extreme_indices(
    daily: pd.DataFrame,
    thresholds: Dict[str, Tuple[float, float]],
) -> pd.DataFrame:
    """
    Calculate exactly the 11 target indices for each station-year.

    ETCCDI-style definitions:
      PRCPTOT: total precipitation on wet days >= 1 mm
      SDII: PRCPTOT / number of wet days
      Rx1day: annual maximum daily precipitation
      Rx5day: annual maximum consecutive 5-day total
      CDD: maximum consecutive days < 1 mm
      CWD: maximum consecutive days >= 1 mm
      R10mm: count >= 10 mm
      R20mm: count >= 20 mm
      R50mm: count >= 50 mm (study-specific)
      R95p: annual total on days > observed P95
      R99p: annual total on days > observed P99

    A calendar year is the unit of aggregation.
    """
    rows = []

    for station in daily.columns:
        sid = normalize_station_id(station)
        s = daily[station]

        p95, p99 = thresholds.get(sid, (np.nan, np.nan))

        for year, grp in s.groupby(s.index.year):
            x = grp.to_numpy(dtype=float)

            # Require at least 90% valid daily values before deriving an
            # annual extreme index. This prevents incomplete years from
            # creating artificial low extremes or shortened spells.
            is_leap = bool(pd.Timestamp(year=year, month=12, day=31).is_leap_year)
            expected_days = 366 if is_leap else 365
            valid_count = int(np.isfinite(x).sum())
            completeness = valid_count / expected_days

            if completeness < MIN_ANNUAL_COMPLETENESS:
                row = {idx: np.nan for idx in INDICES}
                row.update({
                    "Station": sid,
                    "Year": int(year),
                })
                rows.append(row)
                continue

            wet = x[np.isfinite(x) & (x >= WET_THRESHOLD)]

            prcptot = float(wet.sum()) if len(wet) else 0.0
            sdii = float(wet.mean()) if len(wet) else np.nan

            rx1 = float(np.nanmax(x)) if finite.any() else np.nan
            rx5 = rolling_rx5(x)

            cdd = max_consecutive(
                np.isfinite(x) & (x < WET_THRESHOLD)
            )
            cwd = max_consecutive(
                np.isfinite(x) & (x >= WET_THRESHOLD)
            )

            r10 = float(np.sum(np.isfinite(x) & (x >= R10_THRESHOLD)))
            r20 = float(np.sum(np.isfinite(x) & (x >= R20_THRESHOLD)))
            r50 = float(np.sum(np.isfinite(x) & (x >= R50_THRESHOLD)))

            if np.isfinite(p95):
                r95 = float(
                    np.nansum(x[(np.isfinite(x)) & (x > p95)])
                )
            else:
                r95 = np.nan

            if np.isfinite(p99):
                r99 = float(
                    np.nansum(x[(np.isfinite(x)) & (x > p99)])
                )
            else:
                r99 = np.nan

            row = {
                "Station": sid,
                "Year": int(year),
                "PRCPTOT": prcptot,
                "SDII": sdii,
                "Rx1day": rx1,
                "Rx5day": rx5,
                "CDD": cdd,
                "CWD": cwd,
                "R10mm": r10,
                "R20mm": r20,
                "R50mm": r50,
                "R95p": r95,
                "R99p": r99,
            }
            rows.append(row)

    out = pd.DataFrame(rows)
    if not out.empty:
        out = out.set_index(["Station", "Year"]).sort_index()
    return out


def indices_to_wide(
    index_df: pd.DataFrame,
    index_name: str,
) -> pd.DataFrame:
    if index_df.empty:
        return pd.DataFrame()

    s = index_df[index_name].copy()
    wide = s.unstack("Station")
    wide.index = pd.to_datetime(
        wide.index.astype(int).astype(str) + "-12-31"
    )
    wide.index.name = "date"
    return wide


# ---------------------------------------------------------------------
# Mann-Kendall and Sen's slope
# ---------------------------------------------------------------------

def mk_s_statistic(x: np.ndarray) -> Tuple[float, List[int]]:
    x = np.asarray(x, dtype=float)
    x = x[np.isfinite(x)]
    n = len(x)

    if n < 2:
        return np.nan, []

    s = 0
    for i in range(n - 1):
        s += int(np.sign(x[i + 1:] - x[i]).sum())

    _, counts = np.unique(x, return_counts=True)
    ties = [int(c) for c in counts if c > 1]

    return float(s), ties


def mk_variance(n: int, ties: Sequence[int]) -> float:
    if n < 2:
        return np.nan

    tie_term = sum(
        t * (t - 1) * (2 * t + 5)
        for t in ties
    )

    return (
        n * (n - 1) * (2 * n + 5) - tie_term
    ) / 18.0


def z_from_s(s: float, var_s: float) -> float:
    if not np.isfinite(s) or not np.isfinite(var_s) or var_s <= 0:
        return np.nan

    if s > 0:
        return (s - 1.0) / math.sqrt(var_s)
    if s < 0:
        return (s + 1.0) / math.sqrt(var_s)
    return 0.0


def two_sided_p(z: float) -> float:
    if not np.isfinite(z):
        return np.nan
    return float(2.0 * stats.norm.sf(abs(z)))


def kendall_tau_from_s(s: float, n: int) -> float:
    if n < 2 or not np.isfinite(s):
        return np.nan
    return float(s / (0.5 * n * (n - 1)))


def sen_slope(x: np.ndarray) -> float:
    """
    Sen slope using equally spaced annual observations.

    For an annual series, time spacing is one year. Missing years must
    be removed before this function is called; the caller also records
    the actual first/last years so the analysis remains traceable.
    """
    x = np.asarray(x, dtype=float)
    x = x[np.isfinite(x)]
    n = len(x)

    if n < 2:
        return np.nan

    slopes = []
    for i in range(n - 1):
        dx = x[i + 1:] - x[i]
        dt = np.arange(1, n - i, dtype=float)
        slopes.extend((dx / dt).tolist())

    return float(np.median(np.asarray(slopes, dtype=float)))


def sen_slope_with_years(years: np.ndarray, values: np.ndarray) -> float:
    years = np.asarray(years, dtype=float)
    values = np.asarray(values, dtype=float)

    good = np.isfinite(years) & np.isfinite(values)
    years = years[good]
    values = values[good]

    if len(values) < 2:
        return np.nan

    slopes = []
    for i in range(len(values) - 1):
        dy = values[i + 1:] - values[i]
        dt = years[i + 1:] - years[i]
        good_dt = dt != 0
        slopes.extend((dy[good_dt] / dt[good_dt]).tolist())

    return float(np.median(np.asarray(slopes, dtype=float)))


def sen_confidence_interval(
    years: np.ndarray,
    values: np.ndarray,
    alpha: float = ALPHA,
) -> Tuple[float, float]:
    """
    Rank-based Sen slope CI following the standard Gilbert-style
    variance/order-statistic construction.

    This CI is for the slope magnitude and is independent of the
    Yue-Wang variance correction used for MMK significance.
    """
    years = np.asarray(years, dtype=float)
    values = np.asarray(values, dtype=float)

    good = np.isfinite(years) & np.isfinite(values)
    years = years[good]
    values = values[good]
    n = len(values)

    if n < 3:
        return np.nan, np.nan

    slopes = []
    for i in range(n - 1):
        dy = values[i + 1:] - values[i]
        dt = years[i + 1:] - years[i]
        slopes.extend((dy / dt).tolist())

    slopes = np.sort(np.asarray(slopes, dtype=float))
    n_slopes = len(slopes)

    s, ties = mk_s_statistic(values)
    var_s = mk_variance(n, ties)

    zcrit = stats.norm.ppf(1.0 - alpha / 2.0)
    c_alpha = zcrit * math.sqrt(max(var_s, 0.0))

    # N1/N2 order positions. Clamp only to valid array bounds.
    m1 = int(math.ceil((n_slopes - c_alpha) / 2.0))
    m2 = int(math.floor((n_slopes + c_alpha) / 2.0))

    m1 = max(1, min(n_slopes, m1))
    m2 = max(1, min(n_slopes, m2))

    lo = float(slopes[m1 - 1])
    hi = float(slopes[m2 - 1])

    return lo, hi


def lag_autocorrelation(x: np.ndarray, lag: int) -> float:
    x = np.asarray(x, dtype=float)
    if lag <= 0:
        return 1.0
    if len(x) <= lag + 1:
        return np.nan

    a = x[:-lag]
    b = x[lag:]
    good = np.isfinite(a) & np.isfinite(b)
    a = a[good]
    b = b[good]

    if len(a) < 4:
        return np.nan

    a = a - np.mean(a)
    b = b - np.mean(b)

    den = math.sqrt(float(np.sum(a * a) * np.sum(b * b)))
    if den <= 0:
        return np.nan

    return float(np.sum(a * b) / den)


def detrend_by_sen(
    years: np.ndarray,
    values: np.ndarray,
) -> Tuple[np.ndarray, float]:
    years = np.asarray(years, dtype=float)
    values = np.asarray(values, dtype=float)

    beta = sen_slope_with_years(years, values)
    if not np.isfinite(beta):
        return np.full_like(values, np.nan, dtype=float), np.nan

    t0 = float(years[0])
    detrended = values - beta * (years - t0)
    return detrended, beta


def yue_wang_2004_factor(
    years: np.ndarray,
    values: np.ndarray,
) -> Dict[str, object]:
    """
    Yue & Wang (2004) effective-sample-size correction.

    The trend is first removed using Sen's slope. The autocorrelation
    structure is then estimated from the detrended series. The full
    lag structure is used:

        n/n* = 1 + 2 * sum_{k=1}^{n-1} (1-k/n) rho_k

    and

        Var*(S) = Var(S) * (n/n*).

    This is the Yue & Wang (2004) effective-sample-size formulation,
    not the Hamed & Rao (1998) rank-autocorrelation formulation.
    """
    years = np.asarray(years, dtype=float)
    values = np.asarray(values, dtype=float)

    good = np.isfinite(years) & np.isfinite(values)
    years = years[good]
    values = values[good]

    n = len(values)
    if n < 3:
        return {
            "factor": np.nan,
            "n_eff": np.nan,
            "slope_detrend": np.nan,
            "lags_used": 0,
            "rho_sum": np.nan,
        }

    detrended, beta = detrend_by_sen(years, values)

    if not np.isfinite(detrended).all():
        return {
            "factor": np.nan,
            "n_eff": np.nan,
            "slope_detrend": beta,
            "lags_used": 0,
            "rho_sum": np.nan,
        }

    rho_sum = 0.0
    lags_used = 0

    # Yue & Wang (2004): all lags in the ESS formulation.
    for k in range(1, n):
        rho = lag_autocorrelation(detrended, k)
        if not np.isfinite(rho):
            continue

        rho_sum += (1.0 - k / n) * rho
        lags_used += 1

    factor = 1.0 + 2.0 * rho_sum

    # A non-positive factor cannot define a valid variance.
    if not np.isfinite(factor) or factor <= 0:
        return {
            "factor": np.nan,
            "n_eff": np.nan,
            "slope_detrend": beta,
            "lags_used": lags_used,
            "rho_sum": rho_sum,
        }

    n_eff = n / factor

    return {
        "factor": float(factor),
        "n_eff": float(n_eff),
        "slope_detrend": float(beta),
        "lags_used": int(lags_used),
        "rho_sum": float(rho_sum),
    }


def standard_mk(years: np.ndarray, values: np.ndarray) -> Dict[str, object]:
    years = np.asarray(years, dtype=float)
    values = np.asarray(values, dtype=float)

    good = np.isfinite(years) & np.isfinite(values)
    years = years[good]
    values = values[good]

    n = len(values)
    if n < MIN_YEARS:
        return {
            "N": n,
            "S": np.nan,
            "Var_S": np.nan,
            "Z": np.nan,
            "p": np.nan,
            "tau": np.nan,
            "slope": np.nan,
            "slope_lo": np.nan,
            "slope_hi": np.nan,
            "trend": "Insufficient data",
        }

    s, ties = mk_s_statistic(values)
    var_s = mk_variance(n, ties)
    z = z_from_s(s, var_s)
    p = two_sided_p(z)
    tau = kendall_tau_from_s(s, n)
    slope = sen_slope_with_years(years, values)
    lo, hi = sen_confidence_interval(years, values)

    return {
        "N": n,
        "S": s,
        "Var_S": var_s,
        "Z": z,
        "p": p,
        "tau": tau,
        "slope": slope,
        "slope_lo": lo,
        "slope_hi": hi,
        "trend": trend_label(z, p),
    }


def yue_wang_2004_mmk(
    years: np.ndarray,
    values: np.ndarray,
) -> Dict[str, object]:
    """
    Standard MK plus Yue & Wang (2004) ESS variance correction.
    """
    base = standard_mk(years, values)

    if not np.isfinite(base["S"]) or base["N"] < MIN_YEARS:
        base.update(
            {
                "Var_S_adj": np.nan,
                "ESS_factor_n_over_nstar": np.nan,
                "N_eff": np.nan,
                "lags_used": 0,
                "rho_sum": np.nan,
                "MMK_Z": np.nan,
                "MMK_p": np.nan,
                "MMK_trend": "Insufficient data",
            }
        )
        return base

    ess = yue_wang_2004_factor(years, values)
    factor = ess["factor"]

    if not np.isfinite(factor):
        base.update(
            {
                "Var_S_adj": np.nan,
                "ESS_factor_n_over_nstar": np.nan,
                "N_eff": np.nan,
                "lags_used": ess["lags_used"],
                "rho_sum": ess["rho_sum"],
                "MMK_Z": np.nan,
                "MMK_p": np.nan,
                "MMK_trend": "Invalid ESS correction",
            }
        )
        return base

    var_adj = float(base["Var_S"]) * factor
    z_mmk = z_from_s(float(base["S"]), var_adj)
    p_mmk = two_sided_p(z_mmk)

    base.update(
        {
            "Var_S_adj": var_adj,
            "ESS_factor_n_over_nstar": factor,
            "N_eff": ess["n_eff"],
            "lags_used": ess["lags_used"],
            "rho_sum": ess["rho_sum"],
            "MMK_Z": z_mmk,
            "MMK_p": p_mmk,
            "MMK_trend": trend_label(z_mmk, p_mmk),
        }
    )

    return base


# ---------------------------------------------------------------------
# Trend engine
# ---------------------------------------------------------------------

def analyze_index_series(
    series: pd.Series,
    period_label: str,
    station: str,
    index_name: str,
    model: str = "Observed",
    scenario: str = "Historical",
) -> Dict[str, object]:
    years = series.index.year.to_numpy(dtype=float)
    values = series.to_numpy(dtype=float)

    result = yue_wang_2004_mmk(years, values)

    first_year = int(np.nanmin(years)) if len(years) else np.nan
    last_year = int(np.nanmax(years)) if len(years) else np.nan

    return {
        "Period": period_label,
        "Model": model,
        "Scenario": scenario,
        "Station": normalize_station_id(station),
        "Index": index_name,
        "Unit": INDEX_UNITS[index_name],
        "StartYear": first_year,
        "EndYear": last_year,
        "N": result["N"],
        "S": result["S"],
        "Var_S": result["Var_S"],
        "Var_S_adj_YueWang2004": result["Var_S_adj"],
        "Z_MK": result["Z"],
        "p_MK": result["p"],
        "Trend_MK": result["trend"],
        "Z_MMK2004": result["MMK_Z"],
        "p_MMK2004": result["MMK_p"],
        "Trend_MMK2004": result["MMK_trend"],
        "Kendall_tau": result["tau"],
        "Sen_slope": result["slope"],
        "Sen_CI95_low": result["slope_lo"],
        "Sen_CI95_high": result["slope_hi"],
        "n_over_nstar": result["ESS_factor_n_over_nstar"],
        "n_effective": result["N_eff"],
        "ESS_lags_used": result["lags_used"],
        "ESS_weighted_rho_sum": result["rho_sum"],
        "MK_MMK_same_conclusion": (
            result["trend"] == result["MMK_trend"]
            if "trend" in result and "MMK_trend" in result
            else False
        ),
    }


def analyze_index_dataframe(
    index_df: pd.DataFrame,
    period_label: str,
    model: str = "Observed",
    scenario: str = "Historical",
) -> pd.DataFrame:
    rows = []

    if index_df.empty:
        return pd.DataFrame()

    for index_name in INDICES:
        if index_name not in index_df.columns:
            continue

        for station in sorted(index_df.index.get_level_values("Station").unique()):
            s = index_df.xs(station, level="Station")[index_name]
            s = s.sort_index()
            s.index = pd.Index(s.index.astype(int), name="Year")

            rows.append(
                analyze_index_series(
                    s,
                    period_label=period_label,
                    station=station,
                    index_name=index_name,
                    model=model,
                    scenario=scenario,
                )
            )

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------
# Future CMIP6 discovery
# ---------------------------------------------------------------------

def discover_bc_model_files(data_root: Path) -> Dict[Tuple[str, str], Path]:
    """
    Discover only supplied bias-corrected CMIP6 daily files.

    Raw model files are intentionally excluded from this paper.
    """
    files = list(data_root.rglob("*.csv"))
    found = {}

    for path in files:
        stem = path.stem.lower()

        if not stem.startswith("bc_"):
            continue

        model = parse_model_from_filename(path)
        if model is None:
            continue

        scenario = parse_scenario_from_filename(path)

        # Historical files are needed for the model-consistent baseline.
        if "historical" in stem:
            found[(model, "historical")] = path
        elif scenario in SCENARIOS:
            found[(model, scenario)] = path

    return found


def select_station_columns(
    df: pd.DataFrame,
    station_ids: Sequence[str],
) -> pd.DataFrame:
    mapping = {normalize_station_id(c): c for c in df.columns}
    cols = [mapping[s] for s in station_ids if s in mapping]
    out = df.loc[:, cols].copy()
    out.columns = [normalize_station_id(c) for c in out.columns]
    return out


# ---------------------------------------------------------------------
# Future change
# ---------------------------------------------------------------------

def climatological_index_mean(
    index_df: pd.DataFrame,
    start_year: int,
    end_year: int,
) -> pd.DataFrame:
    """
    Return station x index mean for a year interval.
    """
    if index_df.empty:
        return pd.DataFrame()

    tmp = index_df.reset_index()
    tmp = tmp[
        (tmp["Year"] >= start_year)
        & (tmp["Year"] <= end_year)
    ]

    if tmp.empty:
        return pd.DataFrame()

    return tmp.groupby("Station")[INDICES].mean()


def future_change_table(
    future_index_df: pd.DataFrame,
    historical_index_df: pd.DataFrame,
    model: str,
    scenario: str,
) -> pd.DataFrame:
    fut = climatological_index_mean(
        future_index_df, FUTURE_START, FUTURE_END
    )
    base = climatological_index_mean(
        historical_index_df, BASELINE_START, BASELINE_END
    )

    stations = sorted(set(fut.index) & set(base.index))
    rows = []

    for station in stations:
        for index_name in INDICES:
            f = safe_float(fut.loc[station, index_name])
            b = safe_float(base.loc[station, index_name])

            absolute = f - b if np.isfinite(f) and np.isfinite(b) else np.nan

            # Percentage change is not used as the sole metric for R50mm.
            if np.isfinite(f) and np.isfinite(b) and b != 0:
                pct = 100.0 * absolute / b
            else:
                pct = np.nan

            rows.append(
                {
                    "Model": model,
                    "Scenario": scenario,
                    "Station": station,
                    "Index": index_name,
                    "Baseline_Period": f"{BASELINE_START}-{BASELINE_END}",
                    "Future_Period": f"{FUTURE_START}-{FUTURE_END}",
                    "Baseline_mean": b,
                    "Future_mean": f,
                    "Absolute_change": absolute,
                    "Relative_change_pct": pct,
                    "Unit": INDEX_UNITS[index_name],
                }
            )

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------
# IDW
# ---------------------------------------------------------------------

def haversine_km(
    lat1: np.ndarray,
    lon1: np.ndarray,
    lat2: float,
    lon2: float,
) -> np.ndarray:
    r = 6371.0088
    p1 = np.radians(lat1)
    p2 = math.radians(lat2)
    dlat = p2 - p1
    dlon = np.radians(lon2 - lon1)

    a = (
        np.sin(dlat / 2.0) ** 2
        + np.cos(p1) * math.cos(p2) * np.sin(dlon / 2.0) ** 2
    )
    return 2.0 * r * np.arcsin(np.sqrt(np.clip(a, 0.0, 1.0)))


def idw_grid(
    coords: pd.DataFrame,
    values: pd.Series,
    power: float = 2.0,
    nx: int = 300,
    ny: int = 300,
    padding: float = 0.03,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """
    IDW spatial interpolation.

    This is a visualization of station-referenced change, not dynamical
    downscaling and not an independently observed gridded field.
    """
    tmp = coords.copy()
    tmp["value"] = values
    tmp = tmp.dropna(subset=["latitude", "longitude", "value"])

    if len(tmp) < 3:
        raise ValueError("At least 3 valid stations are required for IDW.")

    lat_min = float(tmp["latitude"].min())
    lat_max = float(tmp["latitude"].max())
    lon_min = float(tmp["longitude"].min())
    lon_max = float(tmp["longitude"].max())

    dlat = max(lat_max - lat_min, 0.01)
    dlon = max(lon_max - lon_min, 0.01)

    lat = np.linspace(
        lat_min - padding * dlat,
        lat_max + padding * dlat,
        ny,
    )
    lon = np.linspace(
        lon_min - padding * dlon,
        lon_max + padding * dlon,
        nx,
    )

    grid_lon, grid_lat = np.meshgrid(lon, lat)

    flat_lat = grid_lat.ravel()
    flat_lon = grid_lon.ravel()

    slat = tmp["latitude"].to_numpy(float)
    slon = tmp["longitude"].to_numpy(float)
    sval = tmp["value"].to_numpy(float)

    out = np.empty_like(flat_lat, dtype=float)

    for i in range(len(flat_lat)):
        d = haversine_km(slat, slon, flat_lat[i], flat_lon[i])

        if np.any(d == 0):
            out[i] = sval[np.argmin(d)]
            continue

        w = 1.0 / np.power(d, power)
        out[i] = float(np.sum(w * sval) / np.sum(w))

    return grid_lon, grid_lat, out.reshape(grid_lat.shape)


# ---------------------------------------------------------------------
# Plotting
# ---------------------------------------------------------------------

def setup_publication_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9.5,
            "axes.titlesize": 10.5,
            "axes.labelsize": 9.5,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 8.5,
            "figure.dpi": 150,
            "savefig.dpi": 600,
            "savefig.bbox": "tight",
            "axes.unicode_minus": False,
        }
    )


def plot_observed_trend_map(
    trend_df: pd.DataFrame,
    coords: pd.DataFrame,
    index_name: str,
    out_path: Path,
) -> None:
    sub = trend_df[trend_df["Index"] == index_name].copy()
    sub["Station"] = sub["Station"].map(normalize_station_id)

    m = coords.merge(
        sub[
            [
                "Station",
                "Z_MMK2004",
                "p_MMK2004",
                "Trend_MMK2004",
                "Sen_slope",
            ]
        ],
        on="Station",
        how="inner",
    )

    if m.empty:
        return

    fig, ax = plt.subplots(figsize=(6.8, 5.8))

    inc = m["Trend_MMK2004"].eq("Increasing")
    dec = m["Trend_MMK2004"].eq("Decreasing")
    ns = m["Trend_MMK2004"].eq("No significant trend")

    ax.scatter(
        m.loc[ns, "longitude"],
        m.loc[ns, "latitude"],
        s=45,
        marker="o",
        facecolors="white",
        edgecolors="black",
        linewidths=0.8,
        zorder=4,
        label="No significant trend",
    )
    ax.scatter(
        m.loc[inc, "longitude"],
        m.loc[inc, "latitude"],
        s=65,
        marker="^",
        facecolors="black",
        edgecolors="black",
        linewidths=0.5,
        zorder=5,
        label="Increasing",
    )
    ax.scatter(
        m.loc[dec, "longitude"],
        m.loc[dec, "latitude"],
        s=65,
        marker="v",
        facecolors="black",
        edgecolors="black",
        linewidths=0.5,
        zorder=5,
        label="Decreasing",
    )

    for _, r in m.iterrows():
        ax.text(
            r["longitude"] + 0.008,
            r["latitude"] + 0.008,
            r["Station"],
            fontsize=7,
            ha="left",
            va="bottom",
        )

    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_title(
        f"{index_name} | Observed trend, 1981-2014\n"
        "Yue-Wang (2004) MMK"
    )

    ax.legend(
        loc="upper left",
        bbox_to_anchor=(1.01, 1.0),
        frameon=False,
    )

    ax.grid(alpha=0.18, linewidth=0.6)
    fig.savefig(out_path)
    plt.close(fig)


def plot_future_idw_change(
    change_df: pd.DataFrame,
    coords: pd.DataFrame,
    index_name: str,
    scenario: str,
    out_path: Path,
    power: float = 2.0,
) -> None:
    sub = change_df[
        (change_df["Index"] == index_name)
        & (change_df["Scenario"] == scenario)
    ].copy()

    if sub.empty:
        return

    # Station-wise 7-model median, then IDW.
    station_values = (
        sub.groupby("Station")["Absolute_change"]
        .median()
        .rename("value")
    )

    c = coords.copy()
    c["Station"] = c["Station"].map(normalize_station_id)
    station_values.index = station_values.index.map(normalize_station_id)

    merged = c.merge(
        station_values.reset_index(),
        on="Station",
        how="inner",
    )

    if len(merged) < 3:
        return

    glon, glat, grid = idw_grid(
        merged[["Station", "latitude", "longitude"]],
        merged.set_index("Station")["value"],
        power=power,
    )

    vmax = np.nanmax(np.abs(grid))
    if not np.isfinite(vmax) or vmax == 0:
        vmax = 1.0

    norm = TwoSlopeNorm(vmin=-vmax, vcenter=0.0, vmax=vmax)

    fig, ax = plt.subplots(figsize=(7.0, 5.8))
    im = ax.pcolormesh(
        glon,
        glat,
        grid,
        shading="auto",
        cmap="RdBu_r",
        norm=norm,
    )

    ax.scatter(
        merged["longitude"],
        merged["latitude"],
        s=30,
        facecolors="white",
        edgecolors="black",
        linewidths=0.7,
        zorder=5,
    )

    for _, r in merged.iterrows():
        ax.text(
            r["longitude"] + 0.008,
            r["latitude"] + 0.008,
            r["Station"],
            fontsize=6.8,
            ha="left",
            va="bottom",
        )

    cb = fig.colorbar(im, ax=ax, pad=0.02, shrink=0.88)
    cb.set_label(f"Absolute change ({INDEX_UNITS[index_name]})")

    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_title(
        f"{index_name} | {SCENARIOS.get(scenario, scenario)}\n"
        f"2021-2050 minus {BASELINE_START}-{BASELINE_END} | "
        "7-model median + IDW"
    )
    ax.grid(alpha=0.15, linewidth=0.5)

    fig.savefig(out_path)
    plt.close(fig)


def plot_heatmap(
    trend_df: pd.DataFrame,
    period_label: str,
    out_path: Path,
    scenario: Optional[str] = None,
) -> None:
    sub = trend_df.copy()

    if scenario is not None:
        sub = sub[sub["Scenario"] == scenario]

    if sub.empty:
        return

    # For future model results, summarize by station using the median
    # Sen slope and majority direction across models.
    grouped = []
    for (station, index_name), g in sub.groupby(
        ["Station", "Index"]
    ):
        slopes = pd.to_numeric(g["Sen_slope"], errors="coerce")
        mmk_p = pd.to_numeric(g["p_MMK2004"], errors="coerce")
        mmk_z = pd.to_numeric(g["Z_MMK2004"], errors="coerce")

        med_slope = float(slopes.median())

        # Significance is represented by the direction of the median
        # model Z only for visualization; detailed model results remain
        # in the Excel/CSV outputs.
        med_z = float(mmk_z.median())

        if med_z > 0:
            val = 1
        elif med_z < 0:
            val = -1
        else:
            val = 0

        grouped.append(
            {
                "Station": station,
                "Index": index_name,
                "direction": val,
                "slope": med_slope,
            }
        )

    gdf = pd.DataFrame(grouped)
    pivot = gdf.pivot(
        index="Station",
        columns="Index",
        values="direction",
    )

    pivot = pivot.reindex(
        columns=[x for x in INDICES if x in pivot.columns]
    )

    fig, ax = plt.subplots(
        figsize=(11.0, max(5.0, 0.38 * len(pivot) + 2.5))
    )

    im = ax.imshow(
        pivot.to_numpy(dtype=float),
        aspect="auto",
        cmap="RdBu_r",
        vmin=-1,
        vmax=1,
    )

    ax.set_xticks(np.arange(len(pivot.columns)))
    ax.set_xticklabels(pivot.columns, rotation=45, ha="right")
    ax.set_yticks(np.arange(len(pivot.index)))
    ax.set_yticklabels(pivot.index)

    ax.set_xlabel("Extreme precipitation index")
    ax.set_ylabel("Station")
    ax.set_title(
        f"Trend-direction summary | {period_label}"
        + (
            f" | {SCENARIOS.get(scenario, scenario)}"
            if scenario
            else ""
        )
    )

    cbar = fig.colorbar(im, ax=ax, ticks=[-1, 0, 1], pad=0.02)
    cbar.ax.set_yticklabels(
        ["Decrease", "No trend / mixed", "Increase"]
    )

    fig.savefig(out_path)
    plt.close(fig)


# ---------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------

def write_excel(
    path: Path,
    observed_trend: pd.DataFrame,
    future_trend: pd.DataFrame,
    change_df: pd.DataFrame,
    observed_indices: pd.DataFrame,
    future_indices: pd.DataFrame,
    coords: pd.DataFrame,
    qc_observed: pd.DataFrame,
    qc_models: pd.DataFrame,
    summary: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        observed_trend.to_excel(
            writer, sheet_name="Observed_Trend", index=False
        )
        future_trend.to_excel(
            writer, sheet_name="Future_Trend", index=False
        )
        change_df.to_excel(
            writer, sheet_name="Future_Change", index=False
        )
        observed_indices.to_excel(
            writer, sheet_name="Observed_Indices"
        )
        future_indices.to_excel(
            writer, sheet_name="Future_Indices"
        )
        coords.to_excel(
            writer, sheet_name="Stations", index=False
        )
        qc_observed.to_excel(
            writer, sheet_name="QC_Observed", index=False
        )
        qc_models.to_excel(
            writer, sheet_name="QC_Models", index=False
        )
        summary.to_excel(
            writer, sheet_name="Summary", index=False
        )

    # Basic publication-oriented formatting.
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            max_len = 0
            for row in ws.iter_rows(
                min_col=col,
                max_col=col,
                values_only=True,
            ):
                val = row[0]
                max_len = max(max_len, len(str(val)) if val is not None else 0)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 28)

    wb.save(path)


# ---------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------

def build_summary(
    observed_trend: pd.DataFrame,
    future_trend: pd.DataFrame,
    change_df: pd.DataFrame,
) -> pd.DataFrame:
    rows = []

    def summarize(
        df: pd.DataFrame,
        period: str,
        group_cols: Sequence[str],
    ) -> None:
        if df.empty:
            return

        for keys, g in df.groupby(list(group_cols)):
            if not isinstance(keys, tuple):
                keys = (keys,)

            sig = pd.to_numeric(g["p_MMK2004"], errors="coerce") < ALPHA
            z = pd.to_numeric(g["Z_MMK2004"], errors="coerce")

            inc = int((sig & (z > 0)).sum())
            dec = int((sig & (z < 0)).sum())
            ns = int((~sig).sum())

            row = {
                "Period": period,
                "Index": g["Index"].iloc[0],
                "Increasing_MMK": inc,
                "Decreasing_MMK": dec,
                "No_significant_trend_MMK": ns,
                "Median_Sen_slope": float(
                    pd.to_numeric(g["Sen_slope"], errors="coerce").median()
                ),
            }

            for col, value in zip(group_cols, keys):
                row[col] = value

            rows.append(row)

    summarize(
        observed_trend,
        "1981-2014",
        ["Index"],
    )
    summarize(
        future_trend,
        "2021-2050",
        ["Model", "Scenario", "Index"],
    )

    if not change_df.empty:
        for (scenario, index_name), g in change_df.groupby(
            ["Scenario", "Index"]
        ):
            vals = pd.to_numeric(
                g["Absolute_change"], errors="coerce"
            )
            rows.append(
                {
                    "Period": "2021-2050 change",
                    "Scenario": scenario,
                    "Index": index_name,
                    "Median_future_minus_baseline": float(vals.median()),
                    "IQR_low": float(vals.quantile(0.25)),
                    "IQR_high": float(vals.quantile(0.75)),
                }
            )

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------
# Main analysis
# ---------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Uttaradit 11-index extreme precipitation trend analysis "
            "using MK + Yue-Wang (2004) MMK + Sen slope."
        )
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to extracted dataUttaradit directory or dataUttaradit.rar",
    )
    parser.add_argument(
        "--output",
        default="output_utt_trend",
        help="Output directory",
    )
    parser.add_argument(
        "--idw-power",
        type=float,
        default=2.0,
        help="IDW power parameter (default 2.0)",
    )
    parser.add_argument(
        "--no-figures",
        action="store_true",
        help="Skip figure generation",
    )
    args = parser.parse_args()

    setup_publication_style()

    input_path = Path(args.input)
    output_root = Path(args.output).resolve()
    result_dir = output_root / "results"
    fig_dir = output_root / "figures"

    result_dir.mkdir(parents=True, exist_ok=True)
    fig_dir.mkdir(parents=True, exist_ok=True)

    log("=" * 78)
    log(f"Uttaradit Extreme Precipitation Trend Analysis v{VERSION}")
    log("Core: 11 indices | MK | Yue-Wang (2004) MMK | Sen slope")
    log("Observed: 1981-2014 | Future: 2021-2050")
    log("=" * 78)

    # -------------------------------------------------------------
    # Locate/extract archive
    # -------------------------------------------------------------
    prepared = prepare_input_root(input_path)
    data_root = locate_data_root(prepared)

    log(f"Data root: {data_root}")

    obs_path = (
        data_root
        / "Observed_Rain_daily_198101_201412_Uttaradit.csv"
    )

    coord_candidates = [
        data_root / "station_coordinates.xlsx",
        data_root / "gis" / "station_coordinates.xlsx",
    ]

    coord_path = next(
        (p for p in coord_candidates if p.exists()),
        None,
    )

    if coord_path is None:
        matches = list(
            data_root.rglob("station_coordinates.xlsx")
        )
        if matches:
            coord_path = matches[0]

    if coord_path is None:
        raise FileNotFoundError(
            "station_coordinates.xlsx not found."
        )

    # -------------------------------------------------------------
    # Observed data
    # -------------------------------------------------------------
    log("\n[1/8] Loading observed rainfall ...")
    obs_raw = load_daily_csv(obs_path)
    obs_raw = select_station_columns(
        obs_raw,
        [
            normalize_station_id(c)
            for c in obs_raw.columns
        ],
    )

    obs, qc_obs = qc_daily(
        obs_raw,
        expected_start=f"{OBS_START}-01-01",
        expected_end=f"{OBS_END}-12-31",
    )

    coords = read_coordinates(coord_path)

    observed_station_ids = {
        normalize_station_id(c) for c in obs.columns
    }
    coord_station_ids = set(coords["station"])

    station_ids = sorted(
        observed_station_ids & coord_station_ids
    )

    if len(station_ids) == 0:
        raise RuntimeError(
            "No common stations between observed rainfall and coordinates."
        )

    # Restrict observed to coordinate-supported stations.
    obs = select_station_columns(obs, station_ids)
    coords = coords[
        coords["station"].isin(station_ids)
    ].copy()

    log(f"Observed stations used: {len(station_ids)}")
    log(f"Stations: {', '.join(station_ids)}")

    # -------------------------------------------------------------
    # Observed ETCCDI-based indices
    # -------------------------------------------------------------
    log("\n[2/8] Computing the 11 observed extreme indices ...")
    thresholds = station_thresholds_from_observed(obs)

    obs_indices = annual_extreme_indices(
        obs,
        thresholds,
    )

    observed_trend = analyze_index_dataframe(
        obs_indices,
        period_label="1981-2014",
        model="Observed",
        scenario="Historical",
    )

    observed_trend.to_csv(
        result_dir / "observed_trend_1981_2014.csv",
        index=False,
    )
    obs_indices.to_csv(
        result_dir / "observed_indices_1981_2014.csv"
    )
    qc_obs.to_csv(
        result_dir / "qc_observed.csv",
        index=False,
    )

    log(
        f"Observed trend rows: {len(observed_trend)} "
        f"(expected up to {len(station_ids) * len(INDICES)})"
    )

    # -------------------------------------------------------------
    # Discover CMIP6 supplied bias-corrected files
    # -------------------------------------------------------------
    log("\n[3/8] Discovering supplied bias-corrected CMIP6 files ...")
    bc_files = discover_bc_model_files(data_root)

    for key, path in sorted(bc_files.items()):
        log(f"  {key}: {path.name}")

    models_found = sorted(
        {
            model
            for model, period in bc_files
            if period in {"historical", "ssp245", "ssp585"}
        }
    )

    log(
        f"Models discovered: {len(models_found)} -> "
        f"{', '.join(models_found)}"
    )

    # -------------------------------------------------------------
    # Future model calculations
    # -------------------------------------------------------------
    future_trend_parts = []
    change_parts = []
    qc_model_parts = []
    future_index_parts = []

    log("\n[4/8] Computing model historical baselines and future indices ...")

    model_hist_indices: Dict[str, pd.DataFrame] = {}
    model_future_indices: Dict[Tuple[str, str], pd.DataFrame] = {}

    for model in models_found:
        hist_path = bc_files.get((model, "historical"))

        if hist_path is None:
            log(f"  [SKIP] {model}: no historical BC file")
            continue

        log(f"  Historical BC: {model}")
        hist_raw = load_daily_csv(hist_path)
        hist = select_station_columns(hist_raw, station_ids)
        hist, qc_hist = qc_daily(
            hist,
            expected_start=f"{OBS_START}-01-01",
            expected_end=f"{OBS_END}-12-31",
        )
        qc_hist["Model"] = model
        qc_hist["Series"] = "Historical_BC"
        qc_model_parts.append(qc_hist)

        hist_indices = annual_extreme_indices(
            hist,
            thresholds,
        )
        model_hist_indices[model] = hist_indices

        for scenario in SCENARIOS:
            future_path = bc_files.get((model, scenario))
            if future_path is None:
                log(
                    f"  [SKIP] {model} {SCENARIOS[scenario]}: "
                    "future BC file not found"
                )
                continue

            log(f"  Future BC: {model} {SCENARIOS[scenario]}")

            fut_raw = load_daily_csv(future_path)
            fut = select_station_columns(fut_raw, station_ids)
            fut, qc_fut = qc_daily(
                fut,
                expected_start=f"{FUTURE_START}-01-01",
                expected_end=f"{FUTURE_END}-12-31",
            )
            qc_fut["Model"] = model
            qc_fut["Series"] = scenario
            qc_model_parts.append(qc_fut)

            fut_indices = annual_extreme_indices(
                fut,
                thresholds,
            )
            model_future_indices[(model, scenario)] = fut_indices

            # Future trend: 2021-2050 only.
            fut_trend = analyze_index_dataframe(
                fut_indices,
                period_label="2021-2050",
                model=model,
                scenario=SCENARIOS[scenario],
            )
            future_trend_parts.append(fut_trend)

            # Future minus model-consistent 1995-2014 baseline.
            change = future_change_table(
                fut_indices,
                hist_indices,
                model=model,
                scenario=SCENARIOS[scenario],
            )
            change_parts.append(change)

            # Long-format future indices for audit.
            if not fut_indices.empty:
                tmp = fut_indices.reset_index()
                tmp["Model"] = model
                tmp["Scenario"] = SCENARIOS[scenario]
                future_index_parts.append(tmp)

    future_trend = (
        pd.concat(future_trend_parts, ignore_index=True)
        if future_trend_parts
        else pd.DataFrame()
    )

    change_df = (
        pd.concat(change_parts, ignore_index=True)
        if change_parts
        else pd.DataFrame()
    )

    future_indices_long = (
        pd.concat(future_index_parts, ignore_index=True)
        if future_index_parts
        else pd.DataFrame()
    )

    qc_models = (
        pd.concat(qc_model_parts, ignore_index=True)
        if qc_model_parts
        else pd.DataFrame()
    )

    # -------------------------------------------------------------
    # Save numerical outputs
    # -------------------------------------------------------------
    log("\n[5/8] Writing numerical results ...")

    future_trend.to_csv(
        result_dir / "future_trend_2021_2050.csv",
        index=False,
    )
    change_df.to_csv(
        result_dir / "future_change_2021_2050.csv",
        index=False,
    )
    future_indices_long.to_csv(
        result_dir / "future_indices_2021_2050.csv",
        index=False,
    )
    qc_models.to_csv(
        result_dir / "qc_models.csv",
        index=False,
    )
    coords.to_csv(
        result_dir / "station_coordinates_used.csv",
        index=False,
    )

    # -------------------------------------------------------------
    # Summaries
    # -------------------------------------------------------------
    log("\n[6/8] Building summary tables ...")
    summary = build_summary(
        observed_trend,
        future_trend,
        change_df,
    )
    summary.to_csv(
        result_dir / "summary_by_index.csv",
        index=False,
    )

    if not future_trend.empty:
        future_model_summary = (
            future_trend.groupby(
                ["Model", "Scenario", "Index"],
                dropna=False,
            )
            .agg(
                Median_Sen_slope=("Sen_slope", "median"),
                Median_MMK_Z=("Z_MMK2004", "median"),
                Median_MMK_p=("p_MMK2004", "median"),
                N_series=("Station", "count"),
            )
            .reset_index()
        )
    else:
        future_model_summary = pd.DataFrame()

    future_model_summary.to_csv(
        result_dir / "summary_by_model_scenario.csv",
        index=False,
    )

    # -------------------------------------------------------------
    # Figures
    # -------------------------------------------------------------
    if not args.no_figures:
        log("\n[7/8] Generating publication figures ...")

        for index_name in INDICES:
            plot_observed_trend_map(
                observed_trend,
                coords,
                index_name,
                fig_dir / (
                    f"observed_trends_1981_2014_{index_name}.png"
                ),
            )

        plot_heatmap(
            observed_trend,
            "1981-2014",
            fig_dir / "observed_trend_heatmap.png",
        )

        if not future_trend.empty:
            for scenario in SCENARIOS:
                s = SCENARIOS[scenario]
                for index_name in INDICES:
                    plot_future_idw_change(
                        change_df,
                        coords,
                        index_name,
                        scenario,
                        fig_dir / (
                            f"future_change_IDW_2021_2050_"
                            f"{scenario}_{index_name}.png"
                        ),
                        power=args.idw_power,
                    )

                plot_heatmap(
                    future_trend,
                    "2021-2050",
                    fig_dir / (
                        f"future_trend_heatmap_{scenario}.png"
                    ),
                    scenario=s,
                )

    else:
        log("\n[7/8] Figures skipped by --no-figures")

    # -------------------------------------------------------------
    # Excel
    # -------------------------------------------------------------
    log("\n[8/8] Writing Excel workbook ...")

    write_excel(
        result_dir / "Uttaradit_Trend_Analysis.xlsx",
        observed_trend=observed_trend,
        future_trend=future_trend,
        change_df=change_df,
        observed_indices=obs_indices,
        future_indices=future_indices_long,
        coords=coords,
        qc_observed=qc_obs,
        qc_models=qc_models,
        summary=summary,
    )

    readme = result_dir / "README_results.txt"
    readme.write_text(
        "\n".join(
            [
                f"Uttaradit Extreme Precipitation Trend Analysis v{VERSION}",
                "",
                "Title:",
                "Spatiotemporal Trends of ETCCDI-Based Extreme "
                "Precipitation Indices in Uttaradit Province, Thailand",
                "",
                "Target indices:",
                ", ".join(INDICES),
                "",
                "Observed period: 1981-2014",
                "Future period: 2021-2050",
                "Future baseline: 1995-2014, model-consistent",
                "",
                "Trend methods:",
                "1) Standard Mann-Kendall",
                "2) Yue & Wang (2004) effective-sample-size MMK",
                "3) Sen's slope + 95% CI",
                "",
                "Future spatial visualization:",
                "7-model station-wise median change followed by IDW.",
                "",
                "Excluded:",
                "Pettitt, FDR, Bonferroni, PW, TFPW, seasonal trends, "
                "monthly trends, Taylor diagrams, non-target indices.",
                "",
                "R50mm is treated as a study-specific fixed-threshold index.",
                "",
                "Important:",
                "IDW maps visualize station-referenced change and are not "
                "independent gridded observations or dynamical downscaling.",
            ]
        ),
        encoding="utf-8",
    )

    # -------------------------------------------------------------
    # Console scientific QC
    # -------------------------------------------------------------
    log("\n" + "=" * 78)
    log("ANALYSIS COMPLETE")
    log("=" * 78)
    log(f"Observed trend rows: {len(observed_trend)}")
    log(f"Future trend rows:   {len(future_trend)}")
    log(f"Future change rows:  {len(change_df)}")

    if not observed_trend.empty:
        n_obs_sig = int(
            (
                pd.to_numeric(
                    observed_trend["p_MMK2004"],
                    errors="coerce",
                )
                < ALPHA
            ).sum()
        )
        log(
            f"Observed significant MMK-2004 trends: "
            f"{n_obs_sig}/{len(observed_trend)}"
        )

    if not future_trend.empty:
        n_fut_sig = int(
            (
                pd.to_numeric(
                    future_trend["p_MMK2004"],
                    errors="coerce",
                )
                < ALPHA
            ).sum()
        )
        log(
            f"Future significant MMK-2004 trends: "
            f"{n_fut_sig}/{len(future_trend)}"
        )

    log(f"Results: {result_dir}")
    log(f"Figures: {fig_dir}")
    log("=" * 78)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
