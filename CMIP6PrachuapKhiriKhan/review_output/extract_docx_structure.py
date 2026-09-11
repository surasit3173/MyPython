from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


DOCX = Path(r"C:\MyPython\CMIP6PrachuapKhiriKhan\MK4_Full_Results_Explanation.docx")


def iter_body(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "P", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield "T", Table(child, document)


doc = Document(DOCX)
props = doc.core_properties
print(
    "META",
    {
        "title": props.title,
        "subject": props.subject,
        "author": props.author,
        "created": str(props.created),
        "modified": str(props.modified),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "inline_shapes": len(doc.inline_shapes),
    },
)
for i, section in enumerate(doc.sections, 1):
    print(
        f"SECTION {i}",
        {
            "width_emu": section.page_width,
            "height_emu": section.page_height,
            "top_margin_emu": section.top_margin,
            "bottom_margin_emu": section.bottom_margin,
            "left_margin_emu": section.left_margin,
            "right_margin_emu": section.right_margin,
        },
    )

page = 1
p_index = 0
t_index = 0
for kind, item in iter_body(doc):
    if kind == "P":
        p_index += 1
        xml = item._p.xml
        before = xml.count("lastRenderedPageBreak") + xml.count('w:br w:type="page"')
        text = " ".join(item.text.split())
        if text or before:
            print(f"P{p_index:03d} page~{page} style={item.style.name!r}: {text}")
        page += before
    else:
        t_index += 1
        print(f"TABLE {t_index:02d} page~{page} rows={len(item.rows)} cols={len(item.columns)}")
        for r_index, row in enumerate(item.rows, 1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            print(f"  R{r_index:02d}: " + " || ".join(cells))

with ZipFile(DOCX) as zf:
    names = zf.namelist()
    media = [n for n in names if n.startswith("word/media/")]
    comments = [n for n in names if "comments" in n]
    footnotes = [n for n in names if "footnotes" in n]
    print("PACKAGE", {"media": media, "comments": comments, "footnotes": footnotes})
