"""
Mann–Kendall Trend Test + Sen's Slope Estimator
รายสถานีวัดน้ำฝน – Prachuap Khiri Khan
รวม Z-statistic, p-value, ทิศทาง, 95% CI ของ Sen's slope
และ Mean Annual Rainfall, CV (%)

ผู้เขียน: Climate & Hydrology Analysis Script
อ้างอิง:
  - Mann (1945), Kendall (1975)
  - Sen (1968) – Estimates of the Regression Coefficient Based on Kendall's Tau
  - Gilbert (1987) – Statistical Methods for Environmental Pollution Monitoring
"""

import os
import sys
import math
import warnings
import numpy as np
import pandas as pd
from scipy import stats
from itertools import combinations
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import FORMAT_NUMBER_00

warnings.filterwarnings("ignore")

# ─────────────────────────────────────────────
# 1.  CORE STATISTICAL FUNCTIONS
# ─────────────────────────────────────────────

def mann_kendall_test(x: np.ndarray, alpha: float = 0.05):
    """
    Two-tailed Mann–Kendall trend test with tied-value correction.

    Parameters
    ----------
    x     : 1-D array of observations (chronological order)
    alpha : significance level (default 0.05)

    Returns
    -------
    dict with keys:
        n, S, var_S, Z, p_value, tau, trend, significant
    """
    n = len(x)
    if n < 4:
        raise ValueError("Mann–Kendall requires at least 4 observations.")

    # ── S statistic ──────────────────────────────────────────────────────
    S = 0
    for k in range(n - 1):
        for j in range(k + 1, n):
            diff = x[j] - x[k]
            if diff > 0:
                S += 1
            elif diff < 0:
                S -= 1

    # ── Variance with tied-group correction ──────────────────────────────
    # unique values & their counts
    unique, counts = np.unique(x, return_counts=True)
    tied_sum = np.sum(counts * (counts - 1) * (2 * counts + 5))
    var_S = (n * (n - 1) * (2 * n + 5) - tied_sum) / 18

    # ── Z statistic (continuity correction) ──────────────────────────────
    if S > 0:
        Z = (S - 1) / math.sqrt(var_S)
    elif S < 0:
        Z = (S + 1) / math.sqrt(var_S)
    else:
        Z = 0.0

    # ── Two-tailed p-value ────────────────────────────────────────────────
    p_value = 2 * (1 - stats.norm.cdf(abs(Z)))

    # ── Kendall's tau ─────────────────────────────────────────────────────
    tau = S / (0.5 * n * (n - 1))

    # ── Trend classification ──────────────────────────────────────────────
    if p_value <= alpha:
        trend = "Increasing" if Z > 0 else "Decreasing"
        significant = True
    else:
        trend = "No trend"
        significant = False

    return {
        "n": n,
        "S": S,
        "var_S": var_S,
        "Z": Z,
        "p_value": p_value,
        "tau": tau,
        "trend": trend,
        "significant": significant,
    }


def sens_slope_ci(x: np.ndarray, alpha: float = 0.05):
    """
    Sen's slope estimator with rank-based (1-alpha)×100% confidence interval.

    Parameters
    ----------
    x     : 1-D array of observations (chronological order)
    alpha : significance level for CI (default 0.05 → 95% CI)

    Returns
    -------
    slope    : Sen's slope (median of pairwise slopes)
    ci_lower : lower bound of confidence interval
    ci_upper : upper bound of confidence interval
    """
    n = len(x)
    t = np.arange(1, n + 1, dtype=float)          # time index 1 … n

    # All pairwise slopes  Q_ij = (x_j − x_i) / (j − i)
    slopes = []
    for i in range(n - 1):
        for j in range(i + 1, n):
            slopes.append((x[j] - x[i]) / (t[j] - t[i]))
    slopes = np.sort(slopes)

    N = len(slopes)                                 # = n(n-1)/2

    # Median slope
    slope = np.median(slopes)

    # ── Confidence interval (Gilbert 1987, p. 218) ────────────────────────
    # number of unique pairs for variance of S  (reuse from MK)
    unique, counts = np.unique(x, return_counts=True)
    tied_sum = np.sum(counts * (counts - 1) * (2 * counts + 5))
    var_S = (n * (n - 1) * (2 * n + 5) - tied_sum) / 18

    Z_a2 = stats.norm.ppf(1 - alpha / 2)           # e.g. 1.96 for 95%
    C_alpha = Z_a2 * math.sqrt(var_S)

    M1 = int(round((N - C_alpha) / 2))
    M2 = int(round((N + C_alpha) / 2))

    # Clamp to valid indices
    M1 = max(0, min(M1, N - 1))
    M2 = max(0, min(M2, N - 1))

    ci_lower = slopes[M1]
    ci_upper = slopes[M2]

    return slope, ci_lower, ci_upper


# ─────────────────────────────────────────────
# 2.  LOAD DATA & AGGREGATE TO ANNUAL TOTALS
# ─────────────────────────────────────────────

def load_and_aggregate(csv_path: str):
    """
    Read daily rainfall CSV and return a DataFrame of annual totals
    (rows = years, columns = station IDs).
    Missing / negative values are replaced with NaN.
    """
    df = pd.read_csv(csv_path)
    df.replace(-99, np.nan, inplace=True)          # common missing flag
    df.replace(-999, np.nan, inplace=True)
    df[df < 0] = np.nan

    station_cols = [c for c in df.columns if c not in ("YEAR", "MONTH", "DAY")]

    # Sum daily → annual; skip years with > 20 % missing days
    annual = (
        df.groupby("YEAR")[station_cols]
        .apply(lambda g: g.sum(min_count=int(0.8 * len(g))))
    )
    annual.index.name = "YEAR"
    return annual, station_cols


# ─────────────────────────────────────────────
# 3.  ANALYSE EVERY STATION
# ─────────────────────────────────────────────

def analyse_all_stations(annual: pd.DataFrame, station_cols: list, alpha: float = 0.05):
    """
    Run Mann–Kendall + Sen's slope for each station column.
    Returns a list of result-dict per station.
    """
    results = []
    for stn in station_cols:
        series = annual[stn].dropna().values

        if len(series) < 4:
            print(f"  ⚠  Station {stn}: insufficient data ({len(series)} years) – skipped.")
            continue

        # ── Descriptive statistics ──────────────────────────────────────
        mean_rain = np.mean(series)
        std_rain  = np.std(series, ddof=1)
        cv_pct    = (std_rain / mean_rain) * 100 if mean_rain != 0 else np.nan

        # ── Trend tests ─────────────────────────────────────────────────
        mk  = mann_kendall_test(series, alpha=alpha)
        slope, ci_lo, ci_hi = sens_slope_ci(series, alpha=alpha)

        # ── p-value annotation ──────────────────────────────────────────
        pv = mk["p_value"]
        if pv <= 0.001:
            pv_label = f"{pv:.4f}***"
        elif pv <= 0.01:
            pv_label = f"{pv:.4f}**"
        elif pv <= 0.05:
            pv_label = f"{pv:.4f}*"
        else:
            pv_label = f"{pv:.4f} ns"

        results.append({
            "Station":              stn,
            "N (years)":            mk["n"],
            "Mean Annual RF (mm)":  round(mean_rain, 1),
            "Std Dev (mm)":         round(std_rain, 1),
            "CV (%)":               round(cv_pct, 1),
            "S statistic":          mk["S"],
            "Var(S)":               round(mk["var_S"], 2),
            "Z-statistic":          round(mk["Z"], 3),
            "p-value (raw)":        round(pv, 4),
            "p-value (annot)":      pv_label,
            "Kendall tau":          round(mk["tau"], 4),
            "Sen slope (mm/yr)":    round(slope, 2),
            "95% CI Lower (mm/yr)": round(ci_lo, 2),
            "95% CI Upper (mm/yr)": round(ci_hi, 2),
            "Trend":                mk["trend"],
            "Significant":          mk["significant"],
        })
        print(f"  ✓  Station {stn:>10s}  |  Z={mk['Z']:+.3f}  p={pv:.4f}  "
              f"slope={slope:+.2f} mm/yr  [{mk['trend']}]")

    return results


# ─────────────────────────────────────────────
# 4.  EXPORT TO EXCEL (openpyxl – styled)
# ─────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")   # dark blue
SUBHEAD_FILL  = PatternFill("solid", fgColor="2E75B6")   # medium blue
INC_FILL      = PatternFill("solid", fgColor="C6EFCE")   # light green
DEC_FILL      = PatternFill("solid", fgColor="FFCCCC")   # light red
NO_FILL       = PatternFill("solid", fgColor="FFFFCC")   # light yellow
SIG_FILL      = PatternFill("solid", fgColor="D9E1F2")   # light blue (sig row)

THIN  = Side(style="thin",   color="000000")
THICK = Side(style="medium", color="000000")
thin_border  = Border(left=THIN,  right=THIN,  top=THIN,  bottom=THIN)
thick_border = Border(left=THICK, right=THICK, top=THICK, bottom=THICK)


def style_cell(ws, row, col,
               value=None, bold=False, italic=False,
               font_color="000000", fill=None, align="center",
               border=None, num_fmt=None, font_size=10):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    cell.font = Font(bold=bold, italic=italic, color=font_color,
                     name="Calibri", size=font_size)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                                wrap_text=True)
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if num_fmt:
        cell.number_format = num_fmt
    return cell


def write_excel(results: list, annual: pd.DataFrame,
                station_cols: list, out_path: str):

    wb = Workbook()

    # ──────────────────────────────────────────
    # SHEET 1 : Summary results table
    # ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "MK Results (Summary)"
    ws1.sheet_view.showGridLines = False
    ws1.freeze_panes = "A5"

    # -- Title block --
    ws1.merge_cells("A1:P1")
    style_cell(ws1, 1, 1,
               value="Mann–Kendall Trend Test and Sen's Slope Estimator – Annual Rainfall",
               bold=True, font_color="FFFFFF", fill=HEADER_FILL,
               font_size=13)

    ws1.merge_cells("A2:P2")
    style_cell(ws1, 2, 1,
               value=f"Significance level α = 0.05  |  Two-tailed test  |  "
                     f"Period: {annual.index.min()}–{annual.index.max()}  |  "
                     f"Stations: {len(results)}",
               bold=False, font_color="FFFFFF", fill=SUBHEAD_FILL,
               font_size=10)

    ws1.merge_cells("A3:P3")
    style_cell(ws1, 3, 1,
               value="*** p ≤ 0.001  ** p ≤ 0.01  * p ≤ 0.05  ns = not significant  |  "
                     "CV = Coefficient of Variation  |  95% CI = Confidence Interval of Sen's slope",
               italic=True, font_color="1F4E79", fill=None, font_size=9)

    # -- Column headers --
    headers = [
        "Station", "N\n(years)",
        "Mean Annual\nRainfall (mm)", "Std Dev\n(mm)", "CV\n(%)",
        "S\nStatistic", "Var(S)",
        "Z-\nStatistic", "p-value",
        "Kendall\ntau",
        "Sen's Slope\n(mm yr⁻¹)",
        "95% CI\nLower (mm yr⁻¹)", "95% CI\nUpper (mm yr⁻¹)",
        "Trend",
        "Significant\n(α=0.05)",
    ]
    for c, h in enumerate(headers, 1):
        style_cell(ws1, 4, c, value=h, bold=True,
                   font_color="FFFFFF", fill=HEADER_FILL,
                   border=thin_border, font_size=10)

    ws1.row_dimensions[4].height = 40

    # -- Data rows --
    for r_idx, res in enumerate(results, 5):
        # row background
        sig  = res["Significant"]
        trend = res["Trend"]
        if trend == "Increasing":
            row_fill = INC_FILL
        elif trend == "Decreasing":
            row_fill = DEC_FILL
        else:
            row_fill = NO_FILL if not sig else SIG_FILL

        row_data = [
            res["Station"],
            res["N (years)"],
            res["Mean Annual RF (mm)"],
            res["Std Dev (mm)"],
            res["CV (%)"],
            res["S statistic"],
            res["Var(S)"],
            res["Z-statistic"],
            res["p-value (annot)"],
            res["Kendall tau"],
            res["Sen slope (mm/yr)"],
            res["95% CI Lower (mm/yr)"],
            res["95% CI Upper (mm/yr)"],
            res["Trend"],
            "Yes" if sig else "No",
        ]

        for c_idx, val in enumerate(row_data, 1):
            cell = style_cell(ws1, r_idx, c_idx, value=val,
                              fill=row_fill, border=thin_border,
                              align="center", font_size=10)
            # left-align station name
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="left",
                                            vertical="center")
            # colour-code Trend cell
            if c_idx == 14:
                if val == "Increasing":
                    cell.font = Font(bold=True, color="375623",
                                     name="Calibri", size=10)
                elif val == "Decreasing":
                    cell.font = Font(bold=True, color="9C0006",
                                     name="Calibri", size=10)
                else:
                    cell.font = Font(bold=True, color="7D6608",
                                     name="Calibri", size=10)
        ws1.row_dimensions[r_idx].height = 18

    # -- Column widths --
    col_widths = [14, 8, 18, 12, 8,
                  11, 12, 11, 16, 11,
                  18, 20, 20, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ──────────────────────────────────────────
    # SHEET 2 : Annual Rainfall Time Series
    # ──────────────────────────────────────────
    ws2 = wb.create_sheet("Annual Rainfall (mm)")
    ws2.sheet_view.showGridLines = False

    # header
    ws2.merge_cells(f"A1:{get_column_letter(len(station_cols)+1)}1")
    style_cell(ws2, 1, 1,
               value="Annual Rainfall (mm) – All Stations",
               bold=True, font_color="FFFFFF",
               fill=HEADER_FILL, font_size=12)

    style_cell(ws2, 2, 1, value="YEAR", bold=True,
               font_color="FFFFFF", fill=SUBHEAD_FILL,
               border=thin_border, font_size=10)
    for c, stn in enumerate(station_cols, 2):
        style_cell(ws2, 2, c, value=str(stn), bold=True,
                   font_color="FFFFFF", fill=SUBHEAD_FILL,
                   border=thin_border, font_size=10)

    for r_idx, (yr, row) in enumerate(annual.iterrows(), 3):
        style_cell(ws2, r_idx, 1, value=int(yr),
                   bold=True, border=thin_border, font_size=9)
        for c_idx, stn in enumerate(station_cols, 2):
            val = row[stn]
            cell = style_cell(ws2, r_idx, c_idx,
                              value=None if pd.isna(val) else round(val, 1),
                              border=thin_border, font_size=9)
        ws2.row_dimensions[r_idx].height = 15

    # widths
    ws2.column_dimensions["A"].width = 8
    for c in range(2, len(station_cols) + 2):
        ws2.column_dimensions[get_column_letter(c)].width = 11

    # ──────────────────────────────────────────
    # SHEET 3 : Statistics Reference
    # ──────────────────────────────────────────
    ws3 = wb.create_sheet("Method & References")
    ws3.sheet_view.showGridLines = False
    ws3.column_dimensions["A"].width = 4
    ws3.column_dimensions["B"].width = 28
    ws3.column_dimensions["C"].width = 70

    ws3.merge_cells("A1:C1")
    style_cell(ws3, 1, 1,
               value="Statistical Method Description & References",
               bold=True, font_color="FFFFFF",
               fill=HEADER_FILL, font_size=13)

    info_rows = [
        ("Test", "Mann–Kendall Trend Test",
         "Non-parametric test for monotonic trends in time series. "
         "Null hypothesis H₀: no trend. "
         "Two-tailed test applied (α = 0.05)."),
        ("S statistic", "Kendall's S",
         "S = Σ sgn(xⱼ – xᵢ) for all i < j. "
         "Positive S → upward trend; negative S → downward trend."),
        ("Var(S)", "Variance of S",
         "Var(S) = [n(n–1)(2n+5) – Σtₚ(tₚ–1)(2tₚ+5)] / 18, "
         "where tₚ is the count of ties in tied group p."),
        ("Z", "Z-statistic",
         "Z = (S–1)/√Var(S) if S>0; 0 if S=0; (S+1)/√Var(S) if S<0. "
         "Follows standard normal distribution under H₀."),
        ("p-value", "Two-tailed p-value",
         "p = 2 × [1 – Φ(|Z|)]. "
         "Significance: *** p≤0.001  ** p≤0.01  * p≤0.05  ns p>0.05."),
        ("τ", "Kendall's tau",
         "τ = S / [½n(n–1)]. Ranges from –1 (perfect decreasing) "
         "to +1 (perfect increasing)."),
        ("Slope", "Sen's Slope Estimator",
         "Median of all pairwise slopes Qᵢⱼ = (xⱼ – xᵢ)/(j–i). "
         "Robust to outliers and non-normality (Sen 1968)."),
        ("95% CI", "Confidence Interval of Sen's Slope",
         "Rank-based CI (Gilbert 1987): "
         "M₁ = (N – Cα)/2  and  M₂ = (N + Cα)/2, "
         "where Cα = z_{α/2} √Var(S) and N = n(n–1)/2."),
        ("CV", "Coefficient of Variation (%)",
         "CV = (σ / μ) × 100. Measures inter-annual variability "
         "relative to the mean annual rainfall."),
        ("", "References",
         "Mann, H.B. (1945). Non-parametric tests against trend. "
         "Econometrica 13(3):245–259.\n"
         "Kendall, M.G. (1975). Rank Correlation Methods. 4th ed. "
         "Charles Griffin, London.\n"
         "Sen, P.K. (1968). Estimates of the Regression Coefficient Based on Kendall's Tau. "
         "JASA 63(324):1379–1389.\n"
         "Gilbert, R.O. (1987). Statistical Methods for Environmental Pollution Monitoring. "
         "Van Nostrand Reinhold, New York."),
    ]

    for r_offset, (abbr, term, desc) in enumerate(info_rows, 3):
        fill = PatternFill("solid", fgColor="DEEAF1") if r_offset % 2 == 0 \
               else PatternFill("solid", fgColor="FFFFFF")
        style_cell(ws3, r_offset, 1, value=abbr,
                   bold=True, fill=fill, border=thin_border,
                   font_size=10, align="center")
        style_cell(ws3, r_offset, 2, value=term,
                   bold=True, fill=fill, border=thin_border,
                   font_size=10, align="left")
        cell = style_cell(ws3, r_offset, 3, value=desc,
                          fill=fill, border=thin_border,
                          font_size=10, align="left")
        cell.alignment = Alignment(horizontal="left", vertical="top",
                                    wrap_text=True)
        ws3.row_dimensions[r_offset].height = 45

    # ── Save ────────────────────────────────────────────────────────────
    wb.save(out_path)
    print(f"\n  💾  Saved → {out_path}")


# ─────────────────────────────────────────────
# 5.  MAIN
# ─────────────────────────────────────────────

def main():
    # ── Input path ──────────────────────────────────────────────────────
    # เปลี่ยน input_csv เป็น path ของไฟล์ข้อมูลจริงบนเครื่อง
    input_csv = r"C:\Prachuap Khiri Khan\Observed_Rain_daily_198101_201412_Prachuap_Khiri_Khan.csv"

    if not os.path.isfile(input_csv):
        # ── fallback: ค้นหา csv แรกในโฟลเดอร์เดียวกับ script ──────────
        script_dir = os.path.dirname(os.path.abspath(__file__))
        candidates = [f for f in os.listdir(script_dir)
                      if f.lower().endswith(".csv")]
        if candidates:
            input_csv = os.path.join(script_dir, candidates[0])
            print(f"  ℹ  ไม่พบ path หลัก – ใช้ไฟล์ที่พบ: {input_csv}")
        else:
            sys.exit("  ✗  ไม่พบไฟล์ข้อมูล CSV – กรุณาตรวจสอบ path")

    # ── Output path: ชื่อเดียวกับ input, นามสกุล .xlsx ─────────────────
    base_name   = os.path.splitext(os.path.basename(input_csv))[0]
    output_dir  = os.path.dirname(input_csv)          # โฟลเดอร์เดียวกับ input
    output_xlsx = os.path.join(output_dir, "Output_" + base_name + ".xlsx")

    print("=" * 65)
    print("  Mann–Kendall Trend Test + Sen's Slope – Prachuap Khiri Khan")
    print("=" * 65)
    print(f"  Input  : {input_csv}")
    print(f"  Output : {output_xlsx}")
    print("-" * 65)

    # ── Load & aggregate ─────────────────────────────────────────────────
    print("  กำลังโหลดและรวมข้อมูลรายปี ...")
    annual, station_cols = load_and_aggregate(input_csv)
    print(f"  สถานี {len(station_cols)} สถานี  |  ปี {annual.index.min()}–{annual.index.max()}")
    print("-" * 65)
    print("  วิเคราะห์ Mann–Kendall + Sen's Slope ...")

    # ── Analyse ──────────────────────────────────────────────────────────
    results = analyse_all_stations(annual, station_cols, alpha=0.05)

    # ── Export ───────────────────────────────────────────────────────────
    print("-" * 65)
    print("  กำลังส่งออกผลลัพธ์เป็น Excel ...")
    write_excel(results, annual, station_cols, output_xlsx)

    # ── Console summary table ────────────────────────────────────────────
    print("\n" + "=" * 65)
    print("  SUMMARY TABLE")
    print("=" * 65)
    hdr = f"{'Station':>10} | {'Mean(mm)':>9} | {'CV%':>5} | "
    hdr += f"{'Z':>7} | {'p-value':>10} | {'Slope':>8} | {'Trend':<12}"
    print(hdr)
    print("-" * len(hdr))
    for res in results:
        pv  = res["p-value (raw)"]
        sig = "*" if pv <= 0.05 else " "
        print(f"  {res['Station']:>8} | {res['Mean Annual RF (mm)']:>9.1f} | "
              f"{res['CV (%)']:>5.1f} | "
              f"{res['Z-statistic']:>+7.3f} | "
              f"{pv:>9.4f}{sig} | "
              f"{res['Sen slope (mm/yr)']:>+8.2f} | "
              f"{res['Trend']:<12}")
    print("=" * 65)
    print("  เสร็จสิ้น – ผลลัพธ์บันทึกใน Excel สำเร็จ ✓")


if __name__ == "__main__":
    main()
