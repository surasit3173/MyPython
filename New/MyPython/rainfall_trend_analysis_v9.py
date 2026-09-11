import os
import sys
import math
from datetime import datetime
from pathlib import Path
import numpy as np
import pandas as pd
from scipy import stats as sps
from scipy.stats import norm
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# --- Configuration: Option A: Classic (Journal of Hydrology) ---
ALPHA = 0.05
DPI = 600
COLOR_INC = '#006D2C' # เขียวเข้ม
COLOR_DEC = '#A50F15' # แดงเข้ม
COLOR_NS  = '#525252' # เทาเข้ม
COLOR_CI  = '#F7F7F7' # พื้นหลัง CI เทาอ่อนมาก

plt.rcParams.update({
    "font.family": "serif", "font.serif": ["Times New Roman"],
    "font.size": 10, "axes.titlesize": 12, "axes.labelsize": 10,
    "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": DPI
})

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║  §1  CORE ENGINE (MK, MMK, SEN, PETTITT)                                ║
# ╚══════════════════════════════════════════════════════════════════════════╝

def standard_mk_test(x):
    n = len(x)
    s = sum(np.sign(x[j] - x[i]) for i in range(n - 1) for j in range(i + 1, n))
    var_s = (n * (n - 1) * (2 * n + 5)) / 18.0
    z = (s - 1) / np.sqrt(var_s) if s > 0 else (s + 1) / np.sqrt(var_s) if s < 0 else 0
    p = 2 * (1 - norm.cdf(abs(z)))
    return {'z': z, 'p': p, 'sig': p < ALPHA}

def modified_mk_test(x):
    n = len(x)
    s = sum(np.sign(x[j] - x[i]) for i in range(n - 1) for j in range(i + 1, n))
    var_s = (n * (n - 1) * (2 * n + 5)) / 18.0
    ranks = sps.rankdata(x)
    rho = np.corrcoef(ranks[:-1], ranks[1:])[0, 1] if n > 2 else 0
    if rho > 0.1:
        n_ns = 1 + 2 * (rho**n / (n * (rho - 1)) - rho / (n * (rho - 1)**2))
        var_s *= (n / n_ns)
    z = (s - 1) / np.sqrt(var_s) if s > 0 else (s + 1) / np.sqrt(var_s) if s < 0 else 0
    p = 2 * (1 - norm.cdf(abs(z)))
    return {'z': z, 'p': p, 'sig': p < ALPHA}

def sens_slope_gilbert(x, alpha=0.05):
    n = len(x)
    slopes = [(x[j] - x[i]) / (j - i) for i in range(n - 1) for j in range(i + 1, n)]
    slopes = np.sort(slopes)
    N = len(slopes)
    var_s = (n * (n - 1) * (2 * n + 5)) / 18.0
    C_alpha = norm.ppf(1 - alpha / 2) * math.sqrt(var_s)
    low_idx = max(0, int(np.floor((N - C_alpha) / 2)))
    high_idx = min(N - 1, int(np.ceil((N + C_alpha) / 2)))
    return {'slope': float(np.median(slopes)), 'low_ci': float(slopes[low_idx]), 'high_ci': float(slopes[high_idx])}

def pettitt_test(x):
    n = len(x)
    r = sps.rankdata(x)
    U = [2 * np.sum(r[:t+1]) - (t + 1) * (n + 1) for t in range(n)]
    k_idx = np.argmax(np.abs(U))
    p_val = 2 * np.exp((-6 * np.max(np.abs(U))**2) / (n**3 + n**2))
    return {'idx': k_idx, 'p': p_val, 'sig': p_val < 0.05}

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║  §2  DYNAMIC LUXURY PLOTTING                                            ║
# ╚══════════════════════════════════════════════════════════════════════════╝

def plot_luxury_v6_5(stn, years, data_dict, stats_dict, out_path):
    fig = plt.figure(figsize=(10, 12))
    gs = gridspec.GridSpec(3, 1, hspace=0.32)
    keys = ['annual', 'wet', 'dry']
    titles = ['Annual Total Rainfall', 'Wet Season (May-Oct)', 'Dry Season (Nov-Apr)']

    for i, k in enumerate(keys):
        ax = fig.add_subplot(gs[i])
        y = data_dict[k]
        s = stats_dict[k]
        slp = sens_slope_gilbert(y)
        
        # --- Logic Color Determination ---
        if s['sig']:
            current_color = COLOR_INC if s['z'] > 0 else COLOR_DEC
            status_text = "SIGNIFICANT INCREASE" if s['z'] > 0 else "SIGNIFICANT DECREASE"
        else:
            current_color = COLOR_NS
            status_text = "NON-SIGNIFICANT (NS)"

        # Calculations for Plot
        intercept = np.median(y) - slp['slope'] * np.median(years)
        trend_y = slp['slope'] * years + intercept
        ci_low = slp['low_ci'] * years + (np.median(y) - slp['low_ci'] * np.median(years))
        ci_high = slp['high_ci'] * years + (np.median(y) - slp['high_ci'] * np.median(years))

        # Plotting
        ax.fill_between(years, ci_low, ci_high, color=COLOR_CI, alpha=0.9, label='95% CI (Gilbert)')
        ax.plot(years, y, color=current_color, marker='o', ms=5, mfc='white', mew=1.2, lw=1, alpha=0.4, label='Observed')
        ax.plot(years, trend_y, color=current_color, lw=2.5, label=f"Sen's Slope: {slp['slope']:.2f}")

        # Labeling
        ax.set_title(f"({chr(97+i)}) {titles[i]} - Stn {stn}", loc='left', weight='bold', color='#333333')
        ax.set_ylabel("Rainfall (mm)")
        ax.grid(True, axis='y', ls=':', alpha=0.6)
        
        # Result Box
        stats_box = f"MMK Z: {s['z']:.2f}\np-value: {s['p']:.4f}\n{status_text}"
        ax.text(0.02, 0.95, stats_box, transform=ax.transAxes, va='top', fontsize=9,
                bbox=dict(boxstyle='round,pad=0.4', fc='white', ec='#D5DBDB', alpha=0.9))
        
        if i == 0: ax.legend(loc='upper right', frameon=False, ncol=3, fontsize=8)

    plt.savefig(out_path, dpi=DPI, bbox_inches='tight')
    plt.close()

# ╔══════════════════════════════════════════════════════════════════════════╗
# ║  §3  MAIN PROCESS                                                       ║
# ╚══════════════════════════════════════════════════════════════════════════╝

def main():
    print(f"\nRainfall Trend Analysis v6.5 (Classic Journal Style)\n{'='*55}")
    
    try:
        data_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
        csv_file = next(data_dir.glob("*.csv"))
    except:
        print("Error: No CSV file found."); return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = Path(f"Output_Analysis_v6_5_{timestamp}")
    fig_dir = out_dir / "Luxury_Figures"
    fig_dir.mkdir(parents=True, exist_ok=True)

    # 1. Load Data
    df = pd.read_csv(csv_file)
    df.replace([-99, -999, -9999, 1e20], np.nan, inplace=True)
    df['date'] = pd.to_datetime({'year': df['YEAR'], 'month': df['MONTH'], 'day': df['DAY']})
    df = df.set_index('date').drop(['YEAR','MONTH','DAY'], axis=1).sort_index()
    df_clean = df.interpolate(method='time', limit=3, limit_area='inside')
    
    stns = df_clean.columns.tolist()
    annual = df_clean.resample('YS').sum(min_count=330)
    wet = df_clean[df_clean.index.month.isin(range(5,11))].resample('YS').sum(min_count=150)
    dry = df_clean[df_clean.index.month.isin([11,12,1,2,3,4])].resample('YS').sum(min_count=150)
    years = annual.index.year.values

    # 2. Table Setup
    wb = Workbook()
    sheets = {
        "1_Descriptive": ["Station", "Scale", "Mean", "Std", "CV%", "Max", "Min"],
        "2_Standard_MK": ["Station", "Scale", "Z", "p", "Sig"],
        "3_Modified_MK": ["Station", "Scale", "Z", "p", "Sig"],
        "4_Pettitt_CP": ["Station", "Scale", "CP_Year", "p", "Sig"],
        "5_Sens_Slope": ["Station", "Scale", "Slope", "Low_CI", "High_CI"],
        "6_Comparison": ["Station", "Scale", "MK_Sig", "MMK_Sig", "Agree"],
        "7_Research_Summary": ["Station", "Final_Assessment"]
    }
    
    ws_map = {}
    for i, (name, head) in enumerate(sheets.items()):
        ws = wb.active if i == 0 else wb.create_sheet(name)
        ws.title = name
        ws.append(head)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws_map[name] = ws

    # 3. Processing
    print(f"Processing {len(stns)} stations...")
    for stn in stns:
        stn_stats = {}
        findings = []
        
        for k, frame in zip(['annual', 'wet', 'dry'], [annual, wet, dry]):
            data = frame[stn].dropna()
            if len(data) < 10: continue
            
            mk = standard_mk_test(data.values)
            mmk = modified_mk_test(data.values)
            pet = pettitt_test(data.values)
            slp = sens_slope_gilbert(data.values)
            stn_stats[k] = mmk
            
            ws_map["1_Descriptive"].append([stn, k, data.mean(), data.std(), (data.std()/data.mean())*100, data.max(), data.min()])
            ws_map["2_Standard_MK"].append([stn, k, mk['z'], mk['p'], mk['sig']])
            ws_map["3_Modified_MK"].append([stn, k, mmk['z'], mmk['p'], mmk['sig']])
            ws_map["4_Pettitt_CP"].append([stn, k, years[pet['idx']], pet['p'], pet['sig']])
            ws_map["5_Sens_Slope"].append([stn, k, slp['slope'], slp['low_ci'], slp['high_ci']])
            ws_map["6_Comparison"].append([stn, k, mk['sig'], mmk['sig'], mk['sig']==mmk['sig']])
            
            if mmk['sig']: findings.append(f"{k} ({'Up' if mmk['z']>0 else 'Down'})")

        # Plot All Stations
        print(f"  > Creating Figure for: {stn}")
        data_vals = {k: frame[stn].values for k, frame in zip(['annual', 'wet', 'dry'], [annual, wet, dry])}
        plot_luxury_v6_5(stn, years, data_vals, stn_stats, fig_dir / f"Luxury_Trend_{stn}.png")
        ws_map["7_Research_Summary"].append([stn, ", ".join(findings) if findings else "Stable"])

    # Save
    excel_file = out_dir / f"Rainfall_Report_v6_5_{timestamp}.xlsx"
    wb.save(excel_file)
    print(f"\nDONE! Results in: {out_dir.name}\n")

if __name__ == "__main__":
    main()