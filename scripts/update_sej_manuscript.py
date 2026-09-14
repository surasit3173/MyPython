import os
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def update_docx():
    docx_path = 'Science Essence Journal/Rainfall_Occurrence_Intensity_Heterogeneity_SEJ_READY.docx'
    doc = docx.Document(docx_path)

    # 1. Update Table 2 (Row 12: Station 351012)
    table2 = doc.tables[1]
    # Check headers
    headers = [c.text.strip() for c in table2.rows[0].cells]
    print("Table 2 Headers:", headers)

    for row in table2.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if cells[0] == '351012':
            print("Updating Table 2 Row 351012 before:", cells)
            row.cells[1].text = '365.4'
            row.cells[2].text = '366.9'
            row.cells[3].text = '366.0'
            row.cells[4].text = '0.60'
            row.cells[5].text = '0.448'
            row.cells[6].text = 'Gumbel'
            print("Updating Table 2 Row 351012 after:", [c.text.strip() for c in row.cells])

    # 2. Add Footnote to Table 5
    table5 = doc.tables[4]
    print("Table 5 Headers:", [c.text.strip() for c in table5.rows[0].cells])
    # Add footnote text as paragraph below Table 5 or in caption
    # Let's inspect where Table 5 caption or text resides

    # 3. Update Text Paragraphs
    for i, p in enumerate(doc.paragraphs):
        text = p.text

        # Update Section 3.2 GEV Refit explanation
        if 'Equation (8) selected the log-Pearson type III' in text:
            p.text = (
                "Equation (8) selected the log-Pearson type III at six gauges, the Gumbel at five, "
                "and the generalized extreme value at two (Table 2). The selections were identical to those "
                "obtained with the uncorrected criterion at all 13 gauges. For station 351012, initial "
                "unconstrained GEV optimization suffered solver non-convergence, yielding an inflated AICc (436.3). "
                "Refitting with Nelder-Mead optimization initialized from Gumbel parameters restored global MLE "
                "convergence (GEV AICc = 366.9, ΔAICc = +1.43 relative to Gumbel), satisfying the theoretical "
                "model-nesting bound ΔAICc <= 2.413 for n=34. Gumbel remained the selected distribution for station 351012 "
                "both before and after refit, and selection-inclusive bootstrap replicates evaluate candidate fits independently "
                "per draw, ensuring full validity of all downstream bootstrap confidence intervals."
            )
            print(f"Updated P{i} (Section 3.2 GEV Refit)")

        # Update Section 3.3 Return Levels min/max station IDs
        if 'The ratio of Equation (12) is 3.83 at a 2-year return period' in text:
            p.text = (
                "Between-gauge differences are large and grow with return period (Table 3, Figure 2). "
                "The ratio of Equation (12) is 3.83 at a 2-year return period (ranging from 26.1 mm at Station 351010 to 99.9 mm at Station 351003), "
                "5.56 at 10 years, 7.67 at 25 years, and 12.53 at 100 years, with 100-year estimates spanning 57.8 mm at Station 351010 (Gumbel) "
                "to 724.7 mm at Station 351011 (LP3). These extrapolated return levels represent T-year quantiles and are distinct from "
                "observed daily maximum rainfalls (which span 54.2 mm at 351010 to 298.5 mm at 351012)."
            )
            print(f"Updated P{i} (Section 3.3 Return Levels)")

        # Update Figure 4 caption (linear scale)
        if 'The vertical scale is symmetric-logarithmic' in text:
            p.text = text.replace(
                "The vertical scale is symmetric-logarithmic so that the smaller errors remain legible beside the two largest.",
                "The vertical scale uses a linear axis to display provincial representation error percentages across all gauges."
            )
            print(f"Updated P{i} (Figure 4 Caption)")

        # Update Section 3.7 scenario return periods
        if 'Setting aside the two lowest-S gauges reduces' in text:
            p.text = (
                "Under a 100-year return period (T=100 yr), setting aside the two lowest-S gauges (Scenario B, n=11 gauges) "
                "reduces the between-gauge ratio from 12.53 to 4.70 and raises the network mean from 251.8 to 285.3 mm. "
                "Setting aside five gauges (Scenario C, n=8 gauges) gives exactly the same 100-year ratio, 4.70, and a mean of 290.3 mm, "
                "because the extreme estimates in the reduced network come from the same two gauges in both scenarios (Table 3, lower block)."
            )
            print(f"Updated P{i} (Section 3.7 Scenarios)")

        # Update Acknowledgements (remove anonymous reviewers)
        if 'and the anonymous reviewers for comments that improved the manuscript' in text:
            p.text = text.replace("and the anonymous reviewers for comments that improved the manuscript. ", "")
            print(f"Updated P{i} (Acknowledgements)")

    # Add Footnote text below Table 5
    # Find paragraph after Table 5
    for i, p in enumerate(doc.paragraphs):
        if 'Table 5: Provincial representation error' in p.text:
            p.text = (
                "Table 5: Provincial representation error of Equation (15) at 25- and 100-year return periods, "
                "for four ways of forming a provincial design value.\n"
                "* Note: In Median and 75th percentile rows, exact-zero representation errors where provincial estimates match local gauge values "
                "result in 12 non-zero errors across the 13 gauges, leaving median error statistics fully valid."
            )
            print(f"Updated Table 5 Caption/Footnote at P{i}")

    out_path = 'Science Essence Journal/Rainfall_Occurrence_Intensity_Heterogeneity_SEJ_READY.docx'
    doc.save(out_path)
    print("Saved updated DOCX to:", out_path)

if __name__ == '__main__':
    update_docx()
