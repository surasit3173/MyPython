import docx

def verify_alignment():
    print("========================================================")
    print(" STEP 3: VERIFY MANUSCRIPT ALIGNMENT WITH PIPELINE     ")
    print("========================================================")

    doc_path = 'Science Essence Journal/Rainfall_Occurrence_Intensity_Heterogeneity_SEJ_READY.docx'
    doc = docx.Document(doc_path)

    full_text = []
    for p in doc.paragraphs:
        full_text.append(p.text)

    all_text = "\n".join(full_text)

    checks = []

    # 1. Table 2 Row 351012 contains GEV AICc 366.9
    t2 = doc.tables[1]
    row_351012 = None
    for r in t2.rows[1:]:
        cells = [c.text.strip() for c in r.cells]
        if cells[0] == '351012':
            row_351012 = cells
            break

    print("Table 2 Row 351012:", row_351012)
    t2_pass = (row_351012 is not None and row_351012[2] == '366.9' and row_351012[6] == 'Gumbel')
    checks.append(("Table 2 Row 351012 contains GEV AICc 366.9 & Selected Gumbel", t2_pass))

    # 2. Table 2 does not contain 436.3
    t2_text = " ".join([c.text for r in t2.rows for c in r.cells])
    t2_no_436 = '436.3' not in t2_text
    checks.append(("Table 2 does not contain stale 436.3", t2_no_436))

    # 3. Text explains station 351012 refit (436.3 -> 366.9)
    refit_explained = ('436.3' in all_text) and ('366.9' in all_text) and ('solver non-convergence' in all_text)
    checks.append(("Body text accurately explains station 351012 solver refit", refit_explained))

    # 4. No 'anonymous reviewers' phrasing
    has_reviewers = 'anonymous reviewers' in all_text.lower()
    checks.append(("No 'anonymous reviewers' in text", not has_reviewers))

    # 5. Table 5 footnote present
    has_t5_footnote = 'exact-zero representation errors' in all_text
    checks.append(("Table 5 exact-zero footnote present", has_t5_footnote))

    # 6. Figure 4 caption updated
    has_linear_fig4 = 'linear axis' in all_text
    checks.append(("Figure 4 caption describes linear axis", has_linear_fig4))

    # 7. Station IDs 351010 and 351011 cited for min/max
    has_min_max_ids = 'Station 351010' in all_text and 'Station 351011' in all_text
    checks.append(("Station IDs 351010 & 351011 cited for min/max extremes", has_min_max_ids))

    print("\n--- Alignment Verification Results ---")
    all_pass = True
    for name, status in checks:
        status_str = "PASS" if status else "FAIL"
        print(f"  [{status_str}] {name}")
        if not status:
            all_pass = False

    if all_pass:
        print("\nManuscript Alignment Verification PASSED!")
    else:
        raise RuntimeError("Manuscript Alignment Verification FAILED!")

if __name__ == '__main__':
    verify_alignment()
